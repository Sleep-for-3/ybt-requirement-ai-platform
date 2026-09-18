from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import KnowledgeDocument, MartTable, Project, SourceTable
from test_requirement_snapshot_api import snapshot_api, snapshot_db


def add_document(db, project_id, name, status="indexed"):
    row = KnowledgeDocument(project_id=project_id, file_name=name, file_type="txt",
        source_type="synthetic", storage_path="isolated-test-not-read", document_status=status)
    db.add(row)
    db.flush()
    return row


def test_resource_selection_is_bounded_searchable_and_project_scoped(snapshot_api):
    client, db, project, _, _, membership = snapshot_api
    base = f"/projects/{project.id}/requirement-resources"
    active = add_document(db, project.id, "合成码值100%依据")
    add_document(db, project.id, "已归档依据", "archived")
    other = Project(name="隔离项目")
    db.add(other)
    db.flush()
    add_document(db, other.id, "其他项目依据")
    response = client.get(base)
    assert response.status_code == 200
    options = response.json()
    assert options["document_ids"]["items"] == [{"id": active.id, "name": active.file_name, "code": active.file_name}]
    assert options["source_table_ids"]["items"] and options["mart_table_ids"]["items"]
    assert "storage_path" not in response.text
    assert client.get(base, params={"q": "%"}).json()["document_ids"]["items"][0]["id"] == active.id
    assert client.get(base, params={"q": "不存在"}).json()["document_ids"]["items"] == []
    for index in range(101):
        db.add(MartTable(project_id=project.id, table_code=f"SYNTH_{index}", table_name="合成限量检验"))
    db.flush()
    bounded = client.get(base, params={"q": "合成限量"}).json()["mart_table_ids"]
    assert bounded["truncated"] and len(bounded["items"]) == 100
    assert client.get(f"/projects/{other.id}/requirement-resources").status_code == 404
    membership.project_role = "viewer"
    db.flush()
    restricted = client.get(base).json()["document_ids"]
    assert restricted["items"] == [] and restricted["available"] is False
    membership.status = "inactive"
    db.flush()
    assert client.get(base).status_code == 404


def test_requirement_persists_selected_resources_and_rejects_archived_or_foreign(snapshot_api):
    client, db, project, _, req, _ = snapshot_api
    document = add_document(db, project.id, "合成字典")
    source = db.scalar(select(SourceTable).where(SourceTable.project_id == project.id))
    mart = db.scalar(select(MartTable).where(MartTable.project_id == project.id))
    base = f"/projects/{project.id}/requirements"
    payload = {**req.scope_json, "name": req.name, "expected_version": 1,
               "document_ids": [document.id], "source_table_ids": [source.id], "mart_table_ids": [mart.id]}
    saved = client.put(f"{base}/{req.id}", json=payload)
    assert saved.status_code == 200
    assert saved.json()["version"] == 2
    loaded = client.get(base).json()[0]
    for key in ("document_ids", "source_table_ids", "mart_table_ids"):
        assert loaded[key] == payload[key]
    preview = client.get(f"{base}/{req.id}/document").json()
    assert len(preview["resources"]) == 3
    assert preview["resources"][0]["name"] == "合成字典"
    frozen = client.post(f"{base}/{req.id}/snapshots", json={
        "expected_version": 2, "expected_hash": preview["content_hash"]}).json()
    document.document_status = "archived"
    document.file_name = "后续改名"
    db.flush()
    reassessed = client.get(f"{base}/{req.id}/document").json()
    assert any(gap["id"] == f"scope:document_ids:{document.id}" for gap in reassessed["gaps"])
    assert "后续改名" not in str(reassessed["resources"])
    export = client.get(f"{base}/{req.id}/snapshots/{frozen['id']}/export")
    book = load_workbook(BytesIO(export.content))
    assert book["资料与数据范围"]["B2"].value == "合成字典"
    payload["expected_version"] = 2
    assert client.put(f"{base}/{req.id}", json=payload).status_code == 422
    payload["document_ids"] = []
    payload["source_table_ids"] = [source.id + 99999]
    assert client.put(f"{base}/{req.id}", json=payload).status_code == 422
    assert client.get(base).json()[0]["version"] == 2


def test_business_analyst_can_save_through_shared_guard_but_viewer_cannot(snapshot_api):
    client, db, project, _, req, membership = snapshot_api
    base = f"/projects/{project.id}/requirements"
    payload = {**req.scope_json, "name": "合成业务分析需求", "expected_version": 1}
    membership.project_role = "business_analyst"
    db.flush()
    created = client.post(base, json=payload)
    assert created.status_code == 201
    assert client.put(f"{base}/{req.id}", json=payload).status_code == 200
    membership.project_role = "viewer"
    db.flush()
    payload["expected_version"] = 2
    assert client.post(base, json=payload).status_code == 403
    assert client.put(f"{base}/{req.id}", json=payload).status_code == 403
