from copy import deepcopy
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from sqlalchemy import select

from app.models import (KnowledgeDocumentVersion, LineageEdge, Project, RequirementGenerationInput,
    ScriptFileVersion, TargetTable, TemplateDocument, TemplateVersion)
from app.services.lineage.ingestion import ScriptIngestionService
from app.services.mapping.requirement_input import build_requirement_input
from app.services.requirement_revisions import load_revision
from app.services.requirement_scope import content_digest
from app.services.knowledge_eligibility import validate_frozen_requirement_evidence
from test_requirement_generation_input import add_unit
from test_requirement_resources import add_document
from test_requirement_snapshot_api import snapshot_api, snapshot_db
from test_resource_batch_import import MemoryStorage


def seed(client, db, project, req, field, sql=None, storage=None):
    base = f"/projects/{project.id}/requirements/{req.id}"
    assert client.post(base + "/revisions", json={"expected_version": 1}).status_code == 201
    actor = db.scalar(select(__import__('app.models', fromlist=['User']).User))
    service = ScriptIngestionService(db, storage if storage is not None else MemoryStorage())
    script = service.ingest(project=project, data=(sql or "INSERT INTO bank.reg.report (amount) SELECT a.amount * 2 FROM bank.ods.accounts a WHERE a.amount > 0;").encode(),
        file_name="batch.sql", relative_path="batch.sql", dialect=None, actor_user_id=actor.id, build_revision=False)
    document = TemplateDocument(project_id=project.id, template_code="TEST", display_name="隔离模板", file_name="test.xlsx", file_type="xlsx", storage_path="never-read")
    db.add(document); db.flush()
    table = db.get(TargetTable, req.scope_json["target_table_id"])
    template = TemplateVersion(project_id=project.id, template_document_id=document.id, version_no=1,
        template_code="TEST", file_name="test.xlsx", file_type="xlsx", storage_path="never-read",
        file_hash="synthetic-template", parsed_snapshot_json=[{"table_code": table.table_code}])
    db.add(template); db.flush()
    preview = client.post(f"/projects/{project.id}/requirements/script-preview", json={"script_version_ids": [script.version.id]})
    assert preview.status_code == 200, preview.text
    data = preview.json()
    payload = {"script_version_ids": [script.version.id], "expected_content_version": 1,
        "preview_hash": data["preview_hash"], "template_version_id": template.id,
        "target_key": data["targets"][0]["key"], "field_bindings": {"amount": field.id}}
    return base, service, script, payload, data


def test_selected_versions_freeze_rules_and_do_not_follow_live_scripts(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    base, service, script, payload, preview = seed(client, db, project, req, field)
    excluded = service.ingest(project=project, data=b"INSERT INTO forbidden (x) SELECT secret FROM hidden;",
        file_name="outside.sql", relative_path="outside.sql", dialect=None, actor_user_id=None, build_revision=False)
    result = client.post(base + "/script-basis", json=payload)
    assert result.status_code == 201, result.text
    assert result.json()["revision"]["content_version"] == 2
    frozen = deepcopy(load_revision(db, project.id, req.id, 2).content_json)
    generation = build_requirement_input(db, load_revision(db, project.id, req.id, 2), [field.id], ["lineage"])
    assert "forbidden" not in str(generation) and "hidden" not in str(generation)
    rules = generation["script_basis"]["rules"]
    assert {r["script_version_id"] for r in rules} == {script.version.id}
    assert any(r["transformation_expression"] == "a.amount * 2" and r["filter_condition"] for r in rules)
    assert all(r["statement_id"] and r["source_line_start"] for r in rules)
    service.ingest(project=project, data=b"INSERT INTO bank.reg.report (amount) SELECT 999;",
        file_name="batch.sql", relative_path="batch.sql", dialect=None, actor_user_id=None, build_revision=False)
    assert load_revision(db, project.id, req.id, 2).content_json == frozen
    # Disabled historical edges are valid evidence for the explicitly selected old version.
    old = client.post(f"/projects/{project.id}/requirements/script-preview", json={"script_version_ids": [script.version.id]})
    assert old.json()["rules"] == preview["rules"]
    changes = client.get(base + "/script-basis-impact?content_version=2").json()
    assert changes["pending_review"] and any(c["code"] == f"script:{script.script_file.id}" for c in changes["changes"])
    assert load_revision(db, project.id, req.id, 2).content_json == frozen
    assert client.post(base + "/script-basis", json=payload).status_code == 409


def test_script_confirmation_authorization_scope_staleness_and_manual_preservation(snapshot_api):
    client, db, project, field, req, membership = snapshot_api
    base, _, script, payload, _ = seed(client, db, project, req, field)
    edit = client.put(base + f"/fields/{field.id}", json={"expected_content_version": 1,
        "section": "business", "changes": {"final_content": "人工正文不能被脚本覆盖"}})
    assert edit.status_code == 200, edit.text
    payload["expected_content_version"] = 2
    assert client.post(base + "/script-basis", json={**payload, "field_bindings": {"amount": 99999}}).status_code == 422
    assert client.post(base + "/script-basis", json={**payload, "preview_hash": "0" * 64}).status_code == 409
    other = Project(name="另一个隔离机构项目")
    db.add(other); db.flush()
    assert client.post(f"/projects/{other.id}/requirements/script-preview", json={"script_version_ids": [script.version.id]}).status_code == 404
    membership.project_role = "viewer"; db.flush()
    assert client.post(base + "/script-basis", json=payload).status_code == 403
    membership.project_role = "project_manager"; db.flush()
    attached = client.post(base + "/script-basis", json=payload)
    assert attached.status_code == 201, attached.text
    assert attached.json()["fields"][0]["business"]["final_content"] == "人工正文不能被脚本覆盖"
    assert attached.json()["manual_ownership"][str(field.id)]


def test_incomplete_parse_remains_gap(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    _, _, _, _, preview = seed(client, db, project, req, field,
        "INSERT INTO report SELECT * FROM missing_upstream;")
    assert any("缺少显式目标字段" in g for g in preview["gaps"])
    assert any("星号" in g for g in preview["gaps"])
    assert any("上游" in g for g in preview["gaps"])


@pytest.mark.parametrize("state", ["expired", "future", "withdrawn", "superseded", "draft", "project_scope", "institution_scope", "scenario"])
def test_ineligible_policy_never_enters_generation(snapshot_api, state):
    client, db, project, field, req, _ = snapshot_api
    doc = add_document(db, project.id, "合成制度")
    unit = add_unit(db, project.id, doc, "INELIGIBLE_POLICY_SENTINEL")
    version = db.get(KnowledgeDocumentVersion, unit.document_version_id)
    doc.current_version_id = version.id
    version.lifecycle_status = "active"
    now = datetime.now(timezone.utc)
    if state == "expired": version.expires_at = now - timedelta(days=1)
    elif state == "future": version.effective_at = now + timedelta(days=1)
    elif state in {"withdrawn", "superseded", "draft"}: version.lifecycle_status = state
    elif state == "project_scope": doc.applicable_project_ids_json = [99999]
    elif state == "institution_scope": doc.applicable_institution_names_json = ["其他银行"]
    elif state == "scenario": unit.scenario_id = 99999
    req.scope_json = {**req.scope_json, "document_ids": [doc.id]}; db.flush()
    base = f"/projects/{project.id}/requirements/{req.id}"
    assert client.post(base + "/revisions", json={"expected_version": 1}).status_code == 201
    result = build_requirement_input(db, load_revision(db, project.id, req.id, 1), [field.id], ["business"])
    assert not result["evidence"] and "INELIGIBLE_POLICY_SENTINEL" not in str(result)
    assert "缺少依据" in str(result["limitations"])


def test_frozen_policy_expiration_blocks_retry_without_replacing_evidence(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    doc = add_document(db, project.id, "合成有效制度")
    unit = add_unit(db, project.id, doc, "ACTIVE_POLICY")
    version = db.get(KnowledgeDocumentVersion, unit.document_version_id)
    doc.current_version_id = version.id; version.lifecycle_status = "active"
    req.scope_json = {**req.scope_json, "document_ids": [doc.id]}; db.flush()
    base = f"/projects/{project.id}/requirements/{req.id}"
    client.post(base + "/revisions", json={"expected_version": 1})
    frozen = build_requirement_input(db, load_revision(db, project.id, req.id, 1), [field.id], ["business"])
    assert frozen["evidence"][0]["unit_id"] == unit.id
    before = deepcopy(frozen)
    validate_frozen_requirement_evidence(db, project.id, frozen)
    version.lifecycle_status = "withdrawn"; db.flush()
    with pytest.raises(HTTPException, match="固定制度依据"):
        validate_frozen_requirement_evidence(db, project.id, frozen)
    assert frozen == before


def test_reverse_policy_is_frozen_and_excel_uses_same_snapshot(snapshot_api):
    from io import BytesIO
    from openpyxl import load_workbook
    from app.services.requirement_scope import export_document
    client, db, project, field, req, _ = snapshot_api
    doc = add_document(db, project.id, "选定的制度版本")
    unit = add_unit(db, project.id, doc, "FIXED_POLICY")
    version = db.get(KnowledgeDocumentVersion, unit.document_version_id)
    doc.current_version_id = version.id; version.lifecycle_status = "active"
    req.scope_json = {**req.scope_json, "document_ids": [doc.id]}; db.flush()
    base, _, _, payload, _ = seed(client, db, project, req, field)
    response = client.post(base + "/script-basis", json=payload)
    assert response.status_code == 201, response.text
    fixed = load_revision(db, project.id, req.id, 2)
    snapshot = deepcopy(fixed.content_json)
    assert snapshot["script_basis"]["policy_snapshot"]["evidence"][0]["content"] == "FIXED_POLICY"
    generation = build_requirement_input(db, fixed, [field.id], ["lineage"])
    assert generation["evidence"][0]["content"] == "FIXED_POLICY"
    unit.content = "CHANGED_POLICY"; db.flush()
    with pytest.raises(HTTPException):
        build_requirement_input(db, fixed, [field.id], ["lineage"])
    book = load_workbook(BytesIO(export_document(snapshot)))
    assert book["固定制度依据"]["D2"].value == "FIXED_POLICY"
    assert book["固定脚本事实"].max_row > 1
    assert "CHANGED_POLICY" not in str(snapshot)
    assert client.get(base + "/script-basis-impact?content_version=2").json()["pending_review"]


def test_cross_project_script_version_denied_even_with_both_memberships(snapshot_api):
    from app.models import ProjectMembership
    client, db, project, field, req, membership = snapshot_api
    _, _, script, _, _ = seed(client, db, project, req, field)
    other = Project(name="有权限的第二隔离项目")
    db.add(other); db.flush()
    db.add(ProjectMembership(project_id=other.id, user_id=membership.user_id,
        project_role="project_manager", status="active")); db.flush()
    response = client.post(f"/projects/{other.id}/requirements/script-preview", json={"script_version_ids": [script.version.id]})
    assert response.status_code == 404


def test_scope_revision_keeps_fixed_basis_and_unresolved_script_gaps(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    base, _, _, payload, _ = seed(client, db, project, req, field)
    assert client.post(base + "/script-basis", json=payload).status_code == 201
    before = deepcopy(load_revision(db, project.id, req.id, 2).content_json)
    updated = client.put(base, json={**req.scope_json, "name": req.name, "background": "新的业务背景",
        "expected_version": 1, "expected_content_version": 2})
    assert updated.status_code == 200, updated.text
    current = load_revision(db, project.id, req.id, 3).content_json
    assert current["script_basis"] == before["script_basis"]
    assert {g["id"] for g in before["gaps"] if g["id"].startswith("scope:script:")} <= {g["id"] for g in current["gaps"]}
    assert current["assessment"] == "gaps"


def test_candidate_cannot_cite_unselected_script_rule(snapshot_api, monkeypatch):
    from app.services.requirement_generation_worker import generate_candidate
    client, db, project, field, req, _ = snapshot_api
    base, _, _, payload, _ = seed(client, db, project, req, field)
    assert client.post(base + "/script-basis", json=payload).status_code == 201
    snapshot = build_requirement_input(db, load_revision(db, project.id, req.id, 2), [field.id], ["lineage"])
    async def fake_model(*args, **kwargs):
        return {"final_content": "仅合成测试", "script_rule_ids": ["script-edge-999999"]}
    monkeypatch.setattr("app.services.requirement_generation_worker.execute_runtime_chat", fake_model)
    with pytest.raises(HTTPException, match="范围外脚本规则"):
        generate_candidate(db, SimpleNamespace(input_json=snapshot), SimpleNamespace(field_id=field.id, section="lineage"), project)


def test_legacy_external_profile_uses_safe_requirement_input_budget(snapshot_api, monkeypatch):
    from app.services.llm.prompt_runtime import PromptRuntime
    from app.services.requirement_generation_worker import generate_candidate

    client, db, project, field, req, _ = snapshot_api
    base, _, _, payload, _ = seed(client, db, project, req, field)
    assert client.post(base + "/script-basis", json=payload).status_code == 201
    snapshot = build_requirement_input(db, load_revision(db, project.id, req.id, 2), [field.id], ["lineage"])
    runtime = PromptRuntime(
        prompt_key="requirement_field_candidate",
        version=1,
        system_prompt="system",
        user_template="{target}",
        model_profile_id=None,
        provider_type="openai_compatible",
        base_url="https://provider.example.com/v1",
        model_name="example-model",
        api_key_env_name="MODEL_API_KEY",
        local_only=False,
        config={},
    )

    async def fake_model(*_args, **_kwargs):
        return {"final_content": "外部模型候选"}

    monkeypatch.setattr("app.services.requirement_generation_worker.get_prompt_runtime", lambda *_args: runtime)
    monkeypatch.setattr("app.services.requirement_generation_worker.execute_runtime_chat", fake_model)

    candidate = generate_candidate(
        db,
        SimpleNamespace(input_json=snapshot),
        SimpleNamespace(field_id=field.id, section="lineage"),
        project,
    )
    assert candidate["final_content"] == "外部模型候选"
    assert candidate["runtime"]["provider"] == "openai_compatible"
    assert candidate["runtime"]["test_provider"] is False


def test_requirement_generation_block_reason_preserves_actionable_classification():
    from app.services.requirement_generation_worker import blocked_reason_code

    assert blocked_reason_code(HTTPException(403, "denied")) == "generation_permission_denied"
    assert blocked_reason_code(HTTPException(409, "changed")) == "generation_input_conflict"
    assert blocked_reason_code(HTTPException(422, "invalid candidate")) == "generation_validation_blocked"
    assert blocked_reason_code(HTTPException(503, "provider")) == "policy_or_scope_blocked"
