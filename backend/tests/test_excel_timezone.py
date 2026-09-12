"""Workbook exporters must survive timezone-aware database values.

PostgreSQL returns ``timestamptz`` columns as aware datetimes while SQLite
returns naive ones, and openpyxl refuses aware datetimes outright.  These
regressions keep the SQLite-only test suite honest about production data.
"""

from __future__ import annotations

from datetime import UTC, datetime, timezone
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from app.core.excel import excel_value
from app.services.deliverables.workbook import render_workbook
from app.services.lineage.exporter import _write
from app.services.uat.reporting import build_uat_report


def test_excel_value_normalizes_aware_datetimes_to_utc_without_tzinfo() -> None:
    aware = datetime(2026, 9, 12, 17, 30, tzinfo=timezone.utc)
    normalized = excel_value(aware)
    assert normalized == datetime(2026, 9, 12, 17, 30)
    assert normalized.tzinfo is None
    naive = datetime(2026, 9, 12, 17, 30)
    assert excel_value(naive) is naive
    assert excel_value("普通文本") == "普通文本"
    assert excel_value(None) is None


def test_deliverable_render_writes_timezone_aware_values() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "审核记录"
    sheet.append(["动作", "审核时间"])
    stream = BytesIO()
    workbook.save(stream)
    sheets = [SimpleNamespace(
        id=1, enabled=True, sheet_name="审核记录", business_section="review_record",
        header_row_end=1, data_start_row=2, repeat_direction="vertical",
    )]
    columns = [
        SimpleNamespace(template_sheet_mapping_id=1, business_field="action", excel_column="A", default_value=None, required=True, write_mode="overwrite", merge_strategy="none"),
        SimpleNamespace(template_sheet_mapping_id=1, business_field="approved_at", excel_column="B", default_value=None, required=True, write_mode="overwrite", merge_strategy="none"),
    ]
    approved_at = datetime(2026, 9, 12, 9, 58, 56, tzinfo=UTC)
    rendered, warnings = render_workbook(
        stream.getvalue(), sheets, columns,
        {"review_record": [{"action": "approve", "approved_at": approved_at}]},
    )
    assert warnings == []
    cell = load_workbook(BytesIO(rendered))["审核记录"]["B2"]
    assert cell.value == datetime(2026, 9, 12, 9, 58, 56)
    assert cell.value.tzinfo is None


def test_lineage_workbook_writer_accepts_aware_review_timestamps() -> None:
    workbook = Workbook()
    sheet = workbook.active
    _write(sheet, ["任务 ID", "审核时间"], [[1, datetime(2026, 9, 12, 1, 2, 3, tzinfo=UTC)]])
    stream = BytesIO()
    workbook.save(stream)
    saved = load_workbook(BytesIO(stream.getvalue()))["Sheet"]
    assert saved["B2"].value == datetime(2026, 9, 12, 1, 2, 3)


class _FakeResult:
    def __init__(self, items):
        self._items = items

    def all(self):
        return self._items


class _FakeSession:
    """Minimal session stub so the report builder sees aware datetimes.

    SQLite normalizes ``timestamptz`` values to naive datetimes, so the shared
    test database can never reproduce the production crash; this stub feeds
    PostgreSQL-shaped rows into the exporter instead.
    """

    def __init__(self, suite, cases, results, findings, signoffs):
        self._suite = suite
        self._queue = [cases, results, findings, signoffs]

    def get(self, _model, _ident):
        return self._suite

    def scalars(self, _statement):
        return _FakeResult(self._queue.pop(0))


def test_uat_report_writes_timezone_aware_timestamps() -> None:
    started_at = datetime(2026, 9, 12, 3, 4, 5, tzinfo=UTC)
    resolved_at = datetime(2026, 9, 12, 6, 7, 8, tzinfo=UTC)
    suite = SimpleNamespace(id=1, suite_name="冒烟套件", suite_type="smoke", description="说明")
    cases = [SimpleNamespace(id=1, case_code="C-1", case_name="案例", execution_mode="manual", severity="high")]
    results = [SimpleNamespace(
        id=1, uat_case_id=1, status="passed", duration_ms=12,
        expected_result_json={}, actual_result_json={}, evidence_json={},
    )]
    findings = [SimpleNamespace(
        finding_no=1, finding_type="bug", severity="low", title="标题", status="closed",
        assigned_role="qa", description="描述", resolution_text="已修", resolved_at=resolved_at, verified_at=resolved_at,
    )]
    signoffs = [SimpleNamespace(
        signoff_role="qa", signoff_status="signed", comment="同意", signed_by=1, signed_at=resolved_at,
    )]
    run = SimpleNamespace(
        project_id=1, id=9, run_name="UAT Run", status="passed", uat_suite_id=1,
        environment_name="local", application_version="1.0.0",
        started_at=started_at, completed_at=resolved_at, git_commit_sha="deadbeef", run_no="R-1",
    )
    session = _FakeSession(suite, cases, results, findings, signoffs)
    saved = load_workbook(BytesIO(build_uat_report(session, run)))
    assert saved["环境信息"]["C2"].value == datetime(2026, 9, 12, 3, 4, 5)
    assert saved["环境信息"]["C2"].value.tzinfo is None
    assert saved["签署记录"]["E2"].value == datetime(2026, 9, 12, 6, 7, 8)
    assert saved["修复记录"]["D2"].value == datetime(2026, 9, 12, 6, 7, 8)
