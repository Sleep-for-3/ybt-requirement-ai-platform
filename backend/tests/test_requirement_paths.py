from copy import deepcopy
from io import BytesIO

import pytest
from sqlalchemy import select, delete
from openpyxl import load_workbook
from docx import Document

from app.models import (CatalogColumn, CatalogSchema, CatalogTable, DataSource, KnowledgeDocumentVersion,
    TemplateDocument, TemplateVersion, ReviewTask)
from app.services.requirement_revisions import load_revision
from app.services.requirement_paths import path_preview, path_issues
from app.services.requirement_review import review_readiness, submit_for_review, finalize_formal_delivery
from app.services.requirement_scope import export_document
from app.services.requirement_word import export_word
from test_requirement_snapshot_api import snapshot_api, snapshot_db
from test_requirement_script_basis import seed
from test_requirement_resources import add_document
from test_requirement_generation_input import add_unit


def catalog(db, project, database="bank"):
    source = DataSource(project_id=project.id, name=f"isolated-{database}", db_type="offline", enabled=False)
    db.add(source); db.flush()
    schema = CatalogSchema(project_id=project.id, datasource_id=source.id, schema_name="ods")
    db.add(schema); db.flush()
    table = CatalogTable(project_id=project.id, datasource_id=source.id, catalog_schema_id=schema.id,
        database_name=database, schema_name="ods", table_name="accounts")
    db.add(table); db.flush()
    db.add(CatalogColumn(project_id=project.id, datasource_id=source.id, catalog_table_id=table.id,
        database_name=database, schema_name="ods", table_name="accounts", column_name="amount", data_type="DECIMAL"))
    db.flush()
    return table


def setup_path(sample, multilayer=False, missing=False, storage=None):
    client, db, project, field, req, membership = sample
    from app.models import MartToYbtMapping
    db.execute(delete(MartToYbtMapping).where(MartToYbtMapping.project_id == project.id,
        MartToYbtMapping.target_field_id == field.id))
    if not missing:
        catalog(db, project)
    doc = add_document(db, project.id, "路径核验隔离制度")
    doc.source_category = "regulatory_formal"
    unit = add_unit(db, project.id, doc, "余额取已核验账户期末金额。")
    version = db.get(KnowledgeDocumentVersion, unit.document_version_id)
    version.lifecycle_status = "active"; doc.current_version_id = version.id
    req.scope_json = {**req.scope_json, "document_ids": [doc.id], "mart_table_ids": [], "source_table_ids": []}; db.flush()
    sql = "INSERT INTO bank.reg.report (amount) SELECT amount FROM bank.ods.accounts;"
    if multilayer:
        sql = "INSERT INTO bank.dwd.accounts (amount) SELECT amount FROM bank.ods.accounts;"
    base, service, script, payload, _ = seed(client, db, project, req, field, sql, storage=storage)
    if multilayer:
        for index, (src, dst) in enumerate((("dwd.accounts", "dws.accounts"), ("dws.accounts", "reg.report"))):
            result = service.ingest(project=project,
                data=f"INSERT INTO bank.{dst} (amount) SELECT amount FROM bank.{src};".encode(),
                file_name=f"layer-{index}.sql", relative_path=f"layer-{index}.sql", dialect=None,
                actor_user_id=membership.user_id, build_revision=False)
            payload["script_version_ids"].append(result.version.id)
    template = db.get(TemplateVersion, payload["template_version_id"])
    template.status = "active"
    db.get(TemplateDocument, template.template_document_id).current_version_id = template.id
    db.flush()
    preview = client.post(f"/projects/{project.id}/requirements/script-preview",
        json={"script_version_ids": payload["script_version_ids"]}).json()
    payload["preview_hash"] = preview["preview_hash"]
    payload["target_key"] = next(t["key"] for t in preview["targets"] if t["table_name"] == "report")
    result = client.post(base + "/script-basis", json=payload)
    assert result.status_code == 201, result.text
    return base, unit, payload


def confirm(client, base, version=2):
    view = client.get(base + f"/paths?content_version={version}")
    assert view.status_code == 200, view.text
    return client.post(base + "/paths", json={"expected_content_version": version,
        "preview_hash": view.json()["preview_hash"], "rationale": "已逐条核验来源元数据、完整加工依赖及写入字段"}), view.json()


def test_layer_rename_updates_new_basis_but_preserves_frozen_requirement(snapshot_api):
    from app.services import data_architecture as architecture
    client, db, project, _, req, _ = snapshot_api
    base, _, payload = setup_path(snapshot_api)
    definition = architecture.examples()[0]
    architecture.save_architecture(db, definition, 0, project_id=project.id)
    table = db.scalar(select(CatalogTable).where(CatalogTable.project_id == project.id,
        CatalogTable.table_name == "accounts"))
    architecture.assign_table(db, project.id, table.id, {"layer_key": "layer_1"}, 1)
    preview = client.post(f"/projects/{project.id}/requirements/script-preview",
        json={"script_version_ids": payload["script_version_ids"]}).json()
    payload.update(expected_content_version=2, preview_hash=preview["preview_hash"])
    assert client.post(base + "/script-basis", json=payload).status_code == 201
    frozen = deepcopy(load_revision(db, project.id, req.id, 3).content_json)
    definition["layers"][1]["name"] = "已更名贴源层"
    architecture.save_architecture(db, definition, 1, project_id=project.id)
    current = client.post(f"/projects/{project.id}/requirements/script-preview",
        json={"script_version_ids": payload["script_version_ids"]}).json()
    assert current["metadata"][0]["assignment"]["layer_name"] == "已更名贴源层"
    assert frozen["script_basis"]["metadata"][0]["assignment"]["layer_name"] == "贴源层"
    impact = client.get(base + "/script-basis-impact?content_version=3").json()
    assert impact["pending_review"]
    assert any(change["code"] == "script:metadata" for change in impact["changes"])
    assert load_revision(db, project.id, req.id, 3).content_json == frozen


@pytest.mark.parametrize("multilayer", [False, True])
def test_actual_paths_without_mart_can_complete_review_and_fixed_exports(snapshot_api, multilayer, monkeypatch, tmp_path):
    from app.services.auth.dependencies import Principal
    from app.services.governance.workflow import decide_task
    from app.services.storage import get_storage_service
    from test_requirement_formal_delivery import _reviewers
    client, db, project, field, req, membership = snapshot_api
    base, unit, _ = setup_path(snapshot_api, multilayer)
    original = deepcopy(load_revision(db, project.id, req.id, 2).content_json)
    response, view = confirm(client, base)
    assert not view["paths"][0]["issues"], view
    assert response.status_code == 201, response.text
    content = load_revision(db, project.id, req.id, 3).content_json
    assert not content["fields"][0]["mart_mappings"]
    assert len(content["fields"][0]["confirmed_path"]["rule_ids"]) == (3 if multilayer else 1)
    assert path_issues(content) == []
    graph = client.get(base + "/revisions/3/lineage").json()["graph"]
    assert graph["facts_mode"] == "requirement_fixed_scripts"
    assert len(graph["edges"]) == (3 if multilayer else 1)
    assert len(graph["tables"]) == (4 if multilayer else 2)
    assert all(t["id"].startswith("script-table:") and '"' not in t["id"] for t in graph["tables"])
    assert {n["table_key"] for n in graph["nodes"]} <= {t["id"] for t in graph["tables"]}
    assert all(edge["evidence_refs"][0]["version_no"] == 1 for edge in graph["edges"])
    assert not any(g["id"].endswith(":mapping") for g in content["gaps"])
    assert load_revision(db, project.id, req.id, 2).content_json == original
    policy = client.get(base + "/policy-comparison?content_version=3").json()
    response = client.post(base + "/policy-comparison", json={"expected_content_version": 3,
        "basis_hash": policy["basis_hash"], "decisions": [{"unit_id": unit.id, "status": "matched",
            "rule_ids": [r["rule_id"] for r in policy["rules"]], "rationale": "合成制度与原值取数一致"}]})
    assert response.status_code == 201, response.text
    response = client.put(base + f"/fields/{field.id}", json={"expected_content_version": 4,
        "section": "business", "changes": {"business_definition": "已核验的期末金额", "final_content": "期末账户金额"}})
    assert response.status_code == 200, response.text
    ready = review_readiness(db, project.id, req.id, 5)
    assert ready["eligible"], ready
    monkeypatch.setenv("STORAGE_DIR", str(tmp_path / "path-deliveries")); get_storage_service.cache_clear()
    reviewers = _reviewers(db, project.id)
    revision = load_revision(db, project.id, req.id, 5)
    submission, _ = submit_for_review(db, project_id=project.id, requirement_id=req.id,
        content_version=5, content_hash=revision.content_hash, submitted_by=membership.user_id,
        assignments={role: user.id for role, user in reviewers.items()})
    tasks = {task.step_key: task for task in db.scalars(select(ReviewTask).where(ReviewTask.target_id == submission.id,
        ReviewTask.target_type == "requirement_review_submission"))}
    for step, role in (("business_review", "business_reviewer"), ("technical_review", "technical_reviewer"), ("final_review", "final_reviewer")):
        user = reviewers[role]
        decide_task(db, tasks[step], Principal(user.id, user.username, None), "approved", "隔离路径审核")
    delivery, _ = finalize_formal_delivery(db, project_id=project.id, requirement_id=req.id, submission_id=submission.id)
    book = load_workbook(BytesIO(export_document(delivery.content_json)))
    text = "\n".join(p.text for p in Document(BytesIO(export_word(delivery.content_json))).paragraphs)
    assert "已审核正式交付" in text
    assert book["需求范围"]["B3"].value == 5
    assert book["已确认字段路径"]["B2"].value == "已确认"
    assert "路径核验依据" in text
    db.commit()
    excel_response = client.get(base + f"/formal-deliveries/{delivery.id}/export")
    word_response = client.get(base + f"/formal-deliveries/{delivery.id}/export?format=docx")
    assert excel_response.status_code == word_response.status_code == 200
    assert excel_response.headers["x-requirement-snapshot-hash"] == word_response.headers["x-requirement-snapshot-hash"]
    get_storage_service.cache_clear()


def test_missing_upstream_and_database_ambiguity_are_not_confirmable(snapshot_api):
    client, db, project, _, _, _ = snapshot_api
    catalog(db, project, "another_bank")
    base, _, _ = setup_path(snapshot_api, missing=True)
    result, view = confirm(client, base)
    assert result.status_code == 409
    assert any("缺少上游" in issue for issue in view["paths"][0]["issues"])


def test_path_reconfirmation_after_technical_edit_preserves_old_version(snapshot_api):
    from app.services.requirement_revisions import edit_field
    client, db, project, field, req, membership = snapshot_api
    base, _, _ = setup_path(snapshot_api)
    response, view = confirm(client, base)
    assert response.status_code == 201, response.text
    before = deepcopy(load_revision(db, project.id, req.id, 3).content_json)
    revision = edit_field(db, req, 3, field.id, "lineage", {"processing_logic": "人工补充已核验加工说明"}, membership.user_id)
    db.commit()
    assert path_issues(revision.content_json)
    assert confirm(client, base, 4)[0].status_code == 201
    assert load_revision(db, project.id, req.id, 3).content_json == before


def test_paths_permission_and_optimistic_version(snapshot_api):
    client, db, _, _, _, membership = snapshot_api
    base, _, _ = setup_path(snapshot_api)
    view = client.get(base + "/paths?content_version=2").json()
    payload = {"expected_content_version": 2, "preview_hash": view["preview_hash"], "rationale": "已核验"}
    membership.project_role = "viewer"; db.flush()
    assert client.post(base + "/paths", json=payload).status_code == 403
    membership.project_role = "project_manager"; db.flush()
    assert client.post(base + "/paths", json={**payload, "preview_hash": "0" * 64}).status_code == 409
    assert client.post(base + "/paths", json=payload).status_code == 201
    assert client.post(base + "/paths", json=payload).status_code == 409


@pytest.mark.parametrize("kind", ["cycle", "multiple_writers", "star", "temporary"])
def test_incomplete_paths_remain_blocked(snapshot_api, kind):
    client, db, project, _, req, _ = snapshot_api
    setup_path(snapshot_api)
    content = deepcopy(load_revision(db, project.id, req.id, 2).content_json)
    rule = next(r for r in content["script_basis"]["rules"] if r["target"].get("column_name") == "amount")
    if kind == "cycle":
        rule["source"] = deepcopy(rule["target"])
    elif kind == "multiple_writers":
        other = {**deepcopy(rule), "rule_id": "another-write", "statement_id": 999}
        content["script_basis"]["rules"].append(other)
    elif kind == "star":
        rule["source"]["column_name"] = "*"
    else:
        rule["source"]["temporary"] = True
    assert path_preview(content)["paths"][0]["issues"]
