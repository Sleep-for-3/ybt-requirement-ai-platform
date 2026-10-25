from __future__ import annotations

import json
from io import BytesIO

from docx import Document

from docx_template_fixtures import build_customer_template, customer_config
from test_requirement_snapshot_api import snapshot_api, snapshot_db  # noqa: F401  测试夹具复用
from test_requirement_xlsx_template import build_template as build_xlsx_template
from test_requirement_xlsx_template import SAMPLE_ROW

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_files(template: bytes, name: str = "客户需求模板.xlsx"):
    return {"template": (name, template, XLSX_TYPE)}


def _xlsx_config(template: bytes) -> dict:
    from app.services.requirement_xlsx_template import template_sha256

    return {
        "template_hash": template_sha256(template),
        "config_version": "p1-b-api-test-1",
        "supported_scope": "requirement_document",
        "cell_map": {
            "需求名称": {"named_range": "需求名称", "value_path": "requirement.name", "required": True},
        },
        "row_blocks": {
            "字段清单": {
                "sheet": "字段口径",
                "sample_row": SAMPLE_ROW,
                "columns": {"编码": "A", "名称": "B", "类型": "C", "口径": "D", "来源": "E"},
                "column_paths": {
                    "编码": "field.field_code",
                    "名称": "field.field_name",
                    "类型": "field.field_type",
                    "口径": "business.final_content",
                    "来源": "lineage.source_system_name",
                },
                "max_rows": 50,
                "source": "fields",
            },
        },
        "required_keys": ["需求名称", "字段清单"],
    }


def _files(template: bytes, name: str = "客户需求模板.docx"):
    return {"template": (name, template, DOCX_TYPE)}


def _frozen_snapshot(client, base: str, req) -> int:
    current = client.get(base + "/document").json()
    created = client.post(base + "/snapshots", json={
        "expected_version": req.version, "expected_hash": current["content_hash"]})
    assert created.status_code == 201, created.text
    return created.json()["id"]


def test_template_export_requires_frozen_content(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    base = f"/projects/{project.id}/requirements/{req.id}"
    template = build_customer_template()
    response = client.post(base + "/template-exports", files=_files(template),
                           data={"format": "docx", "config": "{}"})
    assert response.status_code == 409
    detail = response.json()["detail"]
    assert "冻结" in detail and "snapshot_id" in detail


def test_template_export_rejects_bad_config_and_format(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    base = f"/projects/{project.id}/requirements/{req.id}"
    snapshot_id = _frozen_snapshot(client, base, req)
    template = build_customer_template()

    bad_json = client.post(base + "/template-exports", files=_files(template),
                           data={"format": "docx", "config": "{not json", "snapshot_id": snapshot_id})
    assert bad_json.status_code == 422
    assert "合法 JSON" in bad_json.json()["detail"]

    bad_format = client.post(base + "/template-exports", files=_files(template),
                             data={"format": "pdf", "config": "{}", "snapshot_id": snapshot_id})
    assert bad_format.status_code == 422
    assert "unsupported_format" in bad_format.json()["detail"]


def test_check_endpoint_renders_without_returning_file_and_audits(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    base = f"/projects/{project.id}/requirements/{req.id}"
    snapshot_id = _frozen_snapshot(client, base, req)
    template = build_customer_template()
    config = customer_config(template)

    checked = client.post(base + "/template-exports/check", files=_files(template),
                          data={"format": "docx", "config": json.dumps(config, ensure_ascii=False),
                                "snapshot_id": snapshot_id})
    assert checked.status_code == 200, checked.text
    payload = checked.json()
    assert payload["format"] == "docx"
    assert payload["file_size"] > 0
    assert len(payload["template_sha256"]) == 64
    assert payload["renderer_version"]
    assert payload["config_version"] == "customer-p0d-test-1"

    downloaded = client.post(base + "/template-exports", files=_files(template),
                             data={"format": "docx", "config": json.dumps(config, ensure_ascii=False),
                                   "snapshot_id": snapshot_id})
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.headers["content-type"].startswith(DOCX_TYPE)
    assert downloaded.headers["x-requirement-template-sha256"] == payload["template_sha256"]
    assert len(downloaded.headers["x-requirement-snapshot-hash"]) == 64
    assert downloaded.headers["x-requirement-rendered-file-sha256"] == payload["file_sha256"]
    body = Document(BytesIO(downloaded.content))
    assert body.paragraphs  # 生成的是可打开的 docx


def test_xlsx_template_export_renders_with_identity_headers(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    base = f"/projects/{project.id}/requirements/{req.id}"
    snapshot_id = _frozen_snapshot(client, base, req)
    template = build_xlsx_template()
    config = _xlsx_config(template)
    data = {"format": "xlsx", "config": json.dumps(config, ensure_ascii=False),
            "snapshot_id": snapshot_id}

    checked = client.post(base + "/template-exports/check", files=_xlsx_files(template), data=data)
    assert checked.status_code == 200, checked.text
    payload = checked.json()
    assert payload["format"] == "xlsx"
    assert payload["file_size"] > 0
    assert len(payload["template_sha256"]) == 64
    assert payload["config_version"] == "p1-b-api-test-1"

    downloaded = client.post(base + "/template-exports", files=_xlsx_files(template), data=data)
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.headers["content-type"].startswith(XLSX_TYPE)
    assert downloaded.headers["x-requirement-rendered-file-sha256"] == payload["file_sha256"]
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(downloaded.content))
    assert workbook.sheetnames  # 产物是可打开的 xlsx


def test_template_export_is_blocked_without_deliverable_export_permission(snapshot_api):
    client, db, project, field, req, membership = snapshot_api
    base = f"/projects/{project.id}/requirements/{req.id}"
    snapshot_id = _frozen_snapshot(client, base, req)
    template = build_customer_template()
    membership.project_role = "viewer"
    db.flush()
    response = client.post(base + "/template-exports", files=_files(template),
                           data={"format": "docx", "config": "{}", "snapshot_id": snapshot_id})
    assert response.status_code in {401, 403}
