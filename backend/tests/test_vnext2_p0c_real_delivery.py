"""VNEXT2 P0-C：真实监管需求生成 → 候选采用 → 人工修订 → 三级审核 → 正式交付全链路（API 驱动）。

场景（合成，不代表所有银行口径）：
    一表通「EAST客户报送表.CUSTOMER_ID（客户统一标识）」的监管开发需求。
    来源：借据 LOAN_IOU → 合同 CONTRACT → 产品 PRODUCT / 产品期限 PRODUCT_TERM → 客户 CUSTOMER，
    经数据集市 MART.REG_CUSTOMER 报送一表通。合成制度明确：产品维度需按 产品ID+期限 组合业务键
    关联，且产品ID存在重复，无法唯一关联时必须形成待确认问题，不得自定口径。
    业务规则由本测试内置的合成制度文本、集市映射与人工确认结果支撑。

链路（除夹具预置的主数据外全部走 HTTP API）：
    1 创建需求（scope 含目标表/字段/场景/资料来源/集市资料）
    2 冻结内容版本 revisions
    3 generation-inputs → generation-runs（inline 队列真实执行 worker）
    4 轮询 generation-runs / generation-items，读取候选差异、来源与证据
    5 generation-candidates/adopt 真实采用（候选 hash 来自 worker 落库结果）
    6 PUT fields 人工补业务口径（形成新修订）
    7 review-readiness → review-submissions → 业务/技术/终审三级 approve → finalize
    8 导出 xlsx/docx，校验快照哈希头与正文内容
    9 交付后新增待确认问题 → 历史交付不变；对已确认版本发起修订 → 历史交付仍不变

真实性级别：
    - 产品链路（队列/输入冻结/候选/采用/修订/审核/交付/导出）：REAL_COMPONENT（SQLite + TestClient +
      inline 队列，真实 worker 代码路径）。
    - 模型：默认用测试替身（MOCK）验证链路；设置 `VNEXT2_P0C_REAL_MODEL=1` 时使用
      backend/.env 配置的真实模型网关（REAL_MODEL）。
    - 夹具预置：机构/项目/用户/目标表/来源表/集市表/映射/证据/知识文档与知识单元直接 ORM 写入。
      产品操作路径分别由知识上传审核启用（test_knowledge_rag）、元数据导入（test_metadata_catalog）、
      映射确认（test_lineage_paths 等）覆盖；本测试聚焦「生成→采用→审核→交付」链。
"""

from __future__ import annotations

import hashlib
import os
from datetime import datetime, timezone

import pytest
from docx import Document as DocxDocument
from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO
from sqlalchemy import select

from app.api import deliverables, requirements, review_tasks
from app.core.database import get_db
from app.models import (
    BusinessSystem, KnowledgeDocument, KnowledgeDocumentVersion, KnowledgeUnit,
    ProductScenario, ProjectMembership, Requirement, SourceField, SourceTable,
)
from app.services.auth.dependencies import Principal, get_current_principal
from app.services.auth.resource_guard import guard_project_resource
from app.services.storage import get_storage_service
from test_lineage_paths import _seed_assets
from test_requirement_formal_delivery import _reviewers
from test_requirement_snapshot_api import snapshot_db  # noqa: F401  隔离 SQLite 夹具

SCENARIO = "需求：按合成制度从借据/合同/产品及产品期限关联客户信息，报送一表通客户统一标识"
RULE_V1 = (
    "【合成监管制度 2026-v1】第十条 一表通客户统一标识填报口径："
    "报送字段 CUSTOMER_ID（客户统一标识）应取自客户主数据；"
    "需经借据与合同关联到产品信息时，产品维度必须使用“产品ID+期限编号”组合业务键，"
    "不得仅以产品ID关联，避免同一产品不同期限方案造成一对多记录膨胀。"
)
RULE_V2 = (
    "【合成监管制度 2026-v2】第十条（修订）：客户统一标识改为以客户号优先，"
    "产品期限仅用于统计分组，不再作为客户标识的关联键。"
)


@pytest.fixture()
def p0c(snapshot_db, tmp_path, monkeypatch):
    from app.core.settings import get_settings

    monkeypatch.setenv("STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("TASK_QUEUE_PROVIDER", "inline")
    get_storage_service.cache_clear()
    get_settings.cache_clear()
    db = snapshot_db
    seed = _seed_p0c(db)
    holder = {"principal": Principal(seed["owner"].id, seed["owner"].username, None)}
    app = FastAPI()
    app.include_router(requirements.router, dependencies=[Depends(guard_project_resource)])
    app.include_router(deliverables.router)
    app.include_router(review_tasks.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_principal] = lambda: holder["principal"]
    with TestClient(app) as client:
        yield client, db, seed, holder
    get_storage_service.cache_clear()
    get_settings.cache_clear()


def _seed_p0c(db):
    institution, owner, project, target_field, source_field, mart_field = _seed_assets(db)
    db.add(ProjectMembership(project_id=project.id, user_id=owner.id,
                             project_role="project_manager", status="active"))
    scenario = ProductScenario(project_id=project.id, scenario_code="SYNTH",
                               scenario_name="个人消费贷产品方案")
    db.add(scenario)
    db.flush()

    system = db.scalar(select(BusinessSystem).where(BusinessSystem.project_id == project.id))
    loan = SourceTable(project_id=project.id, business_system_id=system.id, table_code="LOAN_IOU",
                       table_name="借据表", table_comment="核心借据", schema_name="SRC",
                       physical_table_name="LOAN_IOU")
    contract = SourceTable(project_id=project.id, business_system_id=system.id, table_code="CONTRACT",
                           table_name="合同表", table_comment="贷款合同", schema_name="SRC",
                           physical_table_name="CONTRACT")
    product = SourceTable(project_id=project.id, business_system_id=system.id, table_code="PRODUCT",
                          table_name="产品表", table_comment="产品主数据（产品ID存在重复）", schema_name="SRC",
                          physical_table_name="PRODUCT")
    term = SourceTable(project_id=project.id, business_system_id=system.id, table_code="PRODUCT_TERM",
                       table_name="产品期限方案表", table_comment="产品期限方案", schema_name="SRC",
                       physical_table_name="PRODUCT_TERM")
    db.add_all([loan, contract, product, term])
    db.flush()
    db.add_all([
        SourceField(project_id=project.id, source_table_id=loan.id, field_code="LOAN_ID",
                    field_name="借据号", field_comment="借据主键", physical_column_name="LOAN_ID"),
        SourceField(project_id=project.id, source_table_id=loan.id, field_code="CUST_NO",
                    field_name="客户号", field_comment="借据上的客户号", physical_column_name="CUST_NO"),
        SourceField(project_id=project.id, source_table_id=loan.id, field_code="PRODUCT_ID",
                    field_name="产品ID", field_comment="借据产品", physical_column_name="PRODUCT_ID"),
        SourceField(project_id=project.id, source_table_id=loan.id, field_code="TERM_NO",
                    field_name="期限编号", field_comment="借据期限方案", physical_column_name="TERM_NO"),
        SourceField(project_id=project.id, source_table_id=contract.id, field_code="CONTRACT_ID",
                    field_name="合同号", field_comment="合同主键", physical_column_name="CONTRACT_ID"),
        SourceField(project_id=project.id, source_table_id=product.id, field_code="PRODUCT_ID",
                    field_name="产品ID", field_comment="产品主数据ID（允许重复）", physical_column_name="PRODUCT_ID"),
        SourceField(project_id=project.id, source_table_id=product.id, field_code="PRODUCT_NAME",
                    field_name="产品名称", field_comment="产品名称", physical_column_name="PRODUCT_NAME"),
        SourceField(project_id=project.id, source_table_id=term.id, field_code="PRODUCT_ID",
                    field_name="产品ID", field_comment="期限方案所属产品", physical_column_name="PRODUCT_ID"),
        SourceField(project_id=project.id, source_table_id=term.id, field_code="TERM_NO",
                    field_name="期限编号", field_comment="期限方案编号", physical_column_name="TERM_NO"),
        SourceField(project_id=project.id, source_table_id=term.id, field_code="TERM_DAYS",
                    field_name="期限天数", field_comment="期限天数", physical_column_name="TERM_DAYS"),
    ])

    document = KnowledgeDocument(
        project_id=project.id, file_name="一表通客户统一标识口径说明-2026v1.txt", file_type="txt",
        source_type="upload", storage_path="p0c/knowledge/customer-id-2026v1.txt",
        knowledge_type="regulatory_formal", knowledge_scope="project", document_status="indexed",
        confidentiality_level="internal", current_version_no=1, source_category="regulatory_formal",
        regulatory_document_no="合成制度-2026-10", publisher="合成监管银行", created_by=owner.username,
    )
    db.add(document)
    db.flush()
    version = KnowledgeDocumentVersion(
        project_id=project.id, document_id=document.id, version_no=1,
        file_name=document.file_name, storage_path=document.storage_path,
        file_hash=hashlib.sha256(RULE_V1.encode("utf-8")).hexdigest(), parse_status="parsed",
        lifecycle_status="active", regulatory_version="2026-v1",
        effective_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
        activated_at=datetime(2026, 1, 2, tzinfo=timezone.utc), reviewed_by="p0c-knowledge-reviewer",
        created_by=owner.username,
    )
    db.add(version)
    db.flush()
    document.current_version_id = version.id
    unit = KnowledgeUnit(
        project_id=project.id, document_id=document.id, document_version_id=version.id,
        knowledge_type="regulatory_formal", knowledge_scope="project", unit_type="clause",
        title="一表通客户统一标识填报口径", content=RULE_V1,
        normalized_content=RULE_V1.replace(" ", ""), source_file_name=document.file_name,
        confidentiality_level="internal", enabled=True,
        content_hash=hashlib.sha256(RULE_V1.encode("utf-8")).hexdigest(),
        target_table_code="EAST_CUSTOMER", target_field_code="CUSTOMER_ID",
        metadata_json={"clause": "第十条"}, tags_json=["一表通", "客户标识"],
    )
    db.add(unit)
    db.flush()
    reviewers = _reviewers(db, project.id)
    db.commit()
    return {
        "institution": institution, "owner": owner, "project": project, "target_field": target_field,
        "source_field": source_field, "mart_field": mart_field, "scenario": scenario,
        "source_tables": [loan, contract, product, term], "document": document,
        "version": version, "unit": unit, "reviewers": reviewers,
    }


def _fake_chat(db, project_id, runtime, prompt, response_model, **kwargs):
    """测试替身：结构合法的候选，引用冻结输入中真实存在的物理来源与证据。"""
    import json

    context = json.loads(prompt)
    physical = context.get("physical_sources", [])
    evidence = context.get("evidence", [])
    references = [{"kind": row["kind"], "table_id": row["table_id"], "field_id": row["field_id"]}
                  for row in physical[:3]]
    return {
        "business_definition": "客户统一标识取自客户主数据，经借据→合同→产品（产品ID+期限编号）关联，"
                               "不得仅以产品ID关联；无法唯一关联时形成待确认问题。",
        "processing_logic": "借据按 LOAN_ID 关联合同，按 PRODUCT_ID+TERM_NO 关联产品期限方案，"
                            "按 CUST_NO 关联客户后取 CUST_ID，空值不参与关联。",
        "final_content": "客户统一标识 = 客户主数据 CUST_ID；产品维度按 产品ID+期限编号 组合业务键关联。",
        "physical_references": references,
        "evidence_unit_ids": [evidence[0]["unit_id"]] if evidence else [],
        "script_rule_ids": [],
        "policy_comparisons": [],
        "gaps": [],
    }


async def _fake_runtime_chat(*args, **kwargs):
    return _fake_chat(*args, **kwargs)


def _use_model(monkeypatch):
    """默认测试替身；VNEXT2_P0C_REAL_MODEL=1 时走真实网关。"""
    if os.getenv("VNEXT2_P0C_REAL_MODEL") == "1":
        return "real_model"
    import app.services.requirement_generation_worker as worker
    monkeypatch.setattr(worker, "execute_runtime_chat", _fake_runtime_chat)
    return "mock_model"


def _create_requirement(client, seed):
    project = seed["project"]
    target_field = seed["target_field"]
    payload = {
        "name": "一表通客户统一标识开发需求",
        "target_table_id": target_field.target_table_id,
        "field_ids": [target_field.id],
        "scenario_id": seed["scenario"].id,
        "objective": "形成可供数据集市开发的一表通客户统一标识需求文档",
        "background": SCENARIO,
        "effective_date": "2026-01-01",
        "inclusion": "个人消费贷借据",
        "exclusion": "已核销借据",
        "document_ids": [seed["document"].id],
        "source_table_ids": [seed["source_tables"][0].id, seed["source_tables"][1].id,
                             seed["source_tables"][2].id, seed["source_tables"][3].id],
        "mart_table_ids": [],
        "expected_version": None,
    }
    response = client.post(f"/projects/{project.id}/requirements", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _run_generation(client, seed, requirement):
    project = seed["project"]
    base = f"/projects/{project.id}/requirements/{requirement['id']}"
    frozen = client.post(base + "/revisions", json={"expected_version": requirement["version"]})
    assert frozen.status_code == 201, frozen.text
    content_version = frozen.json()["revision"]["content_version"]
    prepared = client.post(base + "/generation-inputs", json={
        "expected_content_version": content_version,
        "field_ids": [seed["target_field"].id],
        "sections": ["business", "lineage"],
        "idempotency_key": "p0c-generation-" + str(requirement["id"]),
    })
    assert prepared.status_code == 201, prepared.text
    started = client.post(base + "/generation-runs", json={
        "expected_content_version": content_version,
        "field_ids": [seed["target_field"].id],
        "sections": ["business", "lineage"],
        "idempotency_key": "p0c-generation-run-" + str(requirement["id"]),
    })
    assert started.status_code == 201, started.text
    runs = client.get(base + "/generation-runs").json()
    items = runs[0]["items"] if isinstance(runs, list) else runs.get("items", [])
    assert items, runs
    return content_version, runs, items


def test_generation_adoption_review_and_formal_delivery(p0c, monkeypatch):
    client, db, seed, holder = p0c
    mode = _use_model(monkeypatch)
    requirement = _create_requirement(client, seed)
    content_version, runs, items = _run_generation(client, seed, requirement)
    project = seed["project"]
    base = f"/projects/{project.id}/requirements/{requirement['id']}"

    statuses = {item["section"]: item for item in items}
    assert set(statuses) == {"business", "lineage"}, items
    for item in items:
        assert item["status"] in {"completed", "failed", "blocked"}, item
        assert item["status"] == "completed", item
    detail = client.get(base + f"/generation-items/{statuses['business']['id']}")
    assert detail.status_code == 200, detail.text
    body = detail.json()
    assert len(body["candidate_hash"]) == 64
    assert body["changes"] and body["changes"][0]["proposed"], body
    assert body["stale"] is False
    lineage_detail = client.get(base + f"/generation-items/{statuses['lineage']['id']}")
    assert lineage_detail.status_code == 200, lineage_detail.text
    lineage_hash = lineage_detail.json()["candidate_hash"]

    adopted = client.post(base + "/generation-candidates/adopt", json={
        "expected_content_version": content_version,
        "selections": [
            {"item_id": statuses["business"]["id"],
             "candidate_hash": body["candidate_hash"],
             "selected_fields": ["business_definition", "final_content"],
             "replace_manual_fields": []},
            {"item_id": statuses["lineage"]["id"],
             "candidate_hash": lineage_hash,
             "selected_fields": ["processing_logic", "final_content"],
             "replace_manual_fields": []},
        ],
    })
    assert adopted.status_code == 200, adopted.text
    adopted_body = adopted.json()
    adopted_version = adopted_body["document"]["revision"]["content_version"]
    assert adopted_version == content_version + 1

    edited = client.put(base + f"/fields/{seed['target_field'].id}", json={
        "expected_content_version": adopted_version,
        "section": "business",
        "changes": {"final_content": "客户统一标识 = 客户主数据 CUST_ID（人工确认：按 产品ID+期限编号 组合键关联产品，"
                                     "无法唯一关联时提待确认问题）。"},
    })
    assert edited.status_code == 200, edited.text
    business_version = edited.json()["revision"]["content_version"]

    lineage_confirmed = client.put(base + f"/fields/{seed['target_field'].id}", json={
        "expected_content_version": business_version,
        "section": "lineage",
        "changes": {
            "source_system_name": "客户信息系",
            "source_schema_name": "SRC",
            "source_table_english_name": "SRC.CUSTOMER",
            "source_field_english_name": "CUST_ID",
            "processing_logic": "借据 LOAN_IOU 经 借据号+合同号 关联 CONTRACT，再按 产品ID+期限编号 关联 "
                                "PRODUCT/PRODUCT_TERM；记录数不因产品期限膨胀；无法唯一关联时提待确认问题。",
        },
    })
    assert lineage_confirmed.status_code == 200, lineage_confirmed.text
    revised_version = lineage_confirmed.json()["revision"]["content_version"]

    refreshed = client.post(base + f"/revisions/{revised_version}/refresh-lineage")
    assert refreshed.status_code == 200, refreshed.text
    content_version = refreshed.json()["content_version"]

    # AI 候选自带的风险提示（gap）必须由人工逐条给出依据后关闭，不能自动消失。
    document = client.get(base + f"/revisions/{content_version}")
    assert document.status_code == 200, document.text
    open_gaps = [gap for gap in document.json().get("gaps", [])
                 if gap.get("status") not in {"resolved", "closed"}]
    for gap in open_gaps:
        closed = client.post(base + f"/gaps/{gap['id']}/resolve", json={
            "expected_content_version": content_version,
            "resolution_text": "已人工核对合成制度 2026-v1 第十条与集市映射证据，确认该候选口径可交付。",
        })
        assert closed.status_code == 201, (gap, closed.text)
        content_version = closed.json()["revision"]["content_version"]

    readiness = client.get(base + "/review-readiness", params={"content_version": content_version})
    assert readiness.status_code == 200, readiness.text
    readiness_body = readiness.json()
    assert readiness_body["eligible"] is True, readiness_body.get("reasons")
    content_hash = readiness_body["content_hash"]

    submitted = client.post(base + "/review-submissions", json={
        "expected_content_version": content_version,
        "expected_content_hash": content_hash,
    })
    assert submitted.status_code == 201, submitted.text
    submission_id = submitted.json()["id"]

    blocked = client.post(base + f"/review-submissions/{submission_id}/finalize")
    assert blocked.status_code == 409

    for role in ("business_reviewer", "technical_reviewer", "final_reviewer"):
        holder["principal"] = Principal(seed["reviewers"][role].id, seed["reviewers"][role].username, None)
        tasks = client.get("/me/tasks")
        assert tasks.status_code == 200, tasks.text
        pending = [task for task in tasks.json()
                   if task["target_type"] == "requirement_review_submission" and task["target_id"] == submission_id]
        assert pending, (role, tasks.json())
        approved = client.post(f"/review-tasks/{pending[0]['id']}/approve", json={"comment": f"{role} 通过"})
        assert approved.status_code == 200, approved.text

    holder["principal"] = Principal(seed["owner"].id, seed["owner"].username, None)
    finalized = client.post(base + f"/review-submissions/{submission_id}/finalize")
    assert finalized.status_code == 201, finalized.text
    delivery = finalized.json()
    assert delivery["content_hash"] and delivery["file_hash"]

    deliveries = client.get(base + "/formal-deliveries")
    assert deliveries.status_code == 200, deliveries.text
    delivery_id = deliveries.json()[0]["id"]
    excel = client.get(base + f"/formal-deliveries/{delivery_id}/export", params={"format": "xlsx"})
    assert excel.status_code == 200, excel.text
    snapshot_hash = excel.headers["X-Requirement-Snapshot-Hash"]
    workbook = load_workbook(BytesIO(excel.content))
    text = "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value)
    assert "客户统一标识" in text
    assert "产品ID+期限编号" in text or "产品ID" in text

    word = client.get(base + f"/formal-deliveries/{delivery_id}/export", params={"format": "docx"})
    assert word.status_code == 200, word.text
    document = DocxDocument(BytesIO(word.content))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "客户统一标识" in body
    assert snapshot_hash in body or "固定内容标识" in body

    # 交付后新增待确认问题：历史交付文件不变，当前安全检查可见新风险
    question = client.post(f"/projects/{project.id}/questions", json={
        "target_table_id": seed["target_field"].target_table_id,
        "target_field_id": seed["target_field"].id,
        "question_text": "产品期限方案调整后，客户统一标识是否需要重新报送？",
        "question_type": "business",
    })
    assert question.status_code == 201, question.text
    again = client.get(base + f"/formal-deliveries/{delivery_id}/export", params={"format": "xlsx"})
    assert again.status_code == 200
    assert again.headers["X-Requirement-Snapshot-Hash"] == snapshot_hash
    readiness_after = client.get(base + "/review-readiness", params={"content_version": content_version}).json()
    assert readiness_after["eligible"] is False
    assert any(item["code"].startswith("manual:") for item in readiness_after["reasons"])

    # 交付后修订：新草稿版本不影响历史正式交付
    revised = client.post(base + f"/revisions/{content_version}/revise",
                          json={"expected_content_version": content_version, "reason": "制度 v2 变化复核"})
    assert revised.status_code == 201, revised.text
    assert revised.json()["revision"]["content_version"] == content_version + 1
    still = client.get(base + f"/formal-deliveries/{delivery_id}/export", params={"format": "xlsx"})
    assert still.status_code == 200
    assert still.headers["X-Requirement-Snapshot-Hash"] == snapshot_hash
    assert db.scalar(select(Requirement).where(Requirement.id == requirement["id"])) is not None


def test_generation_failure_is_visible_not_empty_draft(p0c, monkeypatch):
    client, db, seed, holder = p0c
    import app.services.requirement_generation_worker as worker

    async def _broken(*args, **kwargs):
        return {"unexpected": "shape"}

    monkeypatch.setattr(worker, "execute_runtime_chat", _broken)
    requirement = _create_requirement(client, seed)
    content_version, runs, items = _run_generation(client, seed, requirement)
    for item in items:
        assert item["status"] == "failed", item
        assert item["reason_code"] == "generation_failed"
    base = f"/projects/{seed['project'].id}/requirements/{requirement['id']}"
    detail = client.get(base + f"/generation-items/{items[0]['id']}")
    assert detail.status_code == 409
    # 生成失败不会凭空产生正文：草稿仍未取得业务定义/来源/加工，因此仍不可送审。
    readiness = client.get(base + "/review-readiness", params={"content_version": content_version}).json()
    assert readiness["eligible"] is False
    codes = {item["code"] for item in readiness["reasons"]}
    assert {"1:definition", "1:source", "1:transformation"} <= codes
    document = client.get(base + "/document").json()
    record = document["fields"][0]
    assert record["business"] in (None, {}) and record["lineage"] in (None, {})


def test_analysis_gap_cannot_be_closed_by_a_note(p0c, monkeypatch):
    """分析派生的缺口（缺定义/来源/加工）不能靠一句说明关闭，避免用说明绕过口径。"""
    client, db, seed, holder = p0c
    _use_model(monkeypatch)
    requirement = _create_requirement(client, seed)
    content_version, runs, items = _run_generation(client, seed, requirement)
    base = f"/projects/{seed['project'].id}/requirements/{requirement['id']}"
    readiness = client.get(base + "/review-readiness", params={"content_version": content_version}).json()
    gap = next(item for item in readiness["reasons"] if item["code"] == "1:definition")
    closed = client.post(base + f"/gaps/{gap['code']}/resolve", json={
        "expected_content_version": content_version, "resolution_text": "业务口头确认即可",
    })
    assert closed.status_code == 409, closed.text
    assert "补齐字段内容" in closed.json()["detail"]
    empty = client.post(base + f"/gaps/{gap['code']}/resolve", json={
        "expected_content_version": content_version, "resolution_text": "   ",
    })
    assert empty.status_code in {422}, empty.text


def test_new_question_after_freeze_blocks_new_delivery(p0c, monkeypatch):
    client, db, seed, holder = p0c
    _use_model(monkeypatch)
    requirement = _create_requirement(client, seed)
    content_version, runs, items = _run_generation(client, seed, requirement)
    base = f"/projects/{seed['project'].id}/requirements/{requirement['id']}"
    question = client.post(f"/projects/{seed['project'].id}/questions", json={
        "target_table_id": seed["target_field"].target_table_id,
        "question_text": "同一客户存在多笔借据时，客户统一标识是否只取主借据？",
        "question_type": "business",
    })
    assert question.status_code == 201, question.text
    readiness = client.get(base + "/review-readiness", params={"content_version": content_version}).json()
    assert readiness["eligible"] is False
    submitted = client.post(base + "/review-submissions", json={
        "expected_content_version": content_version,
        "expected_content_hash": readiness["content_hash"],
    })
    assert submitted.status_code == 409
    assert "阻断" in submitted.json()["detail"]
