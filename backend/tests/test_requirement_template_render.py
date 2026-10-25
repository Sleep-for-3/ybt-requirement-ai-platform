from __future__ import annotations

import hashlib
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from app.services.requirement_doc_template import template_hash
from app.services.requirement_template_render import (
    DOCX_MEDIA_TYPE,
    RENDERER_VERSION,
    XLSX_MEDIA_TYPE,
    CustomerTemplateRenderError,
    render_requirement_artifact,
)
from app.services.requirement_xlsx_template import template_sha256
from docx_template_fixtures import build_customer_template, customer_config, frozen_content
from test_requirement_xlsx_template import (
    build_config as build_xlsx_config,
    build_content as build_xlsx_content,
    build_template as build_xlsx_template,
)


def _docx_case(n_fields: int = 2):
    template = build_customer_template()
    return template, customer_config(template), frozen_content(n_fields=n_fields)


def test_docx_render_returns_bytes_and_identity():
    template, config, content = _docx_case()
    artifact = render_requirement_artifact(
        format="docx", template_bytes=template, frozen_content=content, config=config,
        base_name="客户需求",
    )
    assert artifact.format == "docx"
    assert artifact.media_type == DOCX_MEDIA_TYPE
    assert artifact.file_name == "客户需求.docx"
    assert artifact.file_sha256 == hashlib.sha256(artifact.content).hexdigest()
    assert artifact.template_sha256 == template_hash(template)
    assert artifact.renderer_version == RENDERER_VERSION
    assert artifact.config_version == "customer-p0d-test-1"
    assert artifact.diagnostics

    body = Document(BytesIO(artifact.content))
    text = "\n".join(paragraph.text for paragraph in body.paragraphs)
    table_text = "\n".join(
        cell.text for table in body.tables for row in table.rows for cell in row.cells
    )
    assert "客户字段1" in table_text or "客户字段1" in text
    assert "{{" not in text and "{{" not in table_text
    # 页眉页脚来自客户模板，不能在渲染中丢失。
    assert any("客户报送需求文档" in part.text for section in body.sections for part in section.header.paragraphs)


def test_xlsx_render_returns_bytes_and_identity():
    template = build_xlsx_template()
    config = build_xlsx_config(template)
    content = build_xlsx_content(3)
    artifact = render_requirement_artifact(
        format="xlsx", template_bytes=template, frozen_content=content, config=config,
        base_name="客户需求.xlsx",
    )
    assert artifact.format == "xlsx"
    assert artifact.media_type == XLSX_MEDIA_TYPE
    assert artifact.file_name == "客户需求.xlsx"
    assert artifact.template_sha256 == template_sha256(template)
    assert artifact.file_sha256 == hashlib.sha256(artifact.content).hexdigest()
    book = load_workbook(BytesIO(artifact.content))
    assert book.sheetnames


def test_unsupported_format_is_rejected_with_code():
    template, config, content = _docx_case()
    for value in ("pdf", "", "doc"):
        try:
            render_requirement_artifact(
                format=value, template_bytes=template, frozen_content=content, config=config,
            )
        except CustomerTemplateRenderError as failure:
            assert failure.code == "unsupported_format"
            assert "docx" in failure.message and "xlsx" in failure.message
        else:  # pragma: no cover - 明确失败才算通过
            raise AssertionError(f"格式 {value!r} 不应被接受")


def test_docx_hash_mismatch_is_reported_before_rendering():
    template, config, content = _docx_case()
    broken = dict(config)
    broken["template_hash"] = "0" * 64
    try:
        render_requirement_artifact(
            format="docx", template_bytes=template, frozen_content=content, config=broken,
        )
    except CustomerTemplateRenderError as failure:
        assert failure.code == "template_hash_mismatch"
    else:  # pragma: no cover
        raise AssertionError("模板 hash 与配置不一致时必须拒绝渲染")


def test_missing_required_anchor_fails_instead_of_partial_output():
    template, config, content = _docx_case()
    broken = dict(config)
    broken["anchors"] = {**config["anchors"], "cover.name": "{{客户模板里没有的锚点}}"}
    broken["required_keys"] = list(broken["anchors"])
    try:
        render_requirement_artifact(
            format="docx", template_bytes=template, frozen_content=content, config=broken,
        )
    except CustomerTemplateRenderError as failure:
        assert failure.code in {"template_anchor_error", "template_hash_mismatch"}
    else:  # pragma: no cover
        raise AssertionError("缺必需锚点时必须失败，不能生成残缺文档")


def test_empty_template_and_empty_content_are_rejected():
    template, config, content = _docx_case()
    cases = (
        ({"template_bytes": b""}, "empty_template"),
        ({"frozen_content": {}}, "empty_content"),
    )
    for override, expected in cases:
        arguments = {"format": "docx", "template_bytes": template,
                     "frozen_content": content, "config": config, **override}
        try:
            render_requirement_artifact(**arguments)
        except CustomerTemplateRenderError as failure:
            assert failure.code == expected
        else:  # pragma: no cover
            raise AssertionError(f"{expected} 未被拦截")


def test_file_name_keeps_single_suffix():
    template, config, content = _docx_case(n_fields=1)
    for base, expected in (("需求说明.docx", "需求说明.docx"), ("需求说明", "需求说明.docx")):
        artifact = render_requirement_artifact(
            format="docx", template_bytes=template, frozen_content=content, config=config,
            base_name=base,
        )
        assert artifact.file_name == expected
