"""一表通调研表式导出：版式（按来源系统重复分组）、取值来源与"不猜测"边界。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.services.requirement_survey_export import SURVEY_DOC_VERSION, export_survey_document

BASE_COLUMNS = 11
BUSINESS_COLUMNS = 6
SOURCE_COLUMNS = 9


def _content(**overrides):
    field = {
        "id": 1,
        "field_code": "H010001",
        "field_name": "借据ID",
        "data_category": "基础数据",
        "data_format": "AN..60",
        "regulatory_original_definition": "填报该笔贷款对应的借据号码。",
        "report_name": "贷款借据",
        "report_field_name": "借据ID",
        "remarks": "制度包版本：R2022100101",
    }
    record = {
        "field": field,
        "business": {
            "business_definition": "借据唯一编号",
            "final_content": "借据编号由许可证号与内部编号组成。",
        },
        "lineage": {"final_content": "取自信贷借据信息表的借据号。"},
        "mart_mappings": [
            {
                "mart_table_summary": "信贷借据信息表(HYF_LN_DUEBILL_INFO)",
                "mart_field_summary": "信贷借据号(DUEBILL_NO)",
                "business_rule": "按借据号直接映射",
            }
        ],
    }
    content = {
        "target_table": {"table_code": "T_8_1", "table_name": "贷款借据"},
        "scenario": {"scenario_name": "一表通监管报送"},
        "requirement": {
            "name": "T_8_1 需求",
            "inclusion": None,
            "objective": None,
            "exclusion": None,
        },
        "fields": [record],
        "script_basis": {
            "rules": [
                {
                    "rule_id": "script-edge-1",
                    "target": {"table_name": "T_8_1", "column_name": "H010001"},
                    "source": {
                        "table_name": "HYF_LN_DUEBILL_INFO_EXT_YBT",
                        "column_name": "DUEBILL_NO",
                    },
                    "transformation_expression": "P1.DUEBILL_NO AS H010001",
                    "transformation_expression_ref": None,
                }
            ],
            "condition_texts": {},
        },
    }
    content.update(overrides)
    return content


def _rows(payload: bytes):
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["需求文档"]
    return sheet, [[cell.value for cell in row] for row in sheet.iter_rows()]


def _header_index(rows):
    return next(index for index, row in enumerate(rows) if row[0] == "数据项编码")


def test_export_writes_title_table_info_group_titles_and_field_rows():
    sheet, rows = _rows(export_survey_document(_content()))
    assert SURVEY_DOC_VERSION == SURVEY_DOC_VERSION
    assert rows[0][0] == "返回目录"

    labels = [row[0] for row in rows[2:8]]
    assert labels == [
        "数据分类",
        "数据主题名称",
        "数据表名称",
        "采集范围",
        "一表通报送范围确认",
        "与EAST报送范围差异",
    ]
    assert rows[2][2] == "待业务补充"  # 表级信息缺失时写占位，不猜测
    assert rows[3][2] == "一表通监管报送"
    assert rows[4][2] == "贷款借据"

    title_row = next(index for index, row in enumerate(rows) if row[0] == "一表通数据基础信息")
    title = rows[title_row]
    assert title[6] == "EAST5.0同源映射列"
    assert title[BASE_COLUMNS] == "业务口径-监管集市"
    assert title[BASE_COLUMNS + BUSINESS_COLUMNS] == "溯源-监管集市"

    header_row = _header_index(rows)
    header = rows[header_row]
    assert header[4].startswith("字段业务定义")
    assert header[BASE_COLUMNS + 4] == "手工补录"
    assert header[BASE_COLUMNS + 5] == "业务口径确认人"
    assert header[BASE_COLUMNS + BUSINESS_COLUMNS] == "来源系统"
    assert header[BASE_COLUMNS + BUSINESS_COLUMNS + 6] == "处理逻辑类型"

    data = rows[header_row + 1]
    assert data[0] == "H010001"
    assert data[1] == "借据ID"
    assert data[4] == "填报该笔贷款对应的借据号码。"
    assert data[BASE_COLUMNS] == "借据唯一编号"
    assert data[BASE_COLUMNS + 4] is None  # 没有补录事实时不得写"否"替业务确认
    source = BASE_COLUMNS + BUSINESS_COLUMNS
    assert data[source] == "监管集市"
    assert data[source + 1] == "HYF_LN_DUEBILL_INFO"
    assert data[source + 2] == "信贷借据信息表"
    assert data[source + 3] == "DUEBILL_NO"
    assert data[source + 4] == "信贷借据号"
    assert data[source + 5] == "按借据号直接映射"
    assert data[source + 6] == "加工"
    # 没有第二个来源系统时不得凭空多出列组
    assert len([value for value in header if value]) == BASE_COLUMNS + BUSINESS_COLUMNS + SOURCE_COLUMNS
    assert sheet.freeze_panes == f"C{header_row + 2}"


def test_second_source_system_adds_another_column_group():
    content = _content()
    content["fields"][0]["mart_mappings"] = [
        {
            "mart_table_summary": "信贷借据信息表(HYF_LN_DUEBILL_INFO)",
            "mart_field_summary": "信贷借据号(DUEBILL_NO)",
            "business_rule": "",
        },
        {
            "mart_table_summary": "补录借据表(BL_YBT_JJB)",
            "mart_field_summary": "借据编号(JJBH)",
            "business_rule": "从补录平台按日导入",
        },
    ]
    _, rows = _rows(export_survey_document(content))
    header_row = _header_index(rows)
    header = rows[header_row]
    second = BASE_COLUMNS + BUSINESS_COLUMNS + SOURCE_COLUMNS
    assert header[BASE_COLUMNS] == "字段业务定义"
    assert header[second] == "字段业务定义"

    title_row = next(index for index, row in enumerate(rows) if row[0] == "一表通数据基础信息")
    assert rows[title_row][BASE_COLUMNS] == "业务口径-监管集市"
    assert rows[title_row][second] == "业务口径-补录平台"
    assert rows[title_row][second + BUSINESS_COLUMNS] == "溯源-补录平台"

    data = rows[header_row + 1]
    assert data[second + BUSINESS_COLUMNS + 1] == "BL_YBT_JJB"
    assert data[second + BUSINESS_COLUMNS + 5] == "从补录平台按日导入"


def test_field_without_candidate_content_uses_placeholder_not_invented():
    content = _content()
    content["fields"][0]["business"] = None
    content["fields"][0]["lineage"] = None
    content["fields"][0]["mart_mappings"] = []
    content["script_basis"] = {}
    _, rows = _rows(export_survey_document(content))
    header_row = _header_index(rows)
    data = rows[header_row + 1]
    # 业务口径列写占位符（交付评审要看到"待业务补充"，空白会被当成漏排），
    # 但不得编造任何口径内容；没有来源映射时溯源列组仍整组留空。
    assert data[BASE_COLUMNS] == "待业务补充"
    source = BASE_COLUMNS + BUSINESS_COLUMNS
    assert not data[source + 1] and not data[source + 3] and not data[source + 5]


def test_condition_refs_are_resolved_from_condition_pool():
    content = _content()
    record = content["fields"][0]
    record["mart_mappings"][0]["business_rule"] = None
    rule = content["script_basis"]["rules"][0]
    rule["transformation_expression"] = None
    rule["transformation_expression_ref"] = "abc123"
    content["script_basis"]["condition_texts"] = {
        "abc123": "P1.DUEBILL_NO AS H010001"
    }
    _, rows = _rows(export_survey_document(content))
    header_row = _header_index(rows)
    source = BASE_COLUMNS + BUSINESS_COLUMNS
    assert "P1.DUEBILL_NO AS H010001" in rows[header_row + 1][source + 5]


def test_fields_without_code_are_skipped():
    content = _content()
    content["fields"].append(
        {"field": {"field_name": "无编码字段"}, "business": {}, "lineage": {}}
    )
    _, rows = _rows(export_survey_document(content))
    codes = [row[0] for row in rows if row[0] == "H010001"]
    assert len(codes) == 1
    assert all(row[1] != "无编码字段" for row in rows)

def test_group_names_only_claim_what_the_evidence_supports():
    """分组名只能来自证据：集市映射前缀、脚本血缘、或"未登记来源"。

    真实事故：早期版本把"没有集市映射"的字段一律归进补录平台组，等于替业务声明
    "只能人工补录"；而其中一类的值其实是跑批脚本直接产出的。这里固定三种来源。
    """
    content = _content()
    market = content["fields"][0]
    script_only = {
        "field": {"id": 2, "field_code": "SFJJGDM", "field_name": "省局分局机构代码"},
        "business": {"final_content": "由脚本按数据拆分生成。"},
        "mart_mappings": [],
    }
    unknown = {
        "field": {"id": 3, "field_code": "SFJJYWCJ", "field_name": "业务场景"},
        "business": {"final_content": "待业务补充"},
        "mart_mappings": [],
    }
    content["fields"] = [market, script_only, unknown]
    content["script_basis"]["rules"].append(
        {
            "rule_id": "script-edge-2",
            "target": {"table_name": "T_8_1", "column_name": "SFJJGDM"},
            "source": {"table_name": "BL_YBT_JGXXB", "column_name": "JGID"},
            "transformation_expression": "BLJG.JGID AS SFJJGDM",
            "transformation_expression_ref": None,
        }
    )

    sheet, rows = _rows(export_survey_document(content))
    title = rows[9]
    names = [value for value in title if isinstance(value, str) and value.startswith("业务口径-")]
    assert names == ["业务口径-监管集市", "业务口径-跑批脚本", "业务口径-未登记来源"]

    header = _header_index(rows)
    # 组在行里的 0 基起止：基础列 0-10，第 i 组（1 基）业务段从 11+15*(i-1) 开始
    def group_span(index: int) -> tuple[int, int]:
        business = BASE_COLUMNS + (BUSINESS_COLUMNS + SOURCE_COLUMNS) * (index - 1)
        return business, business + BUSINESS_COLUMNS

    # 跑批脚本组（第 2 组）：来源表/字段来自脚本血缘，不是编造
    script_business, script_source = group_span(2)
    script_row = next(row for row in rows[header + 1 :] if row[0] == "SFJJGDM")
    assert script_row[script_business] == "由脚本按数据拆分生成。"
    assert script_row[script_source] == "跑批脚本"
    assert script_row[script_source + 1] == "BL_YBT_JGXXB"
    assert script_row[script_source + 3] == "JGID"
    assert "BLJG.JGID" in (script_row[script_source + 5] or "")

    # 未登记来源组（第 3 组）：来源系统必须留空。
    # 真实事故：v6 把分组名直接写进"来源系统"，导出的需求文档里出现"来源系统=未登记来源"，
    # 等于替业务声明存在一个叫这个名字的系统；没有脚本事实也没有集市映射时，来源信息一律留空。
    unknown_business, unknown_source = group_span(3)
    unknown_row = next(row for row in rows[header + 1 :] if row[0] == "SFJJYWCJ")
    assert unknown_row[unknown_business] == "待业务补充"
    assert unknown_row[unknown_source] is None
    assert unknown_row[unknown_source + 1] is None
    assert unknown_row[unknown_source + 5] is None


def test_renderer_version_carries_source_fingerprint() -> None:
    from app.services.requirement_survey_export import (
        RENDERER_FINGERPRINT,
        SURVEY_DOC_VERSION,
        renderer_version,
    )

    version = renderer_version()
    assert version == f"{SURVEY_DOC_VERSION}+{RENDERER_FINGERPRINT}"
    assert len(RENDERER_FINGERPRINT) == 12
    # 源码改了但语义版本号没变时，渲染身份仍必须变化（真实事故见模块 docstring）。
    assert version != SURVEY_DOC_VERSION


def test_completeness_footer_counts_what_is_missing():
    """交付完整度只统计冻结内容里的事实，缺什么写什么，并标明不替代业务确认。"""
    content = _content()
    content["fields"].append(
        {
            "field": {"field_code": "H010999", "field_name": "无来源字段"},
            "business": {},
            "lineage": {},
            "mart_mappings": [],
            "evidence": [],
        }
    )
    _sheet, rows = _rows(export_survey_document(content))
    header = next(index for index, row in enumerate(rows) if str(row[0]).startswith("交付完整度"))
    assert "不替代业务确认" in str(rows[header][0])
    stats = {rows[index][0]: rows[index][1] for index in range(header + 1, len(rows)) if rows[index][0] and rows[index][1] is not None}
    assert stats["字段总数"] == "2"
    assert stats["有业务口径正文"] == "1（待补 1）"
    assert stats["有集市映射"] == "1"
    assert stats["多来源字段（≥2 条映射）"] == "0"
    assert stats["未登记来源字段"] == "1"


def test_table_level_values_come_from_template_meta_when_requirement_is_empty():
    """需求侧没填表级口径时，回退到调研表“表级口径”里的事实，而不是一律写占位。"""

    content = _content()
    content["target_table"] = {
        "table_code": "T_8_1",
        "table_name": "贷款借据",
        "theme": "状态类数据",
        "filing_note": "全量借据",
        "reporting_domain_name": "默认报送数据域",
    }
    content["requirement"] = {"name": "T_8_1", "inclusion": "", "objective": "", "exclusion": ""}

    _, rows = _rows(export_survey_document(content))

    assert rows[3][:3] == ["数据主题名称", None, "状态类数据"]
    assert rows[5][:3] == ["采集范围", None, "全量借据"]
    assert rows[6][:3] == ["一表通报送范围确认", None, "默认报送数据域"]
    # 需求侧填了就以需求侧为准（模板不覆盖人工口径）
    content["requirement"]["inclusion"] = "仅统计已结清借据"
    _, rows2 = _rows(export_survey_document(content))
    assert rows2[5][2] == "仅统计已结清借据"