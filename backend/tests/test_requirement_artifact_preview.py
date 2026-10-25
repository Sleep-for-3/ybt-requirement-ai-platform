"""产物页面预览：读已落库的产物（xlsx 网格 / docx 章节），不重新渲染。

固定四件事：
1. 预览读的是产物文件本身（与下载同一份字节），不是重新渲染；
2. 合并区、冻结窗格、截断与"未缓存公式"都显式表达，不静默省略；
3. docx 产物按章节预览（标题/段落/表格摘要），段落与表格截断都显式标记；
4. 产物被替换/丢失时预览报错，不静默重渲。
"""

from __future__ import annotations

from pathlib import Path

import pytest

from openpyxl import Workbook

from app.services.requirement_artifact_preview import (
    ArtifactPreviewError,
    preview_artifact,
    preview_docx,
    preview_xlsx,
)
from test_requirement_render_artifact import _frozen_snapshot, _use_temp_storage
from test_requirement_snapshot_api import snapshot_api, snapshot_db  # noqa: F401  夹具复用


def _workbook_bytes() -> bytes:
    from io import BytesIO

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "需求文档"
    sheet["A1"] = "数据项编码"
    sheet["B1"] = "数据项名称"
    sheet["A2"] = "H010001"
    sheet["B2"] = "借据ID"
    sheet.merge_cells("A3:C3")
    sheet["A3"] = "一表通数据基础信息"
    sheet.freeze_panes = "C2"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_preview_reads_values_merges_and_freeze_panes():
    preview = preview_xlsx(_workbook_bytes())

    assert preview["preview_version"] == "artifact-preview-v1"
    assert preview["sheet_count"] == 1
    sheet = preview["sheets"][0]
    assert sheet["name"] == "需求文档"
    assert sheet["rows"][0] == ["数据项编码", "数据项名称", None]
    assert sheet["rows"][1][:2] == ["H010001", "借据ID"]
    assert "A3:C3" in sheet["merged"]
    assert sheet["freeze_panes"] == "C2"
    assert sheet["truncated"] is False


def test_preview_marks_truncation_instead_of_silently_dropping():
    long_text = "加工逻辑" * 200
    from io import BytesIO

    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 12):
        sheet.cell(row=row, column=1, value=f"ROW-{row}")
    sheet.cell(row=1, column=2, value=long_text)
    buffer = BytesIO()
    workbook.save(buffer)

    preview = preview_xlsx(buffer.getvalue(), max_rows=5, max_cell_chars=50)
    sheet = preview["sheets"][0]
    assert len(sheet["rows"]) == 5, "只返回前 max_rows 行"
    assert sheet["truncated"] is True
    assert sheet["total_rows"] == 11
    cell = sheet["rows"][0][1]
    assert cell.startswith("加工逻辑") and "已截断，原文 800 字符" in cell
    assert sheet["truncated_cells"] == 1


def test_preview_shows_formula_text_when_cache_is_missing():
    from io import BytesIO

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "合计"
    sheet["A2"] = "=SUM(A1:A1)"
    buffer = BytesIO()
    workbook.save(buffer)

    preview = preview_xlsx(buffer.getvalue())
    sheet = preview["sheets"][0]
    assert sheet["formulas_without_cache"] == 1
    assert sheet["rows"][1][0] == "[公式] =SUM(A1:A1)", "不能把未缓存公式显示成空白，让人误以为漏填"


def _docx_bytes() -> bytes:
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_heading("跑批脚本需求说明", level=1)
    document.add_heading("需求范围与业务说明", level=1)
    document.add_paragraph("本需求覆盖 T_8_1 贷款借据全部 8 个数据项。")
    document.add_paragraph("X" * 40)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "字段代码"
    table.cell(0, 1).text = "字段名称"
    table.cell(1, 0).text = "H010001"
    table.cell(1, 1).text = "借据ID"
    table.cell(2, 0).text = "H010010"
    table.cell(2, 1).text = "借款余额"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_preview_parses_docx_into_sections_with_tables():
    preview = preview_artifact(_docx_bytes(), file_format="docx")
    assert preview["format"] == "docx"
    assert preview["preview_version"] == "artifact-docx-preview-v1"
    assert preview["section_titles"][:2] == ["跑批脚本需求说明", "需求范围与业务说明"]
    section = preview["sections"][1]
    assert any("8 个数据项" in item["text"] for item in section["paragraphs"])
    table = section["tables"][0]
    assert table["header"] == ["字段代码", "字段名称"]
    assert table["rows"] == 3 and table["columns"] == 2
    assert table["sample_rows"][0] == ["H010001", "借据ID"]
    assert table["truncated"] is False


def test_preview_truncates_long_docx_paragraphs_and_tables():
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_heading("章节", level=1)
    document.add_paragraph("Y" * 2000)
    table = document.add_table(rows=60, cols=1)
    buffer = BytesIO()
    document.save(buffer)
    preview = preview_docx(buffer.getvalue(), max_chars=100, max_table_rows=5)
    paragraph = preview["sections"][0]["paragraphs"][0]
    assert paragraph["truncated"] is True
    assert "已截断，原文 2000 字符" in paragraph["text"]
    table_preview = preview["sections"][0]["tables"][0]
    assert table_preview["truncated"] is True
    assert table_preview["rows"] == 60
    assert len(table_preview["sample_rows"]) == 4  # 表头之外最多 4 行


def test_preview_rejects_unsupported_format_and_broken_files():
    with pytest.raises(ArtifactPreviewError) as pdf:
        preview_artifact(b"whatever", file_format="pdf")
    assert pdf.value.code == "artifact_format_unsupported"
    with pytest.raises(ArtifactPreviewError) as broken_docx:
        preview_artifact(b"not-a-docx", file_format="docx")
    assert broken_docx.value.code == "artifact_not_readable"
    with pytest.raises(ArtifactPreviewError) as broken_xlsx:
        preview_artifact(b"not-xlsx", file_format="xlsx")
    assert broken_xlsx.value.code == "artifact_not_readable"


def test_preview_endpoint_serves_the_stored_artifact(tmp_path: Path, monkeypatch, snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    _use_temp_storage(tmp_path, monkeypatch)
    base = f"/projects/{project.id}/requirements/{req.id}"
    snapshot_id = _frozen_snapshot(client, base, req)

    exported = client.get(base + f"/snapshots/{snapshot_id}/export?format=xlsx")
    assert exported.status_code == 200, exported.text
    artifact_id = exported.headers["x-requirement-render-artifact-id"]

    preview = client.get(base + f"/render-artifacts/{artifact_id}/preview")
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["artifact"]["id"] == int(artifact_id)
    assert body["artifact"]["file_sha256"] == exported.headers["x-requirement-rendered-file-sha256"]
    assert body["sheet_count"] >= 1
    flat = [cell for sheet in body["sheets"] for row in sheet["rows"] for cell in row if cell]
    assert any("需求" in cell or "字段" in cell for cell in flat), "预览必须能看到表格内容"


def test_preview_endpoint_serves_docx_sections(tmp_path: Path, monkeypatch, snapshot_api):
    """Word 产物也能在页面上看结构：标题/章节/字段表都在，且不重新渲染。"""

    client, db, project, field, req, _ = snapshot_api
    _use_temp_storage(tmp_path, monkeypatch)
    base = f"/projects/{project.id}/requirements/{req.id}"
    snapshot_id = _frozen_snapshot(client, base, req)

    exported = client.get(base + f"/snapshots/{snapshot_id}/export?format=docx")
    assert exported.status_code == 200, exported.text
    docx_id = exported.headers["x-requirement-render-artifact-id"]

    preview = client.get(base + f"/render-artifacts/{docx_id}/preview")
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["format"] == "docx"
    assert body["preview_version"] == "artifact-docx-preview-v1"
    assert body["artifact"]["format"] == "docx"
    assert body["artifact"]["file_sha256"] == exported.headers["x-requirement-rendered-file-sha256"]
    titles = body["section_titles"]
    assert "字段业务说明" in titles and "待确认事项" in titles
    assert body["truncated_blocks"] == 0

    missing = client.get(base + "/render-artifacts/999999/preview")
    assert missing.status_code == 404
