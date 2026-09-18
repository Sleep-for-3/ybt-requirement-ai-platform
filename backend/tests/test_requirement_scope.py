import pytest
from fastapi import HTTPException
from io import BytesIO
from openpyxl import load_workbook
from app.models import Project, ProductScenario, Requirement
from app.services.requirement_scope import ScopeInput, validate_scope, document_content, export_document
from test_lineage_paths import _seed_assets


def scope_fixture(db):
    _, _, project, field, _, _ = _seed_assets(db)
    scenario = ProductScenario(project_id=project.id, scenario_code="SYNTH", scenario_name="合成场景")
    db.add(scenario)
    db.flush()
    payload = ScopeInput(name="合成需求", target_table_id=field.target_table_id, field_ids=[field.id],
        scenario_id=scenario.id, background="仅用于范围与缺口测试")
    return project, field, payload


def test_scope_rejects_cross_project_fields(db_session):
    project, field, payload = scope_fixture(db_session)
    validate_scope(db_session, project.id, payload)
    other = Project(name="另一个项目")
    db_session.add(other)
    db_session.flush()
    with pytest.raises(HTTPException) as error:
        validate_scope(db_session, other.id, payload)
    assert error.value.status_code == 422


def test_missing_facts_are_gaps_even_without_question_rows_and_export_is_scoped(db_session):
    project, field, payload = scope_fixture(db_session)
    row = Requirement(project_id=project.id, name=payload.name, version=1,
        scope_json=payload.model_dump(mode="json", exclude={"expected_version", "name"}))
    db_session.add(row)
    db_session.flush()
    content = document_content(db_session, row)
    assert content["assessment"] == "gaps"
    assert any(g["id"] == f"{field.id}:source" for g in content["gaps"])
    assert content["status"] == "draft"
    assert document_content(db_session, row)["gaps"] == content["gaps"]
    book = load_workbook(BytesIO(export_document(content)))
    assert book["字段口径"].max_row == 2
    assert book["字段口径"]["A2"].value == field.field_code
    assert book["字段口径"]["D2"].alignment.wrap_text
    scope_values = {
        row[0].value: row[1].value
        for row in book["需求范围"].iter_rows(min_col=1, max_col=2)
    }
    assert scope_values["业务背景"] == payload.background
    assert book["待确认事项"]["A2"].value == f"{field.field_name} ({field.field_code})"


@pytest.mark.parametrize("status,resolution,expected", [
    ("accepted", "已接收待处理", True),
    ("resolved", "", True),
    ("closed", "", True),
    ("resolved", "已核验依据并补齐条件", False),
])
def test_question_resolution_requires_evidence(db_session, monkeypatch, status, resolution, expected):
    from app.services.requirement_workspace_projection import RequirementWorkspaceProjectionService
    project, field, payload = scope_fixture(db_session)
    row = Requirement(project_id=project.id, name=payload.name, version=1,
        scope_json=payload.model_dump(mode="json", exclude={"expected_version", "name"}))
    monkeypatch.setattr(RequirementWorkspaceProjectionService, "projection", lambda *args: {
        "question_summaries": [{"id": 100, "target_field_id": field.id,
            "question_status": status, "resolution_text": resolution, "question_text": "过滤依据待确认"}]
    })
    content = document_content(db_session, row)
    assert any(gap["id"] == "manual:100" for gap in content["gaps"]) is expected


def test_duplicate_scope_fields_do_not_duplicate_export_rows(db_session):
    project, field, payload = scope_fixture(db_session)
    row = Requirement(project_id=project.id, name=payload.name, version=1,
        scope_json={**payload.model_dump(mode="json", exclude={"expected_version", "name"}),
                    "field_ids": [field.id, field.id]})
    content = document_content(db_session, row)
    assert len(content["fields"]) == 1
    assert len({g["id"] for g in content["gaps"]}) == len(content["gaps"])


def test_export_preserves_layer_rules_and_evidence_and_escapes_formulas():
    mapping = {"id": 1, "mapping_name": "合成客户映射", "join_condition": "c.id=a.customer_id",
        "filter_condition": "c.status='ACTIVE'", "code_mapping_rule": "P→个人；C→企业",
        "null_handling_rule": "COALESCE(a.balance,0)", "quality_check_rule": "客户编号唯一"}
    content = {"requirement": {"name": "合成验收", "version": 1}, "gaps": [], "fields": [{
        "field": {"field_code": "CUSTOMER_ID", "field_name": "客户编号"},
        "business": {"final_content": "=SUM(1,2)"}, "lineage": {}, "mart_mappings": [mapping],
        "source_mappings": {"1": [mapping]},
        "evidence": [{"source_name": "合成字典", "location_text": "字段定义第2行", "quoted_content": "客户主键"}],
    }]}
    book = load_workbook(BytesIO(export_document(content)))
    assert book["字段口径"]["D2"].value == "'=SUM(1,2)"
    assert book["双层映射规则"].max_row == 3
    assert book["双层映射规则"]["G2"].value == mapping["join_condition"]
    assert book["双层映射规则"]["I3"].value == mapping["code_mapping_rule"]
    assert book["证据出处"]["C2"].value == "字段定义第2行"
