"""
Simplified end-to-end test for requirement formal delivery workflow.
Tests the core review and delivery mechanics without full content validation.
"""
import pytest
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.requirements import router
from app.core.database import get_db
from app.models import (
    Requirement,
    RequirementRevision,
    RequirementReviewSubmission,
    RequirementFormalDelivery,
    ReviewTask,
    WorkflowInstance,
    ProductScenario,
    ProjectMembership,
)
from app.services.auth.dependencies import Principal, get_current_principal
from test_lineage_paths import _seed_assets


@pytest.fixture
def delivery_client(db_session):
    """Test client with authenticated user."""
    institution, user, project, field, _, _ = _seed_assets(db_session)
    membership = ProjectMembership(
        project_id=project.id,
        user_id=user.id,
        project_role="project_manager",
        status="active"
    )
    db_session.add(membership)
    db_session.flush()

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_db] = lambda: db_session
    app.dependency_overrides[get_current_principal] = lambda: Principal(user.id, user.username, None)

    with TestClient(app) as client:
        yield client, db_session, user, project, field


def test_formal_delivery_api_integration(delivery_client):
    """Test formal delivery API endpoints work correctly."""
    client, db_session, user, project, field = delivery_client

    # Create requirement with minimal valid content
    scenario = ProductScenario(project_id=project.id, scenario_code="TEST", scenario_name="测试")
    db_session.add(scenario)
    db_session.flush()

    requirement = Requirement(
        project_id=project.id,
        name="API集成测试",
        version=1,
        content_version=1,
        scope_json={
            "target_table_id": field.target_table_id,
            "scenario_id": scenario.id,
            "field_ids": [field.id],
        },
    )
    db_session.add(requirement)
    db_session.flush()

    # Create a minimal valid revision
    from app.services.requirement_scope import content_digest
    content = {
        "requirement": {"id": requirement.id, "name": requirement.name},
        "fields": [{
            "field": {
                "id": field.id,
                "field_code": field.field_code,
                "field_name": field.field_name,
            },
            "business": {"business_definition": "测试定义"},
            "lineage": {
                "source_table_english_name": "TEST",
                "source_field_english_name": "ID",
                "processing_logic": "直接映射"
            },
            "evidence": []
        }],
        "resources": [],
        "gaps": [],
        "assessment": "clear",
        "lineage_graph": {"nodes": [], "edges": []}
    }

    revision = RequirementRevision(
        project_id=project.id,
        requirement_id=requirement.id,
        content_version=1,
        scope_version=1,
        parent_version=0,
        status="confirmed",  # Must be confirmed for finalize
        content_json=content,
        content_hash=content_digest(content),
        created_by=user.id,
    )
    db_session.add(revision)
    db_session.flush()

    # Manually create submission (bypassing validation for test)
    from datetime import datetime, UTC
    import hashlib

    # Calculate submission_hash correctly
    submission_hash_value = (
        f"{project.id}:{requirement.id}:{revision.id}:{revision.content_version}:"
        f"{revision.content_hash}:{user.id}"
    )
    submission_hash = hashlib.sha256(submission_hash_value.encode()).hexdigest()

    submission = RequirementReviewSubmission(
        project_id=project.id,
        requirement_id=requirement.id,
        revision_id=revision.id,
        content_version=1,
        content_hash=revision.content_hash,
        submission_hash=submission_hash,
        status="approved",
        submitted_by=user.id,
        submitted_at=datetime.now(UTC),
    )
    db_session.add(submission)
    db_session.flush()

    # Create approved workflow instance
    workflow = WorkflowInstance(
        project_id=project.id,
        workflow_key="requirement_document_review",
        target_type="requirement_review_submission",
        target_id=submission.id,
        status="approved",
        created_by=user.id,
    )
    db_session.add(workflow)
    db_session.commit()

    # Test finalize endpoint
    response = client.post(
        f"/projects/{project.id}/requirements/{requirement.id}/review-submissions/{submission.id}/finalize",
        json={}
    )
    if response.status_code != 201:
        print(f"\nFinalize failed with status {response.status_code}: {response.text}")
    assert response.status_code == 201, f"Finalize failed: {response.text}"
    delivery_data = response.json()
    assert "id" in delivery_data
    assert delivery_data["content_version"] == 1

    # Verify delivery was created
    delivery = db_session.get(RequirementFormalDelivery, delivery_data["id"])
    assert delivery is not None
    assert delivery.content_json is not None

    # Test download endpoint
    response = client.get(
        f"/projects/{project.id}/requirements/{requirement.id}/formal-deliveries/{delivery.id}/export"
    )
    assert response.status_code == 200
    assert "spreadsheet" in response.headers["content-type"]

    # Verify Excel has required sheets
    workbook = load_workbook(BytesIO(response.content))
    assert len(workbook.sheetnames) > 0


def test_cannot_finalize_unapproved_submission(delivery_client):
    """Cannot finalize a submission that hasn't been approved."""
    client, db_session, user, project, field = delivery_client

    scenario = ProductScenario(project_id=project.id, scenario_code="TEST", scenario_name="测试")
    db_session.add(scenario)
    db_session.flush()

    requirement = Requirement(
        project_id=project.id,
        name="未批准测试",
        version=1,
        content_version=1,
        scope_json={
            "target_table_id": field.target_table_id,
            "scenario_id": scenario.id,
            "field_ids": [field.id],
        },
    )
    db_session.add(requirement)
    db_session.flush()

    from app.services.requirement_scope import content_digest
    content = {"fields": [], "gaps": [], "lineage_graph": {}}
    revision = RequirementRevision(
        project_id=project.id,
        requirement_id=requirement.id,
        content_version=1,
        scope_version=1,
        parent_version=0,
        status="draft",
        content_json=content,
        content_hash=content_digest(content),
        created_by=user.id,
    )
    db_session.add(revision)
    db_session.flush()

    # Create unapproved submission
    from datetime import datetime, UTC
    submission = RequirementReviewSubmission(
        project_id=project.id,
        requirement_id=requirement.id,
        revision_id=revision.id,
        content_version=1,
        content_hash=revision.content_hash,
        submission_hash=revision.content_hash,
        status="pending_review",
        submitted_by=user.id,
        submitted_at=datetime.now(UTC),
    )
    db_session.add(submission)
    db_session.commit()

    # Try to finalize - should fail
    response = client.post(
        f"/projects/{project.id}/requirements/{requirement.id}/review-submissions/{submission.id}/finalize",
        json={}
    )
    assert response.status_code in (409, 422), "Should block finalize without approval"
