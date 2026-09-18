from copy import deepcopy
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import ProjectMembership, Requirement, ReviewTask, User
from app.services.auth.dependencies import Principal
from app.services.governance.workflow import decide_task
from app.services.requirement_revisions import append_revision, initialize_content, load_revision
from app.services.requirement_review import (
    finalize_formal_delivery,
    formal_file,
    review_readiness,
    submit_for_review,
)
from app.services.requirement_scope import content_digest
from app.services.storage import get_storage_service
from test_requirement_snapshot_api import snapshot_api, snapshot_db
from test_requirement_scope import scope_fixture


def _complete_revision(db):
    project, field, scope = scope_fixture(db)
    writer = db.scalar(select(User).order_by(User.id))
    requirement = Requirement(
        project_id=project.id,
        name=scope.name,
        version=1,
        scope_json=scope.model_dump(mode="json", exclude={"expected_version", "name"}),
    )
    db.add(requirement)
    db.flush()
    revision = initialize_content(db, requirement, 1, writer.id, include_lineage=True)
    content = deepcopy(revision.content_json)
    record = content["fields"][0]
    record["business"] = {
        **(record.get("business") or {}),
        "business_definition": "监管客户唯一标识",
        "final_content": "按有效客户主数据形成唯一客户编号。",
    }
    record["lineage"] = {
        **(record.get("lineage") or {}),
        "source_system_name": "客户主数据系统",
        "source_table_english_name": "ODS_CUSTOMER",
        "source_field_english_name": "CUSTOMER_ID",
        "processing_logic": "关联有效客户后直接映射；空值阻断。",
    }
    record["mart_mappings"] = [{
        "id": 301,
        "mart_field_id": 201,
        "mapping_status": "draft",
        "join_condition": "M.CUSTOMER_ID = C.CUSTOMER_ID",
        "filter_condition": "C.STATUS = 'ACTIVE'",
        "null_handling_rule": "CUSTOMER_ID IS NULL 时阻断",
    }]
    record["source_mappings"] = {"201": [{
        "id": 101,
        "mapping_status": "draft",
        "join_condition": "M.CUSTOMER_ID = C.CUSTOMER_ID",
        "filter_condition": "C.STATUS = 'ACTIVE'",
    }]}
    record["evidence"] = [{
        "source_name": "隔离合成字段字典",
        "location_text": "客户主键定义第 2 行",
        "quoted_content": "CUSTOMER_ID 为有效客户唯一标识",
    }]
    content["gaps"] = []
    content["assessment"] = "clear"
    if not content.get("lineage_graph") or not content["lineage_graph"].get("edges"):
        content["lineage_graph"] = {
            "nodes": [],
            "tables": [],
            "edges": [{
                "id": "synthetic:101",
                "source_node_id": "asset:source_field:101",
                "target_node_id": f"asset:target_field:{field.id}",
            }],
            "root_ids": [f"asset:target_field:{field.id}"],
            "truncated": False,
        }
    revision.content_json = content
    revision.content_hash = content_digest(content)
    db.commit()
    return project, field, requirement, revision, writer


def _reviewers(db, project_id):
    rows = {}
    for role in ("business_reviewer", "technical_reviewer", "final_reviewer"):
        user = User(username=f"requirement_{role}")
        db.add(user)
        db.flush()
        db.add(ProjectMembership(
            project_id=project_id,
            user_id=user.id,
            project_role=role,
            status="active",
        ))
        rows[role] = user
    db.commit()
    return rows


def test_blocking_gaps_prevent_requirement_review(db_session):
    project, _, scope = scope_fixture(db_session)
    writer = db_session.scalar(select(User).order_by(User.id))
    requirement = Requirement(project_id=project.id, name=scope.name, version=1,
        scope_json=scope.model_dump(mode="json", exclude={"expected_version", "name"}))
    db_session.add(requirement)
    db_session.flush()
    revision = initialize_content(db_session, requirement, 1, writer.id, include_lineage=False)
    readiness = review_readiness(db_session, project.id, requirement.id, revision.content_version)
    assert not readiness["eligible"]
    assert readiness["blocking_count"] >= 1
    assert any(item["code"] == "scope:lineage_graph" for item in readiness["reasons"])
    with pytest.raises(HTTPException) as error:
        submit_for_review(db_session, project_id=project.id, requirement_id=requirement.id,
            content_version=revision.content_version, content_hash=revision.content_hash,
            submitted_by=writer.id)
    assert error.value.status_code == 409


def test_three_stage_review_creates_immutable_formal_excel(db_session, tmp_path, monkeypatch):
    monkeypatch.setenv("STORAGE_DIR", str(tmp_path / "storage"))
    get_storage_service.cache_clear()
    project, _, requirement, revision, writer = _complete_revision(db_session)
    reviewers = _reviewers(db_session, project.id)
    submission, created = submit_for_review(
        db_session,
        project_id=project.id,
        requirement_id=requirement.id,
        content_version=revision.content_version,
        content_hash=revision.content_hash,
        submitted_by=writer.id,
        assignments={role: user.id for role, user in reviewers.items()},
    )
    assert created and submission.status == "pending_review"
    db_session.refresh(revision)
    assert revision.status == "in_review"
    tasks = {task.step_key: task for task in db_session.scalars(select(ReviewTask).where(
        ReviewTask.target_type == "requirement_review_submission",
        ReviewTask.target_id == submission.id,
    ))}
    for step, role in (
        ("business_review", "business_reviewer"),
        ("technical_review", "technical_reviewer"),
        ("final_review", "final_reviewer"),
    ):
        user = reviewers[role]
        decide_task(db_session, tasks[step], Principal(user.id, user.username, None),
                    "approved", f"{step} 通过")
    db_session.refresh(submission)
    db_session.refresh(revision)
    assert submission.status == "approved"
    assert revision.status == "confirmed"

    delivery, created = finalize_formal_delivery(
        db_session,
        project_id=project.id,
        requirement_id=requirement.id,
        submission_id=submission.id,
    )
    db_session.commit()
    assert created and delivery.version_no == 1
    original = formal_file(db_session, delivery)
    book = load_workbook(BytesIO(original))
    assert "正式交付" in book["需求范围"]["B2"].value
    assert book["需求范围"]["B3"].value == revision.content_version
    assert book["审核与交付"].max_row == 4
    assert book["审核与交付"]["C4"].value == reviewers["final_reviewer"].username

    next_content = deepcopy(revision.content_json)
    next_content["fields"][0]["business"]["remarks"] = "后续修订，不得改变历史交付"
    newer = append_revision(db_session, requirement, next_content, revision.content_version, writer.id)
    db_session.commit()
    assert newer.content_version == 2
    assert formal_file(db_session, delivery) == original
    repeated, created_again = finalize_formal_delivery(
        db_session,
        project_id=project.id,
        requirement_id=requirement.id,
        submission_id=submission.id,
    )
    assert not created_again and repeated.id == delivery.id
    get_storage_service.cache_clear()


def test_formal_delivery_http_permissions_and_frozen_download(snapshot_api, tmp_path, monkeypatch):
    monkeypatch.setenv("STORAGE_DIR", str(tmp_path / "api-storage"))
    get_storage_service.cache_clear()
    client, db, project, field, requirement, membership = snapshot_api
    base = f"/projects/{project.id}/requirements/{requirement.id}"
    assert client.post(base + "/revisions", json={"expected_version": 1}).status_code == 201
    revision = load_revision(db, project.id, requirement.id, 1)
    content = deepcopy(revision.content_json)
    record = content["fields"][0]
    record["business"] = {**(record.get("business") or {}),
        "business_definition": "客户唯一编号", "final_content": "有效客户唯一编号。"}
    record["lineage"] = {**(record.get("lineage") or {}),
        "source_table_english_name": "ODS_CUSTOMER", "source_field_english_name": "CUSTOMER_ID",
        "processing_logic": "过滤 STATUS='ACTIVE' 后直接映射。"}
    record["mart_mappings"] = [{"id": 21, "mart_field_id": 22, "mapping_status": "draft"}]
    record["source_mappings"] = {"22": [{"id": 23, "mapping_status": "draft"}]}
    record["evidence"] = [{"source_name": "隔离字典", "location_text": "第 2 行", "quoted_content": "客户编号"}]
    content["gaps"] = []
    content["assessment"] = "clear"
    content["lineage_graph"] = {"edges": [{"id": "e1", "source_node_id": "s1",
        "target_node_id": f"asset:target_field:{field.id}"}], "nodes": [], "tables": [],
        "root_ids": [f"asset:target_field:{field.id}"], "truncated": False}
    revision.content_json = content
    revision.content_hash = content_digest(content)
    db.commit()
    reviewers = _reviewers(db, project.id)
    submit = client.post(base + "/review-submissions", json={
        "expected_content_version": 1,
        "expected_content_hash": revision.content_hash,
        "assignments": {role: user.id for role, user in reviewers.items()},
    })
    assert submit.status_code == 201, submit.text
    submission_id = submit.json()["id"]
    assert submit.json()["workflow"]["current_step"] == "business_review"
    assert client.post(base + f"/review-submissions/{submission_id}/finalize").status_code == 409
    tasks = {task.step_key: task for task in db.scalars(select(ReviewTask).where(
        ReviewTask.target_type == "requirement_review_submission",
        ReviewTask.target_id == submission_id,
    ))}
    for step, role in (("business_review", "business_reviewer"),
                       ("technical_review", "technical_reviewer"),
                       ("final_review", "final_reviewer")):
        reviewer = reviewers[role]
        decide_task(db, tasks[step], Principal(reviewer.id, reviewer.username, None), "approved", "核验通过")
    finalized = client.post(base + f"/review-submissions/{submission_id}/finalize")
    assert finalized.status_code == 201, finalized.text
    delivery_id = finalized.json()["id"]
    repeated = client.post(base + f"/review-submissions/{submission_id}/finalize")
    assert repeated.status_code == 201 and repeated.json()["deduplicated"]
    membership.project_role = "viewer"
    db.commit()
    assert client.get(base + "/formal-deliveries").status_code == 200
    assert client.get(base + f"/formal-deliveries/{delivery_id}/export").status_code == 403
    membership.project_role = "auditor"
    db.commit()
    downloaded = client.get(base + f"/formal-deliveries/{delivery_id}/export")
    assert downloaded.status_code == 200
    assert "正式交付" in load_workbook(BytesIO(downloaded.content))["需求范围"]["B2"].value
    get_storage_service.cache_clear()
