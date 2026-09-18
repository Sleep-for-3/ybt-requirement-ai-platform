"""Synthetic metadata and files shared by isolated resource acceptance tests."""
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.models import (
    BusinessSystem, CatalogColumn, CatalogSchema, CatalogTable, DataSource,
    KnowledgeDocument, KnowledgeDocumentVersion, KnowledgeKeywordIndex, KnowledgeUnit,
    LineageEdge, LineageNode, MartField, MartTable, MartToYbtMapping, ProductScenario,
    Project, ScenarioTechnicalLineage, SourceField, SourceTable, SourceToMartMapping,
    TemplateDocument, TemplateParseResult, TemplateVersion,
)
from app.models import TargetField, TargetTable
from app.services.knowledge_ingestion.parsers import parse_document
from app.services.retrieval.keyword_index import weighted_tokens


def sample_files():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "监管答疑"
    sheet.append(["问题", "监管回复"])
    sheet.append(["客户证件类型报送范围", "合成制度：客户证件类型用于客户识别，适用范围须核验。"])
    excel = BytesIO(); workbook.save(excel); workbook.close()
    doc = Document()
    doc.add_heading("客户识别", level=1)
    doc.add_paragraph("合成制度：客户证件类型用于客户识别，适用范围须核验。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "字段"; table.cell(0, 1).text = "CERT_TYPE"
    word = BytesIO(); doc.save(word)
    writer = PdfWriter()
    page = writer.add_blank_page(width=420, height=595)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 16 Tf 40 540 Td (Synthetic CERT_TYPE policy evidence) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    pdf = BytesIO(); writer.write(pdf)
    scan_writer = PdfWriter(); scan_writer.add_blank_page(width=420, height=595)
    scan = BytesIO(); scan_writer.write(scan)
    return {
        "regulatory.xlsx": (excel.getvalue(), "regulatory_qa"),
        "policy.docx": (word.getvalue(), "regulatory_policy"),
        "policy.pdf": (pdf.getvalue(), "regulatory_policy"),
        "dictionary.txt": ("CERT_TYPE 字段解释：证件类型，来源 ECIF.CUSTOMER.CERT_TYPE。".encode(), "data_dictionary"),
        "notes.md": ("# 合成资料\n\n客户证件类型范围待确认。".encode(), "business_research"),
        "transform.sql": (b"select CERT_TYPE from ECIF.CUSTOMER where ACTIVE_FLAG = 'Y'", "sql_evidence"),
        "scan.pdf": (scan.getvalue(), "regulatory_policy"),
    }


def seed_resources(db):
    def add(model, **values):
        row = model(**values); db.add(row); db.flush(); return row
    project = add(Project, name="资料与数据隔离验收", bank_name="合成机构")
    other = add(Project, name="无访问权限的隔离项目", bank_name="合成机构")
    target_table = add(TargetTable, project_id=project.id, table_code="YBT_CUSTOMER", table_name="监管客户")
    target = add(TargetField, project_id=project.id, target_table_id=target_table.id, field_code="CERT_TYPE", field_name="客户证件类型")
    scenario = add(ProductScenario, project_id=project.id, scenario_code="CUSTOMER", scenario_name="客户报送")
    datasource = add(DataSource, project_id=project.id, name="合成客户目录", db_type="sqlite", database_name="CUSTOMER_DB")
    schema = add(CatalogSchema, project_id=project.id, datasource_id=datasource.id, schema_name="ECIF")
    table = add(CatalogTable, project_id=project.id, datasource_id=datasource.id, catalog_schema_id=schema.id, schema_name="ECIF", table_name="CUSTOMER")
    column = add(CatalogColumn, project_id=project.id, datasource_id=datasource.id, catalog_table_id=table.id, schema_name="ECIF", table_name="CUSTOMER", column_name="CERT_TYPE", column_comment="客户证件类型")
    system = add(BusinessSystem, project_id=project.id, system_code="ECIF", system_name="客户信息系统")
    source_table = add(SourceTable, project_id=project.id, business_system_id=system.id, datasource_id=datasource.id, database_name="CUSTOMER_DB", schema_name="ECIF", table_code="CUSTOMER", physical_table_name="CUSTOMER", table_name="客户")
    source = add(SourceField, project_id=project.id, source_table_id=source_table.id, field_code="CERT_TYPE", physical_column_name="CERT_TYPE", field_name="客户证件类型")
    mart_table = add(MartTable, project_id=project.id, table_code="MART_CUSTOMER", table_name="集市客户")
    mart_field = add(MartField, project_id=project.id, mart_table_id=mart_table.id, field_code="CERT_TYPE", field_name="客户证件类型")
    mapping = add(ScenarioTechnicalLineage, project_id=project.id, target_field_id=target.id, scenario_id=scenario.id,
        source_system_name=system.system_name, source_database_name="CUSTOMER_DB", source_schema_name="ECIF",
        source_table_english_name="CUSTOMER", source_field_english_name="CERT_TYPE", processing_logic="CERT_TYPE 按登记码值映射", tech_confirm_status="draft")
    add(MartToYbtMapping, project_id=project.id, target_field_id=target.id, mart_field_id=mart_field.id, business_rule="直接取已确认的集市字段")
    add(SourceToMartMapping, project_id=project.id, mart_field_id=mart_field.id, source_fields_summary="CUSTOMER.CERT_TYPE",
        join_condition="CUSTOMER.ID = ACCOUNT.CUSTOMER_ID", filter_condition="ACTIVE_FLAG = 'Y'", code_mapping_rule="01 对应身份证")
    source_node = add(LineageNode, project_id=project.id, node_type="column", logical_name="CUSTOMER.CERT_TYPE",
        catalog_column_id=column.id, source_field_id=source.id, table_name="CUSTOMER", column_name="CERT_TYPE", unresolved_flag=False)
    target_node = add(LineageNode, project_id=project.id, node_type="column", logical_name="YBT_CUSTOMER.CERT_TYPE",
        target_field_id=target.id, table_name="YBT_CUSTOMER", column_name="CERT_TYPE", unresolved_flag=False)
    # SQLite test-only metadata edge; no script execution or real repository exists.
    edge = add(LineageEdge, project_id=project.id, script_file_version_id=1, source_node_id=source_node.id, target_node_id=target_node.id,
        edge_type="value", join_condition="CUSTOMER.ID = ACCOUNT.CUSTOMER_ID", transformation_expression="CUSTOMER.CERT_TYPE")
    files, documents = {}, {}
    for name, (data, kind) in sample_files().items():
        doc = add(KnowledgeDocument, project_id=project.id, file_name=name, file_type=name.rsplit(".", 1)[1],
            source_type=kind, storage_path=f"fixture/{name}", knowledge_type=kind, document_status="indexed", parse_status="indexed",
            source_category="regulatory_formal" if name=="regulatory.xlsx" else "business_material",
            publisher="合成监管机构" if name=="regulatory.xlsx" else "合成机构")
        files[doc.storage_path] = data
        version = add(KnowledgeDocumentVersion, project_id=project.id, document_id=doc.id, version_no=1,
            file_name=name, storage_path=doc.storage_path, file_hash=f"{doc.id:064x}"[-64:], parse_status="indexed",
            lifecycle_status="active", regulatory_version="2026正式版" if name=="regulatory.xlsx" else None,
            internal_revision="R1", publisher=doc.publisher)
        doc.current_version_id=version.id;doc.current_version_no=1
        drafts, warnings = parse_document(name, data, kind)
        doc.warnings_json = warnings
        for draft in drafts:
            unit = add(KnowledgeUnit, project_id=project.id, document_id=doc.id, document_version_id=version.id,
                knowledge_type=kind, knowledge_scope="project", unit_type=draft.unit_type, title=draft.title,
                content=draft.content, normalized_content=draft.content, source_file_name=name,
                source_sheet_name=draft.source_sheet_name, source_page_no=draft.source_page_no,
                source_heading=draft.source_heading, source_cell_range=draft.source_cell_range, metadata_json=draft.metadata,
                confidentiality_level="internal", content_hash="0" * 64)
            for token, weight in weighted_tokens(unit.title, unit.content).items():
                add(KnowledgeKeywordIndex, project_id=project.id, knowledge_unit_id=unit.id, token=token, weight=weight)
        documents[name] = doc
    governed_doc=documents["regulatory.xlsx"]
    governed_active=db.get(KnowledgeDocumentVersion,governed_doc.current_version_id)
    governed_candidate=add(KnowledgeDocumentVersion,project_id=project.id,document_id=governed_doc.id,version_no=2,
        file_name="regulatory-2026-revision.xlsx",storage_path=governed_doc.storage_path,file_hash="e"*64,
        parse_status="indexed",lifecycle_status="pending_review",regulatory_version="2026修订版",internal_revision="R2",
        publisher="合成监管机构",replaces_version_id=governed_active.id,change_note="候选新口径，审核激活后才进入当前问答")
    governed_content="问题：客户证件类型报送范围 回复：合成制度新口径：客户证件类型按最新规则报送。"
    governed_unit=add(KnowledgeUnit,project_id=project.id,document_id=governed_doc.id,
        document_version_id=governed_candidate.id,knowledge_type="regulatory_qa",knowledge_scope="project",
        unit_type="qa_pair",title="客户证件类型报送范围",content=governed_content,normalized_content=governed_content,
        source_file_name=governed_candidate.file_name,source_sheet_name="监管答疑",source_cell_range="A2:B2",
        metadata_json={"locator":{"block_id":"sheet-监管答疑-row-2","sheet_name":"监管答疑","cell_range":"A2:B2"}},
        confidentiality_level="internal",enabled=False,content_hash="e"*64)
    for token,weight in weighted_tokens(governed_unit.title,governed_unit.content).items():
        add(KnowledgeKeywordIndex,project_id=project.id,knowledge_unit_id=governed_unit.id,token=token,weight=weight)
    failed_doc=documents["policy.pdf"]
    add(KnowledgeDocumentVersion,project_id=project.id,document_id=failed_doc.id,version_no=2,
        file_name="制度修订稿.pdf",storage_path=failed_doc.storage_path,file_hash="f"*64,
        parse_status="failed",lifecycle_status="draft",change_note="候选修订解析失败，不替换当前生效版本")

    template_file=sample_files()["regulatory.xlsx"][0]
    files["fixture/template-v1.xlsx"]=template_file;files["fixture/template-v2.xlsx"]=template_file
    template=add(TemplateDocument,project_id=project.id,file_name="一表通客户模板-v2.xlsx",file_type="xlsx",
        storage_path="fixture/template-v2.xlsx",sheet_names_json=["客户表"],parse_status="success",
        template_code="YBT_CUSTOMER",display_name="客户信息一表通模板")
    base_rows=[{"row_number":2,"field_code":"CERT_TYPE","field_name":"客户证件类型","field_type":"VARCHAR","required_flag":False,"regulatory_description":"客户证件类型"}]
    next_rows=[{"row_number":2,"field_code":"CERT_TYPE","field_name":"客户证件类型","field_type":"VARCHAR(20)","required_flag":True,"regulatory_description":"按最新监管定义报送"},{"row_number":3,"field_code":"CERT_NO","field_name":"客户证件号码","field_type":"VARCHAR(64)","required_flag":True,"regulatory_description":"客户证件号码"}]
    v1=add(TemplateVersion,template_document_id=template.id,project_id=project.id,version_no=1,
        regulatory_version="2025批次",template_code="YBT_CUSTOMER",publisher="合成监管机构",status="active",
        file_name="一表通客户模板-v1.xlsx",file_type="xlsx",storage_path="fixture/template-v1.xlsx",file_hash="1"*64,
        sheet_names_json=["客户表"],parsed_snapshot_json=[{"sheet_name":"客户表","table_code":"YBT_CUSTOMER","table_name":"监管客户","rows":base_rows}],parse_status="success")
    v2=add(TemplateVersion,template_document_id=template.id,project_id=project.id,version_no=2,
        regulatory_version="2026批次",release_batch="2026-Q3",template_code="YBT_CUSTOMER",publisher="合成监管机构",
        status="pending_review",replaces_version_id=v1.id,change_note="新增证件号码并调整必填性",
        file_name="一表通客户模板-v2.xlsx",file_type="xlsx",storage_path="fixture/template-v2.xlsx",file_hash="2"*64,
        sheet_names_json=["客户表"],parsed_snapshot_json=[{"sheet_name":"客户表","table_code":"YBT_CUSTOMER","table_name":"监管客户","rows":next_rows}],parse_status="success")
    template.current_version_id=v1.id
    for version,rows in ((v1,base_rows),(v2,next_rows)):
        add(TemplateParseResult,template_document_id=template.id,template_version_id=version.id,project_id=project.id,
            sheet_name="客户表",table_code="YBT_CUSTOMER",table_name="监管客户",field_count=len(rows),
            raw_header_json=[],parsed_rows_json=rows,warnings_json=[])
    db.commit()
    return SimpleNamespace(project=project, other=other, target=target, scenario=scenario, column=column,
        source=source, mapping=mapping, edge=edge, documents=documents, files=files, template=template,
        template_versions=(v1,v2))
