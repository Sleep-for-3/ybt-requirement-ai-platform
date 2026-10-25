"""P1-B 客户 XLSX 需求文档模板渲染器的定向测试。

全部使用合成模板（``backend/tests/`` 下自建 fixture），不依赖真实客户文件。
覆盖：命名单元格填充、行块 0/1/100 行与样式复制、合并单元格、打印区域/打印标题、
页眉页脚、公式保留与公式注入防护、必需映射缺失、不能保真的对象（图表等）、
32767 字符上限、公式未重算的口径诊断。
"""
from __future__ import annotations

import io
import zipfile
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from app.services.requirement_xlsx_template import (
    FORMULA_CACHE_NOTE,
    MAX_EXCEL_CELL_CHARS,
    XlsxCellTooLongError,
    XlsxMappingMissingError,
    XlsxRenderValidationError,
    XlsxTemplateConfigError,
    XlsxTemplateHashMismatch,
    XlsxUnsupportedFeatureError,
    analyze_xlsx_template,
    render_xlsx,
    template_sha256,
)

SHEETS = ("封面", "字段口径", "审核与交付")
SAMPLE_ROW = 3
FILL = PatternFill(fill_type="solid", start_color="DDEBF7", end_color="DDEBF7")
THIN = Side(style="thin", color="8EA9DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


# --------------------------------------------------------------------------- #
# 合成模板
# --------------------------------------------------------------------------- #
def build_template(
    *,
    with_chart: bool = False,
    with_rules: bool = False,
) -> bytes:
    """构造一个带命名单元格、合并单元格、公式、打印设置、页眉页脚的合成模板。"""

    wb = Workbook()
    cover = wb.active
    cover.title = "封面"
    fields = wb.create_sheet("字段口径")
    review = wb.create_sheet("审核与交付")

    # ---- 封面 ----
    cover.merge_cells("A1:D1")
    cover["A1"] = "客户需求文档（合成测试模板）"
    cover["A1"].font = Font(bold=True, size=14)
    cover["A1"].alignment = Alignment(horizontal="center", vertical="center")
    cover.row_dimensions[1].height = 30

    labels = [
        ("A3", "需求名称", "B3", "需求名称"),
        ("A4", "业务目标", "B4", "业务目标"),
        ("A5", "业务背景", "B5", "业务背景"),
        ("A6", "交付状态", "B6", "交付状态"),
        ("A7", "内容版本", "B7", "内容版本"),
        ("A8", "固定内容标识", "B8", "内容标识"),
    ]
    for label_cell, label, value_cell, name in labels:
        cover[label_cell] = label
        cover[label_cell].font = Font(bold=True)
        cover[value_cell] = "（待填）"
        cover[value_cell].alignment = WRAP
        wb.defined_names[name] = DefinedName(name, attr_text=f"'封面'!${value_cell[0]}${value_cell[1:]}")

    cover["A9"] = "字段条数"
    cover["B9"] = "=COUNTA('字段口径'!A3:A200)"
    cover.column_dimensions["A"].width = 16
    cover.column_dimensions["B"].width = 40
    cover.print_area = "A1:D9"
    cover.print_title_rows = "1:2"
    cover.oddHeader.center.text = "客户需求文档 & 机密"
    cover.oddFooter.right.text = "第 &P 页 / 共 &N 页"

    # ---- 字段口径 ----
    fields.merge_cells("A1:F1")
    fields["A1"] = "字段口径"
    fields["A1"].font = Font(bold=True, size=12)
    headers = ("A", "B", "C", "D", "E", "F")
    for letter, text in zip(headers, ("字段编码", "字段名称", "字段类型", "业务口径", "来源系统", "校验列")):
        cell = fields[f"{letter}2"]
        cell.value = text
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", start_color="F2F2F2", end_color="F2F2F2")

    sample = {
        "A": "样板编码",
        "B": "样板名称",
        "C": "string",
        "D": "样板口径",
        "E": "样板系统",
    }
    for letter, text in sample.items():
        cell = fields[f"{letter}{SAMPLE_ROW}"]
        cell.value = text
        cell.font = Font(bold=True, name="微软雅黑")
        cell.fill = FILL
        cell.border = BORDER
        cell.alignment = WRAP
    fields[f"F{SAMPLE_ROW}"] = f'=IF(A{SAMPLE_ROW}="","",A{SAMPLE_ROW}&"-"&B{SAMPLE_ROW})'
    fields[f"F{SAMPLE_ROW}"].border = BORDER
    fields.row_dimensions[SAMPLE_ROW].height = 24

    fields.merge_cells("A4:F4")
    fields["A4"] = "注：第 3 行是样板行，会被复制成多行；本合并备注应在插行后下移。"
    fields["A4"].font = Font(italic=True, color="808080")
    fields.print_area = "A1:F200"
    fields.print_title_rows = "1:2"
    fields.print_title_cols = "A:A"
    fields.oddHeader.center.text = "字段口径"
    fields.oddFooter.left.text = "内部资料"

    # ---- 审核与交付 ----
    review["A1"] = "审核与交付"
    review["A2"] = "交付状态"
    review["B2"] = "待审核"

    if with_rules:
        validation = DataValidation(type="list", formula1='"通过,退回"', allow_blank=True)
        validation.add("B2")
        review.add_data_validation(validation)
        review.conditional_formatting.add(
            "A2:B2",
            CellIsRule(operator="equal", formula=['"通过"'], fill=PatternFill(
                fill_type="solid", start_color="C6EFCE", end_color="C6EFCE")),
        )
        review["B2"].comment = Comment("合成批注：仅用于检测批注能力", "tester")

    if with_chart:
        chart = BarChart()
        chart.title = "合成图表"
        chart.add_data(Reference(cover, min_col=2, min_row=3, max_row=5), titles_from_data=False)
        cover.add_chart(chart, "F2")

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def build_macro_package() -> bytes:
    """在合成模板包里伪造一个 vbaProject.bin 与 macroEnabled 主部件（模拟 .xlsm）。"""

    source = build_template()
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as original, \
            zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in original.infolist():
            payload = original.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = payload.decode("utf-8")
                text = text.replace(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
                text = text.replace(
                    "</Types>",
                    '<Override PartName="/xl/vbaProject.bin" '
                    'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                )
                payload = text.encode("utf-8")
            target.writestr(item, payload)
        target.writestr("xl/vbaProject.bin", b"\x00fake-vba-project\x00")
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# 冻结 content 与配置
# --------------------------------------------------------------------------- #
def build_content(field_count: int = 3) -> dict[str, Any]:
    fields = []
    for index in range(1, field_count + 1):
        fields.append({
            "field": {
                "field_code": f"F{index:03d}",
                "field_name": f"字段名称{index}",
                "field_type": "string",
            },
            "business": {"final_content": f"业务口径{index}"},
            "lineage": {"source_system_name": "核心系统", "final_content": "源字段"},
        })
    return {
        "requirement": {
            "name": "反洗钱可疑交易报送口径",
            "objective": "统一报送口径，减少人工核对",
            "background": "监管要求在 2026 年内完成口径收敛",
            "inclusion": "对私客户、对公客户",
            "exclusion": "内部测试账户",
            "effective_date": "2026-11-01",
            "acceptance_criteria": "抽样比对一致率 99% 以上",
        },
        "fields": fields,
        "gaps": [],
        "formal_delivery": {
            "version_no": "V1.2",
            "approved_at": "2026-10-20T10:00:00",
            "content_version": "req-2026.10.20-1",
        },
        "script_basis": {"digest": "abc"},
    }


def build_config(template: bytes, **overrides: Any) -> dict[str, Any]:
    config: dict[str, Any] = {
        "template_hash": template_sha256(template),
        "config_version": "p1-b-test-1",
        "supported_scope": "requirement_document",
        "cell_map": {
            "需求名称": {"named_range": "需求名称", "value_path": "requirement.name", "required": True},
            "业务目标": {"named_range": "业务目标", "value_path": "requirement.objective"},
            "业务背景": {"named_range": "业务背景", "value_path": "requirement.background"},
            "交付状态": {"named_range": "交付状态", "value_path": "document.status"},
            "内容版本": {"named_range": "内容版本", "value_path": "document.content_version"},
            "内容标识": {"named_range": "内容标识", "value_path": "document.digest"},
        },
        "row_blocks": {
            "字段清单": {
                "sheet": "字段口径",
                "sample_row": SAMPLE_ROW,
                "columns": {"编码": "A", "名称": "B", "类型": "C", "口径": "D", "来源": "E"},
                "column_paths": {
                    "编码": "field.field_code",
                    "名称": "field.field_name",
                    "类型": "field.field_type",
                    "口径": "business.final_content",
                    "来源": "lineage.source_system_name",
                },
                "max_rows": 200,
                "source": "fields",
            },
        },
        "required_keys": ["需求名称", "字段清单"],
    }
    config.update(overrides)
    return config


def load_output(payload: bytes):
    return load_workbook(io.BytesIO(payload), data_only=False)


# --------------------------------------------------------------------------- #
# analyze
# --------------------------------------------------------------------------- #
def test_analyze_reports_layout_and_config_binding() -> None:
    template = build_template()
    analysis = analyze_xlsx_template(template, build_config(template))

    assert analysis.template_hash == template_sha256(template)
    assert analysis.template_hash_matches_config is True
    assert analysis.sheet_names == SHEETS
    assert analysis.config_version == "p1-b-test-1"
    assert analysis.blocking_missing == ()

    names = {info.name: info for info in analysis.defined_names}
    assert set(names) >= {"需求名称", "业务目标", "业务背景", "交付状态", "内容版本", "内容标识"}
    assert names["需求名称"].sheet == "封面"
    assert names["需求名称"].ref == "$B$3"
    assert names["需求名称"].is_cell is True

    assert analysis.merged_ranges["封面"] == ("A1:D1",)
    assert set(analysis.merged_ranges["字段口径"]) == {"A1:F1", "A4:F4"}
    assert analysis.print_areas["封面"] == "'封面'!$A$1:$D$9"
    assert analysis.print_titles["封面"] == ("$1:$2", None)
    assert analysis.headers_footers["封面"]["oddHeader"][1] == "客户需求文档 & 机密"
    assert analysis.headers_footers["封面"]["oddFooter"][2] == "第 &P 页 / 共 &N 页"

    assert len(analysis.row_blocks) == 1
    block = analysis.row_blocks[0]
    assert (block.key, block.sheet, block.sample_row, block.first_data_row) == ("字段清单", "字段口径", 3, 3)
    assert block.columns["编码"] == "A"
    assert block.resolved is True

    assert analysis.unsupported_features == ()
    assert analysis.cell_targets["需求名称"].label == "封面!B3"
    assert analysis.summary().startswith("模板 sha256")


def test_analyze_reports_missing_named_range_and_sheet() -> None:
    template = build_template()
    config = build_config(template)
    config["cell_map"]["需求名称"] = {"named_range": "不存在的命名单元格"}
    config["cell_map"]["业务目标"] = {"sheet": "不存在的表", "cell": "B4"}

    analysis = analyze_xlsx_template(template, config)

    missing = {(item.kind, item.key): item for item in analysis.missing}
    assert ("named_range", "需求名称") in missing
    assert "不存在的命名单元格" in missing[("named_range", "需求名称")].detail
    assert "现有命名单元格" in missing[("named_range", "需求名称")].detail
    assert ("sheet", "业务目标") in missing
    assert "不存在的表" in missing[("sheet", "业务目标")].detail
    assert "封面、字段口径、审核与交付" in missing[("sheet", "业务目标")].detail
    assert {item.key for item in analysis.blocking_missing} >= {"需求名称", "业务目标"}


def test_analyze_detects_data_validation_conditional_formatting_comment() -> None:
    template = build_template(with_rules=True)
    analysis = analyze_xlsx_template(template, build_config(template))

    kinds = set(analysis.lossy_kinds)
    assert {"data_validation", "conditional_formatting", "comment"} <= kinds
    assert "数据验证" in analysis.summary()


def test_analyze_detects_vba_macro_package() -> None:
    template = build_macro_package()
    analysis = analyze_xlsx_template(template, build_config(template))

    kinds = set(analysis.lossy_kinds)
    assert "vba" in kinds
    assert any("vbaProject.bin" in feature.location or "工作簿" in feature.location
               for feature in analysis.unsupported_features if feature.kind == "vba")


# --------------------------------------------------------------------------- #
# 命名单元格填充 + 版式保真
# --------------------------------------------------------------------------- #
def test_render_fills_named_cells_and_keeps_layout() -> None:
    template = build_template()
    output, diagnostics = render_xlsx(template, build_content(3), build_config(template))
    analysis = analyze_xlsx_template(template, build_config(template))
    wb = load_output(output)

    cover = wb["封面"]
    assert cover["B3"].value == "反洗钱可疑交易报送口径"
    assert cover["B4"].value == "统一报送口径，减少人工核对"
    assert cover["B6"].value == "正式交付（已审核）"
    assert cover["B7"].value == "req-2026.10.20-1"
    assert len(cover["B8"].value) == 64
    assert cover["A1"].value == "客户需求文档（合成测试模板）"

    # 行列尺寸与单元格样式必须原样保留（写值不能把样式冲掉）
    assert cover.column_dimensions["A"].width == 16
    assert cover.column_dimensions["B"].width == 40
    assert cover.row_dimensions[1].height == 30
    assert cover["A1"].font.bold is True
    assert cover["A1"].font.size == 14
    assert cover["B3"].alignment.wrap_text is True
    assert cover["B3"].alignment.vertical == "top"

    assert tuple(wb.sheetnames) == SHEETS
    # 3 条字段 = 样板行 + 2 条复制行，样板行下方的合并备注由 A4:F4 平移到 A6:F6。
    expected_ranges = {
        "封面": analysis.merged_ranges["封面"],
        "字段口径": ("A1:F1", f"A{SAMPLE_ROW + 3}:F{SAMPLE_ROW + 3}"),
        "审核与交付": analysis.merged_ranges["审核与交付"],
    }
    for sheet in SHEETS:
        assert tuple(sorted(str(rng) for rng in wb[sheet].merged_cells.ranges)) == expected_ranges[sheet]
        assert len(expected_ranges[sheet]) == len(analysis.merged_ranges[sheet])

    assert wb["封面"].print_area == analysis.print_areas["封面"]
    assert wb["封面"].print_title_rows == "$1:$2"
    assert wb["封面"].print_title_cols is None
    assert wb["封面"].oddHeader.center.text == "客户需求文档 & 机密"
    assert wb["封面"].oddFooter.right.text == "第 &P 页 / 共 &N 页"

    assert diagnostics.row_block_counts == {"字段清单": 3}
    assert diagnostics.lossy_rendered is False
    assert diagnostics.output_sha256 == template_sha256(output)
    assert {cell.logical_key for cell in diagnostics.filled_cells} >= {"需求名称", "业务目标", "交付状态"}


def test_render_row_block_zero_rows_clears_sample_values_but_keeps_style() -> None:
    template = build_template()
    config = build_config(template, required_keys=["需求名称"])
    output, diagnostics = render_xlsx(template, build_content(0), config)
    wb = load_output(output)

    ws = wb["字段口径"]
    for letter in ("A", "B", "C", "D", "E"):
        cell = ws[f"{letter}{SAMPLE_ROW}"]
        assert cell.value is None, f"{letter}{SAMPLE_ROW} 应被清空"
        assert cell.fill.fgColor.rgb == "00DDEBF7"
        assert cell.font.name == "微软雅黑"
        assert cell.border.left.style == "thin"

    assert diagnostics.row_block_counts == {"字段清单": 0}
    assert any("已清空" in warning for warning in diagnostics.warnings)
    assert wb["封面"]["B3"].value == "反洗钱可疑交易报送口径"


def test_render_row_block_single_row() -> None:
    template = build_template()
    output, diagnostics = render_xlsx(template, build_content(1), build_config(template))
    wb = load_output(output)

    ws = wb["字段口径"]
    assert ws["A3"].value == "F001"
    assert ws["B3"].value == "字段名称1"
    assert ws["C3"].value == "string"
    assert ws["D3"].value == "业务口径1"
    assert ws["E3"].value == "核心系统"
    assert ws["F3"].value == '=IF(A3="","",A3&"-"&B3)'
    assert ws["F3"].data_type == "f"
    assert ws["A4"].value == "注：第 3 行是样板行，会被复制成多行；本合并备注应在插行后下移。"
    assert diagnostics.row_block_counts == {"字段清单": 1}


def test_render_row_block_hundred_rows_copies_values_styles_and_shifts_merges() -> None:
    template = build_template()
    analysis = analyze_xlsx_template(template, build_config(template))
    output, diagnostics = render_xlsx(template, build_content(100), build_config(template))
    wb = load_output(output)

    ws = wb["字段口径"]
    assert diagnostics.row_block_counts == {"字段清单": 100}
    for index in range(1, 101):
        row = SAMPLE_ROW + index - 1
        assert ws[f"A{row}"].value == f"F{index:03d}"
        assert ws[f"B{row}"].value == f"字段名称{index}"
        assert ws[f"C{row}"].value == "string"
        assert ws[f"D{row}"].value == f"业务口径{index}"
        assert ws[f"E{row}"].value == "核心系统"
        style_cell = ws[f"A{row}"]
        assert style_cell.fill.fgColor.rgb == "00DDEBF7"
        assert style_cell.font.name == "微软雅黑"
        assert style_cell.border.left.style == "thin"
        assert style_cell.alignment.wrap_text is True
        assert ws.row_dimensions[row].height == 24
        assert ws[f"F{row}"].data_type == "f"
        assert ws[f"F{row}"].value == f'=IF(A{row}="","",A{row}&"-"&B{row})'

    assert ws[f"A{SAMPLE_ROW + 100}"].value == "注：第 3 行是样板行，会被复制成多行；本合并备注应在插行后下移。"
    # 合并区域数量必须与模板一致；样板行下方的合并备注应随插行整体平移到第 103 行。
    shifted_note = f"A{SAMPLE_ROW + 100}:F{SAMPLE_ROW + 100}"
    rendered_ranges = tuple(sorted(str(rng) for rng in ws.merged_cells.ranges))
    assert len(rendered_ranges) == len(analysis.merged_ranges["字段口径"]), (
        f"合并区域数量应保持 {len(analysis.merged_ranges['字段口径'])} 个，实际 {rendered_ranges}")
    assert set(rendered_ranges) == {"A1:F1", shifted_note}, (
        f"合并区域应为标题 A1:F1 与平移后的备注 {shifted_note}，实际 {rendered_ranges}")
    assert ws.print_area == analysis.print_areas["字段口径"]
    assert ws.print_title_rows == "$1:$2"
    assert ws.print_title_cols == "$A:$A"
    assert ws.oddHeader.center.text == "字段口径"
    assert ws.oddFooter.left.text == "内部资料"


def test_render_keeps_template_formula_but_text_escapes_business_formula() -> None:
    template = build_template()
    content = build_content(3)
    content["requirement"]["background"] = "=SUM(1,2)"
    output, diagnostics = render_xlsx(template, content, build_config(template))
    wb = load_output(output)

    cover = wb["封面"]
    assert cover["B9"].value == "=COUNTA('字段口径'!A3:A200)"
    assert cover["B9"].data_type == "f"

    escaped = cover["B5"]
    assert escaped.value == "=SUM(1,2)"
    assert escaped.data_type == "s"
    assert escaped.quotePrefix is True

    background = next(cell for cell in diagnostics.filled_cells if cell.logical_key == "业务背景")
    assert background.escaped_as_text is True
    assert any("强制写成文本" in warning for warning in diagnostics.warnings)


@pytest.mark.parametrize("payload", ["=1+1", "+86-10", "-42", "@SUM(A1)"])
def test_render_escapes_every_formula_injection_prefix(payload: str) -> None:
    template = build_template()
    content = build_content(1)
    content["requirement"]["name"] = payload
    output, _ = render_xlsx(template, content, build_config(template))
    cell = load_output(output)["封面"]["B3"]
    assert cell.value == payload
    assert cell.data_type == "s"
    assert cell.quotePrefix is True


def test_render_keeps_plain_text_unmodified() -> None:
    template = build_template()
    output, diagnostics = render_xlsx(template, build_content(1), build_config(template))
    cell = load_output(output)["封面"]["B3"]
    assert cell.value == "反洗钱可疑交易报送口径"
    assert cell.data_type == "s"
    assert cell.quotePrefix in (None, False)
    assert all(not filled.escaped_as_text for filled in diagnostics.filled_cells)


# --------------------------------------------------------------------------- #
# 失败路径
# --------------------------------------------------------------------------- #
def test_render_missing_required_named_range_fails() -> None:
    template = build_template()
    config = build_config(template)
    config["cell_map"]["需求名称"] = {"named_range": "不存在的命名单元格", "required": True}

    with pytest.raises(XlsxMappingMissingError) as excinfo:
        render_xlsx(template, build_content(3), config)

    message = str(excinfo.value)
    assert "不存在的命名单元格" in message
    assert "需求名称" in message
    assert excinfo.value.code == "XLSX_MAPPING_MISSING"


def test_render_missing_sheet_mapping_fails_with_sheet_names() -> None:
    template = build_template()
    config = build_config(template)
    config["cell_map"]["业务目标"] = {"sheet": "不存在的表", "cell": "B4"}

    with pytest.raises(XlsxMappingMissingError) as excinfo:
        render_xlsx(template, build_content(3), config)

    assert "不存在的表" in str(excinfo.value)
    assert "封面、字段口径、审核与交付" in str(excinfo.value)


def test_render_required_row_block_without_data_fails() -> None:
    template = build_template()
    with pytest.raises(XlsxMappingMissingError) as excinfo:
        render_xlsx(template, build_content(0), build_config(template))
    assert "字段清单" in str(excinfo.value)
    assert "没有数据" in str(excinfo.value)


def test_render_rejects_template_hash_mismatch() -> None:
    template = build_template()
    config = build_config(template, template_hash="0" * 64)
    with pytest.raises(XlsxTemplateHashMismatch):
        render_xlsx(template, build_content(1), config)


def test_config_rejects_bad_hash_and_unknown_required_key() -> None:
    with pytest.raises(XlsxTemplateConfigError) as excinfo:
        analyze_xlsx_template(b"PK\x03\x04", {"template_hash": "not-a-sha"})
    assert "64 位十六进制" in str(excinfo.value)

    template = build_template()
    config = build_config(template, required_keys=["不存在的键"])
    with pytest.raises(XlsxTemplateConfigError) as excinfo2:
        analyze_xlsx_template(template, config)
    assert "不存在的键" in str(excinfo2.value)


def test_render_rejects_unsupported_chart_by_default_and_allows_when_lossy() -> None:
    template = build_template(with_chart=True)
    config = build_config(template)
    analysis = analyze_xlsx_template(template, config)
    assert "chart" in analysis.lossy_kinds

    with pytest.raises(XlsxUnsupportedFeatureError) as excinfo:
        render_xlsx(template, build_content(3), config)
    assert "图表" in str(excinfo.value)
    assert "allow_lossy" in str(excinfo.value)
    assert excinfo.value.code == "XLSX_UNSUPPORTED_FEATURE"

    output, diagnostics = render_xlsx(template, build_content(3), build_config(template, allow_lossy=True))
    assert diagnostics.lossy_rendered is True
    assert any(feature.kind == "chart" for feature in diagnostics.unsupported_features)
    assert any("[allow_lossy]" in warning and "图表" in warning for warning in diagnostics.warnings)
    wb = load_output(output)
    assert wb["封面"]["B3"].value == "反洗钱可疑交易报送口径"


def test_render_rejects_cell_text_over_excel_limit() -> None:
    template = build_template()
    content = build_content(1)
    content["requirement"]["name"] = "长" * (MAX_EXCEL_CELL_CHARS + 1)

    with pytest.raises(XlsxCellTooLongError) as excinfo:
        render_xlsx(template, content, build_config(template))

    message = str(excinfo.value)
    assert str(MAX_EXCEL_CELL_CHARS) in message
    assert "需求名称" in message
    assert excinfo.value.code == "XLSX_CELL_TOO_LONG"


def test_render_warns_when_text_exceeds_configured_max_cell_chars() -> None:
    template = build_template()
    content = build_content(1)
    content["requirement"]["name"] = "长" * 1200
    output, diagnostics = render_xlsx(template, content, build_config(template, max_cell_chars=1000))
    assert any("max_cell_chars=1000" in warning for warning in diagnostics.warnings)
    assert load_output(output)["封面"]["B3"].value == "长" * 1200


def test_render_validation_error_aborts_without_output() -> None:
    """残留占位符必须被读回校验拦下，且不产出字节。"""

    template = build_template()
    content = build_content(1)
    content["requirement"]["name"] = "{{需求名称}}"

    with pytest.raises(XlsxRenderValidationError) as excinfo:
        render_xlsx(template, content, build_config(template))

    message = str(excinfo.value)
    assert "残留占位符" in message
    assert "封面!B3" in message
    assert excinfo.value.code == "XLSX_RENDER_VALIDATION_FAILED"


# --------------------------------------------------------------------------- #
# 公式未重算口径
# --------------------------------------------------------------------------- #
def test_diagnostics_never_claim_formula_recalculated() -> None:
    template = build_template()
    output, diagnostics = render_xlsx(template, build_content(1), build_config(template))

    assert diagnostics.formula_recalculated is False
    assert diagnostics.formula_cache_verified is False
    assert diagnostics.formula_cache_note == FORMULA_CACHE_NOTE
    assert "不重算" in diagnostics.formula_cache_note
    assert diagnostics.recalculation_tool_available in (None, "libreoffice")
    assert diagnostics.recalculation_note
    assert "重算" in diagnostics.recalculation_note

    cached = load_workbook(io.BytesIO(output), data_only=True)["封面"]
    # 没有 Excel/LibreOffice 重算，缓存值就是空；这里只断言「不宣称有值」。
    if cached["B9"].value is not None:
        assert isinstance(cached["B9"].value, (int, float))
    assert load_output(output)["封面"]["B9"].data_type == "f"
