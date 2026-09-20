from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile
import hashlib

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import select, func

from app.models import (BackgroundJob, CatalogColumn, CatalogTable, Institution, LineageNode, LineageRevision,
                        LineageRevisionNode, Project, ResourceImportBatch, ScriptFileVersion, User)
from app.models.data_architecture import CatalogClassification
from app.services.metadata import batch_import as service
from app.services.metadata.batch_parser import parse_file
from app.services import data_architecture


class MemoryStorage:
    def __init__(self):
        self.files = {}
        self.fail = False

    def save(self, data, *, file_name, project_id=None):
        key = hashlib.sha256(data).hexdigest()
        self.files[key] = data
        return SimpleNamespace(storage_key=key, content_hash=key, byte_size=len(data))

    def read(self, key):
        if self.fail:
            raise OSError("synthetic unavailable storage")
        return self.files[key]


@pytest.fixture
def sample(db_session, monkeypatch):
    db = db_session
    institution = Institution(institution_code="isolated-batch-bank", institution_name="隔离银行")
    db.add(institution); db.flush()
    project = Project(name="隔离批量导入", institution_id=institution.id)
    user = User(username="isolated-import-user")
    db.add_all([project, user]); db.flush()
    storage = MemoryStorage()
    monkeypatch.setattr(service, "get_storage_service", lambda: storage)
    return db, project, user, storage


def make_batch(sample, files, options=None, key="one"):
    db, project, user, storage = sample
    return service.create_batch(db, project, user.id, key, files, options or {}, storage)


def apply_batch(sample, batch, key="job"):
    db, project, user, _ = sample
    batch.status = "queued"
    job = BackgroundJob(project_id=project.id, created_by=user.id, idempotency_key=key,
        job_type="resource_batch_import", status="running", payload_summary_json={"batch_id": batch.id, "preview_version": batch.version})
    db.add(job); db.commit()
    result = service.batch_import_handler(db, job)
    return result, job


def dictionary():
    workbook = Workbook()
    for index, sheet in enumerate([workbook.active, workbook.create_sheet("第二表")]):
        sheet.append(["库名", "Schema", "表名", "表中文名", "字段名", "字段类型", "字段中文名"])
        sheet.append(["bank", "ods", f"excel_{index}", "人工字典", "id", "BIGINT", "客户标识"])
    output = BytesIO(); workbook.save(output)
    return output.getvalue()


def test_ddl_materializes_columns_comments_database_identity_and_offline_sources(sample):
    db, _, _, _ = sample
    batch = make_batch(sample, [("a.ddl", b"CREATE TABLE a.s.same (id BIGINT NOT NULL);"),
        ("b.ddl", b"CREATE TABLE b.s.same (id VARCHAR(20));"), ("dictionary.xlsx", dictionary())])
    assert db.scalar(select(func.count()).select_from(CatalogTable)) == 0
    result, _ = apply_batch(sample, batch)
    assert result["success_count"] == 3 and not result["failed_count"]
    tables = list(db.scalars(select(CatalogTable)))
    assert len(tables) == 4
    assert {(x.database_name, x.schema_name, x.table_name) for x in tables} >= {("a", "s", "same"), ("b", "s", "same")}
    assert db.scalar(select(CatalogColumn).where(CatalogColumn.table_name == "excel_0")).column_comment == "客户标识"


def test_duplicate_request_and_repeated_apply_do_not_duplicate_assets(sample):
    db, _, _, _ = sample
    files = [("a.ddl", b"CREATE TABLE a.s.t (id INT);")]
    first = make_batch(sample, files)
    assert make_batch(sample, files).id == first.id
    with pytest.raises(HTTPException) as error:
        make_batch(sample, [("a.ddl", b"CREATE TABLE a.s.t (other INT);")])
    assert error.value.status_code == 409
    _, job = apply_batch(sample, first)
    service.batch_import_handler(db, job)
    second = make_batch(sample, files, key="two")
    assert second.preview_json["items"][0]["tables"][0]["operation"] == "duplicate"
    apply_batch(sample, second, "job-two")
    assert db.scalar(select(func.count()).select_from(CatalogTable)) == 1
    assert db.scalar(select(func.count()).select_from(CatalogColumn)) == 1


def test_conflicts_require_explicit_merge_and_preserve_existing_annotations(sample):
    db, _, _, _ = sample
    first = make_batch(sample, [("a.ddl", b"CREATE TABLE bank.s.t (id INT);")])
    apply_batch(sample, first)
    table = db.scalar(select(CatalogTable)); table.table_comment = "人工确认说明"
    column = db.scalar(select(CatalogColumn)); column.column_comment = "人工字段注释"; db.commit()
    second = make_batch(sample, [("b.ddl", b"CREATE TABLE bank.s.t (id VARCHAR(30), new_col DATE);")], key="two")
    plan = second.preview_json["items"][0]["tables"][0]
    assert plan["operation"] == "conflict"
    second = service.save_options(db, second, 1, {"table_options": {plan["key"]: {"decision": "merge"}}})
    result, _ = apply_batch(sample, second, "two-job")
    assert result["success_count"] == 1
    db.refresh(table); db.refresh(column)
    assert table.table_comment == "人工确认说明" and column.column_comment == "人工字段注释" and column.data_type == "INT"
    assert db.scalar(select(func.count()).select_from(CatalogColumn)) == 2


def test_partial_failures_and_storage_retry_are_truthful(sample):
    db, _, _, storage = sample
    batch = make_batch(sample, [("ok.ddl", b"CREATE TABLE b.s.t (id INT);"), ("bad.ddl", b"not valid ddl")])
    result, job = apply_batch(sample, batch)
    assert result["success_count"] == 1 and result["failed_count"] == 1 and batch.status == "partial"
    result = service.batch_import_handler(db, job)
    assert result["success_count"] == 1 and result["failed_count"] == 1
    assert db.scalar(select(func.count()).select_from(CatalogTable)) == 1
    retry = make_batch(sample, [("retry.ddl", b"CREATE TABLE b.s.retry (id INT);")], key="retry")
    storage.fail = True
    result, job = apply_batch(sample, retry, "retry-job")
    assert result["failed_count"] == 1
    storage.fail = False
    assert service.batch_import_handler(db, job)["success_count"] == 1


def test_preview_change_is_detected_and_does_not_overwrite_user_data(sample):
    db, _, _, _ = sample
    first = make_batch(sample, [("a.ddl", b"CREATE TABLE b.s.t (id INT);")])
    apply_batch(sample, first)
    second = make_batch(sample, [("a.ddl", b"CREATE TABLE b.s.t (id INT, added DATE);")], key="two")
    table = db.scalar(select(CatalogTable)); table.table_comment = "预览后的编辑"; db.commit()
    result, _ = apply_batch(sample, second, "two-job")
    assert result["failed_count"] == 1
    assert db.scalar(select(func.count()).select_from(CatalogColumn)) == 1


def test_file_layer_override_and_existing_classification_protected(sample):
    db, project, _, _ = sample
    data_architecture.save_architecture(db, data_architecture.examples()[0], 0, project_id=project.id)
    batch = make_batch(sample, [("a.ddl", b"CREATE TABLE b.s.a (id INT);"), ("b.ddl", b"CREATE TABLE b.s.b (id INT);")],
        {"defaults": {"layer_key": "layer_0"}, "file_options": {"b.ddl": {"layer_key": "layer_2"}}})
    result, _ = apply_batch(sample, batch)
    assert not result["failed_count"]
    bindings = list(db.scalars(select(CatalogClassification).order_by(CatalogClassification.id)))
    assert [x.assignment_json["layer_key"] for x in bindings] == ["layer_0", "layer_2"]
    duplicate = make_batch(sample, [("a.ddl", b"CREATE TABLE b.s.a (id INT);")], {"defaults": {"layer_key": "layer_1"}}, key="two")
    apply_batch(sample, duplicate, "two-job")
    db.refresh(bindings[0]); assert bindings[0].assignment_json["layer_key"] == "layer_0"


def test_cross_layer_script_reads_never_inherit_batch_default(sample):
    db, project, _, _ = sample
    data_architecture.save_architecture(db, data_architecture.examples()[0], 0, project_id=project.id)
    tables = make_batch(sample, [("tables.ddl", b"CREATE TABLE b.s.source (id INT); CREATE TABLE b.s.target (id INT);")])
    apply_batch(sample, tables)
    batch = make_batch(sample, [("batch.sql", b"INSERT INTO b.s.target (id) SELECT id FROM b.s.source;")],
        {"defaults": {"layer_key": "layer_2"}}, key="sql")
    refs = batch.preview_json["items"][0]["references"]
    assert any(x["role"] == "read" and x["proposed_layer_key"] is None for x in refs)
    result, job = apply_batch(sample, batch, "sql-job")
    assert not result["failed_count"]
    assert db.scalar(select(func.count()).select_from(ScriptFileVersion)) == 1
    service.batch_import_handler(db, job)
    assert db.scalar(select(func.count()).select_from(ScriptFileVersion)) == 1
    bindings = list(db.scalars(select(CatalogClassification)))
    assert len(bindings) == 1 and db.get(CatalogTable, bindings[0].catalog_table_id).table_name == "target"


def test_batch_reresolves_earlier_sql_after_later_ddl_is_applied(sample):
    db, project, _, _ = sample
    # The SQL item is applied first and initially cannot see the tables created
    # by the following DDL item. The final batch pass must rebind those nodes
    # before building the immutable lineage revision.
    batch = make_batch(sample, [
        ("01_load.sql", b"INSERT INTO b.s.target (id) SELECT id FROM b.s.source;"),
        ("02_tables.ddl", b"CREATE TABLE b.s.source (id INT); CREATE TABLE b.s.target (id INT);"),
    ])
    result, _ = apply_batch(sample, batch, "interleaved-metadata-job")

    assert result["success_count"] == 2 and not result["failed_count"]
    columns = list(db.scalars(select(LineageNode).where(
        LineageNode.project_id == project.id,
        LineageNode.node_type == "column",
    )))
    assert columns and all(node.catalog_column_id is not None for node in columns)

    revision = db.scalar(select(LineageRevision).where(LineageRevision.project_id == project.id))
    snapshots = list(db.scalars(select(LineageRevisionNode).where(LineageRevisionNode.revision_id == revision.id)))
    assert snapshots
    assert any(row.snapshot_json.get("catalog_column_id") is not None for row in snapshots)

    db.refresh(batch)
    sql_references = next(item for item in batch.preview_json["items"] if item["path"] == "01_load.sql")["references"]
    assert sql_references
    assert all(reference["catalog_table_id"] is not None for reference in sql_references)


def test_script_references_match_case_folded_identity_without_cross_schema_merge(sample):
    db, project, _, _ = sample
    batch = make_batch(sample, [
        ("01_load.sql", b"INSERT INTO B.S1.TARGET (ID) SELECT ID FROM B.S2.TARGET;"),
        ("02_tables.ddl", b"CREATE TABLE b.s1.target (id INT); CREATE TABLE b.s2.target (id INT);"),
    ])
    result, _ = apply_batch(sample, batch, "case-folded-identity-job")

    assert result["success_count"] == 2 and not result["failed_count"]
    db.refresh(batch)
    references = next(item for item in batch.preview_json["items"] if item["path"] == "01_load.sql")["references"]
    resolved = {reference["table_name"].lower(): reference["catalog_table_id"] for reference in references}
    assert resolved["target"] is not None
    assert len({reference["catalog_table_id"] for reference in references}) == 2


@pytest.mark.parametrize("path", ["../escape.sql", "/absolute.sql", "..\\escape.sql"])
def test_zip_path_traversal_rejected_without_assets(sample, path):
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr(path, "SELECT 1")
    with pytest.raises(HTTPException):
        make_batch(sample, [("archive.zip", stream.getvalue())])
    assert sample[0].scalar(select(func.count()).select_from(ResourceImportBatch)) == 0


def test_static_ddl_comments_and_unknown_sql_are_not_executed():
    _, data = parse_file("CREATE TABLE b.s.t (id BIGINT); COMMENT ON COLUMN b.s.t.id IS '编号';".encode(), "a.ddl")
    assert data["tables"][0]["columns"][0]["column_comment"] == "编号"
    _, data = parse_file(b"EXECUTE IMMEDIATE 'DROP TABLE real_data';", "a.sql")
    assert not data["complete"] and data["warnings"]


def test_missing_insert_column_list_is_a_gap():
    _, data = parse_file(b"INSERT INTO b.s.t SELECT id FROM b.s.source;", "a.sql")
    assert not data["complete"] and any("目标字段" in x for x in data["warnings"])


def test_excel_uses_explicit_batch_database_and_schema_defaults(sample):
    workbook = Workbook(); workbook.active.append(["表名", "字段名", "字段类型"])
    workbook.active.append(["customer", "id", "INT"])
    stream = BytesIO(); workbook.save(stream)
    batch = make_batch(sample, [("plain.xlsx", stream.getvalue())], {"defaults": {"database_name": "bank_a", "schema_name": "custom"}})
    result, _ = apply_batch(sample, batch)
    assert result["success_count"] == 1
    table = sample[0].scalar(select(CatalogTable))
    assert (table.database_name, table.schema_name) == ("bank_a", "custom")


def test_zip_limits_and_supported_multifile_sql(sample):
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        for index in range(101):
            archive.writestr(f"{index}.sql", "SELECT 1")
    with pytest.raises(HTTPException):
        make_batch(sample, [("too-many.zip", stream.getvalue())])
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("upstream.sql", "INSERT INTO b.s.mid (id) SELECT id FROM b.s.source;")
        archive.writestr("downstream.sql", "INSERT INTO b.s.target (id) SELECT id FROM b.s.mid;")
    batch = make_batch(sample, [("safe.zip", stream.getvalue())])
    result, _ = apply_batch(sample, batch)
    assert result["success_count"] == 2 and not result["failed_count"]
    assert all(x.result_json["script_version_id"] for x in service.items(sample[0], batch))


def test_real_http_permissions_queue_retry_and_reopen(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from sqlalchemy.pool import StaticPool
    from app.core.database import Base, get_db
    from app.api import batch_imports
    from app.models import ProjectMembership
    from app.services.auth.dependencies import Principal, get_current_principal
    from app.services.task_queue.inline import InlineTaskQueue
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = Session(engine, expire_on_commit=False)
    bank = Institution(institution_code="http-test-bank", institution_name="HTTP 隔离银行")
    db.add(bank); db.flush()
    project = Project(name="HTTP测试", institution_id=bank.id)
    foreign = Project(name="另一项目", institution_id=bank.id)
    user = User(username="isolated-http-import")
    viewer = User(username="isolated-http-viewer")
    db.add_all([project, foreign, user, viewer]); db.flush()
    db.add_all([ProjectMembership(project_id=project.id, user_id=user.id, project_role="project_manager"),
                ProjectMembership(project_id=project.id, user_id=viewer.id, project_role="viewer")]); db.commit()
    storage = MemoryStorage()
    monkeypatch.setattr(service, "get_storage_service", lambda: storage)
    monkeypatch.setattr(batch_imports, "get_task_queue", lambda: InlineTaskQueue())
    active = {"user": Principal(user.id, user.username, None)}
    app = FastAPI(); app.include_router(batch_imports.router, prefix="/api")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_principal] = lambda: active["user"]
    client = TestClient(app)
    root = f"/api/projects/{project.id}/resource-imports"
    try:
        response = client.post(root, data={"idempotency_key": "http-one", "options": "{}"},
            files=[("files", ("a.ddl", b"CREATE TABLE bank.s.t (id INT);", "text/plain")),
                   ("files", ("bad.ddl", b"invalid structure", "text/plain"))])
        assert response.status_code == 201, response.text
        batch_id = response.json()["id"]
        response = client.post(f"{root}/{batch_id}/apply", json={"expected_version": 1})
        assert response.status_code == 200, response.text
        assert response.json()["status"] == "partial"
        job_id = response.json()["job_id"]
        assert db.get(BackgroundJob, job_id).status == "partially_completed"
        assert client.post(f"{root}/{batch_id}/apply", json={"expected_version": 1}).json()["job_id"] == job_id
        assert client.post(f"{root}/{batch_id}/retry").json()["status"] == "partial"
        reopened = client.post(f"{root}/{batch_id}/reopen").json()
        options = {"file_options": {"bad.ddl": {"decision": "skip"}}}
        updated = client.put(f"{root}/{batch_id}/preview", json={"expected_version": reopened["version"], "options": options})
        assert updated.status_code == 200, updated.text
        completed = client.post(f"{root}/{batch_id}/apply", json={"expected_version": updated.json()["version"]})
        assert completed.json()["status"] == "completed", completed.text
        assert db.scalar(select(func.count()).select_from(CatalogTable)) == 1
        assert client.get(f"/api/projects/{foreign.id}/resource-imports/{batch_id}").status_code in (403, 404)
        active["user"] = Principal(viewer.id, viewer.username, None)
        assert client.get(f"{root}/{batch_id}").status_code == 403
        assert client.post(f"{root}/{batch_id}/apply", json={"expected_version": 1}).status_code == 403
    finally:
        db.close(); engine.dispose()


def test_batch_migration_is_additive():
    import importlib.util
    from pathlib import Path
    from sqlalchemy import create_engine, inspect
    from sqlalchemy.orm import Session
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from app.core.database import Base
    engine = create_engine("sqlite://")
    added = {"resource_import_batches", "resource_import_items"}
    Base.metadata.create_all(engine, tables=[table for table in Base.metadata.sorted_tables if table.name not in added])
    with Session(engine) as db:
        db.add(Project(name="迁移前的项目")); db.commit()
    spec = importlib.util.spec_from_file_location("batch_migration", Path(__file__).resolve().parents[1] / "alembic/versions/202609170035_resource_import_batches.py")
    migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(migration)
    with engine.begin() as connection:
        migration.op = Operations(MigrationContext.configure(connection)); migration.upgrade()
        assert added <= set(inspect(connection).get_table_names())
    with Session(engine) as db:
        assert db.scalar(select(Project)).name == "迁移前的项目"
    engine.dispose()


def test_sparse_workbook_cannot_amplify_parser_memory():
    workbook = Workbook(); workbook.active["ZZ10000"] = "sparse"
    stream = BytesIO(); workbook.save(stream)
    with pytest.raises(ValueError, match="稀疏"):
        parse_file(stream.getvalue(), "sparse.xlsx")


def test_old_job_cannot_apply_a_new_preview_version(sample):
    batch = make_batch(sample, [("a.ddl", b"CREATE TABLE bank.s.t (id INT);")])
    _, job = apply_batch(sample, batch)
    batch.version += 1; sample[0].commit()
    with pytest.raises(ValueError):
        service.batch_import_handler(sample[0], job)


def test_preview_reports_missing_comment_fill_as_modification(sample):
    first = make_batch(sample, [("a.ddl", b"CREATE TABLE bank.s.t (id INT);")])
    apply_batch(sample, first)
    second = make_batch(sample, [("b.ddl", "CREATE TABLE bank.s.t (id INT); COMMENT ON COLUMN bank.s.t.id IS '编号';".encode())], key="comment-fill")
    assert second.preview_json["items"][0]["tables"][0]["operation"] == "modify"
    result, _ = apply_batch(sample, second, "comment-job")
    assert not result["failed_count"]
    assert sample[0].scalar(select(CatalogColumn)).column_comment == "编号"
