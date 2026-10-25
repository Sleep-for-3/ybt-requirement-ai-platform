"""P0-D 客户模板注册：注册一次、导出多次。

守的边界：

- 注册时必须能跑通渲染器自己的分析器（缺锚点/命名单元格直接 422，不存半可用模板）；
- 同名多版本、同文件同配置幂等（不产生重复版本）；
- 导出可以只给 ``template_id``，且与现场上传模板得到**同一份字节**（同一渲染身份）；
- 已注册模板不接受"换一份配置"或"再叠一个文件"，避免注册结论与实际渲染不一致。
"""

from __future__ import annotations

import hashlib
import io
import json

import pytest
from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.customer_templates import router as templates_router
from app.api.requirements import router as requirements_router
from app.core.database import get_db
from app.models import ProjectMembership, User
from app.services import requirement_delivery as service
from app.services.auth.dependencies import Principal, get_current_principal
from app.services.auth.resource_guard import guard_project_resource
from app.services.requirement_scope import content_digest
from app.services.storage.factory import get_storage_service
from docx_template_fixtures import CUSTOMER_ANCHORS, build_customer_template, customer_config
from test_requirement_draft_snapshots import fixture
from test_requirement_snapshot_api import snapshot_db  # noqa: F401
from test_requirement_xlsx_template import build_config as build_xlsx_config
from test_requirement_xlsx_template import build_template as build_xlsx_template

REGISTER_URL = "/projects/{pid}/customer-templates"
LIST_URL = "/projects/{pid}/customer-templates"
DETAIL_URL = "/projects/{pid}/customer-templates/{tid}"
UNIFIED_EXPORT_URL = "/projects/{pid}/requirements/{rid}/exports"
TEMPLATE_EXPORT_URL = "/projects/{pid}/requirements/{rid}/template-exports"


@pytest.fixture
def registry_api(snapshot_db, tmp_path, monkeypatch):  # noqa: F811
    monkeypatch.setenv("STORAGE_DIR", str(tmp_path / "storage"))
    get_storage_service.cache_clear()
    db_session = snapshot_db
    project, field, requirement, content = fixture(db_session)
    user = db_session.scalar(select(User))
    membership = ProjectMembership(project_id=project.id, user_id=user.id,
                                   project_role="project_manager", status="active")
    db_session.add(membership)
    db_session.flush()
    snapshot, _created = service.freeze_draft(db_session, project.id, requirement.id,
                                             requirement.version, content_digest(content), None)
    db_session.commit()
    app = FastAPI()
    app.include_router(templates_router, dependencies=[Depends(guard_project_resource)])
    app.include_router(requirements_router, dependencies=[Depends(guard_project_resource)])
    app.dependency_overrides[get_db] = lambda: db_session
    app.dependency_overrides[get_current_principal] = lambda: Principal(user.id, user.username, None)
    with TestClient(app) as client:
        yield client, db_session, project, membership, requirement, snapshot
    get_storage_service.cache_clear()


def _register(client, project_id: int, *, name="客户需求文档模板", fmt="docx",
              template_bytes=None, config=None, notes=""):
    template_bytes = template_bytes or build_customer_template()
    config = config if config is not None else customer_config(template_bytes)
    return client.post(
        REGISTER_URL.format(pid=project_id),
        data={"name": name, "format": fmt,
              "config": json.dumps(config, ensure_ascii=False), "notes": notes},
        files={"template": (f"template.{fmt}", template_bytes, "application/octet-stream")},
    )


def _export_with_template_id(client, project_id, requirement_id, snapshot_id, template_id, fmt="docx"):
    return client.post(
        UNIFIED_EXPORT_URL.format(pid=project_id, rid=requirement_id),
        data={"format": fmt, "source_kind": "draft_snapshot",
              "source_id": str(snapshot_id), "template_id": str(template_id)},
    )


def test_register_is_idempotent_for_same_file_and_config(registry_api):
    client, db, project, _membership, _req, _snap = registry_api
    template_bytes = build_customer_template()
    first = _register(client, project.id, template_bytes=template_bytes)
    assert first.status_code == 201, first.text
    body = first.json()
    assert body["created"] is True
    assert body["template"]["version_no"] == 1
    assert body["template"]["supported_scope"]["supported"] is True
    assert body["template"]["file_sha256"] == hashlib.sha256(template_bytes).hexdigest()

    second = _register(client, project.id, template_bytes=template_bytes)
    assert second.status_code == 201
    assert second.json()["created"] is False
    assert second.json()["template"]["id"] == body["template"]["id"]
    assert client.get(LIST_URL.format(pid=project.id)).json()["count"] == 1


def test_register_new_version_when_config_changes(registry_api):
    client, db, project, *_ = registry_api
    template_bytes = build_customer_template()
    first = _register(client, project.id, template_bytes=template_bytes).json()["template"]
    relaxed = customer_config(
        template_bytes,
        required_keys=[key for key in CUSTOMER_ANCHORS if key != "table.review"],
        optional_keys=("table.review",),
    )
    assert relaxed != customer_config(template_bytes)
    second = _register(client, project.id, template_bytes=template_bytes, config=relaxed)
    assert second.status_code == 201, second.text
    assert second.json()["created"] is True, second.text
    assert second.json()["template"]["version_no"] == 2, second.text
    latest = client.get(LIST_URL.format(pid=project.id)).json()["items"]
    assert [item["version_no"] for item in latest] == [2]
    all_versions = client.get(LIST_URL.format(pid=project.id),
                             params={"latest_only": False}).json()["items"]
    assert [item["version_no"] for item in all_versions] == [1, 2]
    assert first["id"] != latest[0]["id"]


def test_register_rejects_template_missing_anchors(registry_api):
    client, db, project, *_ = registry_api
    template_bytes = build_customer_template()
    config = customer_config(template_bytes)
    broken = build_customer_template(omit_key="section.policy")
    response = _register(client, project.id, template_bytes=broken,
                         config=customer_config(broken))
    assert response.status_code == 422, response.text
    assert "template_anchor_missing" in response.json()["detail"]
    assert "section.policy" in response.json()["detail"]
    assert client.get(LIST_URL.format(pid=project.id)).json()["count"] == 0


def test_register_rejects_invalid_config_and_unknown_format(registry_api):
    client, db, project, *_ = registry_api
    template_bytes = build_customer_template()
    legacy = dict(customer_config(template_bytes))
    legacy["anchors"] = {key: value for key, value in legacy["anchors"].items()
                         if key != "section.policy"}
    bad = _register(client, project.id, template_bytes=template_bytes, config=legacy)
    assert bad.status_code == 422
    assert "template_config_invalid" in bad.json()["detail"]

    pdf = _register(client, project.id, fmt="pdf")
    assert pdf.status_code == 422
    assert "unsupported_format" in pdf.json()["detail"]


def test_registry_detail_exposes_config_and_scope(registry_api):
    client, db, project, *_ = registry_api
    registered = _register(client, project.id).json()["template"]
    detail = client.get(DETAIL_URL.format(pid=project.id, tid=registered["id"]))
    assert detail.status_code == 200
    payload = detail.json()
    assert payload["name"] == "客户需求文档模板"
    assert "cover.title" in payload["config"]["anchors"]
    assert payload["supported_scope"]["missing"] == []
    assert client.get(DETAIL_URL.format(pid=project.id, tid=9999)).status_code == 404


def test_export_by_template_id_matches_upload_path(registry_api):
    client, db, project, _membership, requirement, snapshot = registry_api
    template_bytes = build_customer_template()
    registered = _register(client, project.id, template_bytes=template_bytes).json()["template"]

    by_id = _export_with_template_id(client, project.id, requirement.id, snapshot.id,
                                     registered["id"])
    assert by_id.status_code == 200, by_id.text
    assert by_id.headers["x-requirement-template-sha256"] == registered["file_sha256"]

    config = customer_config(template_bytes)
    by_upload = client.post(
        TEMPLATE_EXPORT_URL.format(pid=project.id, rid=requirement.id),
        data={"format": "docx", "snapshot_id": str(snapshot.id),
              "config": json.dumps(config, ensure_ascii=False)},
        files={"template": ("template.docx", template_bytes, "application/octet-stream")},
    )
    assert by_upload.status_code == 200, by_upload.text
    assert hashlib.sha256(by_id.content).hexdigest() == hashlib.sha256(by_upload.content).hexdigest()
    assert (by_id.headers["x-requirement-render-artifact-id"]
            == by_upload.headers["x-requirement-render-artifact-id"])


def test_export_by_template_id_rejects_conflicting_inputs(registry_api):
    client, db, project, _membership, requirement, snapshot = registry_api
    template_bytes = build_customer_template()
    registered = _register(client, project.id, template_bytes=template_bytes).json()["template"]
    base = {"format": "docx", "source_kind": "draft_snapshot", "source_id": str(snapshot.id)}
    url = UNIFIED_EXPORT_URL.format(pid=project.id, rid=requirement.id)

    both = client.post(url, data={**base, "template_id": str(registered["id"])},
                       files={"template": ("template.docx", template_bytes,
                                          "application/octet-stream")})
    assert both.status_code == 422
    assert "不要再上传模板文件" in both.json()["detail"]

    mismatched = client.post(
        url, data={**base, "template_id": str(registered["id"]),
                   "config": json.dumps(customer_config(template_bytes,
                                                        optional_keys=("table.review",)),
                                        ensure_ascii=False)})
    assert mismatched.status_code == 422
    assert "已注册模板不一致" in mismatched.json()["detail"]

    assert client.post(url, data={**base, "template_id": "9999"}).status_code == 404

    neither = client.post(url, data={**base, "config": "{}"})
    assert neither.status_code == 200


def test_registered_xlsx_template_exports(registry_api):
    client, db, project, _membership, requirement, snapshot = registry_api
    template_bytes = build_xlsx_template()
    registered = _register(client, project.id, name="客户 Excel 模板", fmt="xlsx",
                           template_bytes=template_bytes,
                           config=build_xlsx_config(template_bytes)).json()["template"]
    response = _export_with_template_id(client, project.id, requirement.id, snapshot.id,
                                        registered["id"], fmt="xlsx")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml")
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames


def test_viewer_cannot_register(registry_api):
    client, db, project, membership, _req, _snap = registry_api
    membership.project_role = "viewer"
    db.commit()
    assert _register(client, project.id).status_code == 403
