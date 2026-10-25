"""内置 Word 版式的结构自检：标题/元信息/章节/字段块/固定内容标识都必须是可判定结论。"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from test_requirement_snapshot_api import snapshot_api, snapshot_db  # noqa: F401  夹具注册

from app.services.requirement_docx_structure_check import (
    BUILTIN_DOCX_RENDERER,
    DOCX_CHECK_VERSION,
    DocxStructureCheckError,
    check_docx_structure,
    contract_for_docx_renderer,
)

SECTIONS = (
    "需求范围与业务说明",
    "字段业务说明",
    "来源系统与分层路径",
    "加工规则与脚本证据",
    "制度依据与差异",
    "待确认事项",
)


def _doc(payload: bytes) -> Document:
    return Document(BytesIO(payload))


def _save(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _paragraph(document: Document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    return None


@pytest.fixture()
def builtin_docx(tmp_path, monkeypatch, snapshot_api) -> bytes:
    """走真实导出链路拿内置版式 Word 字节（不另造内容夹具，避免与导出器契约漂移）。"""

    from test_requirement_export_unified import _frozen_snapshot, _use_temp_storage

    client, db, project, field, req, _ = snapshot_api
    _use_temp_storage(tmp_path, monkeypatch)
    base = f"/projects/{project.id}/requirements/{req.id}"
    snapshot_id = _frozen_snapshot(client, base, req)
    response = client.post(
        f"{base}/exports",
        data={
            "format": "docx",
            "source_kind": "draft_snapshot",
            "source_id": str(snapshot_id),
            "layout": "standard",
        },
    )
    assert response.status_code == 200, response.text
    return response.content


def test_builtin_export_passes_docx_structure_check(builtin_docx):
    result = check_docx_structure(builtin_docx, renderer_version=BUILTIN_DOCX_RENDERER)
    assert result["check_version"] == DOCX_CHECK_VERSION
    assert result["verdict"] == "pass"
    assert result["failed_keys"] == []
    assert result["section_titles"][: len(SECTIONS)] == list(SECTIONS)
    assert result["field_section_count"] >= 1
    checks = {item["key"]: item for item in result["checks"]}
    for key in ("title", "meta", "sections", "fields", "content_id"):
        assert checks[key]["ok"] is True, key
    assert checks["placeholders"]["severity"] == "info"
    assert result["contract"]["renderer_version"] == BUILTIN_DOCX_RENDERER


def test_missing_section_fails_sections_only(builtin_docx):
    document = _doc(builtin_docx)
    paragraph = _paragraph(document, "制度依据与差异")
    assert paragraph is not None
    paragraph.runs[0].text = "制度依据"
    for run in list(paragraph.runs)[1:]:
        run.text = ""
    result = check_docx_structure(_save(document), renderer_version=BUILTIN_DOCX_RENDERER)
    assert result["verdict"] == "fail"
    assert result["failed_keys"] == ["sections"]


def test_broken_content_id_fails_content_id_check(builtin_docx):
    document = _doc(builtin_docx)
    paragraph = None
    for item in document.paragraphs:
        if item.text.startswith("固定内容标识"):
            paragraph = item
            break
    assert paragraph is not None
    paragraph.runs[0].text = "固定内容标识：not-a-hash"
    for run in list(paragraph.runs)[1:]:
        run.text = ""
    result = check_docx_structure(_save(document), renderer_version=BUILTIN_DOCX_RENDERER)
    assert "content_id" in result["failed_keys"]


def test_duplicate_section_title_fails_sections_check(builtin_docx):
    document = _doc(builtin_docx)
    paragraph = _paragraph(document, "待确认事项")
    assert paragraph is not None
    paragraph.runs[0].text = "来源系统与分层路径"
    for run in list(paragraph.runs)[1:]:
        run.text = ""
    result = check_docx_structure(_save(document), renderer_version=BUILTIN_DOCX_RENDERER)
    assert result["verdict"] == "fail"
    assert "sections" in result["failed_keys"]


def test_customer_template_renderer_is_refused_not_scored():
    with pytest.raises(DocxStructureCheckError) as caught:
        check_docx_structure(b"whatever", renderer_version="customer-template-render-v1")
    assert caught.value.code == "artifact_not_builtin_layout"
    assert "不是平台内置版式" in caught.value.message


def test_contract_lookup_and_unreadable_payload():
    contract = contract_for_docx_renderer("builtin-docx-v1")
    assert contract is not None
    assert contract["required_keys"] == ["title", "meta", "sections", "fields", "content_id"]
    assert contract_for_docx_renderer("customer-template-render-v1") is None
    with pytest.raises(DocxStructureCheckError) as caught:
        check_docx_structure(b"not-a-docx")
    assert caught.value.code == "artifact_not_readable"
    with pytest.raises(DocxStructureCheckError) as empty:
        check_docx_structure(b"")
    assert empty.value.code == "artifact_empty"
