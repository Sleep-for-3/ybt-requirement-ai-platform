"""P0-D 测试夹具：合成「客户需求文档模板」与冻结内容 JSON。

这里刻意不使用 ``requirement_word.py`` 的内置版式，而是模拟一份经客户美工排版的
.docx：宋体正文、深蓝标题、A4 纸张、两个分节、页眉页脚、自定义三级编号，以及
中文 ``{{...}}`` 锚点。用它验证渲染器「原位填充、保留客户版式」的能力。
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CUSTOMER_HEADER = "客户报送需求文档（内部资料）"
CUSTOMER_FOOTER = "版本以冻结内容标识为准"

CUSTOMER_ANCHORS: dict[str, str] = {
    "cover.title": "{{文档标题}}",
    "cover.name": "{{需求名称}}",
    "cover.status": "{{文档状态}}",
    "cover.content_version": "{{内容版本}}",
    "cover.content_digest": "{{内容标识}}",
    "cover.summary": "{{文档摘要}}",
    "cover.frozen_at": "{{冻结时间}}",
    "section.scope": "{{需求范围}}",
    "section.fields_intro": "{{字段说明}}",
    "section.lineage": "{{分层路径}}",
    "section.script_basis": "{{脚本依据}}",
    "section.policy": "{{制度依据}}",
    "section.gaps": "{{待确认}}",
    "section.review": "{{审核说明}}",
    "table.fields": "{{字段表}}",
    "table.review": "{{审核表}}",
}

SECTION_TOKENS: tuple[tuple[str, str], ...] = (
    ("需求范围与业务说明", "section.scope"),
    ("字段业务说明", "section.fields_intro"),
    ("来源系统与分层路径", "section.lineage"),
    ("加工规则与脚本证据", "section.script_basis"),
    ("制度依据与差异", "section.policy"),
    ("待确认事项", "section.gaps"),
    ("审核与正式交付记录", "section.review"),
)

_CUSTOMER_NUM_ID = 900
_CUSTOMER_ABSTRACT_ID = 900


# ---------------------------------------------------------------------------
# 模板生成
# ---------------------------------------------------------------------------


def _configure_customer_styles(document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    for name, size in (("Title", 18), ("Heading 1", 16), ("Heading 2", 13)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def _configure_page(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)


def _add_customer_numbering(document) -> None:
    """在 numbering.xml 里新增一套三级编号，并挂到一级章节标题上。"""
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(_CUSTOMER_ABSTRACT_ID))
    for level, lvl_text in enumerate(("%1.", "%1.%2.", "%1.%2.%3.")):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), "decimal")
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), lvl_text)
        lvl.append(start)
        lvl.append(fmt)
        lvl.append(text)
        abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(_CUSTOMER_NUM_ID))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(_CUSTOMER_ABSTRACT_ID))
    num.append(abstract_ref)
    numbering.append(num)


def _numbered_heading(document, text: str):
    paragraph = document.add_heading(text, level=1)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = _CUSTOMER_NUM_ID
    return paragraph


def _label_token(document, label: str, token: str):
    paragraph = document.add_paragraph()
    paragraph.add_run(label + "：").bold = True
    paragraph.add_run(token)
    return paragraph


def _split_token_paragraph(document, label: str) -> None:
    """构造被 Word 拆成 3 个 run 的占位符：``{{需求`` / ``名称}`` / ``}``。"""
    paragraph = document.add_paragraph()
    paragraph.add_run(label + "：").bold = True
    for chunk in ("{{需求", "名称}", "}"):
        paragraph.add_run(chunk)


def _header_row(table, labels: tuple[str, ...]) -> None:
    for index, text in enumerate(labels):
        cell = table.rows[0].cells[index]
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True


def _sample_row(table, first_cell_token: str, *, bold: bool = True) -> None:
    """样板行：首格放锚点，各格都带字符格式，用于验证复制样板行格式。"""
    row = table.rows[1]
    for index, cell in enumerate(row.cells):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        if index == 0:
            run = paragraph.add_run(first_cell_token)
        else:
            run = paragraph.add_run("")
        run.bold = bold
        run.font.size = Pt(9)


def build_customer_template(*, split_anchor: bool = False,
                            duplicate_key: str | None = None,
                            omit_key: str | None = None,
                            section_count: int = 2) -> bytes:
    """生成一份合成客户模板；参数用于制造异常场景。"""
    document = Document()
    _configure_customer_styles(document)
    _configure_page(document.sections[0])
    _add_customer_numbering(document)

    document.sections[0].header.paragraphs[0].text = CUSTOMER_HEADER
    document.sections[0].footer.paragraphs[0].text = CUSTOMER_FOOTER

    document.add_paragraph(CUSTOMER_ANCHORS["cover.title"], "Title")
    if split_anchor:
        _split_token_paragraph(document, "需求名称")
    else:
        _label_token(document, "需求名称", CUSTOMER_ANCHORS["cover.name"])
    _label_token(document, "文档状态", CUSTOMER_ANCHORS["cover.status"])
    _label_token(document, "固定内容版本", CUSTOMER_ANCHORS["cover.content_version"])
    _label_token(document, "固定内容标识", CUSTOMER_ANCHORS["cover.content_digest"])
    _label_token(document, "冻结时间", CUSTOMER_ANCHORS["cover.frozen_at"])
    document.add_paragraph(CUSTOMER_ANCHORS["cover.summary"])

    for title, key in SECTION_TOKENS:
        _numbered_heading(document, title)
        if omit_key == key:
            continue
        document.add_paragraph(CUSTOMER_ANCHORS[key])
        if key == "section.fields_intro":
            field_table = document.add_table(rows=2, cols=9)
            _header_row(field_table, ("序号", "字段名称", "字段代码", "业务定义", "业务正文",
                                      "技术正文", "来源系统", "加工规则", "路径核验"))
            token = "" if omit_key == "table.fields" else CUSTOMER_ANCHORS["table.fields"]
            _sample_row(field_table, token)
        if key == "section.review":
            review_table = document.add_table(rows=2, cols=5)
            _header_row(review_table, ("审核步骤", "结论", "审核人", "时间", "意见"))
            token = "" if omit_key == "table.review" else CUSTOMER_ANCHORS["table.review"]
            _sample_row(review_table, token)

    if section_count > 1:
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("附录：核对清单", "Heading 2")
        document.add_paragraph("本附录由客户模板自带，渲染器不得删除或重排。")
        for extra in range(section_count - 2):
            document.add_section(WD_SECTION.NEW_PAGE)
            document.add_paragraph(f"附件 {extra + 1}：补充说明")

    if duplicate_key:
        document.add_paragraph("重复锚点：" + CUSTOMER_ANCHORS[duplicate_key])

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def customer_config(template_bytes: bytes, *, anchors: dict[str, str] | None = None,
                    required_keys: tuple[str, ...] | list[str] | None = None,
                    optional_keys: tuple[str, ...] | list[str] = (),
                    table_columns: dict[str, dict[str, int]] | None = None,
                    supported_scope: dict[str, Any] | None = None) -> dict[str, Any]:
    """构造模板适配配置 dict；默认 16 个锚点全部必需。"""
    from app.services.requirement_doc_template import DEFAULT_TABLE_COLUMNS, template_hash

    resolved_anchors = dict(CUSTOMER_ANCHORS if anchors is None else anchors)
    if required_keys is None:
        required_keys = tuple(resolved_anchors)
    return {
        "template_hash": template_hash(template_bytes),
        "config_version": "customer-p0d-test-1",
        "supported_scope": supported_scope or {
            "min_fields": 0,
            "max_fields": 1000,
            "sections": [title for title, _key in SECTION_TOKENS],
        },
        "anchors": resolved_anchors,
        "required_keys": list(required_keys),
        "optional_keys": list(optional_keys),
        "table_columns": table_columns or {
            key: dict(value) for key, value in DEFAULT_TABLE_COLUMNS.items()
        },
    }


# ---------------------------------------------------------------------------
# 冻结内容 JSON
# ---------------------------------------------------------------------------


def frozen_content(*, n_fields: int = 1, formal: bool = True,
                   long_scope_chars: int = 0, special_chars: bool = False,
                   script_basis: bool = True) -> dict[str, Any]:
    """构造与 ``requirement_scope`` 冻结快照同形的 content dict。

    ``script_basis=False`` 时省略脚本依据：既有 ``export_word`` 会据此跳过
    ``requirement_paths.path_preview``（避免测试夹具去构造完整目录元数据），
    同时也让「无脚本证据」分支可被覆盖。
    """
    background = "客户存款余额按日跑批，用于监管报送。"
    if long_scope_chars:
        background = "数据口径说明" * (long_scope_chars // 6 + 1)
        background = background[:long_scope_chars]
    if special_chars:
        background = '特殊字符 & < > " 中文引号“测试” 全角（括号）100% 反斜杠\\ 制表\t结束'

    fields: list[dict[str, Any]] = []
    for index in range(n_fields):
        number = index + 1
        fields.append({
            "field": {"id": number, "field_name": f"客户字段{number}", "field_code": f"F{number:03d}"},
            "business": {
                "business_definition": f"业务含义{number}",
                "final_content": f"业务口径{number}",
                "ai_generated_content": f"AI业务候选{number}",
            },
            "lineage": {
                "final_content": f"技术口径{number}",
                "source_system_name": "核心存款系统",
                "processing_logic": f"加工逻辑{number}",
                "ai_generated_content": f"AI技术候选{number}",
            },
        })

    content: dict[str, Any] = {
        "requirement": {
            "name": "对公客户存款余额跑批",
            "background": background,
            "objective": "支撑监管报送",
            "inclusion": "对公客户",
            "exclusion": "同业客户",
            "effective_date": "2026-09-30",
            "acceptance_criteria": "余额与核心系统一致",
            "content_version": 3,
        },
        "target_table": {"table_name": "dwd_customer_deposit", "table_code": "DM_F_001"},
        "fields": fields,
        "gaps": [{"message": "待确认：口径以核心系统为准"}, {"message": "待确认：报送时间窗口"}],
        "manual_ownership": {},
    }
    if script_basis:
        content["script_basis"] = {
            "rules": [{
                "rule_id": "R1",
                "script_version_id": "sv1",
                "source_line_start": 10,
                "source_line_end": 12,
                "statement_id": "S1",
                "source": {"database_name": "ods", "schema_name": "src", "table_name": "deposit", "column_name": "bal"},
                "target": {"database_name": "dwd", "schema_name": "dm", "table_name": "customer_deposit", "column_name": "bal"},
                "transformation_expression": "SUM(bal)",
            }],
            "versions": [{"id": "sv1", "path": "scripts/deposit.sql", "version_no": 2,
                          "file_hash": "f" * 64}],
            "metadata": [{"id": 1, "database_name": "dwd", "schema_name": "dm",
                          "table_name": "customer_deposit",
                          "assignment": {"layer_name": "明细层", "business_system_name": "核心存款系统"}}],
            "policy_snapshot": {"evidence": [{
                "unit_id": 1, "title": "客户存款报送管理办法", "document_version_id": 5,
                "regulatory_version": "2026-01", "source_category": "监管制度",
                "content": "存款余额应按日统计。",
            }]},
        }
        content["policy_comparisons"] = {"decisions": {"1": {
            "status": "matched", "rule_ids": ["R1"], "rationale": "已实现",
            "difference": "无", "confirmed_by": "审核员甲", "confirmed_at": "2026-09-16T10:00:00",
        }}}
    if formal:
        content["formal_delivery"] = {
            "version_no": 7,
            "content_version": 3,
            "approved_at": "2026-09-17T00:00:00Z",
            "review_record": [
                {"step": "业务审核", "decision": "通过", "reviewer": "业务主管",
                 "decided_at": "2026-09-16T09:00:00", "comment": "口径确认"},
                {"step": "技术审核", "decision": "通过", "reviewer": "技术主管",
                 "decided_at": "2026-09-16T15:30:00", "comment": "脚本核对"},
                {"step": "合规审核", "decision": "通过", "reviewer": "合规主管",
                 "decided_at": "2026-09-17T08:00:00", "comment": "同意交付"},
            ],
        }
    return content
