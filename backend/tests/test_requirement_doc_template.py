"""P0-D 客户 DOCX 需求文档模板渲染器定向测试。"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import hashlib

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

from app.services.requirement_doc_template import (
    DEFAULT_SECTION_TITLES,
    FrozenContentScopeError,
    TemplateAnchorError,
    TemplateConfigError,
    TemplateHashMismatch,
    analyze_template,
    build_default_template,
    default_config,
    render,
    template_hash,
)
from app.services.requirement_word import export_word, export_word_with_template
from docx_template_fixtures import (
    CUSTOMER_ANCHORS,
    CUSTOMER_FOOTER,
    CUSTOMER_HEADER,
    SECTION_TOKENS,
    build_customer_template,
    customer_config,
    frozen_content,
)


# ---------------------------------------------------------------------------
# 辅助
# ---------------------------------------------------------------------------


def _document_text(document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


def _render_customer(template_bytes: bytes, content: dict | None = None, **config_kwargs):
    config = customer_config(template_bytes, **config_kwargs)
    return render(template_bytes, content if content is not None else frozen_content(), config)


# ---------------------------------------------------------------------------
# 分析器
# ---------------------------------------------------------------------------


def test_analyze_reports_anchors_structure_header_footer_styles_and_numbering():
    template_bytes = build_customer_template()
    config = customer_config(template_bytes)

    analysis = analyze_template(template_bytes, config)

    assert analysis.ok, analysis.to_dict()
    assert set(analysis.found_anchors) == set(CUSTOMER_ANCHORS)
    assert all(len(locations) == 1 for locations in analysis.found_anchors.values())
    assert analysis.missing_keys == ()
    assert analysis.duplicate_keys == {}
    assert analysis.unsupported_keys == ()
    assert analysis.template_hash == template_hash(template_bytes)
    assert analysis.section_count == 2
    assert analysis.header_text == (CUSTOMER_HEADER,)
    assert analysis.footer_text == (CUSTOMER_FOOTER,)
    assert "Normal" in analysis.style_names and "Heading 1" in analysis.style_names
    assert "900" in analysis.numbering_abstract_ids
    assert "900" in analysis.numbering_num_ids
    assert analysis.max_numbering_levels == 3
    assert analysis.table_count == 2
    assert [table.columns for table in analysis.tables] == [9, 5]
    assert analysis.paragraph_count > 20
    assert analysis.core_modified is not None


def test_analyze_reports_missing_and_duplicate_without_touching_template():
    missing_template = build_customer_template(omit_key="section.gaps")
    missing_analysis = analyze_template(missing_template, customer_config(missing_template))
    assert not missing_analysis.ok
    assert "section.gaps" in missing_analysis.missing_keys

    duplicated_template = build_customer_template(duplicate_key="cover.name")
    duplicated_analysis = analyze_template(duplicated_template, customer_config(duplicated_template))
    assert not duplicated_analysis.ok
    assert len(duplicated_analysis.duplicate_keys["cover.name"]) == 2


# ---------------------------------------------------------------------------
# 规模：0 / 1 / 100 字段
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("n_fields", [0, 1, 100])
def test_render_field_scales_zero_one_and_hundred(n_fields):
    template_bytes = build_customer_template()
    content = frozen_content(n_fields=n_fields)

    output, diagnostics = _render_customer(template_bytes, content)
    document = Document(BytesIO(output))

    field_table = document.tables[0]
    assert len(field_table.rows) == n_fields + 1, "表头 1 行 + 每字段 1 行"
    assert diagnostics.table_rows["table.fields"] == n_fields
    assert len(document.tables[1].rows) == 4, "审核表：表头 + 3 条审核记录"

    text = _document_text(document)
    assert "{{" not in text and "}}" not in text
    assert content["requirement"]["name"] in text
    assert "已审核正式交付" in text

    rendered_names = [row.cells[1].text for row in field_table.rows[1:]]
    assert rendered_names == [f"客户字段{index}" for index in range(1, n_fields + 1)]
    if n_fields:
        assert field_table.rows[1].cells[0].text == "1"
    if n_fields == 100:
        assert field_table.rows[100].cells[1].text == "客户字段100"
        assert field_table.rows[100].cells[2].text == "F100"
        assert field_table.rows[100].cells[2].text == "F100"
        # 序号列必须是 1..100 且无重复
        assert [row.cells[0].text for row in field_table.rows[1:]] == [
            str(index) for index in range(1, 101)]


def test_render_draft_without_formal_delivery_has_empty_review_rows():
    template_bytes = build_customer_template()
    content = frozen_content(n_fields=2, formal=False)

    output, diagnostics = _render_customer(template_bytes, content)
    document = Document(BytesIO(output))

    assert len(document.tables[1].rows) == 1
    assert diagnostics.table_rows["table.review"] == 0
    text = _document_text(document)
    assert "草稿 待审核 不作为正式交付" in text
    assert "{{" not in text


# ---------------------------------------------------------------------------
# 失败路径：缺失 / 重复 / 必需映射缺失 / 哈希不一致
# ---------------------------------------------------------------------------


def test_missing_anchor_fails_with_actionable_message():
    template_bytes = build_customer_template(omit_key="section.gaps")

    with pytest.raises(TemplateAnchorError) as excinfo:
        _render_customer(template_bytes, frozen_content(n_fields=1))

    message = str(excinfo.value)
    assert "section.gaps" in message
    assert "{{待确认}}" in message
    assert "缺失必需锚点" in message


def test_optional_anchor_missing_is_reported_but_tolerated():
    template_bytes = build_customer_template(omit_key="section.gaps")
    required = tuple(key for key in CUSTOMER_ANCHORS if key != "section.gaps")

    output, diagnostics = _render_customer(
        template_bytes, frozen_content(), required_keys=required,
        optional_keys=("section.gaps",))

    assert "section.gaps" in diagnostics.skipped_keys
    assert any("可选锚点" in warning and "section.gaps" in warning
               for warning in diagnostics.warnings)
    assert "{{" not in _document_text(Document(BytesIO(output)))


def test_duplicate_anchor_fails_with_both_locations():
    template_bytes = build_customer_template(duplicate_key="cover.name")

    with pytest.raises(TemplateAnchorError) as excinfo:
        _render_customer(template_bytes)

    message = str(excinfo.value)
    assert "锚点重复" in message
    assert "cover.name" in message
    assert message.count("表格第") == 0 or message.count("->") >= 1


def test_required_key_missing_from_anchors_fails_before_any_output(tmp_path):
    template_bytes = build_customer_template()
    anchors = dict(CUSTOMER_ANCHORS)
    anchors.pop("section.policy")

    with pytest.raises(TemplateConfigError) as excinfo:
        render(template_bytes, frozen_content(), customer_config(
            template_bytes, anchors=anchors, required_keys=tuple(CUSTOMER_ANCHORS)))

    message = str(excinfo.value)
    assert "section.policy" in message
    assert "必需映射" in message
    assert list(tmp_path.iterdir()) == []


def test_missing_required_anchor_at_render_produces_no_file(tmp_path):
    template_bytes = build_customer_template(omit_key="section.gaps")
    config = customer_config(template_bytes)
    output_path = tmp_path / "rendered.docx"

    with pytest.raises(TemplateAnchorError) as excinfo:
        render(template_bytes, frozen_content(), config)

    assert "section.gaps" in str(excinfo.value)
    assert not output_path.exists()


def test_template_hash_mismatch_is_rejected():
    template_bytes = build_customer_template()
    other_template = build_customer_template(split_anchor=True)
    config = customer_config(template_bytes)

    with pytest.raises(TemplateHashMismatch) as excinfo:
        render(other_template, frozen_content(), config)

    message = str(excinfo.value)
    assert template_hash(template_bytes) in message
    assert template_hash(other_template) in message


def test_config_validation_rejects_unknown_ambiguous_and_bad_table_columns():
    template_bytes = build_customer_template()
    base = customer_config(template_bytes)

    unknown = dict(base)
    unknown_anchors = dict(base["anchors"])
    unknown_anchors["unknown.key"] = "{{未知}}"
    unknown["anchors"] = unknown_anchors
    with pytest.raises(TemplateConfigError) as excinfo:
        render(template_bytes, frozen_content(), unknown)
    assert "unknown.key" in str(excinfo.value)
    assert "cover.name" in str(excinfo.value), "错误消息要列出支持键"

    ambiguous = dict(base)
    ambiguous_anchors = dict(base["anchors"])
    ambiguous_anchors["cover.status"] = ambiguous_anchors["cover.name"]
    ambiguous["anchors"] = ambiguous_anchors
    with pytest.raises(TemplateConfigError) as excinfo:
        render(template_bytes, frozen_content(), ambiguous)
    assert "cover.name" in str(excinfo.value) and "cover.status" in str(excinfo.value)

    # 列下标超出样板行宽度是模板事实，只能在渲染期报 TemplateAnchorError。
    bad_columns = dict(base)
    bad_columns["table_columns"] = {"table.fields": {"name": 99}}
    with pytest.raises(TemplateAnchorError) as excinfo:
        render(template_bytes, frozen_content(), bad_columns)
    assert "99" in str(excinfo.value) and "table.fields" in str(excinfo.value)

    # 负下标、同列重复占用属于配置错误，应在渲染前直接拒绝。
    for invalid_columns in ({"table.fields": {"name": -1}},
                            {"table.fields": {"name": 1, "code": 1}}):
        invalid = dict(base)
        invalid["table_columns"] = invalid_columns
        with pytest.raises(TemplateConfigError):
            render(template_bytes, frozen_content(), invalid)

    bad_hash = dict(base)
    bad_hash["template_hash"] = "NOT-A-SHA256"
    with pytest.raises(TemplateConfigError) as excinfo:
        render(template_bytes, frozen_content(), bad_hash)
    assert "template_hash" in str(excinfo.value)


def test_supported_scope_rejects_too_many_fields():
    template_bytes = build_customer_template()
    config = customer_config(template_bytes, supported_scope={
        "min_fields": 0, "max_fields": 10,
        "sections": [title for title, _key in SECTION_TOKENS],
    })

    with pytest.raises(FrozenContentScopeError) as excinfo:
        render(template_bytes, frozen_content(n_fields=11), config)

    assert "max_fields=10" in str(excinfo.value)
    assert "11" in str(excinfo.value)


# ---------------------------------------------------------------------------
# 跨 run 拆分
# ---------------------------------------------------------------------------


def test_placeholder_split_across_runs_is_merged_and_replaced():
    template_bytes = build_customer_template(split_anchor=True)
    config = customer_config(template_bytes)
    analysis = analyze_template(template_bytes, config)
    assert analysis.ok

    output, _diagnostics = _render_customer(template_bytes, frozen_content())
    document = Document(BytesIO(output))
    text = _document_text(document)

    assert "对公客户存款余额跑批" in text
    assert "{{" not in text
    assert any(paragraph.text == "需求名称：对公客户存款余额跑批"
               for paragraph in document.paragraphs)


# ---------------------------------------------------------------------------
# 长文本 / 特殊字符 / 无残留
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("content_kwargs", [
    {"long_scope_chars": 2400},
    {"special_chars": True},
])
def test_long_text_and_special_characters_round_trip_without_residue(content_kwargs):
    template_bytes = build_customer_template()
    content = frozen_content(n_fields=2, **content_kwargs)
    expected = content["requirement"]["background"]
    assert len(expected) > 2000 or content_kwargs.get("special_chars")

    output, _diagnostics = _render_customer(template_bytes, content)
    document = Document(BytesIO(output))
    text = _document_text(document)

    assert expected in text
    assert "{{" not in text and "}}" not in text
    if content_kwargs.get("special_chars"):
        assert "&amp;" not in text and "&lt;" not in text, "特殊字符不能被二次转义"
        assert '"' in text and "中文引号“测试”" in text


# ---------------------------------------------------------------------------
# 原位保留：页眉/页脚/分节/样式/页面设置/编号
# ---------------------------------------------------------------------------


def test_header_footer_sections_styles_page_setup_and_numbering_are_preserved():
    template_bytes = build_customer_template()
    output, diagnostics = _render_customer(template_bytes, frozen_content(n_fields=5))
    document = Document(BytesIO(output))

    assert diagnostics.section_count == 2
    assert len(document.sections) == 2

    def _close(actual, expected):
        # Word 以 twip 存储页面尺寸，python-docx 读回时会与 Cm() 有 <1 twip 的偏差。
        return abs(actual - expected) <= Emu(700)

    for section in document.sections:
        assert _close(section.page_width, Cm(21))
        assert _close(section.page_height, Cm(29.7))
        assert _close(section.top_margin, Cm(2.5))
        assert _close(section.bottom_margin, Cm(2.5))
        assert _close(section.left_margin, Cm(2.0))
        assert _close(section.right_margin, Cm(2.0))
        assert CUSTOMER_HEADER in "\n".join(p.text for p in section.header.paragraphs)
        assert CUSTOMER_FOOTER in "\n".join(p.text for p in section.footer.paragraphs)
    assert diagnostics.header_text == (CUSTOMER_HEADER,)
    assert diagnostics.footer_text == (CUSTOMER_FOOTER,)

    for style_name, size in (("Normal", 10.5), ("Title", 18), ("Heading 1", 16)):
        style = document.styles[style_name]
        assert style.font.name == "Times New Roman", "不得把客户样式覆盖成 Arial"
        assert style.font.size == Pt(size)
    assert document.styles["Heading 1"].font.color.rgb == RGBColor(0x1F, 0x49, 0x7D)
    assert document.styles["Normal"].element.get_or_add_rPr().rFonts.get(qn("w:eastAsia")) == "宋体"

    numbering = document.part.numbering_part.element
    custom = [abstract for abstract in numbering.findall(qn("w:abstractNum"))
              if abstract.get(qn("w:abstractNumId")) == "900"]
    assert len(custom) == 1
    levels = custom[0].findall(qn("w:lvl"))
    assert len(levels) == 3, "模板自带三级编号定义必须保留"
    assert [level.get(qn("w:ilvl")) for level in levels] == ["0", "1", "2"]

    headings = [paragraph for paragraph in document.paragraphs
                if paragraph.style.name == "Heading 1"]
    assert [paragraph.text for paragraph in headings] == [title for title, _key in SECTION_TOKENS]
    for heading in headings:
        num_pr = heading._p.pPr.find(qn("w:numPr"))
        assert num_pr is not None
        assert num_pr.find(qn("w:numId")).get(qn("w:val")) == "900"


# ---------------------------------------------------------------------------
# 确定性
# ---------------------------------------------------------------------------


def test_two_renders_are_byte_identical_and_core_properties_are_frozen():
    template_bytes = build_customer_template()
    content = frozen_content(n_fields=100)
    config = customer_config(template_bytes)

    first, first_diagnostics = render(
        template_bytes, content, config, frozen_at="2026-10-01T08:00:00+08:00")
    second, second_diagnostics = render(
        template_bytes, content, config, frozen_at="2026-10-01T08:00:00+08:00")

    assert first == second
    assert hashlib.sha256(first).hexdigest() == hashlib.sha256(second).hexdigest()
    assert first_diagnostics.output_sha256 == second_diagnostics.output_sha256
    assert first_diagnostics.frozen_at == "2026-10-01 00:00:00"

    reopened = Document(BytesIO(first))
    expected = datetime(2026, 10, 1, 0, 0, 0, tzinfo=timezone.utc)
    assert reopened.core_properties.created == expected
    assert reopened.core_properties.modified == expected


def test_two_renders_are_identical_when_frozen_at_comes_from_content():
    template_bytes = build_customer_template()
    content = frozen_content(n_fields=3)
    config = customer_config(template_bytes)

    first, _ = render(template_bytes, content, config)
    second, _ = render(template_bytes, content, config)

    assert first == second
    reopened = Document(BytesIO(first))
    assert reopened.core_properties.modified == datetime(2026, 9, 17, 0, 0, 0, tzinfo=timezone.utc)


# ---------------------------------------------------------------------------
# 内置默认模板：与既有 export_word 骨架等价、可渲染
# ---------------------------------------------------------------------------


def test_default_template_is_deterministic_and_matches_legacy_skeleton():
    first = build_default_template()
    second = build_default_template()
    assert first == second

    legacy = Document(BytesIO(export_word(frozen_content(n_fields=2, script_basis=False))))
    default = Document(BytesIO(first))

    def headings(document):
        return [paragraph.text for paragraph in document.paragraphs
                if paragraph.style.name == "Heading 1"]

    assert headings(legacy) == list(DEFAULT_SECTION_TITLES)
    assert headings(default) == list(DEFAULT_SECTION_TITLES)
    assert default.styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr")) is None
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        assert default.styles[style_name].font.name == legacy.styles[style_name].font.name == "Arial"
        assert default.styles[style_name].font.size == legacy.styles[style_name].font.size
    for attribute in ("page_width", "page_height", "top_margin", "bottom_margin",
                      "left_margin", "right_margin"):
        assert getattr(default.sections[0], attribute) == getattr(legacy.sections[0], attribute)


def test_default_template_renders_frozen_content():
    template_bytes = build_default_template()
    config = default_config(template_bytes)
    assert config.template_hash == template_hash(template_bytes)

    analysis = analyze_template(template_bytes, config)
    assert analysis.ok, analysis.to_dict()
    assert analysis.missing_keys == () and analysis.duplicate_keys == {}

    output, diagnostics = render(template_bytes, frozen_content(n_fields=3), config)
    document = Document(BytesIO(output))
    text = _document_text(document)

    assert "对公客户存款余额跑批" in text
    assert "已审核正式交付" in text
    assert "固定内容版本：3" in text
    assert "{{" not in text
    assert len(document.tables[0].rows) == 4
    assert diagnostics.table_rows["table.fields"] == 3

    # 默认模板 + 默认配置的便捷入口，返回与直接 render 完全一致的字节。
    wrapped, wrapped_diagnostics = export_word_with_template(frozen_content(n_fields=3))
    assert wrapped == output
    assert wrapped_diagnostics.output_sha256 == diagnostics.output_sha256
