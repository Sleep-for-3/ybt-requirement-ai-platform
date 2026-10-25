"""P0-C: hidden XLSX parse loss must surface as a located, actionable finding.

Before this change the workbook reader called ``load_workbook(data_only=True)``
and returned ``None`` for every formula that had never been recalculated by
Excel, with no warning at all -- the requirement document was then generated
without that business definition and nobody could tell.  These tests pin the
three readers (knowledge ingestion / regulatory template / traceability
template) and the requirement generation gate to the diagnostic contract.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook

from app.models import KnowledgeDocumentVersion
from app.services.knowledge_ingestion.parsers import parse_document, parse_document_with_report
from app.services.parse_diagnostics import (
    CELL_BLANK,
    CELL_FORMULA_CACHED,
    CELL_FORMULA_UNCACHED,
    CELL_VALUE,
    classify_cell,
)
from app.services.template_parser.excel_parser import ExcelTemplateParser
from app.services.template_parser.traceability_excel_parser import TraceabilityExcelParser
from test_knowledge_rag import _client, _get, _post, _upload
from test_requirement_resources import add_document
from test_requirement_snapshot_api import snapshot_api, snapshot_db  # noqa: F401  (fixtures)
from test_traceability_excel import _workbook_bytes

FORMULA = '=IF(1=1,"只报送正常贷款余额","")'
FORMULA_ROWS = [["字段代码", "字段名称", "口径说明"], ["CUST_NO", "客户号", None]]


def _base_xlsx(rows: list[list[object]], sheet_title: str = "指标") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _patch_cell(data: bytes, coord: str, formula: str, cached: str | None) -> bytes:
    """Write a real ``<f>`` plus an optional cached ``<v>``, as Excel itself would.

    openpyxl only ever writes the formula, so a cached result has to be injected
    into the sheet XML to model a workbook that Excel already recalculated.
    """

    body = escape(formula[1:])
    if cached is None:
        replacement = f'<c r="{coord}" t="str"><f>{body}</f></c>'
    else:
        replacement = f'<c r="{coord}" t="str"><f>{body}</f><v>{escape(cached)}</v></c>'
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = payload.decode("utf-8")
                text, count = re.subn(
                    rf'<c r="{coord}"[^>]*>.*?</c>|<c r="{coord}"[^>]*/>', replacement, text, count=1, flags=re.S
                )
                if count == 0:
                    row_no = re.match(r"[A-Z]+(\d+)", coord).group(1)
                    text, count = re.subn(
                        rf'(<row r="{row_no}"[^>]*>)(.*?)(</row>)',
                        lambda match: match.group(1) + match.group(2) + replacement + match.group(3),
                        text,
                        count=1,
                        flags=re.S,
                    )
                assert count == 1, (coord, text[:400])
                payload = text.encode("utf-8")
            zout.writestr(item, payload)
    return out.getvalue()


def test_cell_classifier_judges_cache_presence_only() -> None:
    assert classify_cell(None, None) == CELL_BLANK
    assert classify_cell(None, "文本") == CELL_VALUE
    assert classify_cell("=A1", "文本") == CELL_FORMULA_CACHED
    assert classify_cell("=A1", None) == CELL_FORMULA_UNCACHED
    assert classify_cell("=A1", "") == CELL_FORMULA_UNCACHED


def test_two_tuple_parse_document_contract_is_preserved() -> None:
    units, warnings = parse_document("口径.xlsx", _base_xlsx(FORMULA_ROWS), "regulatory_qa")
    assert [unit.content for unit in units] == ["字段代码：CUST_NO；字段名称：客户号"]
    assert warnings == []


def test_uncached_formula_is_blocking_instead_of_silently_dropped() -> None:
    data = _patch_cell(_base_xlsx(FORMULA_ROWS), "C2", FORMULA, None)

    units, warnings, report = parse_document_with_report("口径.xlsx", data, "regulatory_qa")

    assert [unit.content for unit in units] == ["字段代码：CUST_NO；字段名称：客户号"]
    assert report.has_blocking
    assert report.formula_total == 1 and report.formula_uncached == 1 and report.formula_cached == 0
    assert report.formula_cells[0]["sheet"] == "指标"
    assert report.formula_cells[0]["cell"] == "C2"
    assert report.formula_cells[0]["formula"] == FORMULA
    assert any(item.code == "formula_cache_missing" and item.severity == "blocking" for item in report.findings)
    assert "C2" in warnings[0] and "定位" in warnings[0] and "建议" in warnings[0]
    assert report.to_dict()["blocking_count"] >= 1


def test_cached_formula_keeps_the_value_and_stays_clean() -> None:
    data = _patch_cell(_base_xlsx(FORMULA_ROWS), "C2", FORMULA, "只报送正常贷款余额")

    units, warnings, report = parse_document_with_report("口径.xlsx", data, "regulatory_qa")

    assert [unit.content for unit in units] == ["字段代码：CUST_NO；字段名称：客户号；口径说明：只报送正常贷款余额"]
    assert warnings == []
    assert not report.has_blocking
    assert report.formula_cached == 1 and report.formula_uncached == 0
    assert report.to_dict()["blocking_count"] == 0


def test_zero_value_and_blank_rows_are_reported_not_misread() -> None:
    rows = [["字段代码", "字段名称"], [0, "客户号"], ["", ""], ["CUST_NO", "客户号"]]

    units, warnings, report = parse_document_with_report("零值.xlsx", _base_xlsx(rows), "regulatory_qa")

    assert [unit.target_field_code for unit in units] == ["0", "CUST_NO"]
    assert report.row_skipped_blank == 1
    assert report.row_total == 3 and report.row_emitted == 2
    assert not report.has_blocking


def test_header_below_the_legacy_fifty_row_window_is_found_and_flagged() -> None:
    rows = [["说明"]] * 50 + [["字段代码", "字段名称"], ["CUST_NO", "客户号"]]

    units, warnings, report = parse_document_with_report("深表头.xlsx", _base_xlsx(rows), "regulatory_qa")

    assert [unit.content for unit in units] == ["字段代码：CUST_NO；字段名称：客户号"]
    assert report.has_blocking is False
    assert [(item.code, item.severity) for item in report.findings] == [("header_beyond_legacy_scan", "info")]
    # An info-only finding must not turn a clean parse into "parsed_with_warnings".
    assert warnings == []


def test_duplicate_and_merged_headers_stay_visible() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "指标"
    sheet.append(["字段代码", "字段名称", "字段名称", None])
    sheet.merge_cells("C1:D1")
    sheet.append(["CUST_NO", "客户号", "别名", "别名二"])
    buffer = BytesIO()
    workbook.save(buffer)

    units, warnings, report = parse_document_with_report("合并表头.xlsx", buffer.getvalue(), "regulatory_qa")

    # 重复列名的覆盖行为本身是历史语义（后出现的列覆盖先出现的列），测试固定的是
    # 「这次覆盖必须被计为 warning 并可定位」，而不是把覆盖结果当成正确解析。
    assert [unit.content for unit in units] == ["字段代码：CUST_NO；字段名称：别名二"]
    assert report.merged_header_cells == 1
    # 三列同名（含合并展开出来的那一列），每个冲突都要能被指到具体列号。
    assert [item["name"] for item in report.duplicate_headers] == ["字段名称", "字段名称"]
    assert [item["columns"] for item in report.duplicate_headers] == [[2, 3], [2, 4]]
    assert {"duplicate_header", "merged_header"} <= {item.code for item in report.findings}
    assert not report.has_blocking
    assert any("重复列名" in item for item in warnings)


def test_unrecognisable_sheet_is_reported_as_blocking_sheet_loss() -> None:
    units, warnings, report = parse_document_with_report(
        "无表头.xlsx", _base_xlsx([["甲", "乙"], ["1", "2"]]), "regulatory_qa"
    )

    assert units == []
    assert len(report.sheets_skipped) == 1 and report.sheet_parsed == 0
    assert report.has_blocking
    assert any("未识别表头" in item and "整张" in item for item in warnings)


def test_regulatory_template_keeps_legacy_warning_first_and_adds_diagnostics() -> None:
    rows = [
        ["表代码", "表名称", "字段代码", "字段名称"],
        ["YBT_A", "甲表", "CUST_NO", "客户号"],
        ["YBT_A", "甲表", "", "客户号"],
    ]

    output = ExcelTemplateParser().parse(BytesIO(_base_xlsx(rows, sheet_title="YBT_A")))

    result = output.results[0]
    # Existing contract: the legacy text stays first in the list.
    assert "缺少字段代码" in result.warnings[0]
    assert result.parse_report is not None
    assert result.parse_report["strategy_version"] == "xlsx-parse-diagnostics-v1"
    assert output.parse_report["reader"] == "regulatory-template"
    assert output.parse_report["config_version"] == "xlsx-template-config-v1"


def test_regulatory_template_missing_critical_column_is_blocking() -> None:
    rows = [["表代码", "表名称", "备注"], ["YBT_B", "乙表", "说明"]]

    output = ExcelTemplateParser().parse(BytesIO(_base_xlsx(rows, sheet_title="YBT_B")))

    report = output.parse_report
    assert report["blocking_count"] >= 1
    assert any(item["code"] == "critical_column_missing" for item in report["findings"])


def test_regulatory_template_uncached_formula_is_blocking() -> None:
    rows = [["表代码", "表名称", "字段代码", "字段名称", "字段定义"], ["YBT_A", "甲表", "CUST_NO", "客户号", None]]
    data = _patch_cell(_base_xlsx(rows, sheet_title="YBT_A"), "E2", FORMULA, None)

    output = ExcelTemplateParser().parse(BytesIO(data))

    report = output.parse_report
    assert report["formula_uncached"] == 1
    assert report["blocking_count"] >= 1
    assert any(item["code"] == "formula_cache_missing" for item in report["findings"])


def test_traceability_template_reports_merged_headers_without_blocking() -> None:
    output = TraceabilityExcelParser().parse(BytesIO(_workbook_bytes()))

    report = output.parse_report
    assert report["reader"] == "traceability-template"
    assert report["sheet_parsed"] == 1
    assert report["merged_header_cells"] > 0
    assert report["blocking_count"] == 0
    assert any(item["code"] == "merged_header" for item in report["findings"])


def test_traceability_template_uncached_formula_is_blocking() -> None:
    workbook = load_workbook(BytesIO(_workbook_bytes()))
    workbook["业务口径"]["D4"] = FORMULA
    buffer = BytesIO()
    workbook.save(buffer)

    output = TraceabilityExcelParser().parse(BytesIO(buffer.getvalue()))

    report = output.parse_report
    assert report["formula_uncached"] >= 1
    assert report["blocking_count"] >= 1


def test_response_models_expose_the_parse_report() -> None:
    from app.schemas.api import (
        TemplateUploadResponse,
        TemplateVersionRead,
        TraceabilityTemplateUploadResponse,
    )

    assert "parse_report" in TemplateUploadResponse.model_fields
    assert "parse_report_json" in TemplateVersionRead.model_fields
    assert "parse_report" in TraceabilityTemplateUploadResponse.model_fields


def test_knowledge_upload_exposes_parse_block_and_its_locations(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("STORAGE_DIR", str(tmp_path))
    data = _patch_cell(_base_xlsx(FORMULA_ROWS), "C2", FORMULA, None)

    with _client() as client:
        project = _post(client, "/api/projects", {"name": "解析缺失项目", "bank_name": "甲银行"})
        document = _upload(client, project["id"], "口径.xlsx", data, "regulatory_qa")
        versions = _get(client, f"/api/knowledge/documents/{document['id']}/versions?project_id={project['id']}")

    version = versions[0]
    summary = version["parse_summary_json"]
    assert summary["parse_blocked"] is True
    assert summary["parse_report"]["blocking_count"] >= 1
    assert summary["parse_report"]["formula_cells"][0]["cell"] == "C2"
    assert any("[阻断]" in item for item in version["warnings_json"])


def test_generation_input_refuses_documents_with_blocking_parse_loss(snapshot_api) -> None:
    client, db, project, field, req, _membership = snapshot_api
    document = add_document(db, project.id, "含未缓存公式的口径资料")
    version = KnowledgeDocumentVersion(
        project_id=project.id,
        document_id=document.id,
        version_no=1,
        file_name=document.file_name,
        storage_path="isolated-not-read",
        file_hash="synthetic",
        parse_summary_json={"parse_blocked": True, "parse_report": {"blocking_count": 1}},
    )
    db.add(version)
    db.flush()
    document.current_version_id = version.id
    db.commit()

    base = f"/projects/{project.id}/requirements/{req.id}"
    updated = client.put(
        base, json={**req.scope_json, "name": req.name, "expected_version": 1, "document_ids": [document.id]}
    )
    assert updated.status_code == 200, updated.text
    assert client.post(base + "/revisions", json={"expected_version": 2}).status_code in {200, 201}
    payload = {
        "expected_content_version": 1,
        "field_ids": [field.id],
        "sections": ["business"],
        "idempotency_key": "blocked-parse",
    }
    blocked = client.post(base + "/generation-inputs", json=payload)
    assert blocked.status_code == 409, blocked.text
    assert "选定资料存在未处理的解析缺失" in blocked.json()["detail"]
    assert "口径资料" in blocked.json()["detail"]

    version.parse_summary_json = {"parse_blocked": False, "parse_report": {"blocking_count": 0}}
    db.commit()
    cleared = client.post(base + "/generation-inputs", json={**payload, "idempotency_key": "clear-parse"})
    assert cleared.status_code == 201, cleared.text
