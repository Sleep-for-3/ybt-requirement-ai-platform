from copy import deepcopy
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from app.services.requirement_scope import content_digest, export_document
from app.services.requirement_word import export_word
from test_requirement_policy_comparison import setup_comparison, decision
from test_requirement_snapshot_api import snapshot_api, snapshot_db


def paragraphs(data):
    return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)


def test_word_excel_export_same_frozen_revision_after_live_changes(snapshot_api):
    client, db, project, field, req, _ = snapshot_api
    base, unit, _, view, _ = setup_comparison(snapshot_api)
    response = client.post(base + "/policy-comparison", json=decision(view, unit.id))
    assert response.status_code == 201, response.text
    current = client.get(base + "/document").json()
    result = client.post(base + "/snapshots", json={"expected_version": req.version,
        "expected_hash": current["content_hash"]})
    assert result.status_code == 201, result.text
    snapshot_id = result.json()["id"]
    field.field_name = "LIVE_NAME_MUST_NOT_ENTER_EXPORT"
    unit.content = "LIVE_POLICY_MUST_NOT_ENTER_EXPORT"; db.flush()
    xlsx = client.get(base + f"/snapshots/{snapshot_id}/export")
    word = client.get(base + f"/snapshots/{snapshot_id}/export?format=docx")
    assert xlsx.status_code == word.status_code == 200
    assert xlsx.headers["x-requirement-snapshot-hash"] == word.headers["x-requirement-snapshot-hash"]
    book = load_workbook(BytesIO(xlsx.content))
    content_hash = dict(book["需求范围"].values)["固定内容标识"]
    text = paragraphs(word.content)
    assert content_hash in text and "固定内容版本：3" in text
    assert "LIVE_NAME" not in text and "LIVE_POLICY" not in text
    assert "脚本将余额乘以二" in text and "存在冲突" in text
    assert book["制度逐条对照"]["D2"].value == "存在冲突"
    assert "草稿 待审核" in text
    assert word.headers["content-disposition"].endswith('.docx"')


def test_word_export_permissions_and_cross_project_isolation(snapshot_api):
    client, db, project, field, req, membership = snapshot_api
    base = f"/projects/{project.id}/requirements/{req.id}"
    current = client.get(base + "/document").json()
    snapshot = client.post(base + "/snapshots", json={"expected_version": req.version,
        "expected_hash": current["content_hash"]}).json()
    membership.project_role = "viewer"; db.flush()
    assert client.get(base + f"/snapshots/{snapshot['id']}/export?format=docx").status_code == 403
    assert client.get(f"/projects/99999/requirements/{req.id}/snapshots/{snapshot['id']}/export?format=docx").status_code == 404


def test_formal_label_and_review_record_come_only_from_fixed_snapshot(snapshot_api):
    client, _, project, _, req, _ = snapshot_api
    base = f"/projects/{project.id}/requirements/{req.id}"
    content = client.get(base + "/document").json()
    content["formal_delivery"] = {"version_no": 7, "content_version": 3,
        "approved_at": "2026-09-17T00:00:00Z", "review_record": [{"step": "final_review",
            "decision": "approved", "reviewer": "隔离审核人", "decided_at": "2026-09-17", "comment": "合成审核记录"}]}
    before = deepcopy(content)
    data = export_word(content)
    text = paragraphs(data)
    document = Document(BytesIO(data))
    assert "已审核正式交付" in text and "隔离审核人" in text
    assert "固定内容版本：3" in text and content_digest(content) in text
    assert document.styles["Title"].element.get_or_add_pPr().find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr") is None
    assert content == before
