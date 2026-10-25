"""把已冻结的导出产物（xlsx 网格 / docx 章节）解析成可在页面上查看的内容（只读、不重新渲染）。

为什么需要它：需求文档的验收动作是"看这张表填得对不对"，而在此之前平台只能
下载 .xlsx 再看——评审者无法确认"平台导出的到底是不是监管字段标准模板的版式"，
也无法在同一个界面里对照内容与产物身份（文件 sha256 / 渲染器版本）。

因此这里只做两件事：读取**已落库的产物字节**（不重新渲染，避免"预览与下载不一致"），
解析成行列网格与合并区信息。任何截断都显式标记，绝不静默省略。
"""

from dataclasses import dataclass
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

PREVIEW_VERSION = "artifact-preview-v1"
# 页面渲染上限：超过部分显式标记 truncated，不静默丢弃。
MAX_ROWS = 300
MAX_COLUMNS = 60
# 单元格展示上限：真实脚本的加工逻辑可能很长，超出后截断并计数。
MAX_CELL_CHARS = 500
SUPPORTED_FORMATS = ("xlsx", "docx")
DOCX_PREVIEW_VERSION = "artifact-docx-preview-v1"
# Word 预览上限：正文段落按块截断、表格只给表头与前若干行，都显式标记，不静默省略。
MAX_DOCX_BLOCKS = 600
MAX_DOCX_CHARS = 1200
MAX_DOCX_TABLE_ROWS = 40


class ArtifactPreviewError(RuntimeError):
    """产物无法预览（格式不支持 / 文件损坏）。"""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class Cell:
    value: str | None
    formula: str | None = None


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text


def _load(data: bytes):
    try:
        values = load_workbook(BytesIO(data), data_only=True)
        formulas = load_workbook(BytesIO(data), data_only=False)
    except Exception as exc:  # noqa: BLE001 - 任何解析异常都归为"产物不可预览"
        raise ArtifactPreviewError("artifact_not_readable", f"产物无法解析为 xlsx：{type(exc).__name__}") from exc
    return values, formulas


def preview_xlsx(
    data: bytes,
    *,
    max_rows: int = MAX_ROWS,
    max_columns: int = MAX_COLUMNS,
    max_cell_chars: int = MAX_CELL_CHARS,
) -> dict:
    """解析 xlsx 产物为网格。

    取值优先级：公式的缓存结果（``data_only=True``）→ 公式原文。真实链路里
    openpyxl 写出的公式没有缓存值，此时显示公式原文并计入 ``formulas_without_cache``，
    让评审者知道"这是公式不是结果"，而不是看到空白单元格误以为漏填。
    """

    values, formulas = _load(data)
    sheets: list[dict] = []
    for name in values.sheetnames:
        value_sheet = values[name]
        formula_sheet = formulas[name]
        total_rows = value_sheet.max_row or 0
        total_columns = value_sheet.max_column or 0
        rows: list[list[str | None]] = []
        truncated_cells = 0
        formulas_without_cache = 0
        for row_index in range(1, min(total_rows, max_rows) + 1):
            cells: list[str | None] = []
            for column_index in range(1, min(total_columns, max_columns) + 1):
                cached = value_sheet.cell(row=row_index, column=column_index).value
                formula_value = formula_sheet.cell(row=row_index, column=column_index).value
                text = _text(cached)
                if text is None and isinstance(formula_value, str) and formula_value.startswith("="):
                    text = f"[公式] {formula_value}"
                    formulas_without_cache += 1
                if text is not None and len(text) > max_cell_chars:
                    text = text[:max_cell_chars] + f"…（已截断，原文 {len(text)} 字符）"
                    truncated_cells += 1
                cells.append(text)
            rows.append(cells)
        sheets.append({
            "name": name,
            "rows": rows,
            "total_rows": total_rows,
            "total_columns": total_columns,
            "merged": [str(rng) for rng in value_sheet.merged_cells.ranges],
            "freeze_panes": value_sheet.freeze_panes,
            "truncated": total_rows > max_rows or total_columns > max_columns,
            "truncated_cells": truncated_cells,
            "formulas_without_cache": formulas_without_cache,
        })
    if not sheets:
        raise ArtifactPreviewError("artifact_empty", "产物没有可预览的工作表")
    return {
        "preview_version": PREVIEW_VERSION,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "max_rows": max_rows,
        "max_columns": max_columns,
        "max_cell_chars": max_cell_chars,
        "note": "预览读取的是已落库的产物字节（不重新渲染）；截断与未缓存公式均有显式标记。",
    }


def _docx_paragraph(paragraph) -> dict:
    text = paragraph.text or ""
    original = len(text)
    style = paragraph.style.name if paragraph.style is not None else ""
    if original > MAX_DOCX_CHARS:
        text = text[:MAX_DOCX_CHARS] + f"…（已截断，原文 {original} 字符）"
    return {"style": style, "text": text, "chars": original, "truncated": original > MAX_DOCX_CHARS}


def _docx_table(table, *, max_rows: int = MAX_DOCX_TABLE_ROWS) -> dict:
    rows = list(table.rows)
    total_rows = len(rows)
    shown = rows[:max_rows]
    grid: list[list[str]] = []
    for row in shown:
        grid.append([(cell.text or "").strip() for cell in row.cells])
    columns = max((len(item) for item in grid), default=0)
    header = grid[0] if grid else []
    return {
        "rows": total_rows,
        "columns": columns,
        "header": header,
        "sample_rows": grid[1:] if len(grid) > 1 else [],
        "truncated": total_rows > max_rows,
    }


def preview_docx(
    data: bytes,
    *,
    max_blocks: int = MAX_DOCX_BLOCKS,
    max_chars: int = MAX_DOCX_CHARS,
    max_table_rows: int = MAX_DOCX_TABLE_ROWS,
) -> dict:
    """把 docx 产物解析成"标题 + 段落 + 表格摘要"的章节视图。

    为什么不是整段文本：内置版式 Word 的验收点是"章节在不在、每节写了什么、
    字段块有没有丢"，整段拼接会把这些结构抹平；这里按 Heading 切章节，
    段落保留样式名（评审能看到"这是 Heading 1 还是正文"），表格只给表头与前若干行。
    """

    try:
        document = Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - 任何解析异常都归为"产物不可预览"
        raise ArtifactPreviewError("artifact_not_readable", f"产物无法解析为 docx：{type(exc).__name__}") from exc

    body = document.element.body
    # 先建一次映射：按 body 顺序遍历时用 id(child) 取对象，避免每个块都线性扫一遍 paragraphs（大文档会退化成 O(n²)）。
    paragraph_map = {id(paragraph._p): paragraph for paragraph in document.paragraphs}  # noqa: SLF001
    table_map = {id(table._tbl): table for table in document.tables}  # noqa: SLF001
    sections: list[dict] = []
    current: dict | None = None
    block_count = 0
    truncated_blocks = 0

    def ensure_section(title: str, style: str) -> dict:
        section = {"title": title, "style": style, "paragraphs": [], "tables": [], "chars": 0}
        sections.append(section)
        return section

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            paragraph = paragraph_map.get(id(child))
            if paragraph is None:
                continue
            item = _docx_paragraph(paragraph)
            block_count += 1
            if block_count > max_blocks:
                truncated_blocks += 1
                continue
            style = item["style"]
            if style.startswith("Heading") or style == "Title":
                current = ensure_section(item["text"] or style, style)
                current["chars"] += item["chars"]
                continue
            if current is None:
                current = ensure_section("（未分节内容）", "")
            current["paragraphs"].append(item)
            current["chars"] += item["chars"]
        elif tag == "tbl":
            table = table_map.get(id(child))
            if table is None:
                continue
            if current is None:
                current = ensure_section("（未分节内容）", "")
            current["tables"].append(_docx_table(table, max_rows=max_table_rows))

    if not sections:
        raise ArtifactPreviewError("artifact_empty", "产物没有可预览的正文内容")

    return {
        "preview_version": DOCX_PREVIEW_VERSION,
        "format": "docx",
        "section_count": len(sections),
        "section_titles": [section["title"] for section in sections],
        "sections": sections,
        "max_blocks": max_blocks,
        "max_chars": max_chars,
        "max_table_rows": max_table_rows,
        "truncated_blocks": truncated_blocks,
        "note": "预览读取的是已落库的产物字节（不重新渲染）；段落与表格的截断均有显式标记。",
    }


def preview_xlsx_document(data: bytes, **kwargs) -> dict:
    result = preview_xlsx(data, **kwargs)
    result["format"] = "xlsx"
    return result


def preview_artifact(data: bytes, *, file_format: str, **kwargs) -> dict:
    if file_format == "xlsx":
        return preview_xlsx_document(data, **kwargs)
    if file_format == "docx":
        return preview_docx(data)
    raise ArtifactPreviewError(
        "artifact_format_unsupported",
        f"只有 xlsx 与 docx 产物支持页面预览（当前 {file_format}）。",
    )
