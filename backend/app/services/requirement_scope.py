"""Project-scoped requirements, deterministic gaps, and frozen exports."""
from datetime import date
from io import BytesIO
import hashlib
import json

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import select

from app.models import (
    Requirement, TargetField, TargetTable, ProductScenario, KnowledgeDocument,
    SourceTable, MartTable, ScenarioBusinessMapping, ScenarioTechnicalLineage,
)
from app.services.requirement_workspace_projection import RequirementWorkspaceProjectionService


class ScopeInput(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    target_table_id: int
    field_ids: list[int] = Field(min_length=1, max_length=500)
    scenario_id: int
    objective: str = Field(default="", max_length=10000)
    background: str = Field(default="", max_length=20000)
    effective_date: date | None = None
    inclusion: str = Field(default="", max_length=10000)
    exclusion: str = Field(default="", max_length=10000)
    document_ids: list[int] = Field(default_factory=list, max_length=200)
    source_table_ids: list[int] = Field(default_factory=list, max_length=200)
    mart_table_ids: list[int] = Field(default_factory=list, max_length=200)
    expected_version: int | None = None
    expected_content_version: int | None = None


def validate_scope(db, project_id, payload):
    for model, ids in (
        (TargetTable, [payload.target_table_id]),
        (ProductScenario, [payload.scenario_id]),
        (TargetField, payload.field_ids),
        (KnowledgeDocument, payload.document_ids),
        (SourceTable, payload.source_table_ids),
        (MartTable, payload.mart_table_ids),
    ):
        query = select(model.id).where(model.project_id == project_id, model.id.in_(ids))
        if model is TargetField:
            query = query.where(TargetField.target_table_id == payload.target_table_id)
        if model is KnowledgeDocument:
            query = query.where(KnowledgeDocument.document_status != "archived")
        if set(db.scalars(query)) != set(ids):
            raise HTTPException(422, "所选范围包含不属于当前项目或目标表的对象")
    if not payload.name.strip():
        raise HTTPException(422, "请输入需求名称")


def requirement_dict(row):
    return {"id": row.id, "project_id": row.project_id, "version": row.version,
            "content_version": row.content_version or 0,
            **row.scope_json, "name": row.name}


def load_requirement(db, project_id, requirement_id):
    row = db.scalar(select(Requirement).where(
        Requirement.id == requirement_id, Requirement.project_id == project_id))
    if not row:
        raise HTTPException(404, "需求不存在或不可见")
    return row


def document_content(db, row):
    if row.content_version:
        from app.services.requirement_revisions import revision_document, load_revision
        return revision_document(load_revision(db, row.project_id, row.id, row.content_version))
    return live_document_content(db, row)


def live_document_content(db, row):
    scope = row.scope_json
    service = RequirementWorkspaceProjectionService(db)
    fields, gaps = [], []
    resources = []
    for key, model, label, name, code in (
        ("document_ids", KnowledgeDocument, "知识资料", KnowledgeDocument.file_name, KnowledgeDocument.file_name),
        ("source_table_ids", SourceTable, "来源表", SourceTable.table_name, SourceTable.table_code),
        ("mart_table_ids", MartTable, "集市表", MartTable.table_name, MartTable.table_code),
    ):
        ids = list(dict.fromkeys(scope.get(key, [])))
        query = select(model.id, name.label("name"), code.label("code")).where(
            model.project_id == row.project_id, model.id.in_(ids))
        if model is KnowledgeDocument:
            query = query.where(KnowledgeDocument.document_status != "archived")
        visible = {item.id: item for item in db.execute(query)}
        for resource_id in ids:
            item = visible.get(resource_id)
            resources.append({"kind": label, "id": resource_id,
                "name": item.name if item else "已失效或不可用的选择",
                "code": item.code if item else "", "status": "selected" if item else "unavailable"})
            if item is None:
                gaps.append({"id": f"scope:{key}:{resource_id}", "field_id": None, "origin": "analysis",
                    "status": "open", "message": f"所选{label}已失效或不可用，请重新选择并核对依据"})
    projection = service.projection(row.project_id, scope["target_table_id"], scope["scenario_id"])
    for field_id in dict.fromkeys(scope["field_ids"]):
        record = service.field_detail(row.project_id, field_id, scope["scenario_id"])
        evidence = service.field_evidence(row.project_id, field_id, scope["scenario_id"])
        from app.services.requirement_gaps import field_gaps
        gaps.extend(field_gaps({**record, "evidence": evidence}))
        fields.append({**record, "evidence": evidence})
    for question in projection.get("question_summaries", []):
        if question.get("target_field_id") in scope["field_ids"] or (
            not question.get("target_field_id") and question.get("target_table_id") == scope["target_table_id"]
        ):
            resolved = question.get("question_status") in {"resolved", "closed", "rejected"}
            if not resolved or not (question.get("resolution_text") or "").strip():
                gaps.append({"id": f"manual:{question['id']}", "field_id": question.get("target_field_id"),
                             "origin": "manual", "status": question.get("question_status"),
                             "message": ("问题已标记处理，但缺少解决依据：" if resolved else "") + (question.get("question_text") or question.get("question_title") or "人工待确认问题")})
    approved = bool(fields) and all(
        (f.get("business") or {}).get("business_confirm_status") in {"confirmed", "approved"}
        and (f.get("lineage") or {}).get("tech_confirm_status") in {"confirmed", "approved"}
        and bool(f.get("mart_mappings"))
        and all(m.get("mapping_status") in {"confirmed", "approved"} for m in f["mart_mappings"])
        and all(
            bool(f.get("source_mappings", {}).get(str(m.get("mart_field_id"))))
            and all(s.get("mapping_status") in {"confirmed", "approved"}
                    for s in f["source_mappings"][str(m.get("mart_field_id"))])
            for m in f["mart_mappings"]
        )
        for f in fields
    )
    return {"requirement": requirement_dict(row), "fields": fields, "gaps": gaps, "resources": resources,
            "target_table": next((t for t in projection.get("tables", []) if t["id"] == scope["target_table_id"]), None),
            "scenario": next((s for s in projection.get("scenarios", []) if s["id"] == scope["scenario_id"]), None),
            "mart_tables": projection.get("mart_tables", []), "mart_fields": projection.get("mart_fields", []),
            "assessment": "unassessed" if not fields else "gaps" if gaps else "clear",
            "status": "deliverable" if not gaps and approved else "draft"}


def content_digest(content):
    return hashlib.sha256(json.dumps(content, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def export_document(content):
    workbook = Workbook()
    overview = workbook.active
    overview.title = "需求范围"
    req = content["requirement"]
    formal = content.get("formal_delivery")
    overview.append(["需求名称", req["name"]])
    overview.append(["交付状态", "正式交付（已审核）" if formal else "草稿（待审核交付）"])
    overview.append(["内容版本", formal.get("content_version") if formal else req.get("content_version") or "尚未建立独立内容"])
    if formal:
        overview.append(["正式交付版本", formal["version_no"]])
        overview.append(["审核通过时间", formal["approved_at"]])
    for key, label in (("version", "需求版本"), ("objective", "业务目标"), ("background", "业务背景"),
                       ("effective_date", "口径生效日期"), ("inclusion", "纳入范围"), ("exclusion", "排除条件")):
        overview.append([label, req.get(key) or "待确认"])
    overview.append(["固定内容标识", content_digest(content)])
    resource_sheet = workbook.create_sheet("资料与数据范围")
    resource_sheet.append(["类别", "资料或业务表名", "技术名称", "选择状态"])
    for resource in content.get("resources", []):
        resource_sheet.append([resource["kind"], resource["name"], resource["code"],
                               "已选（不代表已核验）" if resource["status"] == "selected" else "不可用，待重新选择"])
    if not content.get("resources"):
        resource_sheet.append(["待确认", "未指定资料范围，不代表允许使用全部资料", "", "待指定"])
    if formal:
        review_sheet = workbook.create_sheet("审核与交付")
        review_sheet.append(["审核步骤", "决定", "审核人", "时间", "意见"])
        for record in formal.get("review_record", []):
            review_sheet.append([
                record.get("step"),
                record.get("decision"),
                record.get("reviewer"),
                record.get("decided_at"),
                record.get("comment") or "未填写意见",
            ])
    sheet = workbook.create_sheet("字段口径")
    sheet.append(["字段", "名称", "业务定义", "业务口径", "来源系统", "来源表", "来源字段",
                  "加工规则", "技术口径", "证据"])
    for record in content["fields"]:
        f, b, t = record["field"], record.get("business") or {}, record.get("lineage") or {}
        sheet.append([f["field_code"], f["field_name"], b.get("business_definition"),
                      b.get("final_content") or b.get("ai_generated_content") or "待确认",
                      t.get("source_system_name"), t.get("source_table_english_name"), t.get("source_field_english_name"),
                      t.get("processing_logic"), t.get("final_content") or t.get("ai_generated_content") or "待确认",
                      "\n".join(str(e.get("source_name") or e.get("quoted_content") or "已绑定依据") for e in record["evidence"])])
    rules = workbook.create_sheet("双层映射规则")
    rules.append(["目标字段", "加工层级", "映射名称", "来源表", "来源字段", "业务规则",
                  "关联条件", "过滤条件", "码值映射", "空值处理", "优先级", "合并规则",
                  "异常处理", "质量与校验", "报送条件", "待确认事项"])
    evidence_sheet = workbook.create_sheet("证据出处")
    evidence_sheet.append(["目标字段", "出处", "位置", "原文", "证据说明"])
    for record in content["fields"]:
        code = record["field"]["field_code"]
        mappings = [("集市 → 监管", m) for m in record.get("mart_mappings", [])]
        seen = set()
        for rows in record.get("source_mappings", {}).values():
            for mapping in rows:
                if mapping["id"] not in seen:
                    seen.add(mapping["id"])
                    mappings.append(("来源 → 集市", mapping))
        for layer, mapping in mappings:
            rules.append([code, layer, mapping.get("mapping_name"),
                mapping.get("source_tables_summary") or mapping.get("mart_table_summary"),
                mapping.get("source_fields_summary") or mapping.get("mart_field_summary"),
                mapping.get("business_rule"), mapping.get("join_condition"), mapping.get("filter_condition"),
                mapping.get("code_mapping_rule"), mapping.get("null_handling_rule"),
                mapping.get("priority_rule"), mapping.get("merge_rule"), mapping.get("exception_rule"),
                mapping.get("quality_check_rule") or mapping.get("validation_rule"),
                mapping.get("reporting_condition"), mapping.get("open_questions")])
        for evidence in record.get("evidence", []):
            evidence_sheet.append([code, evidence.get("source_name"), evidence.get("location_text"),
                                   evidence.get("quoted_content"), evidence.get("evidence_summary")])
    basis = content.get("script_basis")
    if basis:
        path_sheet = workbook.create_sheet("已确认字段路径")
        path_sheet.append(["目标字段", "确认状态", "规则编号", "确认人", "确认时间", "核验依据"])
        from app.services.requirement_paths import path_issues
        blocked = {i["field_id"] for i in path_issues(content)}
        for record in content["fields"]:
            path = record.get("confirmed_path") or {}
            path_sheet.append([record["field"]["field_code"], "待核验" if record["field"]["id"] in blocked else "已确认",
                "、".join(path.get("rule_ids", [])), path.get("confirmed_by"), path.get("confirmed_at"), path.get("rationale")])
        script_sheet = workbook.create_sheet("固定脚本事实")
        script_sheet.append(["规则编号", "脚本版本", "语句编号", "起始行", "结束行", "来源", "目标",
            "加工表达式", "关联条件", "过滤条件", "聚合", "码值转换"])
        for rule in basis.get("rules", []):
            labels = [".".join(str(rule[side].get(key) or "") for key in
                ("database_name", "schema_name", "table_name", "column_name")) for side in ("source", "target")]
            script_sheet.append([rule["rule_id"], rule["script_version_id"], rule["statement_id"],
                rule["source_line_start"], rule["source_line_end"], *labels, rule["transformation_expression"],
                rule["join_condition"], rule["filter_condition"], rule["aggregation_rule"], rule["code_mapping_rule"]])
        policy_sheet = workbook.create_sheet("固定制度依据")
        policy_sheet.append(["条款编号", "资料版本", "监管版本", "条款", "来源类别", "对照状态"])
        for unit in basis.get("policy_snapshot", {}).get("evidence", []):
            policy_sheet.append([unit["unit_id"], unit["document_version_id"], unit.get("regulatory_version"),
                unit["content"], unit.get("source_category"), "核验结果见制度逐条对照，不代表已经正式审核"])
        if policy_sheet.max_row == 1:
            policy_sheet.append([None, None, None, "缺少依据", None, "脚本现状不是制度要求"])
        comparisons = workbook.create_sheet("制度逐条对照")
        comparisons.append(["条款编号", "资料版本", "关联规则", "人工判断", "核验理由", "差异", "确认人", "确认时间"])
        decisions = content.get("policy_comparisons", {}).get("decisions", {})
        labels = {"matched": "匹配", "conflict": "存在冲突", "missing_implementation": "缺少实现", "pending": "待确认"}
        for unit in basis.get("policy_snapshot", {}).get("evidence", []):
            row = decisions.get(str(unit["unit_id"])) or {}
            comparisons.append([unit["unit_id"], unit["document_version_id"], "、".join(row.get("rule_ids", [])),
                labels.get(row.get("status"), "未确认"), row.get("rationale"), row.get("difference"),
                row.get("confirmed_by"), row.get("confirmed_at")])
        manifest = workbook.create_sheet("脚本版本清单")
        manifest.append(["脚本", "版本编号", "版本号", "文件哈希", "解析状态"])
        for version in basis["versions"]:
            manifest.append([version["path"], version["id"], version["version_no"], version["file_hash"], version["parse_status"]])
    gap_sheet = workbook.create_sheet("待确认事项")
    gap_sheet.append(["字段", "发现来源", "事项", "状态"])
    field_labels = {record["field"].get("id"): f'{record["field"]["field_name"]} ({record["field"]["field_code"]})'
                    for record in content["fields"] if record["field"].get("id") is not None}
    for gap in content["gaps"]:
        gap_sheet.append([field_labels.get(gap["field_id"], "需求范围"), "分析" if gap["origin"] == "analysis" else "人工", gap["message"], "待确认"])
    for ws in workbook:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 34
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245B51")
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
