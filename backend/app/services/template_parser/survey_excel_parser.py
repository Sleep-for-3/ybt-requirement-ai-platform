"""一表通监管报送口径需求调研表解析器。

真实客户模板（``一表通报送口径需求调研表-一表通项目.xlsx``）与研发早期约定的
"单行表头 + 表编号/表名称/字段代码/字段名称" 版式不同：

* 前 11 个 sheet 是说明性内容（业务口径填写示例、注意事项、版本信息、目录……）；
* 每个报送表一个 sheet，命名 ``T_<n>_<m>-<中文表名>``；
* 第 1 行是 "返回" 链接；
* 第 2 行是合并的分区表头（数据表信息 / 业务口径 / 技术口径 / 监管集市信息 …）；
* **第 3 行才是列名**（数据项编码、数据项名称、数据项监管要求、集市表英文名 …）；
* 第 4 行是填写说明（稀疏）；
* 数据自第 5 行起，一行一个数据项。

本模块只做确定性解析：不推断缺失值、不执行公式、不猜测列语义。识别不到列名就
把整个 sheet 记为跳过并留下可定位的告警，而不是静默产出空表。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.parse_diagnostics import (
    SEVERITY_BLOCKING,
    SEVERITY_INFO,
    SEVERITY_WARNING,
    ParseFinding,
    ParseReport,
    cell_reference,
)

SURVEY_CONFIG_VERSION = "xlsx-survey-config-v1"
LEGACY_HEADER_SCAN_ROWS = 20
#: 报送表 sheet 名，例如 ``T_8_1-贷款借据``、``T_10_2-存款状态``。
TABLE_SHEET_PATTERN = re.compile(r"^T_\d+_\d+(?:[-_]|$)")
#: 数据项编码形如 ``H010001``；真实模板里个别表用 ``A040001`` 这类监管码。
FIELD_CODE_PATTERN = re.compile(r"^[A-Z]{1,3}\d{4,}$")
#: 数据项编码至少应当是 ASCII 标识符。真实模板里数据区中间还夹着"辅助字段""数据项编码"
#: 这类分段标记行（表编号/表名称列为空），它们不是数据项；含中文的编码一律当标记跳过，
#: 否则会把"辅助字段"建成一个监管字段并在导出文档里出现空行。
CODE_IDENTIFIER_PATTERN = re.compile(r"^[A-Za-z][A-Za-z0-9_]{1,63}$")
HEADER_MARKERS = ("数据项编码", "数据项名称")
#: 调研表里独立的一张“表级口径”sheet：91 行，一行一张监管表，
#: 带 主题 / 填报说明 / 报送数据域编码 / 报送数据域名称。
#: 导出文档第 3-8 行的“数据主题名称 / 采集范围 / 一表通报送范围确认”应当引用它，
#: 而不是一律写“待业务补充”。
TABLE_LEVEL_SHEET = "表级口径"
#: 表级口径 sheet 的表头只在前 10 行内找（真实表头在第 1 行）
TABLE_LEVEL_HEADER_SCAN_ROWS = 10
TABLE_META_ALIASES: dict[str, tuple[str, ...]] = {
    "theme": ("主题", "数据主题", "数据主题名称"),
    "table_code": ("表编号", "表代码"),
    "table_name": ("数据表名称", "表名称"),
    "filing_note": ("填报说明", "采集范围", "填报范围说明"),
    "reporting_domain_code": ("报送数据域编码",),
    "reporting_domain_name": ("报送数据域名称",),
    "internal_caliber": ("行内业务口径描述",),
}

#: 规范键 -> 中文列名别名。多个同名列按出现次序加 ``#2``/``#3`` 后缀。
SURVEY_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "table_code": ("表编号",),
    "table_name": ("表名称",),
    "field_code": ("数据项编码",),
    "field_name": ("数据项名称",),
    "package_version": ("制度包版本",),
    "data_category": ("数据类别",),
    "data_format": ("数据格式",),
    "version_note": ("版本说明",),
    "regulatory_requirement": ("数据项监管要求",),
    "required_flag": ("是否必填",),
    "domain_code": ("报送数据域编码",),
    "domain_name": ("报送数据域名称",),
    "field_owner": ("数据项主管部门",),
    "business_definition": ("取值业务口径描述",),
    "source_system": ("取值来源业务系统",),
    "remark": ("备注",),
    "source_system_cn": ("源系统中文名",),
    "source_system_en": ("源系统英文名",),
    "source_table_cn": ("源表中文名",),
    "source_table_en": ("源表英文名",),
    "source_field_cn": ("源表字段中文名",),
    "source_field_en": ("源表字段英文名",),
    "source_logic": ("数据项取数逻辑描述",),
    "confirmed_date": ("确认日期",),
    "mart_table_cn": ("集市表中文名",),
    "mart_table_en": ("集市表英文名",),
    "mart_field_cn": ("集市字段中文名",),
    "mart_field_en": ("集市字段英文名",),
    "mart_logic": ("数据项取数逻辑",),
    "is_manual": ("是否补录",),
    "manual_confirmer": ("确认人",),
    "manual_confirmed_date": ("补录确认日期",),
    "status": ("状态",),
}
_REQUIRED_KEYS = ("field_code", "field_name")
#: 出现多次的列名 -> 依次对应到规范键。
_MULTI_OCCURRENCE_KEYS = {"备注": ("remark", "remark_technical", "remark_other")}
_REQUIRED_FLAG_VALUES = {"是", "y", "yes", "true", "1", "必填", "必输", "not null", "非空"}


@dataclass
class SurveySheetResult:
    sheet_name: str
    table_code: str | None
    table_name: str | None
    field_count: int
    raw_header: list[str]
    parsed_rows: list[dict[str, Any]]
    warnings: list[str] = field(default_factory=list)
    parse_report: dict[str, Any] | None = None


@dataclass
class SurveyTemplateParseOutput:
    sheet_count: int
    table_count: int
    field_count: int
    sheet_names: list[str]
    results: list[SurveySheetResult]
    warnings: list[str]
    parse_report: dict[str, Any] | None = None
    #: 表级口径：{表编号: {theme, filing_note, reporting_domain_code, reporting_domain_name, ...}}
    table_meta: dict[str, dict[str, str]] = field(default_factory=dict)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text.lower() == "none" else text


def _norm(value: Any) -> str:
    return re.sub(r"\s+", "", _text(value)).replace("（", "(").replace("）", ")")


def is_survey_workbook(source: Any) -> bool:
    """判断是否是"一表通调研表"版式（第 3 行左右出现数据项编码/数据项名称）。"""

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
                values = {_norm(item) for item in row if item not in (None, "")}
                if all(marker in values for marker in HEADER_MARKERS):
                    return True
            if any(TABLE_SHEET_PATTERN.match(sheet.title or "") for _ in (0,)):
                continue
        return False
    finally:
        workbook.close()


#: 监管集市列：真实模板在同一个 sheet 里会**重复多组**（“集市表中文名/集市表英文名/
#: 集市字段中文名/集市字段英文名/数据项取数逻辑”出现 2~7 次，Excel 里表现为同名列），
#: 每组对应一个不同的集市表映射。只认第一组会静默丢掉其余映射，因此按出现次序分组。
MART_GROUP_KEYS: tuple[str, ...] = ("mart_table_cn", "mart_table_en", "mart_field_cn",
                                   "mart_field_en", "mart_logic")


def _column_map(raw_header: list[str]) -> tuple[dict[str, int], list[str], list[dict[str, int]]]:
    """返回 `(主列映射, 未识别列, 追加集市列组)`。

    追加集市列组按“第几组”排序（`mart_table_cn#2` 属于第 2 组），每组是
    `{规范键: 列号}`；列名里不带 `#N` 的第一组直接进 `mapping`（兼容原有单组行为）。
    """

    mapping: dict[str, int] = {}
    unmapped: list[str] = []
    seen: dict[str, int] = {}
    mart_occurrence: dict[str, int] = {}
    extra_groups: dict[int, dict[str, int]] = {}
    for index, item in enumerate(raw_header, start=1):
        name = _norm(item)
        if not name:
            continue
        occurred = seen.get(name, 0) + 1
        seen[name] = occurred
        key = None
        if name in _MULTI_OCCURRENCE_KEYS:
            # 同名列（真实模板里“备注”出现三次，分别属于业务口径/技术口径/其他）按出现次序
            # 依次落到 remark / remark_technical / remark_other，不能只保留最后一列。
            alternates = _MULTI_OCCURRENCE_KEYS[name]
            key = alternates[min(occurred, len(alternates)) - 1]
        else:
            for canonical, aliases in SURVEY_COLUMN_ALIASES.items():
                if name in {_norm(alias) for alias in aliases}:
                    key = canonical
                    break
        if key is None:
            unmapped.append(name)
            continue
        if key in MART_GROUP_KEYS:
            position = mart_occurrence.get(key, 0) + 1
            mart_occurrence[key] = position
            if position > 1:
                extra_groups.setdefault(position, {})[key] = index
                continue
        if key in mapping:
            # 同一规范键重复出现时保留第一列，后续列不覆盖（并记入未识别列以便定位）。
            unmapped.append(f"{name}#{occurred}")
            continue
        mapping[key] = index
    return mapping, unmapped, [extra_groups[position] for position in sorted(extra_groups)]


def _parse_table_level(rows: list[list[Any]]) -> dict[str, dict[str, str]]:
    """解析“表级口径”sheet：返回 {表编号: {规范键: 文本}}。

    只取该 sheet 上真正存在的列；缺失的键不写，避免用空串冒充“已确认”。
    """

    header_offset: int | None = None
    for index, row in enumerate(rows[:TABLE_LEVEL_HEADER_SCAN_ROWS]):
        values = {_text(cell) for cell in row}
        if any(alias in values for alias in TABLE_META_ALIASES["table_code"]) and any(
            alias in values for alias in TABLE_META_ALIASES["theme"]
        ):
            header_offset = index
            break
    if header_offset is None:
        return {}
    raw_header = [_text(cell) for cell in rows[header_offset]]
    column_map: dict[int, str] = {}
    for position, name in enumerate(raw_header):
        if not name:
            continue
        for key, aliases in TABLE_META_ALIASES.items():
            if name in aliases and key not in column_map.values():
                column_map[position] = key
                break
    result: dict[str, dict[str, str]] = {}
    for row in rows[header_offset + 1 :]:
        values: dict[str, str] = {}
        for position, key in column_map.items():
            cell = _text(row[position]) if position < len(row) else ""
            if cell:
                values[key] = cell
        table_code = values.get("table_code")
        if not table_code:
            continue
        values.pop("table_code", None)
        if values:
            result[table_code.upper()] = values
    return result


def _find_header_row(rows: list[list[Any]]) -> int | None:
    """定位列名行。

    只要出现任一标记就认为是表头（优先两个都出现的行）：真实模板存在缺“数据项编码”列的
    情况，必须让解析继续走到“缺少关键列 → blocking”，而不是把它当说明性 sheet 静默跳过。
    """

    best: tuple[int, int] | None = None
    for offset, row in enumerate(rows[:LEGACY_HEADER_SCAN_ROWS]):
        values = {_norm(item) for item in row if item not in (None, "")}
        hits = sum(1 for marker in HEADER_MARKERS if marker in values)
        if hits == 0:
            continue
        if hits == len(HEADER_MARKERS):
            return offset
        if best is None or hits > best[0]:
            best = (hits, offset)
    return best[1] if best else None


def _row_payload(values: list[Any], mapping: dict[str, int],
                 extra_groups: list[dict[str, int]] | None = None) -> dict[str, Any]:
    payload: dict[str, Any] = {}
    for key, column in mapping.items():
        index = column - 1
        payload[key] = _text(values[index]) if index < len(values) else ""
    if extra_groups:
        groups = [payload]
        for group in extra_groups:
            item = {key: _text(values[column - 1]) if column - 1 < len(values) else ""
                    for key, column in group.items()}
            if any(item.values()):
                groups.append(item)
        payload["mart_group_payloads"] = groups
    return payload


def _required(value: Any) -> bool:
    return _text(value).lower() in _REQUIRED_FLAG_VALUES


class SurveyTemplateParser:
    """解析一表通监管报送口径调研表。"""

    reader = "survey-template"
    config_version = SURVEY_CONFIG_VERSION

    def parse(self, file_path: Any) -> SurveyTemplateParseOutput:
        payload: bytes | None = None
        target: Any = file_path
        if isinstance(file_path, (str, Path)):
            with open(file_path, "rb") as handle:
                payload = handle.read()
            target = Path(file_path)
        elif hasattr(file_path, "read"):
            payload = file_path.read()
        workbook = load_workbook(target if payload is None else __import__("io").BytesIO(payload), data_only=True)
        sheet_names = [sheet.title for sheet in workbook.worksheets]
        report = ParseReport(
            config_version=self.config_version,
            reader=self.reader,
            file_suffix=".xlsx",
            sheet_total=len(workbook.worksheets),
        )
        results: list[SurveySheetResult] = []
        try:
            table_meta: dict[str, dict[str, str]] = {}
            for sheet in workbook.worksheets:
                if (sheet.title or "").strip() == TABLE_LEVEL_SHEET:
                    table_meta.update(
                        _parse_table_level([list(row) for row in sheet.iter_rows(values_only=True)])
                    )
                    continue
                result = self._parse_sheet(sheet, report)
                if result is not None:
                    results.append(result)
        finally:
            workbook.close()
        warnings = report.warnings()
        return SurveyTemplateParseOutput(
            sheet_count=len(sheet_names),
            table_count=len(results),
            field_count=sum(item.field_count for item in results),
            sheet_names=sheet_names,
            results=results,
            warnings=warnings,
            parse_report=report.to_dict(),
            table_meta=table_meta,
        )

    def _parse_sheet(self, sheet: Any, report: ParseReport) -> SurveySheetResult | None:
        start = len(report.findings)
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        header_offset = _find_header_row(rows)
        if header_offset is None:
            report.sheets_skipped.append({"sheet": sheet.title, "reason": "no_survey_header"})
            report.add(ParseFinding(
                code="sheet_skipped",
                severity=SEVERITY_INFO,
                message=f"{sheet.title} 不是报送表 sheet（未找到“数据项编码/数据项名称”表头），已跳过",
                locator={"sheet": sheet.title},
                suggestion="说明性 sheet 无需导入；如为报送表，请确认第 3 行列名未被改名。",
            ))
            report.sheet_parsed += 0
            return None
        report.sheet_parsed += 1
        raw_header = [_text(item) for item in rows[header_offset]]
        mapping, unmapped, extra_mart_groups = _column_map(raw_header)
        header_row_number = header_offset + 1
        report.header_rows.append({"sheet": sheet.title, "row": header_row_number})
        report.merged_header_cells += sum(
            1
            for merged in sheet.merged_cells.ranges
            if merged.min_row <= header_row_number <= merged.max_row
        )
        missing = [key for key in _REQUIRED_KEYS if key not in mapping]
        warnings: list[str] = []
        if missing:
            labels = {"field_code": "数据项编码", "field_name": "数据项名称"}
            message = f"{sheet.title} 缺少关键列：" + "、".join(labels[key] for key in missing)
            report.sheets_skipped.append({"sheet": sheet.title, "reason": "missing_critical_column"})
            report.add(ParseFinding(
                code="critical_column_missing",
                severity=SEVERITY_BLOCKING,
                message=message + "，无法导入该表",
                locator={"sheet": sheet.title, "row": header_row_number},
                suggestion="请确认表头行未被改名或合并覆盖。",
            ))
            # 缺关键列的 sheet 不产出表：否则 Apply 会建出 0 字段的空监管表。
            return None
        if unmapped:
            report.columns_unrecognized.extend(unmapped)
            report.add(ParseFinding(
                code="unrecognized_column",
                severity=SEVERITY_INFO,
                message=f"{sheet.title} 有 {len(unmapped)} 列不在已知口径内，已忽略",
                locator={"sheet": sheet.title, "row": header_row_number, "columns": unmapped[:10]},
                suggestion="如需保留这些列，请把它们加入调研表解析配置。",
            ))
        table_code = None
        table_name = None
        parsed_rows: list[dict[str, Any]] = []
        seen_codes: set[str] = set()
        blank_streak = 0
        for offset in range(header_offset + 1, len(rows)):
            values = rows[offset]
            excel_row = offset + 1
            payload = _row_payload(values, mapping, extra_mart_groups)
            field_code = payload.get("field_code", "")
            field_name = payload.get("field_name", "")
            if not field_code and not field_name:
                blank_streak += 1
                report.row_skipped_blank += 1
                if blank_streak >= 20 and parsed_rows:
                    break
                continue
            blank_streak = 0
            report.row_total += 1
            table_code = table_code or payload.get("table_code") or sheet.title.split("-")[0]
            table_name = table_name or payload.get("table_name") or sheet.title.split("-", 1)[-1]
            if not field_code or not field_name:
                report.add(ParseFinding(
                    code="row_missing_key",
                    severity=SEVERITY_WARNING,
                    message=f"{sheet.title} 第 {excel_row} 行数据项编码或名称为空，已跳过",
                    locator={"sheet": sheet.title, "row": excel_row},
                    suggestion="补齐数据项编码与名称后重新导入。",
                ))
                continue
            if not CODE_IDENTIFIER_PATTERN.match(field_code):
                report.add(ParseFinding(
                    code="row_marker_skipped",
                    severity=SEVERITY_INFO,
                    message=f"{sheet.title} 第 {excel_row} 行的数据项编码 {field_code!r} 不是字段编码，按分段标记跳过",
                    locator={"sheet": sheet.title, "row": excel_row},
                    suggestion="分段标记行（如“辅助字段”）不需要填写数据项编码；如确为字段请改用 ASCII 编码。",
                ))
                continue
            if not FIELD_CODE_PATTERN.match(field_code):
                report.add(ParseFinding(
                    code="unrecognized_column",
                    severity=SEVERITY_INFO,
                    message=f"{sheet.title} 第 {excel_row} 行的数据项编码 {field_code} 不符合常见编码格式，仍按原值导入",
                    locator={"sheet": sheet.title, "row": excel_row, "cell": cell_reference(excel_row, mapping["field_code"])},
                    suggestion="如编码确实特殊，可忽略；否则检查该单元格是否被手工改动。",
                ))
            if field_code in seen_codes:
                # 真实调研表会把同一数据项块重复排版（打印分页），编码相同的行只保留首行，
                # 否则重复字段会在需求文档里出现多份互相矛盾的正文。
                report.row_skipped_duplicate += 1
                report.add(ParseFinding(
                    code="row_skipped_by_config",
                    severity=SEVERITY_INFO,
                    message=f"{sheet.title} 第 {excel_row} 行的数据项编码 {field_code} 与前面重复，保留首行",
                    locator={"sheet": sheet.title, "row": excel_row},
                    suggestion="同一数据项在一张表里只应出现一次；重复块通常是排版遗留，可删除。",
                ))
                continue
            seen_codes.add(field_code)
            parsed_rows.append(_target_payload(payload, excel_row, table_code, table_name, sheet.title))
        report.row_emitted += len(parsed_rows)
        if not parsed_rows:
            report.add(ParseFinding(
                code="row_missing_key",
                severity=SEVERITY_WARNING,
                message=f"{sheet.title} 表头之后没有可导入的数据行",
                locator={"sheet": sheet.title, "row": header_row_number + 1},
                suggestion="确认该表已填写数据项，或从模板中删除空 sheet。",
            ))
        if mapping.get("mart_table_en") and mapping.get("mart_field_en"):
            report.add(ParseFinding(
                code="merged_header",
                severity=SEVERITY_INFO,
                message=f"{sheet.title} 已识别监管集市列（集市表/集市字段/取数逻辑），将写入集市映射",
                locator={"sheet": sheet.title, "row": header_row_number},
                suggestion="导入后可在“集市口径”中复核。",
            ))
        return SurveySheetResult(sheet.title, table_code, table_name, len(parsed_rows), raw_header,
                                 parsed_rows, report.warnings()[start:], _sheet_report(report, start))


def _target_payload(payload: dict[str, Any], excel_row: int, table_code: str, table_name: str,
                    sheet_name: str) -> dict[str, Any]:
    """把调研表一行映射为 ``TargetField`` 语义载荷（含扩展列与集市映射）。

    映射规则（写进文档，避免后人猜）：

    * ``field_definition``/``regulatory_original_definition`` ← 数据项监管要求（监管原始口径）
    * ``regulatory_refined_definition`` ← 取值业务口径描述（业务细化口径）
    * ``internal_definition`` ← 取值业务口径描述（行内口径缺失时以业务口径代替，并保留原文）
    * ``report_name`` ← sheet 中文表名；``report_field_name`` ← 数据项名称
    * ``remarks`` ← 备注 + 来源系统/制度包/补录标记等（无法落库的列以键值形式保留）
    """

    remark_parts = []
    for key in ("remark", "remark_technical", "remark_other"):
        if payload.get(key):
            remark_parts.append(payload[key])
    meta = []
    for label, key in (("制度包版本", "package_version"), ("报送数据域编码", "domain_code"),
                       ("报送数据域名称", "domain_name"), ("数据项主管部门", "field_owner"),
                       ("取值来源业务系统", "source_system"), ("是否补录", "is_manual"),
                       ("确认人", "manual_confirmer"), ("补录确认日期", "manual_confirmed_date"),
                       ("状态", "status")):
        if payload.get(key):
            meta.append(f"{label}：{payload[key]}")
    if payload.get("source_table_en") or payload.get("source_field_en"):
        meta.append("源系统口径：" + " / ".join(filter(None, [
            payload.get("source_system_cn") or payload.get("source_system_en"),
            payload.get("source_table_cn") or payload.get("source_table_en"),
            payload.get("source_field_cn") or payload.get("source_field_en"),
        ])))
    if meta:
        remark_parts.append("；".join(meta))
    requirement = payload.get("regulatory_requirement") or ""
    business = payload.get("business_definition") or ""
    return {
        "row_number": excel_row,
        "sheet_name": sheet_name,
        "table_code": table_code,
        "table_name": table_name,
        "field_code": payload.get("field_code") or "",
        "field_name": payload.get("field_name") or "",
        "field_type": payload.get("data_format") or "",
        "required_flag": "是" if _required(payload.get("required_flag")) else "否",
        "field_definition": requirement,
        "regulatory_description": requirement,
        "data_category": payload.get("data_category") or "",
        "data_format": payload.get("data_format") or "",
        "regulatory_original_definition": requirement,
        "regulatory_refined_definition": business,
        "report_name": table_name,
        "report_field_name": payload.get("field_name") or "",
        "east_definition": "",
        "internal_definition": business,
        "remarks": "；".join(remark_parts),
        "mart_table_name": payload.get("mart_table_cn") or "",
        "mart_table_code": payload.get("mart_table_en") or "",
        "mart_field_name": payload.get("mart_field_cn") or "",
        "mart_field_code": payload.get("mart_field_en") or "",
        "mart_logic": payload.get("mart_logic") or "",
        "mart_mappings": _mart_mappings(payload),
        "source_system": payload.get("source_system") or "",
        "source_table_name": payload.get("source_table_en") or "",
        "source_field_name": payload.get("source_field_en") or "",
        "is_manual": payload.get("is_manual") or "",
    }


def _mart_mappings(payload: dict[str, Any]) -> list[dict[str, str]]:
    """把一组或多组集市列还原成映射列表（第 1 组来自主列，其余来自重复列组）。

    真实调研表一个字段可能对应多个集市表（例如信贷借据表 + 贴现票据表、借记卡 + 贷记卡
    参数表），这些映射都要保留；只取第一组会让需求文档漏掉一半溯源信息。
    """

    groups = payload.get("mart_group_payloads") or [payload]
    mappings: list[dict[str, str]] = []
    for group in groups:
        table_code = (group.get("mart_table_en") or "").strip()
        field_code = (group.get("mart_field_en") or "").strip()
        if not table_code and not field_code:
            continue
        mappings.append({
            "mart_table_name": (group.get("mart_table_cn") or "").strip(),
            "mart_table_code": table_code,
            "mart_field_name": (group.get("mart_field_cn") or "").strip(),
            "mart_field_code": field_code,
            "mart_logic": (group.get("mart_logic") or "").strip(),
        })
    return mappings


def _sheet_report(report: ParseReport, start: int) -> dict[str, Any]:
    findings = report.findings[start:]
    return {
        "strategy_version": report.strategy_version,
        "config_version": report.config_version,
        "findings": [item.to_dict() for item in findings],
        "complete": not any(item.severity == SEVERITY_BLOCKING for item in findings),
    }


__all__ = [
    "SURVEY_CONFIG_VERSION",
    "SURVEY_COLUMN_ALIASES",
    "SurveySheetResult",
    "SurveyTemplateParseOutput",
    "SurveyTemplateParser",
    "is_survey_workbook",
]
