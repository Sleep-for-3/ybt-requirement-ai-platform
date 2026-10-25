from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.parse_diagnostics import (
    CELL_FORMULA_UNCACHED,
    SEVERITY_BLOCKING,
    SEVERITY_INFO,
    SEVERITY_WARNING,
    ParseFinding,
    ParseReport,
    cell_reference,
    scan_formulas,
)


COLUMN_ALIASES = {
    "table_code": {"表编号", "表代码", "数据表编号", "监管表编号", "table_code"},
    "table_name": {"表名称", "表名", "数据表名称", "监管表名称", "table_name"},
    "field_code": {"字段编号", "字段代码", "字段名", "字段英文名", "字段标识", "field_code"},
    "field_name": {"字段名称", "字段中文名", "中文名称", "字段说明", "field_name"},
    "field_type": {"字段类型", "数据类型", "类型", "field_type"},
    "required_flag": {"是否必填", "必填", "是否为空", "是否可空", "required", "required_flag"},
    "field_definition": {"字段定义", "业务定义", "口径说明", "填报说明", "采集口径", "field_definition"},
    "regulatory_description": {"监管说明", "监管口径", "报送说明", "校验规则", "校验规则说明", "regulatory_description"},
}


HEADER_SCAN_ROWS = 200
LEGACY_HEADER_SCAN_ROWS = 20
TEMPLATE_CONFIG_VERSION = "xlsx-template-config-v1"
CRITICAL_COLUMNS = ("field_code", "field_name")
CRITICAL_COLUMN_LABELS = {"field_code": "字段代码", "field_name": "字段名称"}
ALL_ALIASES = {alias for aliases in COLUMN_ALIASES.values() for alias in aliases}


@dataclass
class ExcelSheetParseResult:
    sheet_name: str
    table_code: str | None
    table_name: str | None
    field_count: int
    raw_header: list[str]
    parsed_rows: list[dict[str, Any]]
    warnings: list[str] = field(default_factory=list)
    parse_report: dict[str, Any] | None = None


@dataclass
class ExcelTemplateParseOutput:
    sheet_count: int
    table_count: int
    field_count: int
    sheet_names: list[str]
    results: list[ExcelSheetParseResult]
    warnings: list[str]
    parse_report: dict[str, Any] | None = None


class ExcelTemplateParser:
    def parse(self, file_path: Any) -> ExcelTemplateParseOutput:
        payload: bytes | None = None
        target: Any = file_path
        if isinstance(file_path, (str, Path)):
            path = Path(file_path)
            if path.suffix.lower() == ".xls":
                raise ValueError("暂不支持 .xls，请另存为 .xlsx 后上传。")
            if path.suffix.lower() != ".xlsx":
                raise ValueError("MVP 阶段只支持 .xlsx 一表通模板。")
            target = path
        else:
            payload = file_path.read() if hasattr(file_path, "read") else bytes(file_path)
            target = BytesIO(payload)

        # Two views of the same workbook: cached values for content, formula
        # text so a cache miss can be reported instead of silently dropped.
        if payload is None:
            values_workbook = load_workbook(target, data_only=True)
            formula_workbook = load_workbook(target, data_only=False)
        else:
            values_workbook = load_workbook(BytesIO(payload), data_only=True)
            formula_workbook = load_workbook(BytesIO(payload), data_only=False)

        report = ParseReport(
            config_version=TEMPLATE_CONFIG_VERSION,
            reader="regulatory-template",
            file_suffix=".xlsx",
            sheet_total=len(values_workbook.worksheets),
        )
        results: list[ExcelSheetParseResult] = []
        all_warnings: list[str] = []
        for sheet in values_workbook.worksheets:
            formula_sheet = formula_workbook[sheet.title] if sheet.title in formula_workbook.sheetnames else None
            result = self._parse_sheet(sheet, formula_sheet, report)
            results.append(result)
            all_warnings.extend(result.warnings)
        report.row_emitted = sum(result.field_count for result in results)
        return ExcelTemplateParseOutput(
            sheet_count=len(results),
            table_count=sum(1 for result in results if result.field_count > 0),
            field_count=sum(result.field_count for result in results),
            sheet_names=[sheet.title for sheet in values_workbook.worksheets],
            results=results,
            warnings=all_warnings,
            parse_report=report.to_dict(),
        )

    def _parse_sheet(self, sheet, formula_sheet, report: ParseReport) -> ExcelSheetParseResult:
        start = len(report.findings)
        rows = list(sheet.iter_rows(values_only=True))
        kinds = scan_formulas(
            sheet.title, formula_sheet, lambda row, column: sheet.cell(row, column).value, report
        )
        header_index = _find_header_index(rows)
        if header_index is None:
            has_content = any(_cell_to_text(value) for row in rows for value in row)
            report.sheets_skipped.append({"sheet": sheet.title, "row_total": sheet.max_row})
            report.add(
                ParseFinding(
                    "sheet_skipped",
                    SEVERITY_BLOCKING if has_content else SEVERITY_INFO,
                    f"Sheet {sheet.title} 未识别到表头",
                    locator={"sheet": sheet.title, "scanned_rows": len(rows)},
                )
            )
            warning = f"Sheet {sheet.title} 未识别到表头"
            return ExcelSheetParseResult(
                sheet.title, None, sheet.title, 0, [], [], [warning], _sheet_report(report, start)
            )

        report.sheet_parsed += 1
        if header_index >= LEGACY_HEADER_SCAN_ROWS:
            report.add(
                ParseFinding(
                    "header_beyond_legacy_scan",
                    SEVERITY_INFO,
                    f"Sheet {sheet.title} 的表头在第 {header_index + 1} 行，超出历史扫描范围（前 {LEGACY_HEADER_SCAN_ROWS} 行）",
                    locator={"sheet": sheet.title, "row": header_index + 1},
                )
            )
        raw_header = [_cell_to_text(value) for value in rows[header_index]]
        column_map = _build_column_map(raw_header)
        header_cells = [
            cell_reference(header_index + 1, index + 1)
            for index in range(len(raw_header))
            if kinds.get((header_index + 1, index + 1)) == CELL_FORMULA_UNCACHED
        ]
        if header_cells:
            report.add(
                ParseFinding(
                    "formula_in_header",
                    SEVERITY_BLOCKING,
                    f"Sheet {sheet.title} 表头含无缓存公式，列名可能为空",
                    locator={"sheet": sheet.title, "cells": ",".join(header_cells[:5])},
                )
            )
        has_rows = len(rows) > header_index + 1
        _record_header_findings(sheet.title, raw_header, column_map, report, has_rows)

        parsed_rows: list[dict[str, Any]] = []
        warnings: list[str] = []
        for excel_row_number, values in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if not any(_cell_to_text(value) for value in values):
                continue
            parsed = _parse_row(values, column_map)
            if not parsed.get("field_code"):
                warnings.append(f"Sheet {sheet.title} 第 {excel_row_number} 行缺少字段代码，apply 时将跳过")
            if not parsed.get("field_name"):
                warnings.append(f"Sheet {sheet.title} 第 {excel_row_number} 行缺少字段名称，apply 时将跳过")
            missing_cells = [
                cell_reference(excel_row_number, index + 1)
                for index in column_map.values()
                if kinds.get((excel_row_number, index + 1)) == CELL_FORMULA_UNCACHED
            ]
            if missing_cells:
                report.add(
                    ParseFinding(
                        "formula_cache_missing",
                        SEVERITY_BLOCKING,
                        f"Sheet {sheet.title} 第 {excel_row_number} 行有 {len(missing_cells)} 列是无缓存公式，apply 时会写入空值",
                        locator={
                            "sheet": sheet.title,
                            "row": excel_row_number,
                            "cells": ",".join(missing_cells[:5]),
                        },
                    )
                )
            parsed["row_number"] = excel_row_number
            parsed_rows.append(parsed)

        first_with_table = next((row for row in parsed_rows if row.get("table_code") or row.get("table_name")), {})
        table_code = first_with_table.get("table_code") or _guess_table_code(sheet.title)
        table_name = first_with_table.get("table_name") or sheet.title
        valid_field_count = sum(1 for row in parsed_rows if row.get("field_code") and row.get("field_name"))
        # Legacy messages stay first: existing callers and tests index them by position.
        diagnostics = [
            item.as_warning()
            for item in report.findings[start:]
            if item.severity != SEVERITY_INFO
        ]
        return ExcelSheetParseResult(
            sheet_name=sheet.title,
            table_code=table_code,
            table_name=table_name,
            field_count=valid_field_count,
            raw_header=raw_header,
            parsed_rows=parsed_rows,
            warnings=[*warnings, *diagnostics],
            parse_report=_sheet_report(report, start),
        )


def _find_header_index(rows: list[tuple]) -> int | None:
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        headers = {_cell_to_text(value) for value in row}
        if headers & ALL_ALIASES:
            return index
    return None


def _record_header_findings(
    sheet_name: str,
    raw_header: list[str],
    column_map: dict[str, int],
    report: ParseReport,
    has_rows: bool,
) -> None:
    seen: dict[str, int] = {}
    for index, name in enumerate(raw_header):
        if not name:
            continue
        if name in seen:
            report.duplicate_headers.append(
                {"sheet": sheet_name, "name": name, "columns": [seen[name] + 1, index + 1]}
            )
            report.add(
                ParseFinding(
                    "duplicate_header",
                    SEVERITY_WARNING,
                    f"Sheet {sheet_name} 重复列名“{name}”，只有一列会被采用，另一列的数据会丢失",
                    locator={"sheet": sheet_name, "column": get_column_letter(index + 1)},
                )
            )
        else:
            seen[name] = index
    unmapped = sorted({name for name in raw_header if name and name not in ALL_ALIASES})
    if unmapped:
        report.columns_unrecognized.append({"sheet": sheet_name, "columns": unmapped})
        report.add(
            ParseFinding(
                "unrecognized_column",
                SEVERITY_INFO,
                f"Sheet {sheet_name} 有 {len(unmapped)} 个未映射列：{'、'.join(unmapped[:10])}",
                locator={"sheet": sheet_name},
            )
        )
    for canonical in CRITICAL_COLUMNS:
        if canonical not in column_map:
            report.columns_critical_missing.append({"sheet": sheet_name, "column": canonical})
            report.add(
                ParseFinding(
                    "critical_column_missing",
                    SEVERITY_BLOCKING if has_rows else SEVERITY_WARNING,
                    f"Sheet {sheet_name} 缺少关键列 {CRITICAL_COLUMN_LABELS[canonical]}",
                    locator={"sheet": sheet_name},
                )
            )


def _sheet_report(report: ParseReport, start: int) -> dict[str, Any]:
    findings = report.findings[start:]
    return {
        "strategy_version": report.strategy_version,
        "config_version": report.config_version,
        "findings": [item.to_dict() for item in findings],
        "complete": all(item.severity == SEVERITY_INFO for item in findings),
    }


def _build_column_map(raw_header: list[str]) -> dict[str, int]:
    column_map: dict[str, int] = {}
    for index, name in enumerate(raw_header):
        normalized = name.strip()
        for canonical, aliases in COLUMN_ALIASES.items():
            if normalized in aliases and canonical not in column_map:
                column_map[canonical] = index
    return column_map


def _parse_row(values: tuple, column_map: dict[str, int]) -> dict[str, Any]:
    parsed: dict[str, Any] = {}
    for key in COLUMN_ALIASES:
        index = column_map.get(key)
        parsed[key] = _cell_to_text(values[index]) if index is not None and index < len(values) else ""
    parsed["required_flag"] = _parse_required_flag(parsed.get("required_flag", ""))
    return parsed


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_required_flag(value: str) -> bool:
    normalized = value.strip().lower()
    if normalized in {"是", "y", "yes", "true", "1", "必填", "not null", "非空"}:
        return True
    if normalized in {"否", "n", "no", "false", "0", "可空", "nullable"}:
        return False
    if "不可空" in normalized or "不能为空" in normalized:
        return True
    return False


def _guess_table_code(sheet_name: str) -> str | None:
    cleaned = sheet_name.strip()
    return cleaned if cleaned.isascii() and cleaned else None
