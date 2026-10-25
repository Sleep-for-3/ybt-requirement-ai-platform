"""客户 XLSX 需求文档模板渲染器（内部代号 P1-B）。

本模块只做一件事：把客户给定的 Excel 版式当作「模板」，用冻结的需求 JSON
（``backend/app/services/requirement_scope.py`` 里 ``export_document`` 消费的那份
``content``）按配置声明的「命名单元格」和「重复行块」填充，并尽量保持模板的
版式特征：sheet 顺序、合并单元格、行列尺寸、单元格样式、公式、打印区域、
打印标题、页眉页脚。

与监管报送模板（``TargetTable`` + 字段列映射，见 ``app/services/template_parser/``）
是两套不同的东西，本模块不复用也不修改它。

设计边界
--------
* 纯函数：``analyze_xlsx_template`` / ``render_xlsx``。不建数据库模型、不落盘、
  不发起网络请求、不调用外部程序（包括 LibreOffice）。
* **不重算公式**：openpyxl 只写公式文本，不计算也不刷新缓存值。输出文件里的
  公式单元格在 Excel/LibreOffice 打开并重算之前，``data_only=True`` 读到的缓存值
  仍为空或旧值。「保存成功」不等于「公式已重算」，诊断对象里对此有明确字段。
* 来自业务文本的单元格值如果以 ``=`` ``+`` ``-`` ``@`` 开头，一律写成**文本**
  （``data_type='s'`` + ``quotePrefix=True``），不会被 Excel 当成公式执行；
  模板自带的可信公式原样保留为公式（行块内还会按相对行号翻译）。

不能保真的对象（检测到后按 ``allow_lossy`` 拒绝渲染或渲染并记录警告）
------------------------------------------------------------------
VBA/宏（.xlsm）、图表、数据透视表、图片、数据验证、条件格式、批注、外部链接。
检测方式见 :func:`_detect_unsupported_features`，清单见
:data:`UNSUPPORTED_FEATURE_LABELS`。
"""
from __future__ import annotations

from dataclasses import dataclass, field
import copy
import hashlib
import io
import re
import shutil
import zipfile
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter
from openpyxl.worksheet.cell_range import CellRange

__all__ = [
    "MAX_EXCEL_CELL_CHARS",
    "DEFAULT_MAX_CELL_CHARS",
    "FORMULA_CACHE_NOTE",
    "UNSUPPORTED_FEATURE_LABELS",
    "XlsxTemplateError",
    "XlsxTemplateConfigError",
    "XlsxTemplateHashMismatch",
    "XlsxUnsupportedFeatureError",
    "XlsxMappingMissingError",
    "XlsxCellTooLongError",
    "XlsxRenderValidationError",
    "CellMapping",
    "RowBlock",
    "XlsxTemplateConfig",
    "UnsupportedFeature",
    "MissingItem",
    "SheetInfo",
    "DefinedNameInfo",
    "RowBlockInfo",
    "XlsxTemplateAnalysis",
    "FilledCell",
    "RenderDiagnostics",
    "template_sha256",
    "document_values",
    "analyze_xlsx_template",
    "render_xlsx",
]

#: Excel 单元格文本硬上限（字符）。
MAX_EXCEL_CELL_CHARS = 32767
#: 默认「提前告警」阈值；超过只告警不截断（截断会静默丢业务内容）。
DEFAULT_MAX_CELL_CHARS = 32000

#: 公式缓存值的统一口径说明，写在 RenderDiagnostics 里，防止调用方误宣称已重算。
FORMULA_CACHE_NOTE = (
    "本模块不重算公式：openpyxl 只写入公式文本，不计算公式、也不刷新缓存值。"
    "输出文件中的公式单元格在 Excel/LibreOffice 打开并重算之前，"
    "data_only=True 读到的缓存值仍为空或旧值；保存成功不代表公式已重算。"
)

UNSUPPORTED_FEATURE_LABELS: dict[str, str] = {
    "vba": "VBA 宏（.xlsm）",
    "chart": "图表",
    "pivot": "数据透视表/透视缓存",
    "image": "图片",
    "data_validation": "数据验证",
    "conditional_formatting": "条件格式",
    "comment": "批注",
    "external_link": "外部链接",
}

UNSUPPORTED_FEATURE_HINTS: dict[str, str] = {
    "vba": "openpyxl 以 keep_vba=False 打开并重写工作簿时会丢弃 xl/vbaProject.bin，"
           "宏会真正丢失；请先用 Excel「另存为 .xlsx（不含宏）」重新制作模板，"
           "或改为在 .xlsm 上手工维护。",
    "chart": "图表锚点不随行块插入移动，图表的引用区域也不会跟着扩展；"
             "请把图表移出被填充的区域，或在 Excel 中确认后再登记模板。",
    "pivot": "透视表不保证能被字节级保留，透视源区域也不会随插行调整；"
             "请在模板中删除透视表，或单独维护透视表版本。",
    "image": "图片锚点不随行块插入移动，可能与新数据重叠；请把图片移出被填充的区域。",
    "data_validation": "数据验证范围（sqref）不会随插行调整，校验规则可能落到错误的行；"
                       "请用配置化校验或填充后手工修正。",
    "conditional_formatting": "条件格式范围不会随插行调整，格式可能退化为常量样式；"
                              "请在插行后手工修正规则范围。",
    "comment": "批注锚点不随插行移动；请删除批注或改为填充后人工添加。",
    "external_link": "外部链接部件不会被本模块维护，且重写工作簿后可能丢失；"
                     "请先用 Excel「编辑链接」断开外部引用。",
}

#: 默认的「残留占位符」形态：{{key}} / ${key} / <<key>>。
DEFAULT_PLACEHOLDER_PATTERNS: tuple[str, ...] = (
    r"\{\{[^{}]{1,200}\}\}",
    r"\$\{[^{}]{1,200}\}",
    r"<<[^<>]{1,200}>>",
)

#: 中文逻辑键别名 → 需求 content 的字段路径。适配人员可以直接用中文键写 cell_map。
STANDARD_VALUE_ALIASES: dict[str, str] = {
    "需求名称": "requirement.name",
    "业务目标": "requirement.objective",
    "需求目的": "requirement.objective",
    "业务背景": "requirement.background",
    "纳入范围": "requirement.inclusion",
    "排除范围": "requirement.exclusion",
    "排除条件": "requirement.exclusion",
    "口径生效日期": "requirement.effective_date",
    "验收要求": "requirement.acceptance_criteria",
    "交付状态": "document.status",
    "内容版本": "document.content_version",
    "固定内容标识": "document.digest",
    "需求版本": "requirement.version",
}

_BUILTIN_DEFINED_NAME_PREFIXES = ("_xlnm.",)
_CELL_REF_RE = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]*$")
_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_ILLEGAL_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# 包部件 → 不能保真的对象。按 OOXML 包结构探测，比对象模型更可靠
# （对象模型只覆盖 openpyxl 认识的部分）。
_PART_FEATURE_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("vba", re.compile(r"^xl/vbaProject\.bin$", re.IGNORECASE)),
    ("chart", re.compile(r"^xl/charts/", re.IGNORECASE)),
    ("pivot", re.compile(r"^xl/pivot(Tables|Cache)/", re.IGNORECASE)),
    ("image", re.compile(r"^xl/media/", re.IGNORECASE)),
    ("comment", re.compile(r"^xl/(threaded)?comments/", re.IGNORECASE)),
    ("external_link", re.compile(r"^xl/externalLinks/", re.IGNORECASE)),
)


# --------------------------------------------------------------------------- #
# 异常
# --------------------------------------------------------------------------- #
class XlsxTemplateError(Exception):
    """模板渲染的统一异常基类；``code`` 供 API 层映射到业务错误码。"""

    code = "XLSX_TEMPLATE_ERROR"

    def __init__(self, message: str, *, code: str | None = None, details: Sequence[str] = ()):
        self.message = message
        self.details = tuple(details)
        if code:
            self.code = code
        full = message
        if self.details:
            full = message + "\n" + "\n".join(f"- {item}" for item in self.details)
        super().__init__(full)

    def as_dict(self) -> dict[str, Any]:
        return {"code": self.code, "message": self.message, "details": list(self.details)}


class XlsxTemplateConfigError(XlsxTemplateError):
    """配置本身不合法（缺字段、类型错误、哈希格式不对）。"""

    code = "XLSX_CONFIG_INVALID"


class XlsxTemplateHashMismatch(XlsxTemplateError):
    """模板字节的 sha256 与配置声明的 template_hash 不一致。"""

    code = "XLSX_TEMPLATE_HASH_MISMATCH"


class XlsxUnsupportedFeatureError(XlsxTemplateError):
    """模板含不能保真的对象，且 allow_lossy=False。"""

    code = "XLSX_UNSUPPORTED_FEATURE"


class XlsxMappingMissingError(XlsxTemplateError):
    """配置声明了模板/内容里不存在的映射或必需键。"""

    code = "XLSX_MAPPING_MISSING"


class XlsxCellTooLongError(XlsxTemplateError):
    """单元格文本超过 Excel 32767 字符上限。"""

    code = "XLSX_CELL_TOO_LONG"


class XlsxRenderValidationError(XlsxTemplateError):
    """渲染后读回校验失败；此时不产出文件。"""

    code = "XLSX_RENDER_VALIDATION_FAILED"


# --------------------------------------------------------------------------- #
# 配置模型（纯 dataclass，不建数据库模型）
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class CellTarget:
    """一个具体的单元格位置。"""

    sheet: str
    coordinate: str

    @property
    def label(self) -> str:
        return f"{self.sheet}!{self.coordinate}"


@dataclass(frozen=True)
class CellMapping:
    """``cell_map`` 的一项：逻辑键 → 命名单元格 / 坐标 (+ 值来源路径)。"""

    key: str
    target: CellTarget | None = None
    named_range: str | None = None
    value_path: str | None = None
    required: bool = False

    @property
    def description(self) -> str:
        if self.named_range:
            return f"命名单元格「{self.named_range}」"
        if self.target:
            return self.target.label
        return "未声明目标"

    @classmethod
    def from_dict(cls, key: str, raw: Any) -> "CellMapping":
        if isinstance(raw, str):
            raw = _parse_target_shorthand(raw, key)
        if not isinstance(raw, Mapping):
            raise XlsxTemplateConfigError(
                f"cell_map 的键「{key}」必须是字符串或对象，当前类型为 {type(raw).__name__}。",
                details=["示例：\"需求名称\": {\"named_range\": \"需求名称\", \"value_path\": \"requirement.name\"}",
                         "或简写：\"需求名称\": \"封面!B2\""],
            )
        named = raw.get("named_range")
        if named is not None:
            named = str(named).strip()
            if not named:
                raise XlsxTemplateConfigError(f"cell_map 的键「{key}」的 named_range 不能为空字符串。")
        target = None
        sheet = raw.get("sheet")
        cell = raw.get("cell") or raw.get("coordinate") or raw.get("at")
        if isinstance(cell, str) and "!" in cell:
            # 形如 "封面!B2" 的合写
            sheet_part, _, cell_part = cell.strip().rpartition("!")
            if sheet_part:
                sheet, cell = sheet_part.strip("'"), cell_part
        if cell is not None and named is None:
            if sheet is None:
                raise XlsxTemplateConfigError(
                    f"cell_map 的键「{key}」声明了坐标 {cell!r} 但没有给出 sheet。",
                    details=["请写成 {\"sheet\": \"封面\", \"cell\": \"B2\"} 或简写 \"封面!B2\"。"],
                )
            sheet_text = str(sheet).strip()
            cell_text = str(cell).strip().upper()
            if not _CELL_REF_RE.match(cell_text):
                raise XlsxTemplateConfigError(
                    f"cell_map 的键「{key}」的坐标 {cell_text!r} 不是合法单元格坐标（应形如 B2）。")
            target = CellTarget(sheet=sheet_text, coordinate=cell_text)
        if named is None and target is None:
            raise XlsxTemplateConfigError(
                f"cell_map 的键「{key}」必须声明 named_range 或 sheet+cell 其中之一。")
        value_path = raw.get("value_path")
        return cls(
            key=key,
            target=target,
            named_range=named,
            value_path=None if value_path in (None, "") else str(value_path),
            required=bool(raw.get("required", False)),
        )


@dataclass(frozen=True)
class RowBlock:
    """``row_blocks`` 的一项：逻辑列表键 → 一个可重复的样板行块。"""

    key: str
    sheet: str
    columns: Mapping[str, str]
    sample_row: int | None = None
    header_row: int | None = None
    anchor_cell: str | None = None
    sample_rows: int = 1
    max_rows: int = 1000
    source: str | None = None
    column_paths: Mapping[str, str] = field(default_factory=dict)
    required: bool = False

    @property
    def first_data_row(self) -> int | None:
        if self.sample_row is not None:
            return int(self.sample_row)
        if self.anchor_cell:
            match = re.match(r"^[A-Za-z]{1,3}([1-9][0-9]*)$", self.anchor_cell.strip())
            if match:
                return int(match.group(1))
        if self.header_row is not None:
            return int(self.header_row) + 1
        return None

    @property
    def column_letters(self) -> tuple[str, ...]:
        return tuple(sorted(self.columns.values(), key=column_index_from_string))

    @classmethod
    def from_dict(cls, key: str, raw: Any) -> "RowBlock":
        if not isinstance(raw, Mapping):
            raise XlsxTemplateConfigError(
                f"row_blocks 的键「{key}」必须是对象，当前类型为 {type(raw).__name__}。")
        sheet = raw.get("sheet")
        if not sheet:
            raise XlsxTemplateConfigError(f"row_blocks 的键「{key}」必须给出 sheet。")
        columns = raw.get("columns") or {}
        if not isinstance(columns, Mapping) or not columns:
            raise XlsxTemplateConfigError(
                f"row_blocks 的键「{key}」必须给出非空 columns（逻辑列 → 列字母），如 {{\"字段编码\": \"A\"}}。")
        normalized: dict[str, str] = {}
        for column_key, letter in columns.items():
            text = str(letter).strip().upper()
            if not re.match(r"^[A-Z]{1,3}$", text):
                raise XlsxTemplateConfigError(
                    f"row_blocks 的键「{key}」的列「{column_key}」列字母 {letter!r} 不合法，应为 A..XFD。")
            normalized[str(column_key)] = text
        sample_row = raw.get("sample_row")
        header_row = raw.get("header_row")
        anchor_cell = raw.get("anchor_cell")
        if sample_row is None and not anchor_cell and header_row is None:
            raise XlsxTemplateConfigError(
                f"row_blocks 的键「{key}」必须给出 sample_row、anchor_cell、header_row 三者之一来定位样板行。")
        sample_rows = int(raw.get("sample_rows", 1))
        if sample_rows < 1:
            raise XlsxTemplateConfigError(f"row_blocks 的键「{key}」的 sample_rows 必须 >= 1。")
        max_rows = int(raw.get("max_rows", 1000))
        if max_rows < 1:
            raise XlsxTemplateConfigError(f"row_blocks 的键「{key}」的 max_rows 必须 >= 1。")
        column_paths = raw.get("column_paths") or {}
        if not isinstance(column_paths, Mapping):
            raise XlsxTemplateConfigError(
                f"row_blocks 的键「{key}」的 column_paths 必须是对象（逻辑列 → 行内字段路径）。")
        block = cls(
            key=key,
            sheet=str(sheet),
            columns=normalized,
            sample_row=None if sample_row is None else int(sample_row),
            header_row=None if header_row is None else int(header_row),
            anchor_cell=None if anchor_cell in (None, "") else str(anchor_cell).strip().upper(),
            sample_rows=sample_rows,
            max_rows=max_rows,
            source=None if raw.get("source") in (None, "") else str(raw.get("source")),
            column_paths={str(k): str(v) for k, v in column_paths.items()},
            required=bool(raw.get("required", False)),
        )
        if block.first_data_row is None:
            raise XlsxTemplateConfigError(
                f"row_blocks 的键「{key}」无法确定样板行行号：anchor_cell={block.anchor_cell!r} 不合法。")
        if block.first_data_row < 1:
            raise XlsxTemplateConfigError(f"row_blocks 的键「{key}」的样板行行号必须 >= 1。")
        return block


@dataclass(frozen=True)
class XlsxTemplateConfig:
    """模板配置。``template_hash`` 是登记模板时的 sha256，必须与字节一致。"""

    template_hash: str
    cell_map: Mapping[str, CellMapping] = field(default_factory=dict)
    row_blocks: Mapping[str, RowBlock] = field(default_factory=dict)
    required_keys: tuple[str, ...] = ()
    config_version: str = "1"
    supported_scope: str = "requirement_document"
    allow_lossy: bool = False
    max_cell_chars: int = DEFAULT_MAX_CELL_CHARS
    placeholder_patterns: tuple[str, ...] = DEFAULT_PLACEHOLDER_PATTERNS

    @classmethod
    def from_dict(cls, raw: Mapping[str, Any]) -> "XlsxTemplateConfig":
        if not isinstance(raw, Mapping):
            raise XlsxTemplateConfigError(
                f"模板配置必须是 dict，当前类型为 {type(raw).__name__}。")
        template_hash = str(raw.get("template_hash") or "").strip().lower()
        if not template_hash:
            raise XlsxTemplateConfigError(
                "模板配置缺少 template_hash（模板字节的 sha256），无法保证「填的是登记过的版式」。",
                details=["可用 requirement_xlsx_template.template_sha256(模板字节) 计算后写入配置。"],
            )
        if not _SHA256_RE.match(template_hash):
            raise XlsxTemplateConfigError(
                f"template_hash 必须是 64 位十六进制 sha256，当前为 {template_hash!r}。")
        raw_cell_map = raw.get("cell_map") or {}
        raw_row_blocks = raw.get("row_blocks") or {}
        if not isinstance(raw_cell_map, Mapping):
            raise XlsxTemplateConfigError("cell_map 必须是对象（逻辑键 → 目标单元格）。")
        if not isinstance(raw_row_blocks, Mapping):
            raise XlsxTemplateConfigError("row_blocks 必须是对象（逻辑键 → 行块定义）。")
        cell_map = {str(k): CellMapping.from_dict(str(k), v) for k, v in raw_cell_map.items()}
        row_blocks = {str(k): RowBlock.from_dict(str(k), v) for k, v in raw_row_blocks.items()}
        overlap = set(cell_map) & set(row_blocks)
        if overlap:
            raise XlsxTemplateConfigError(
                f"cell_map 与 row_blocks 使用了同名逻辑键：{'、'.join(sorted(overlap))}。两者必须唯一。")
        required_keys = tuple(str(k) for k in (raw.get("required_keys") or ()))
        unknown_required = [k for k in required_keys if k not in cell_map and k not in row_blocks]
        if unknown_required:
            raise XlsxTemplateConfigError(
                f"required_keys 引用了未在 cell_map / row_blocks 中声明的逻辑键：{'、'.join(unknown_required)}。")
        max_cell_chars = int(raw.get("max_cell_chars", DEFAULT_MAX_CELL_CHARS))
        if not 1 <= max_cell_chars <= MAX_EXCEL_CELL_CHARS:
            raise XlsxTemplateConfigError(
                f"max_cell_chars 必须在 1..{MAX_EXCEL_CELL_CHARS} 之间，当前为 {max_cell_chars}。")
        patterns = tuple(raw.get("placeholder_patterns") or DEFAULT_PLACEHOLDER_PATTERNS)
        for pattern in patterns:
            try:
                re.compile(pattern)
            except re.error as exc:
                raise XlsxTemplateConfigError(
                    f"placeholder_patterns 里的正则 {pattern!r} 无法编译：{exc}") from exc
        return cls(
            template_hash=template_hash,
            cell_map=cell_map,
            row_blocks=row_blocks,
            required_keys=required_keys,
            config_version=str(raw.get("config_version") or "1"),
            supported_scope=str(raw.get("supported_scope") or "requirement_document"),
            allow_lossy=bool(raw.get("allow_lossy", False)),
            max_cell_chars=max_cell_chars,
            placeholder_patterns=patterns,
        )


def _parse_target_shorthand(raw: str, key: str) -> dict[str, Any]:
    text = raw.strip()
    if not text:
        raise XlsxTemplateConfigError(f"cell_map 的键「{key}」的目标为空字符串。")
    if text.startswith("@"):
        return {"named_range": text[1:].strip()}
    if "!" in text:
        sheet_part, _, cell_part = text.rpartition("!")
        return {"sheet": sheet_part.strip().strip("'"), "cell": cell_part.strip()}
    if _CELL_REF_RE.match(text):
        raise XlsxTemplateConfigError(
            f"cell_map 的键「{key}」只给了坐标 {text!r} 没给 sheet，无法定位（同一坐标在不同 sheet 上是不同单元格）。",
            details=["请写成 \"封面!B2\" 或 {\"sheet\": \"封面\", \"cell\": \"B2\"}；"
                     "若这是命名单元格的名字，请写成 \"@名字\"。"],
        )
    return {"named_range": text}


def _coerce_config(config: Mapping[str, Any] | XlsxTemplateConfig) -> XlsxTemplateConfig:
    if isinstance(config, XlsxTemplateConfig):
        return config
    return XlsxTemplateConfig.from_dict(config)


# --------------------------------------------------------------------------- #
# 分析结果
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class UnsupportedFeature:
    kind: str
    location: str
    detail: str
    source: str = "package"  # package | object

    @property
    def label(self) -> str:
        return UNSUPPORTED_FEATURE_LABELS.get(self.kind, self.kind)

    @property
    def message(self) -> str:
        return f"{self.label}（{self.location}）：{self.detail}"


@dataclass(frozen=True)
class MissingItem:
    kind: str  # sheet | named_range | cell | row_block | required_key | template_hash
    key: str
    detail: str
    sheet: str | None = None

    @property
    def message(self) -> str:
        location = f"[{self.sheet}] " if self.sheet else ""
        return f"{location}{self.key}：{self.detail}"


@dataclass(frozen=True)
class SheetInfo:
    title: str
    state: str
    max_row: int
    max_column: int
    merged_ranges: tuple[str, ...]
    print_area: str | None
    print_title_rows: str | None
    print_title_cols: str | None
    header_footer: Mapping[str, tuple[str | None, str | None, str | None]]
    freeze_panes: str | None
    data_validation_count: int
    conditional_formatting_count: int
    chart_count: int
    image_count: int
    comment_count: int


@dataclass(frozen=True)
class DefinedNameInfo:
    name: str
    attr_text: str
    sheet: str | None
    ref: str | None
    is_cell: bool
    is_builtin: bool = False


@dataclass(frozen=True)
class RowBlockInfo:
    key: str
    sheet: str
    columns: Mapping[str, str]
    header_row: int | None
    sample_row: int
    sample_rows: int
    first_data_row: int
    max_rows: int
    source: str
    resolved: bool
    note: str = ""


@dataclass(frozen=True)
class XlsxTemplateAnalysis:
    """``analyze_xlsx_template`` 的返回值。"""

    template_hash: str
    template_hash_matches_config: bool
    config_version: str
    supported_scope: str
    sheet_names: tuple[str, ...]
    sheets: tuple[SheetInfo, ...]
    defined_names: tuple[DefinedNameInfo, ...]
    merged_ranges: Mapping[str, tuple[str, ...]]
    print_areas: Mapping[str, str | None]
    print_titles: Mapping[str, tuple[str | None, str | None]]
    headers_footers: Mapping[str, Mapping[str, tuple[str | None, str | None, str | None]]]
    row_blocks: tuple[RowBlockInfo, ...]
    cell_targets: Mapping[str, CellTarget | None]
    unsupported_features: tuple[UnsupportedFeature, ...]
    missing: tuple[MissingItem, ...]
    warnings: tuple[str, ...]

    @property
    def blocking_missing(self) -> tuple[MissingItem, ...]:
        return tuple(
            item for item in self.missing
            if item.kind in {"sheet", "named_range", "cell", "row_block", "required_key", "template_hash"}
        )

    @property
    def lossy_kinds(self) -> tuple[str, ...]:
        return tuple(dict.fromkeys(feature.kind for feature in self.unsupported_features))

    def summary(self) -> str:
        lines = [f"模板 sha256 {self.template_hash[:12]}…，{len(self.sheet_names)} 个 sheet："
                 f"{'、'.join(self.sheet_names)}"]
        if self.unsupported_features:
            lines.append("不能保真的对象：" + "；".join(
                f"{UNSUPPORTED_FEATURE_LABELS.get(k, k)}" for k in self.lossy_kinds))
        if self.missing:
            lines.append("配置声明但模板中缺失：" + "；".join(item.message for item in self.missing))
        return "\n".join(lines)


@dataclass(frozen=True)
class FilledCell:
    logical_key: str
    sheet: str
    coordinate: str
    value: str
    escaped_as_text: bool


@dataclass(frozen=True)
class RenderDiagnostics:
    """``render_xlsx`` 的第二个返回值：给 API/前端看的真实渲染诊断。"""

    template_hash: str
    config_version: str
    supported_scope: str
    output_sha256: str
    filled_cells: tuple[FilledCell, ...]
    row_block_counts: Mapping[str, int]
    warnings: tuple[str, ...]
    unsupported_features: tuple[UnsupportedFeature, ...]
    lossy_rendered: bool
    # 公式口径：这三个字段永远如实反映「没有重算」。
    formula_recalculated: bool = False
    formula_cache_verified: bool = False
    formula_cache_note: str = FORMULA_CACHE_NOTE
    recalculation_tool_available: str | None = None
    recalculation_note: str = ""

    def summary(self) -> str:
        return (
            f"已渲染 {len(self.filled_cells)} 个命名单元格、"
            f"{sum(self.row_block_counts.values())} 行行块数据"
            f"（{self.row_block_counts}）；输出 sha256 {self.output_sha256[:12]}…。"
            + (f"\n公式：{self.formula_cache_note}" if not self.formula_recalculated else "")
        )


# --------------------------------------------------------------------------- #
# 包结构扫描
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class _PackageScan:
    namelist: tuple[str, ...]
    has_vba_content_type: bool


def template_sha256(template_bytes: bytes) -> str:
    """模板字节的 sha256（登记模板时写进配置的 template_hash）。"""

    return hashlib.sha256(bytes(template_bytes)).hexdigest()


def _scan_package(data: bytes) -> _PackageScan:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = tuple(archive.namelist())
            content_types = ""
            if "[Content_Types].xml" in names:
                content_types = archive.read("[Content_Types].xml").decode("utf-8", "replace")
    except zipfile.BadZipFile as exc:
        raise XlsxTemplateError(
            "模板不是有效的 xlsx（OOXML zip）包，无法读取包结构。"
            "请确认上传的是 .xlsx；.xls / .xlsb / 加密文件需要先另存为 .xlsx。",
            code="XLSX_TEMPLATE_NOT_OOXML",
        ) from exc
    return _PackageScan(
        namelist=names,
        has_vba_content_type="macroEnabled" in content_types,
    )


def _detect_unsupported_features(
    data: bytes, workbook, package: _PackageScan
) -> tuple[UnsupportedFeature, ...]:
    """按 OOXML 包部件 + openpyxl 对象模型双重探测不能保真的对象。"""

    found: dict[tuple[str, str], UnsupportedFeature] = {}

    for name in package.namelist:
        for kind, pattern in _PART_FEATURE_PATTERNS:
            if pattern.search(name):
                found.setdefault((kind, name), UnsupportedFeature(
                    kind=kind,
                    location=name,
                    detail=f"包部件 {name}；{UNSUPPORTED_FEATURE_HINTS[kind]}",
                    source="package",
                ))
                break
    if package.has_vba_content_type and not any(k == "vba" for k, _ in found):
        found[("vba", "工作簿")] = UnsupportedFeature(
            kind="vba",
            location="工作簿",
            detail="[Content_Types].xml 声明了 macroEnabled 主部件（宏工作簿）；"
                   f"{UNSUPPORTED_FEATURE_HINTS['vba']}",
            source="package",
        )

    for ws in workbook.worksheets:
        def add(kind: str, detail: str, sheet: str | None = None) -> None:
            location = sheet or ws.title
            found.setdefault((kind, location), UnsupportedFeature(
                kind=kind, location=location, detail=detail, source="object"))

        charts = list(getattr(ws, "_charts", []) or [])
        if charts:
            add("chart", f"openpyxl 读到 {len(charts)} 个图表对象；{UNSUPPORTED_FEATURE_HINTS['chart']}")
        images = list(getattr(ws, "_images", []) or [])
        if images:
            add("image", f"openpyxl 读到 {len(images)} 张图片；{UNSUPPORTED_FEATURE_HINTS['image']}")
        pivots = list(getattr(ws, "_pivots", []) or [])
        if pivots:
            add("pivot", f"openpyxl 读到 {len(pivots)} 个透视表；{UNSUPPORTED_FEATURE_HINTS['pivot']}")
        validations = list(getattr(ws.data_validations, "dataValidation", []) or [])
        if validations:
            ranges = "、".join(str(item.sqref) for item in validations[:5])
            add("data_validation",
                f"存在 {len(validations)} 条数据验证（范围 {ranges}）；"
                f"{UNSUPPORTED_FEATURE_HINTS['data_validation']}")
        conditional = list(ws.conditional_formatting)
        if conditional:
            ranges = "、".join(str(item.sqref) for item in conditional[:5])
            add("conditional_formatting",
                f"存在 {len(conditional)} 组条件格式（范围 {ranges}）；"
                f"{UNSUPPORTED_FEATURE_HINTS['conditional_formatting']}")
        comments = [cell.coordinate for row in ws.iter_rows() for cell in row if cell.comment]
        if comments:
            add("comment", f"存在 {len(comments)} 条批注（{ '、'.join(comments[:5]) }）；"
                          f"{UNSUPPORTED_FEATURE_HINTS['comment']}")

    external_links = list(getattr(workbook, "_external_links", []) or [])
    if external_links:
        found.setdefault(("external_link", "工作簿"), UnsupportedFeature(
            kind="external_link",
            location="工作簿",
            detail=f"openpyxl 读到 {len(external_links)} 条外部链接；{UNSUPPORTED_FEATURE_HINTS['external_link']}",
            source="object",
        ))
    return tuple(found.values())


def _feature_refusal_message(features: Sequence[UnsupportedFeature]) -> str:
    kinds = "、".join(dict.fromkeys(UNSUPPORTED_FEATURE_LABELS.get(f.kind, f.kind) for f in features))
    return (
        f"模板包含本模块不能保证保真的对象（{kinds}），已按 allow_lossy=False 拒绝渲染。"
        "请在模板中移除这些对象后重新登记模板，或在配置里显式设置 allow_lossy=True "
        "并接受「这些对象不会被校验、可能失真」的风险。"
    )


# --------------------------------------------------------------------------- #
# 工作簿读取与小工具
# --------------------------------------------------------------------------- #
def _load_workbook(data: bytes, *, data_only: bool = False):
    if not isinstance(data, (bytes, bytearray)):
        raise XlsxTemplateError(
            f"模板内容必须是 bytes，当前类型为 {type(data).__name__}。", code="XLSX_TEMPLATE_NOT_BYTES")
    payload = bytes(data)
    if not payload:
        raise XlsxTemplateError("模板内容为空字节，无法渲染。", code="XLSX_TEMPLATE_EMPTY")
    if not payload.startswith(b"PK"):
        raise XlsxTemplateError(
            "模板不是有效的 xlsx（OOXML zip）包：文件头不是 PK。"
            "请确认上传的是 .xlsx；.xls / .xlsb / 加密文件需要先另存为 .xlsx。",
            code="XLSX_TEMPLATE_NOT_OOXML",
        )
    try:
        return load_workbook(io.BytesIO(payload), data_only=data_only, keep_vba=False)
    except Exception as exc:  # openpyxl 的坏包异常类型不稳定，统一收口
        raise XlsxTemplateError(
            f"读取模板工作簿失败：{type(exc).__name__}: {exc}。"
            "请用 Excel 重新另存为 .xlsx 后再登记模板。",
            code="XLSX_TEMPLATE_UNREADABLE",
        ) from exc


def _header_footer_snapshot(ws) -> dict[str, tuple[str | None, str | None, str | None]]:
    snapshot: dict[str, tuple[str | None, str | None, str | None]] = {}
    for kind in ("oddHeader", "evenHeader", "firstHeader", "oddFooter", "evenFooter", "firstFooter"):
        holder = getattr(ws, kind, None)
        parts: list[str | None] = []
        for part in ("left", "center", "right"):
            item = getattr(holder, part, None) if holder is not None else None
            parts.append(getattr(item, "text", None))
        snapshot[kind] = (parts[0], parts[1], parts[2])
    return snapshot


def _normalize_ref(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1]
    return re.sub(r"[\s'$]", "", text).lower()


def _defined_name_infos(workbook) -> tuple[DefinedNameInfo, ...]:
    infos: list[DefinedNameInfo] = []
    for name in list(workbook.defined_names.keys()):
        defined = workbook.defined_names[name]
        attr_text = str(getattr(defined, "attr_text", "") or "")
        is_builtin = str(name).startswith(_BUILTIN_DEFINED_NAME_PREFIXES)
        sheet: str | None = None
        ref: str | None = None
        try:
            destinations = list(defined.destinations)
        except Exception:
            destinations = []
        if destinations:
            sheet, ref = destinations[0]
            sheet = str(sheet)
            ref = str(ref)
        infos.append(DefinedNameInfo(
            name=str(name),
            attr_text=attr_text,
            sheet=sheet,
            ref=ref,
            is_cell=bool(ref and ":" not in ref and _CELL_REF_RE.match(ref.replace("$", ""))),
            is_builtin=is_builtin,
        ))
    return tuple(infos)


def _cell_from_ref(ref: str) -> str | None:
    text = ref.replace("$", "").strip()
    return text.upper() if _CELL_REF_RE.match(text) else None


def _split_cell_ref(coordinate: str) -> tuple[int, int]:
    """'B12' → (2, 12)，即 (列号, 行号)；坐标已在配置校验阶段保证合法。"""

    row, column = coordinate_to_tuple(coordinate.upper())
    return column, row


def _write_text_cell(cell, text: str, *, escaped: bool) -> None:
    """把业务文本写进单元格；以 = + - @ 开头时强制文本，避免被当作公式执行。"""

    cell.value = text
    if escaped:
        cell.data_type = "s"      # 显式声明字符串，openpyxl 不会写成 <f>
        cell.quotePrefix = True   # Excel 视为文本，编辑时不会改成公式


def _writable_cell(ws, row: int, column: int) -> tuple[Any | None, str | None]:
    """合并单元格区域不能直接写值；返回可写单元格（合并区域左上角）与提示。"""

    cell = ws.cell(row=row, column=column)
    if not isinstance(cell, MergedCell):
        return cell, None
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= column <= rng.max_col:
            if rng.min_row == row:
                return ws.cell(row=rng.min_row, column=rng.min_col), (
                    f"列 {get_column_letter(column)} 属于合并区域 {rng.coord}，已写入合并区域左上角")
            return None, (
                f"列 {get_column_letter(column)} 位于合并区域 {rng.coord} 的非首行，无法写入业务值")
    return None, f"单元格 {get_column_letter(column)}{row} 是合并单元格但找不到合并区域"


# --------------------------------------------------------------------------- #
# 值来源：冻结 content → 逻辑键 → 值
# --------------------------------------------------------------------------- #
def _flatten_paths(node: Any, prefix: str = "", out: dict[str, Any] | None = None) -> dict[str, Any]:
    if out is None:
        out = {}
    if isinstance(node, Mapping):
        if prefix:
            out[prefix] = node
        for key, value in node.items():
            child = f"{prefix}.{key}" if prefix else str(key)
            _flatten_paths(value, child, out)
    elif isinstance(node, (list, tuple)):
        if prefix:
            out[prefix] = node
        for index, value in enumerate(node):
            _flatten_paths(value, f"{prefix}.{index}" if prefix else str(index), out)
    else:
        if prefix:
            out[prefix] = node
    return out


def document_values(frozen_content: Mapping[str, Any],
                    extra_values: Mapping[str, Any] | None = None
                    ) -> dict[str, Any]:
    """把冻结 content 摊平成「逻辑键 → 值」，并补充交付状态等派生键。"""

    values = _flatten_paths(frozen_content or {})
    requirement = (frozen_content or {}).get("requirement") or {}
    formal = (frozen_content or {}).get("formal_delivery")
    values["document.status"] = "正式交付（已审核）" if formal else "草稿（待审核交付）"
    values["document.formal_delivery"] = bool(formal)
    if formal:
        values["document.version_no"] = formal.get("version_no")
        values["document.approved_at"] = formal.get("approved_at")
        values["document.content_version"] = formal.get("content_version")
    else:
        values["document.content_version"] = (
            requirement.get("content_version") if isinstance(requirement, Mapping) else None) or "尚未建立独立内容"
        values["document.version_no"] = None
        values["document.approved_at"] = None
    values["document.digest"] = _content_digest(frozen_content)
    for alias, path in STANDARD_VALUE_ALIASES.items():
        if path in values:
            values[alias] = values[path]
    if extra_values:
        values.update({str(k): v for k, v in extra_values.items()})
    return values


def _content_digest(content: Any) -> str:
    import json

    return hashlib.sha256(
        json.dumps(content, ensure_ascii=False, sort_keys=True, default=str).encode()
    ).hexdigest()


def _row_items(frozen_content: Mapping[str, Any], source: str) -> list[Any] | None:
    node: Any = frozen_content
    for part in source.split("."):
        if isinstance(node, Mapping):
            if part not in node:
                return None
            node = node[part]
        elif isinstance(node, (list, tuple)) and part.isdigit():
            index = int(part)
            if index >= len(node):
                return None
            node = node[index]
        else:
            return None
    if node is None:
        return []
    if isinstance(node, (list, tuple)):
        return list(node)
    return None


# --------------------------------------------------------------------------- #
# analyze
# --------------------------------------------------------------------------- #
def analyze_xlsx_template(
    template_bytes: bytes,
    config: Mapping[str, Any] | XlsxTemplateConfig,
) -> XlsxTemplateAnalysis:
    """分析客户模板：sheet 顺序、命名单元格、合并区域、打印区域、页眉页脚、行块。

    同时报告配置声明但模板中缺失的项，以及不能保真的对象。纯读，不修改模板。
    """

    cfg = _coerce_config(config)
    data = bytes(template_bytes)
    digest = template_sha256(data)
    package = _scan_package(data)
    workbook = _load_workbook(data)
    warnings: list[str] = []
    missing: list[MissingItem] = []

    sheets: list[SheetInfo] = []
    merged_ranges: dict[str, tuple[str, ...]] = {}
    print_areas: dict[str, str | None] = {}
    print_titles: dict[str, tuple[str | None, str | None]] = {}
    headers_footers: dict[str, Mapping[str, tuple[str | None, str | None, str | None]]] = {}
    for ws in workbook.worksheets:
        ranges = tuple(sorted(str(rng) for rng in ws.merged_cells.ranges))
        merged_ranges[ws.title] = ranges
        print_areas[ws.title] = ws.print_area
        print_titles[ws.title] = (ws.print_title_rows, ws.print_title_cols)
        headers_footers[ws.title] = _header_footer_snapshot(ws)
        sheets.append(SheetInfo(
            title=ws.title,
            state=str(ws.sheet_state),
            max_row=int(ws.max_row or 0),
            max_column=int(ws.max_column or 0),
            merged_ranges=ranges,
            print_area=ws.print_area,
            print_title_rows=ws.print_title_rows,
            print_title_cols=ws.print_title_cols,
            header_footer=headers_footers[ws.title],
            freeze_panes=ws.freeze_panes,
            data_validation_count=len(list(getattr(ws.data_validations, "dataValidation", []) or [])),
            conditional_formatting_count=len(list(ws.conditional_formatting)),
            chart_count=len(list(getattr(ws, "_charts", []) or [])),
            image_count=len(list(getattr(ws, "_images", []) or [])),
            comment_count=sum(1 for row in ws.iter_rows() for cell in row if cell.comment),
        ))

    defined_names = _defined_name_infos(workbook)
    name_index = {info.name: info for info in defined_names if not info.is_builtin}

    # ---- 命名单元格 / 坐标映射 ----
    cell_targets: dict[str, CellTarget | None] = {}
    for key, mapping in cfg.cell_map.items():
        target: CellTarget | None = None
        if mapping.named_range:
            info = name_index.get(mapping.named_range)
            if info is None:
                missing.append(MissingItem(
                    kind="named_range", key=key,
                    detail=f"配置把「{key}」映射到命名单元格「{mapping.named_range}」，"
                           f"但模板里没有这个名称。模板现有命名单元格："
                           f"{'、'.join(sorted(name_index)) or '（无）'}",
                ))
            elif not info.is_cell or not info.ref or not info.sheet:
                missing.append(MissingItem(
                    kind="named_range", key=key, sheet=info.sheet,
                    detail=f"命名单元格「{mapping.named_range}」不是单个单元格（{info.attr_text}），"
                           f"cell_map 只接受单格命名单元格。",
                ))
            else:
                coordinate = _cell_from_ref(info.ref)
                if coordinate is None:
                    missing.append(MissingItem(
                        kind="named_range", key=key,
                        detail=f"命名单元格「{mapping.named_range}」的引用 {info.attr_text!r} 无法解析成单元格坐标。"))
                else:
                    target = CellTarget(sheet=info.sheet, coordinate=coordinate)
                    if mapping.target and (mapping.target.sheet != target.sheet
                                           or mapping.target.coordinate != target.coordinate):
                        warnings.append(
                            f"逻辑键「{key}」同时声明了命名单元格「{mapping.named_range}」和坐标 "
                            f"{mapping.target.label}，以命名单元格解析结果 {target.label} 为准。")
        elif mapping.target:
            target = mapping.target
            if target.sheet not in workbook.sheetnames:
                missing.append(MissingItem(
                    kind="sheet", key=key, sheet=target.sheet,
                    detail=f"配置把「{key}」映射到 {target.label}，但模板没有 sheet「{target.sheet}」。"
                           f"模板现有 sheet：{'、'.join(workbook.sheetnames)}",
                ))
                target = None
            else:
                ws = workbook[target.sheet]
                column_index, row_index = _split_cell_ref(target.coordinate)
                if column_index > max(int(ws.max_column or 0), 1) or \
                        row_index > max(int(ws.max_row or 0), 1):
                    warnings.append(
                        f"逻辑键「{key}」的目标 {target.label} 在模板使用区域（{ws.dimensions}）之外，"
                        f"写入后可能扩大 sheet 范围，请确认版式。")
        cell_targets[key] = target
        if target is None and not any(item.key == key for item in missing):
            missing.append(MissingItem(
                kind="cell", key=key, detail=f"逻辑键「{key}」无法解析到任何单元格目标。"))

    # ---- 行块 ----
    row_block_infos: list[RowBlockInfo] = []
    for key, block in cfg.row_blocks.items():
        if block.sheet not in workbook.sheetnames:
            missing.append(MissingItem(
                kind="sheet", key=key, sheet=block.sheet,
                detail=f"行块「{key}」指向的 sheet「{block.sheet}」不存在。"
                       f"模板现有 sheet：{'、'.join(workbook.sheetnames)}",
            ))
            continue
        ws = workbook[block.sheet]
        first = block.first_data_row
        note = ""
        resolved = True
        if first is not None and ws.max_row and first > int(ws.max_row):
            resolved = False
            note = (f"样板行 {first} 超过模板使用区域（{ws.dimensions}）；"
                    f"该行没有任何样板格式可复制。")
            missing.append(MissingItem(
                kind="row_block", key=key, sheet=block.sheet,
                detail=f"行块「{key}」的样板行 {first} 超过模板使用区域 {ws.dimensions}。",
            ))
        sheet_max_column = max(int(ws.max_column or 0), 1)
        for column_key, letter in block.columns.items():
            index = column_index_from_string(letter)
            if index > sheet_max_column:
                warnings.append(
                    f"行块「{key}」的列「{column_key}」={letter} 超出模板使用区域（{ws.dimensions}），"
                    f"该列没有样板样式可复制。")
        row_block_infos.append(RowBlockInfo(
            key=key,
            sheet=block.sheet,
            columns=dict(block.columns),
            header_row=block.header_row,
            sample_row=int(first or 0),
            sample_rows=block.sample_rows,
            first_data_row=int(first or 0),
            max_rows=block.max_rows,
            source=block.source or key,
            resolved=resolved,
            note=note,
        ))

    # ---- 必需键 ----
    for key in cfg.required_keys:
        if key not in cfg.cell_map and key not in cfg.row_blocks:
            missing.append(MissingItem(
                kind="required_key", key=key,
                detail="required_keys 声明了该键，但 cell_map / row_blocks 里都没有它。"))
        elif key in cfg.cell_map and cell_targets.get(key) is None:
            missing.append(MissingItem(
                kind="required_key", key=key,
                detail=f"必需键「{key}」无法定位到模板单元格，渲染必然缺值。"))

    # ---- 模板哈希 ----
    hash_matches = digest == cfg.template_hash
    if not hash_matches:
        missing.append(MissingItem(
            kind="template_hash", key="template_hash",
            detail=f"模板字节 sha256={digest} 与配置声明的 template_hash={cfg.template_hash} 不一致；"
                   f"模板已被改动或配置登记在别的版式上。",
        ))

    unsupported = _detect_unsupported_features(data, workbook, package)

    return XlsxTemplateAnalysis(
        template_hash=digest,
        template_hash_matches_config=hash_matches,
        config_version=cfg.config_version,
        supported_scope=cfg.supported_scope,
        sheet_names=tuple(workbook.sheetnames),
        sheets=tuple(sheets),
        defined_names=defined_names,
        merged_ranges=merged_ranges,
        print_areas=print_areas,
        print_titles=print_titles,
        headers_footers=headers_footers,
        row_blocks=tuple(row_block_infos),
        cell_targets=cell_targets,
        unsupported_features=unsupported,
        missing=tuple(missing),
        warnings=tuple(warnings),
    )


# --------------------------------------------------------------------------- #
# 渲染
# --------------------------------------------------------------------------- #
@dataclass
class _ExpectedCell:
    sheet: str
    coordinate: str
    expected: str
    logical_key: str
    escaped: bool


def _strip_illegal_chars(text: str) -> tuple[str, int]:
    cleaned = _ILLEGAL_CONTROL_RE.sub("", text)
    return cleaned, len(text) - len(cleaned)


def _normalize_cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return str(value)


def _write_business_value(
    cell,
    value: Any,
    *,
    logical_key: str,
    sheet: str,
    config: XlsxTemplateConfig,
    warnings: list[str],
) -> tuple[str, bool]:
    """写一个业务值；返回 (落库后的文本, 是否按文本转义)。"""

    if value is None:
        cell.value = None
        return "", False
    if not isinstance(value, str):
        if isinstance(value, bool) or isinstance(value, (int, float)):
            cell.value = value          # 真数值保持数值类型，不当文本
            return _normalize_cell_value(value), False
        text = str(value)
    else:
        text = value
    cleaned, removed = _strip_illegal_chars(text)
    if removed:
        warnings.append(
            f"「{logical_key}」→{sheet}!{cell.coordinate} 的文本含 {removed} 个 Excel 不允许的控制字符，"
            f"已剔除（否则无法保存成 xlsx）。")
    if len(cleaned) > MAX_EXCEL_CELL_CHARS:
        raise XlsxCellTooLongError(
            f"单元格文本超过 Excel 上限：逻辑键「{logical_key}」→ {sheet}!{cell.coordinate} "
            f"共 {len(cleaned)} 字符，上限 {MAX_EXCEL_CELL_CHARS} 字符。"
            "请拆分内容（例如长正文放到「字段口径」的独立列或附件），不要截断。",
            details=[f"超限逻辑键：{logical_key}", f"超限单元格：{sheet}!{cell.coordinate}"],
        )
    if len(cleaned) > config.max_cell_chars:
        warnings.append(
            f"「{logical_key}」→{sheet}!{cell.coordinate} 文本长度 {len(cleaned)} 字符，"
            f"超过配置阈值 max_cell_chars={config.max_cell_chars}（未截断，请关注可读性与打印分页）。")
    probe = cleaned.lstrip("\t\r\n ")
    escaped = probe[:1] in {"=", "+", "-", "@"} if probe else False
    if escaped:
        warnings.append(
            f"「{logical_key}」→{sheet}!{cell.coordinate} 的文本以 {probe[:1]!r} 开头，"
            f"已强制写成文本（data_type='s' + quotePrefix），不会被 Excel 当作公式执行。")
    _write_text_cell(cell, cleaned, escaped=escaped)
    return cleaned, escaped


def _copy_row_style(ws, source_row: int, target_row: int) -> None:
    # 注意：openpyxl 3.1.5 的 RowDimension.customHeight 是只读属性（由 ht 是否存在推导），
    # 不能赋值，只需复制 height / hidden。
    source = ws.row_dimensions.get(source_row)
    if source is not None and source.height:
        target = ws.row_dimensions[target_row]
        target.height = source.height
    if source is not None and source.hidden:
        ws.row_dimensions[target_row].hidden = True


def _copy_template_row(ws, source_row: int, target_row: int, mapped_columns: set[int],
                       max_column: int, warnings: list[str]) -> None:
    """把样板行的样式（以及未映射列上的模板常量/公式）复制到目标行。"""

    _copy_row_style(ws, source_row, target_row)
    for column in range(1, max_column + 1):
        src = ws.cell(row=source_row, column=column)
        dst = ws.cell(row=target_row, column=column)
        if isinstance(dst, MergedCell) or isinstance(src, MergedCell):
            if column not in mapped_columns and not isinstance(dst, MergedCell) and src.has_style:
                dst._style = copy.copy(src._style)
            continue
        if src.has_style:
            dst._style = copy.copy(src._style)
        if column in mapped_columns:
            continue  # 该列稍后由业务值覆盖
        if isinstance(src.value, str) and src.data_type == "f" and src.value.startswith("="):
            origin = f"{get_column_letter(column)}{source_row}"
            destination = f"{get_column_letter(column)}{target_row}"
            try:
                dst.value = Translator(src.value, origin=origin).translate_formula(destination)
            except Exception as exc:  # 公式无法翻译时明确告警，不静默丢公式
                warnings.append(
                    f"样板行公式 {origin} 无法翻译到 {destination}（{type(exc).__name__}: {exc}），"
                    f"该单元格保持模板原样，请人工确认。")
        elif src.value is not None:
            dst.value = src.value


def _shift_rows_meta(ws, insert_at: int, delta: int, warnings: list[str]) -> None:
    """openpyxl 的 insert_rows 不移动合并区域/行高，这里补齐；其余对象告警。"""

    # 不能调用 ws.unmerge_cells(old)：insert_rows 已经把该区域的非锚点 MergedCell 下移，
    # unmerge_cells 会 del 已不存在的 _cells[(row, col)] 而抛 KeyError。
    # 因此直接改写 merged_cells 集合：丢弃旧范围，登记平移后的新范围（单元格已被下移，无需重建）。
    for rng in [r for r in ws.merged_cells.ranges if r.min_row >= insert_at]:
        old = str(rng)
        new = CellRange(
            min_col=rng.min_col, min_row=rng.min_row + delta,
            max_col=rng.max_col, max_row=rng.max_row + delta,
        )
        ws.merged_cells.ranges.discard(rng)
        ws.merged_cells.ranges.add(new)
        for row, column in new.cells:
            if (row, column) == (new.min_row, new.min_col):
                continue  # 锚点单元格保留原值/样式
            existing = ws._cells.get((row, column))
            if not isinstance(existing, MergedCell):
                ws._cells[(row, column)] = MergedCell(ws, row=row, column=column)
        warnings.append(
            f"插入行后合并区域 {old} 已手工平移到 {new}（openpyxl 不会自动移动合并区域）。")
    heights = {
        index: dimension
        for index, dimension in ws.row_dimensions.items()
        if index >= insert_at and (dimension.height or dimension.hidden)
    }
    for index in sorted(heights):
        ws.row_dimensions.pop(index, None)
    for index, dimension in heights.items():
        target = ws.row_dimensions[index + delta]
        target.height = dimension.height
        target.hidden = dimension.hidden
    if heights:
        warnings.append(
            f"插入 {delta} 行后已平移 {len(heights)} 个行高设置（openpyxl 不会自动移动行高）。")
    if len(list(getattr(ws.data_validations, "dataValidation", []) or [])) or \
            len(list(ws.conditional_formatting)):
        warnings.append(
            "插入行后数据验证/条件格式的范围没有随之平移，请人工修正这两类规则的范围。")
    if list(getattr(ws, "_charts", []) or []) or list(getattr(ws, "_images", []) or []):
        warnings.append("插入行后图表/图片的锚点没有随之移动，请人工核对位置。")


def _resolve_row_paths(block: RowBlock) -> dict[str, str]:
    return {key: block.column_paths.get(key, key) for key in block.columns}


def render_xlsx(
    template_bytes: bytes,
    frozen_content: Mapping[str, Any],
    config: Mapping[str, Any] | XlsxTemplateConfig,
    *,
    extra_values: Mapping[str, Any] | None = None,
) -> tuple[bytes, RenderDiagnostics]:
    """按模板填充，返回 (输出字节, 诊断)。校验不通过时抛异常且不产出文件。"""

    cfg = _coerce_config(config)
    data = bytes(template_bytes)
    analysis = analyze_xlsx_template(data, cfg)
    warnings: list[str] = list(analysis.warnings)

    if not analysis.template_hash_matches_config:
        raise XlsxTemplateHashMismatch(
            f"模板哈希不匹配，拒绝渲染：系统登记的是 template_hash={cfg.template_hash}，"
            f"本次传入的模板字节 sha256={analysis.template_hash}。"
            "请重新登记模板配置，或换回登记的模板文件。",
            details=[item.message for item in analysis.missing if item.kind == "template_hash"],
        )
    if analysis.blocking_missing:
        raise XlsxMappingMissingError(
            "模板映射不可用，拒绝渲染（未产出文件）。请按下列提示修正配置或模板：",
            details=[item.message for item in analysis.blocking_missing],
        )
    if analysis.unsupported_features:
        if not cfg.allow_lossy:
            raise XlsxUnsupportedFeatureError(
                _feature_refusal_message(analysis.unsupported_features),
                details=[feature.message for feature in analysis.unsupported_features],
            )
        for feature in analysis.unsupported_features:
            warnings.append(
                f"[allow_lossy] 模板含{feature.label}（{feature.location}）：本模块不会校验其保真性，"
                f"渲染结果可能失真。{feature.detail}")

    workbook = _load_workbook(data)
    values = document_values(frozen_content, extra_values)
    expected: list[_ExpectedCell] = []
    filled: list[FilledCell] = []
    overwritten: set[tuple[str, str]] = set()
    template_formulas = _formula_snapshot(workbook)

    # ---- 命名单元格 ----
    for key, mapping in cfg.cell_map.items():
        target = analysis.cell_targets.get(key)
        if target is None:  # analyze 已把缺失项放进 blocking_missing，这里只是兜底
            raise XlsxMappingMissingError(
                f"必需映射缺失：逻辑键「{key}」找不到模板目标（{mapping.description}）。")
        ws = workbook[target.sheet]
        value = _lookup_value(values, mapping.value_path or key)
        column_index, row_index = _split_cell_ref(target.coordinate)
        cell, note = _writable_cell(ws, row_index, column_index)
        if cell is None:
            raise XlsxMappingMissingError(
                f"逻辑键「{key}」的目标 {target.label} 不可写入：{note}")
        if note:
            warnings.append(f"「{key}」→{target.label}：{note}")
        text, escaped = _write_business_value(
            cell, value, logical_key=key, sheet=target.sheet, config=cfg, warnings=warnings)
        overwritten.add((target.sheet, cell.coordinate))
        expected.append(_ExpectedCell(target.sheet, cell.coordinate, text, key, escaped))
        filled.append(FilledCell(key, target.sheet, cell.coordinate, text, escaped))

    # ---- 行块 ----
    row_block_counts: dict[str, int] = {}
    shifts: dict[str, list[tuple[int, int]]] = {}
    for key, block in cfg.row_blocks.items():
        ws = workbook[block.sheet]
        source = block.source or key
        items = _row_items(frozen_content, source)
        if items is None:
            raise XlsxMappingMissingError(
                f"行块「{key}」的数据来源「{source}」在冻结内容里不是列表，无法展开成多行。"
                f"请检查 row_blocks[{key!r}].source。")
        if len(items) > block.max_rows:
            raise XlsxTemplateError(
                f"行块「{key}」要写 {len(items)} 行，超过配置上限 max_rows={block.max_rows}。"
                f"请调大 max_rows 或确认模板是否该拆页。",
                code="XLSX_ROW_BLOCK_OVERFLOW",
                details=[f"sheet：{block.sheet}", f"数据来源：{source}"],
            )
        written = _fill_row_block(ws, block, items, cfg, warnings, expected, overwritten, shifts)
        row_block_counts[key] = written
        if block.key in cfg.required_keys and written == 0:
            raise XlsxMappingMissingError(
                f"必需行块「{key}」没有数据：冻结内容 {source} 为空，渲染会产出空表，已拒绝。"
                f"如允许空表请在 required_keys 中移除该键。")

    # ---- 必需键取值校验 ----
    required = set(cfg.required_keys) | {k for k, m in cfg.cell_map.items() if m.required} | \
        {k for k, b in cfg.row_blocks.items() if b.required}
    for record in expected:
        if record.logical_key in required and not record.expected.strip():
            raise XlsxMappingMissingError(
                f"必需键「{record.logical_key}」没有值：目标 {record.sheet}!{record.coordinate} 的取值为空。"
                f"请补齐冻结内容或从 required_keys 中移除该键。")

    stream = io.BytesIO()
    workbook.save(stream)
    output = stream.getvalue()

    # ---- 读回校验（data_only=False，确认公式仍是公式）----
    issues = _validate_output(
        output,
        analysis=analysis,
        expected=expected,
        template_formulas=_translate_formula_coordinates(template_formulas, shifts),
        overwritten=overwritten,
        config=cfg,
        row_block_counts=row_block_counts,
    )
    if issues:
        raise XlsxRenderValidationError(
            "渲染结果读回校验失败，已丢弃输出（未产出文件）。请把下列问题反馈给模板适配人员：",
            details=issues,
        )

    tool = _recalculation_tool()
    recalculation_warning = (
        "环境中未发现 LibreOffice（soffice），公式重算未验证；请用 Excel 打开输出文件确认公式结果。"
        if tool is None
        else "环境中发现 LibreOffice，但本模块不会调用它重算公式；公式重算仍需人工执行并验证。"
    )
    warnings.append(recalculation_warning)
    diagnostics = RenderDiagnostics(
        template_hash=analysis.template_hash,
        config_version=cfg.config_version,
        supported_scope=cfg.supported_scope,
        output_sha256=hashlib.sha256(output).hexdigest(),
        filled_cells=tuple(filled),
        row_block_counts=dict(row_block_counts),
        warnings=tuple(warnings),
        unsupported_features=analysis.unsupported_features,
        lossy_rendered=bool(analysis.unsupported_features and cfg.allow_lossy),
        recalculation_tool_available=tool,
        recalculation_note=recalculation_warning,
    )
    return output, diagnostics


def _lookup_value(values: Mapping[str, Any], path: str) -> Any:
    if path in values:
        return values[path]
    node: Any = values
    for part in path.split("."):
        if isinstance(node, Mapping) and part in node:
            node = node[part]
        else:
            return None
    return node


def _fill_row_block(
    ws,
    block: RowBlock,
    items: Sequence[Any],
    config: XlsxTemplateConfig,
    warnings: list[str],
    expected: list[_ExpectedCell],
    overwritten: set[tuple[str, str]],
    shifts: dict[str, list[tuple[int, int]]],
) -> int:
    first = block.first_data_row
    if first is None:
        raise XlsxMappingMissingError(f"行块「{block.key}」无法确定样板行行号，拒绝渲染。")
    span = block.sample_rows
    needed = len(items)
    max_column = max(int(ws.max_column or 0), 1, max(column_index_from_string(letter)
                                                for letter in block.columns.values()))
    mapped_columns = {column_index_from_string(letter) for letter in block.columns.values()}
    paths = _resolve_row_paths(block)

    if needed > span:
        delta = needed - span
        insert_at = first + span
        ws.insert_rows(insert_at, delta)
        _shift_rows_meta(ws, insert_at, delta, warnings)
        shifts.setdefault(block.sheet, []).append((insert_at, delta))

    if needed < span:
        for row in range(first + needed, first + span):
            for letter in block.columns.values():
                cell, _ = _writable_cell(ws, row, column_index_from_string(letter))
                if cell is not None:
                    cell.value = None
                    overwritten.add((block.sheet, cell.coordinate))
            warnings.append(
                f"行块「{block.key}」模板有 {span} 个样板行，实际只有 {needed} 行数据，"
                f"第 {row} 行的映射列已清空（样式保留）。")

    missing_value_keys: set[str] = set()
    for index, item in enumerate(items):
        target_row = first + index
        if index > 0:
            _copy_template_row(ws, first, target_row, mapped_columns, max_column, warnings)
        flat = _flatten_paths(item)
        for column_key, letter in block.columns.items():
            column = column_index_from_string(letter)
            cell, note = _writable_cell(ws, target_row, column)
            if cell is None:
                if note:
                    warnings.append(f"行块「{block.key}」第 {index + 1} 条数据：{note}")
                continue
            if note:
                warnings.append(f"行块「{block.key}」第 {index + 1} 条数据：{note}")
            value = _lookup_value(flat, paths[column_key])
            if value is None and isinstance(item, Mapping) and column_key in item:
                value = item[column_key]
            if value in (None, "") and column_key not in missing_value_keys:
                missing_value_keys.add(column_key)
                warnings.append(
                    f"行块「{block.key}」的列「{column_key}」在冻结内容里取不到值"
                    f"（列路径 {paths[column_key]!r}），已写空。")
            text, escaped = _write_business_value(
                cell, value, logical_key=f"{block.key}.{column_key}",
                sheet=block.sheet, config=config, warnings=warnings)
            overwritten.add((block.sheet, cell.coordinate))
            expected.append(_ExpectedCell(block.sheet, cell.coordinate, text,
                                          f"{block.key}.{column_key}", escaped))
    return needed


def _formula_snapshot(workbook) -> dict[tuple[str, str], str]:
    formulas: dict[tuple[str, str], str] = {}
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas[(ws.title, cell.coordinate)] = cell.value
    return formulas


def _translate_formula_coordinates(
    formulas: Mapping[tuple[str, str], str],
    shifts: Mapping[str, Sequence[tuple[int, int]]],
) -> dict[tuple[str, str], str]:
    """把插行前的公式坐标平移到插行后的行号，用于读回校验。"""

    translated: dict[tuple[str, str], str] = {}
    for (sheet, coordinate), formula in formulas.items():
        column_index, row_index = _split_cell_ref(coordinate)
        for insert_at, delta in shifts.get(sheet, ()):
            if row_index >= insert_at:
                row_index += delta
        translated[(sheet, f"{get_column_letter(column_index)}{row_index}")] = formula
    return translated


def _validate_output(
    output: bytes,
    *,
    analysis: XlsxTemplateAnalysis,
    expected: Sequence[_ExpectedCell],
    template_formulas: Mapping[tuple[str, str], str],
    overwritten: set[tuple[str, str]],
    config: XlsxTemplateConfig,
    row_block_counts: Mapping[str, int],
) -> list[str]:
    """用 data_only=False 读回输出，逐项核对保真与填充结果。"""

    issues: list[str] = []
    try:
        wb = _load_workbook(output, data_only=False)
    except XlsxTemplateError as exc:
        return [f"输出文件无法被 openpyxl 重新打开：{exc}"]

    if tuple(wb.sheetnames) != analysis.sheet_names:
        issues.append(
            f"sheet 顺序被改变：模板为 {'、'.join(analysis.sheet_names)}，"
            f"输出为 {'、'.join(wb.sheetnames)}。")

    for info in analysis.sheets:
        if info.title not in wb.sheetnames:
            issues.append(f"输出缺少 sheet「{info.title}」。")
            continue
        ws = wb[info.title]
        output_ranges = sorted(str(rng) for rng in ws.merged_cells.ranges)
        if len(output_ranges) < len(info.merged_ranges):
            issues.append(
                f"合并区域减少：sheet「{info.title}」模板 {len(info.merged_ranges)} 个，"
                f"输出 {len(output_ranges)} 个。")
        if _normalize_ref(ws.print_area) != _normalize_ref(info.print_area):
            issues.append(
                f"打印区域不一致：sheet「{info.title}」模板 {info.print_area!r}，输出 {ws.print_area!r}。")
        if (_normalize_ref(ws.print_title_rows) != _normalize_ref(info.print_title_rows)
                or _normalize_ref(ws.print_title_cols) != _normalize_ref(info.print_title_cols)):
            issues.append(
                f"打印标题（重复行/列）不一致：sheet「{info.title}」模板 "
                f"{info.print_title_rows}/{info.print_title_cols}，"
                f"输出 {ws.print_title_rows}/{ws.print_title_cols}。")
        output_hf = _header_footer_snapshot(ws)
        if output_hf != info.header_footer:
            issues.append(f"页眉页脚不一致：sheet「{info.title}」。")

    # 公式仍在
    for (sheet, coordinate), formula in template_formulas.items():
        if (sheet, coordinate) in overwritten:
            continue
        if sheet not in wb.sheetnames:
            continue
        cell = wb[sheet][coordinate]
        if cell.data_type != "f" or not (isinstance(cell.value, str) and cell.value.startswith("=")):
            issues.append(
                f"模板公式丢失：{sheet}!{coordinate} 原为公式 {formula!r}，"
                f"输出为 {cell.value!r}（data_type={cell.data_type}），公式可能被写成了值。")

    # 值读回
    for record in expected:
        if record.sheet not in wb.sheetnames:
            issues.append(f"取值校验失败：输出缺少 sheet「{record.sheet}」。")
            continue
        cell = wb[record.sheet][record.coordinate]
        actual = _normalize_cell_value(cell.value)
        if actual != record.expected:
            preview = lambda text: (text[:40] + "…") if len(text) > 40 else text
            issues.append(
                f"取值回读不一致：{record.sheet}!{record.coordinate}"
                f"（逻辑键「{record.logical_key}」）期望 {preview(record.expected)!r}，"
                f"实际 {preview(actual)!r}。")
        elif record.escaped and cell.data_type != "s":
            issues.append(
                f"公式注入防护失效：{record.sheet}!{record.coordinate}"
                f" 的文本以 =/+/-/@ 开头，但 data_type={cell.data_type}，可能被 Excel 当公式执行。")

    # 残留占位符
    patterns = [re.compile(p) for p in config.placeholder_patterns]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for pattern in patterns:
                        if pattern.search(cell.value):
                            issues.append(
                                f"存在残留占位符：{ws.title}!{cell.coordinate} 的值 "
                                f"{cell.value[:60]!r} 命中占位符形态 {pattern.pattern!r}。")
                            break
    return issues


def _recalculation_tool() -> str | None:
    for candidate in ("soffice", "libreoffice"):
        if shutil.which(candidate):
            return "libreoffice"
    return None
