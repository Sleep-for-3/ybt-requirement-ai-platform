"""导出产物与监管标准模板的取值比对（round 58）。

解决的问题：结构自检只能证明"这份文件长得像标准模板"，不能证明"文件里的模板列与模板一致"。
调研表的监管原始口径、报表名称、字段名称等是业务方给定的标准值，导出时必须一字不差；
一旦需求冻结后有人改过正文、或模板重新应用过，文件里就会出现与当前模板不一致的值。
本模块把这件事做成可执行比对：逐字段、逐模板列核对，并**只报告差异，不改任何内容**。
"""

from __future__ import annotations

import re
from collections.abc import Mapping
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.models import TargetField
from app.services.requirement_render_artifact import RenderArtifactError, read_artifact_bytes

FIDELITY_VERSION = "requirement-template-fidelity-v1"

#: 导出文档里能追溯到标准模板的列：(列序号(1-based), 显示名, TargetField 属性名)。
#:
#: 只比对模板真正提供的列：数据类别/数据格式/监管原始口径/监管定义细化/报表名称/字段名称/
#: EAST 口径/行内定义/备注。业务口径、溯源与确认人来自需求侧或脚本，不属于模板值。
FIDELITY_COLUMNS: tuple[tuple[int, str, str], ...] = (
    (3, "数据类别", "data_category"),
    (4, "数据格式", "data_format"),
    (5, "监管原始口径", "regulatory_original_definition"),
    (6, "监管定义细化", "regulatory_refined_definition"),
    (7, "报表名称", "report_name"),
    (8, "字段名称", "report_field_name"),
    (9, "EAST口径", "east_definition"),
    (10, "行内定义", "internal_definition"),
    (11, "备注", "remarks"),
)

MAX_SAMPLES = 5
_FIELD_CODE = re.compile(r"^[A-Za-z][A-Za-z0-9_]{1,31}$")


class FidelityError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def _normalize(value: Any) -> str:
    """比对口径：去掉 Excel 里的人工换行与多余空白，其余必须一字不差。"""

    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n")
    return re.sub(r"\s+", " ", text).strip()


def _rows(payload: bytes) -> tuple[dict[str, dict[int, str]], str]:
    try:
        workbook = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - 任何解析异常都归为"不是一表通调研表式"
        # Word 产物同样是 zip，openpyxl 会以 "File contains no valid workbook part" 失败；
        # 对用户来说这不是"文件坏了"，而是"这份产物没有可比的模板列"。
        raise FidelityError(
            "artifact_not_survey",
            "这份产物不是一表通调研表式（xlsx），没有可与标准模板比对的模板列；"
            f"Word 产物请用内置版式人工核对或先做结构自检。（原始错误：{exc}）",
        ) from exc
    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    header_index = -1
    for index, row in enumerate(rows):
        cells = [_normalize(cell) for cell in row or ()]
        if "数据项编码" in cells:
            header_index = index
            break
    if header_index < 0:
        raise FidelityError(
            "artifact_not_survey",
            "这份产物里没有「数据项编码」表头，不是一表通调研表式，无法与标准模板比对。",
        )
    data: dict[str, dict[int, str]] = {}
    for row in rows[header_index + 1 :]:
        values = list(row or ())
        code = _normalize(values[0] if values else "")
        if not _FIELD_CODE.match(code):
            # 数据行之外的完整度/来源说明区块按代码格式排除，避免把统计行当字段。
            continue
        data[code] = {column: _normalize(values[column - 1] if len(values) >= column else "") for column, _, _ in FIDELITY_COLUMNS}
    return data, _normalize(sheet.title)


def check_template_fidelity(
    db: Any,
    *,
    project_id: int,
    artifact_id: int,
    requirement_id: int | None = None,
) -> dict[str, Any]:
    """比对产物里的模板列与当前项目模板值。只读：不写库、不改文件。"""

    from app.models import Requirement, RequirementRenderArtifact, TargetField  # 局部导入，避免循环依赖

    artifact = db.get(RequirementRenderArtifact, artifact_id)
    if artifact is None or artifact.project_id != project_id:
        raise FidelityError("artifact_missing", "这份产物不存在或不属于当前项目。")
    if requirement_id is not None and artifact.requirement_id != requirement_id:
        raise FidelityError("artifact_missing", "这份产物不属于当前需求。")
    try:
        payload = read_artifact_bytes(artifact)
    except RenderArtifactError as exc:
        raise FidelityError(exc.code, exc.message) from exc
    rows, sheet_name = _rows(payload)

    # 字段代码不是全局唯一的：SFJJGDM 这类系统字段在 90 张表里都有同名行。
    # 只按 field_code 取模板值会拿到别的表的值，把正确文档判成“不一致”——
    # 必须限定在需求自己的目标表范围内。
    requirement = db.get(Requirement, artifact.requirement_id)
    target_table_id = (requirement.scope_json or {}).get("target_table_id") if requirement else None
    query = db.query(TargetField).filter(TargetField.project_id == project_id)
    if target_table_id:
        query = query.filter(TargetField.target_table_id == target_table_id)
    fields = {row.field_code: row for row in query.all()}
    columns: list[dict[str, Any]] = []
    missing_in_document: list[str] = []
    extra_in_document: list[str] = []
    unexpected_values: list[dict[str, str]] = []
    for order, label, attribute in FIDELITY_COLUMNS:
        matched = 0
        mismatched: list[dict[str, str]] = []
        mismatch_count = 0
        template_empty = 0
        blank_in_document = 0
        for code, cells in sorted(rows.items()):
            field = fields.get(code)
            if field is None:
                if code not in extra_in_document:
                    extra_in_document.append(code)
                continue
            expected = _normalize(getattr(field, attribute, ""))
            actual = cells.get(order, "")
            if not expected:
                # 模板本身没有值：不能判"不一致"，只能记"模板没给"。文档里却写了内容时，
                # 单独列为"模板无值但文档有值"，让人知道这是需求侧补充而非模板事实。
                template_empty += 1
                if actual and len(unexpected_values) < MAX_SAMPLES * len(FIDELITY_COLUMNS):
                    unexpected_values.append(
                        {"label": label, "field_code": code, "actual": actual[:120]}
                    )
                continue
            if not actual:
                blank_in_document += 1
                missing_in_document.append(code)
                continue
            if actual == expected:
                matched += 1
                continue
            mismatch_count += 1
            if len(mismatched) < MAX_SAMPLES:
                mismatched.append({"field_code": code, "expected": expected[:200], "actual": actual[:200]})
        checked = len(rows) - template_empty
        columns.append(
            {
                "key": attribute,
                "label": label,
                "checked_count": max(checked, 0),
                "matched_count": matched,
                "mismatched_count": mismatch_count,
                "blank_in_document_count": blank_in_document,
                "template_empty_count": template_empty,
                "samples": mismatched,
            }
        )
    blocking = [
        column
        for column in columns
        if column["mismatched_count"] or column["blank_in_document_count"]
    ]
    return {
        "fidelity_version": FIDELITY_VERSION,
        "artifact": {
            "id": artifact.id,
            "format": artifact.format,
            "file_name": artifact.file_name,
            "file_size": artifact.file_size,
            "file_sha256": artifact.file_sha256,
            "renderer_version": artifact.renderer_version,
        },
        "sheet_name": sheet_name,
        "field_count": len(rows),
        "columns": columns,
        "missing_in_document": sorted(set(missing_in_document)),
        "extra_in_document": extra_in_document,
        "unexpected_values": unexpected_values,
        "verdict": "fail" if blocking else "pass",
        "failed_labels": [column["label"] for column in blocking],
        "note": "只比对模板提供的列（监管原始口径/报表名称等），业务口径与溯源列来自需求与脚本，不在此列；比对只读，不修改文件或需求。",
    }
