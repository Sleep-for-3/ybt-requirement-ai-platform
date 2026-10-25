from io import BytesIO
from pathlib import Path
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from app.services.sql_parser import parse_sql
from app.services.parse_diagnostics import (
    CELL_FORMULA_UNCACHED,
    SEVERITY_BLOCKING,
    SEVERITY_INFO,
    SEVERITY_WARNING,
    ParseFinding,
    ParseReport,
    cell_reference,
    scan_formulas,
)
from app.core.settings import get_settings
from .base import KnowledgeUnitDraft
from .ocr import IMAGE_SUFFIXES

# A clause number written into the unit's own text, mirroring
# ``agent.tools.evidence.INSCRIBED_CLAUSE_PATTERN``: the ingestion side records what
# the citation side will accept, and neither may drift.
CLAUSE_NUMBER_PATTERN = re.compile(r"^(?:第[零一二三四五六七八九十百千\d]{1,12}条|\d+(?:\.\d+){1,4})(?![\d.])")
CHAPTER_NUMBER_PATTERN = re.compile(r"^(第[零一二三四五六七八九十百千\d]{1,12}章)")
QUESTION={"问题","答疑问题","监管问题","银行问题"};ANSWER={"回复","监管回复","答疑结果","处理意见"};SUGGESTION={"同业建议"};FIELD_CODE={"字段代码","数据项编码"};FIELD_NAME={"字段名称","字段中文名","数据项名称"};TABLE={"一表通表","表代码"};SCENARIO={"业务场景","产品场景","场景名称"};BUSINESS_SYSTEM={"来源系统","业务系统","系统名称"};SOURCE_TABLE={"来源表","来源表英文名","源表","源表英文名"};SOURCE_FIELD={"来源字段","来源字段英文名","源字段","源字段英文名"};MART_TABLE={"集市表","监管集市表"};MART_FIELD={"集市字段","监管集市字段"}
HEADER_MARKERS=QUESTION|ANSWER|FIELD_CODE|FIELD_NAME
ALL_ALIASES=(QUESTION|ANSWER|SUGGESTION|FIELD_CODE|FIELD_NAME|TABLE|SCENARIO|BUSINESS_SYSTEM
             |SOURCE_TABLE|SOURCE_FIELD|MART_TABLE|MART_FIELD)

# Versioned parsing rules.  Changing any value here requires bumping the version:
# a frozen requirement keeps pointing at the rules that produced its snapshot, and
# a configuration upgrade must never rewrite facts that are already frozen.
XLSX_PARSE_CONFIG_VERSION="xlsx-knowledge-config-v1"
XLSX_PARSE_CONFIG={
    # Historical code scanned the first 50 rows only; a header on row 51 was an
    # invisible whole-sheet loss.  The window is a documented, versioned setting.
    "header_scan_rows":200,
    "legacy_header_scan_rows":50,
    "column_aliases":sorted(ALL_ALIASES),
}

def parse_document(file_name,content,knowledge_type):
    units, warnings, _report = parse_document_with_report(file_name, content, knowledge_type)
    return units, warnings


def parse_document_with_report(file_name,content,knowledge_type):
    """Parse a document and return the loss-visible report alongside the units.

    The report is ``None`` for formats whose reader cannot lose cell-level data in
    a way this module can describe (docx/pdf/text/sql); XLSX always returns one.
    """
    units, warnings, report = _parse_document(file_name, content, knowledge_type)
    for index, unit in enumerate(units, 1):
        locator = unit.metadata.setdefault("locator", {})
        locator.setdefault("block_id", f"parsed-block-{index}")
        for key, value in (("page_no", unit.source_page_no), ("sheet_name", unit.source_sheet_name),
                           ("cell_range", unit.source_cell_range), ("heading", unit.source_heading),
                           ("row_no", unit.metadata.get("row")), ("paragraph_index", unit.metadata.get("paragraph")),
                           ("chapter_no", unit.metadata.get("chapter_no")), ("clause_no", unit.metadata.get("clause_no"))):
            if value is not None:
                locator.setdefault(key, value)
    return units, warnings, report


def _text_units(text, file_name, unit_type="paragraph"):
    units = []
    for index, match in enumerate(re.finditer(r"[^\n]+(?:\n(?!\n)[^\n]+)*", text), 1):
        if not match.group().strip():
            continue
        units.append(KnowledgeUnitDraft(unit_type, file_name, match.group(), metadata={"locator": {
            "paragraph_index": index, "char_start": match.start(), "char_end": match.end(),
            "line_start": text.count("\n", 0, match.start()) + 1,
            "line_end": text.count("\n", 0, match.end()) + 1,
        }}))
    return units


def _parse_document(file_name,content,knowledge_type):
    suffix=Path(file_name).suffix.lower()
    SUPPORTED_SUFFIXES = {".xlsx", ".docx", ".pdf", ".txt", ".md", ".sql"} | IMAGE_SUFFIXES
    if suffix not in SUPPORTED_SUFFIXES:raise ValueError(f"Unsupported knowledge document format: {suffix or 'missing extension'}")
    if suffix==".xlsx":return _excel(content,knowledge_type)
    if suffix==".docx":units,warnings=_docx(content);return units,warnings,None
    if suffix==".pdf":units,warnings=_pdf(content);return units,warnings,None
    if suffix in IMAGE_SUFFIXES:
        # Images are only a knowledge document once OCR can read them; otherwise the
        # historical "unsupported format" contract is preserved byte-for-byte.
        if not getattr(get_settings(),"ocr_enabled",False):raise ValueError(f"Unsupported knowledge document format: {suffix or 'missing extension'}")
        units,warnings=_image(content,suffix);return units,warnings,None
    text=content.decode("utf-8",errors="replace").replace("\r\n","\n").replace("\r","\n")
    if suffix==".sql":
        parsed=parse_sql(text);summary=f"来源表: {', '.join(parsed.source_tables)}\n来源字段: {', '.join(parsed.selected_fields)}\n关联条件: {'; '.join(parsed.joins)}\n过滤条件: {'; '.join(parsed.where_conditions)}"
        return [KnowledgeUnitDraft("sql_summary",file_name,summary,metadata={"raw_sql":text,"parsed_success":parsed.parsed_success,"locator":{"line_start":1,"line_end":len(text.splitlines()) or 1}}), *_text_units(text,file_name,"sql_text")],[parsed.error_message] if parsed.error_message else [],None
    return _text_units(text,file_name),[],None

def _cell_text(value):
    if value is None:return ""
    if isinstance(value,str):return value.strip()
    return str(value).strip()


def _first_text(values,keys):
    """First mapped value that is actually present; ``0`` is a value, not absence."""
    for key in keys:
        if key not in values:continue
        text=_cell_text(values[key])
        if text:return text
    return None


def _record_header_findings(sheet,headers,header_row,report):
    seen={};duplicates=[]
    for column,name in headers.items():
        if not name:continue
        if name in seen:duplicates.append({"sheet":sheet.title,"row":header_row,"name":name,"columns":[seen[name],column]})
        else:seen[name]=column
    for item in duplicates:
        report.duplicate_headers.append(item)
        report.add(ParseFinding("duplicate_header",SEVERITY_WARNING,
                                f"{sheet.title} 第 {header_row} 行重复列名“{item['name']}”，后出现的列会覆盖先出现的列",
                                locator={"sheet":sheet.title,"row":header_row,"column":get_column_letter(item["columns"][1])}))
    unmapped=sorted({name for name in headers.values() if name and name not in ALL_ALIASES})
    if unmapped:
        report.columns_unrecognized.append({"sheet":sheet.title,"row":header_row,"columns":unmapped})
        report.add(ParseFinding("unrecognized_column",SEVERITY_INFO,
                                f"{sheet.title} 有 {len(unmapped)} 个未映射列：{'、'.join(unmapped[:10])}",
                                locator={"sheet":sheet.title,"row":header_row},
                                suggestion="列内容仍写入单元文本，但不作为结构化字段；如需纳入口径请补齐列别名。"))
    merged=[item for item in sheet.merged_cells.ranges if item.min_row<=header_row<=item.max_row]
    if merged:
        report.merged_header_cells+=len(merged)
        report.add(ParseFinding("merged_header",SEVERITY_INFO,
                                f"{sheet.title} 表头含 {len(merged)} 个合并单元格，其值已展开到合并区域",
                                locator={"sheet":sheet.title,"row":header_row}))


def _excel(content,knowledge_type):
    values_workbook=load_workbook(BytesIO(content),data_only=True)
    # A second, formula-view workbook is what makes an uncached formula visible:
    # data_only=True silently returns None for a formula Excel never recalculated.
    formula_workbook=load_workbook(BytesIO(content),data_only=False)
    report=ParseReport(config_version=XLSX_PARSE_CONFIG_VERSION,reader="knowledge-ingestion",
                       file_suffix=".xlsx",sheet_total=len(values_workbook.worksheets))
    units=[]
    for sheet in values_workbook.worksheets:
        matrix=_expanded_matrix(sheet)
        formula_sheet=formula_workbook[sheet.title] if sheet.title in formula_workbook.sheetnames else None
        def _cached(row,column,_matrix=matrix):
            if row<=len(_matrix) and column<=len(_matrix[row-1]):return _matrix[row-1][column-1]
            return None
        kinds=scan_formulas(sheet.title,formula_sheet,_cached,report)
        header_row=None;headers={}
        scan_rows=min(len(matrix),XLSX_PARSE_CONFIG["header_scan_rows"])
        for row in range(1,scan_rows+1):
            values={column:_cell_text(matrix[row-1][column-1]) for column in range(1,sheet.max_column+1)}
            if any(value in HEADER_MARKERS for value in values.values()):header_row=row;headers=values;break
        if not header_row:
            report.sheets_skipped.append({"sheet":sheet.title,"row_total":sheet.max_row})
            lost=any(_cell_text(cell) for line in matrix for cell in line)
            report.add(ParseFinding("sheet_skipped",SEVERITY_BLOCKING if lost else SEVERITY_INFO,
                                    f"{sheet.title} 未识别表头，整张 sheet 未导入",
                                    locator={"sheet":sheet.title,"scanned_rows":scan_rows}))
            continue
        report.sheet_parsed+=1
        report.header_rows.append({"sheet":sheet.title,"row":header_row})
        if header_row>XLSX_PARSE_CONFIG["legacy_header_scan_rows"]:
            report.add(ParseFinding("header_beyond_legacy_scan",SEVERITY_INFO,
                                    f"{sheet.title} 表头在第 {header_row} 行，超出历史扫描范围（前 {XLSX_PARSE_CONFIG['legacy_header_scan_rows']} 行）",
                                    locator={"sheet":sheet.title,"row":header_row}))
        _record_header_findings(sheet,headers,header_row,report)
        for (row_index,column_index),kind in kinds.items():
            if row_index==header_row and kind==CELL_FORMULA_UNCACHED:
                report.add(ParseFinding("formula_in_header",SEVERITY_BLOCKING,
                                        f"{sheet.title} 第 {row_index} 行表头含无缓存公式，表头识别可能失败",
                                        locator={"sheet":sheet.title,"cell":cell_reference(row_index,column_index)}))
        uncached_rows={}
        for (row_index,column_index),kind in kinds.items():
            if kind==CELL_FORMULA_UNCACHED:uncached_rows.setdefault(row_index,[]).append(column_index)
        for row in range(header_row+1,sheet.max_row+1):
            report.row_total+=1
            values={headers[column]:matrix[row-1][column-1] for column in headers if headers[column]}
            missing_cells=uncached_rows.get(row,[])
            if not any(value not in (None,"") for value in values.values()):
                report.row_skipped_blank+=1
                if missing_cells:
                    report.add(ParseFinding("formula_cache_missing",SEVERITY_BLOCKING,
                                            f"{sheet.title} 第 {row} 行只有无缓存公式，被当作空行跳过",
                                            locator={"sheet":sheet.title,"row":row,"cells":",".join(cell_reference(row,column) for column in missing_cells[:5])}))
                continue
            missing=[headers[column] for column in headers if headers[column] and kinds.get((row,column))==CELL_FORMULA_UNCACHED]
            if missing:
                report.add(ParseFinding("formula_cache_missing",SEVERITY_BLOCKING,
                                        f"{sheet.title} 第 {row} 行有 {len(missing)} 列是无缓存公式（{'、'.join(missing[:5])}），这些值未被解析",
                                        locator={"sheet":sheet.title,"row":row,"cells":",".join(cell_reference(row,column) for column in missing_cells[:5])}))
            question=_first_text(values,QUESTION) or "";answer=_first_text(values,ANSWER) or "";suggestion=_first_text(values,SUGGESTION) or ""
            field_code=_first_text(values,FIELD_CODE);field_name=_first_text(values,FIELD_NAME);table=_first_text(values,TABLE);scenario_name=_first_text(values,SCENARIO);business_system=_first_text(values,BUSINESS_SYSTEM);source_table=_first_text(values,SOURCE_TABLE);source_field=_first_text(values,SOURCE_FIELD);mart_table=_first_text(values,MART_TABLE);mart_field=_first_text(values,MART_FIELD)
            content_text="\n".join(filter(None,[f"问题：{question}" if question else "",f"回复：{answer}" if answer else "",f"同业建议：{suggestion}" if suggestion else "", "；".join(f"{key}：{_cell_text(value)}" for key,value in values.items() if value not in (None,"") and key not in QUESTION|ANSWER|SUGGESTION)]))
            unit_type="qa" if question or knowledge_type=="regulatory_qa" else "field_mapping" if field_code else "table_row"
            units.append(KnowledgeUnitDraft(unit_type,question or field_name or f"{sheet.title} 第{row}行",content_text,source_sheet_name=sheet.title,source_cell_range=f"A{row}:{get_column_letter(sheet.max_column)}{row}",target_table_code=table,target_field_code=field_code,target_field_name=field_name,scenario_name=scenario_name,source_table_name=source_table,source_field_name=source_field,metadata={"row":row,"scenario_name":scenario_name,"business_system_name":business_system,"mart_table_name":mart_table,"mart_field_name":mart_field}))
    report.row_emitted=len(units)
    return units,report.warnings(),report

def _expanded_matrix(sheet):
    matrix=[[sheet.cell(row,column).value for column in range(1,sheet.max_column+1)] for row in range(1,sheet.max_row+1)]
    for merged in sheet.merged_cells.ranges:
        value=matrix[merged.min_row-1][merged.min_col-1]
        for row in range(merged.min_row,merged.max_row+1):
            for column in range(merged.min_col,merged.max_col+1):matrix[row-1][column-1]=value
    return matrix

def _docx(content):
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc=Document(BytesIO(content));units=[];heading=None;chapter=None
    paragraph_index=0;table_index=0
    for element in doc.element.body.iterchildren():
        if element.tag.endswith("}p"):
            paragraph_index+=1;p=Paragraph(element,doc);text=p.text.strip()
            if not text:continue
            if p.style and p.style.name.lower().startswith("heading"):
                heading=text
                chapter_match=CHAPTER_NUMBER_PATTERN.match(text)
                # A heading without a chapter number (``附录``, ``总则``) ends the chapter
                # scope of the previous one: clauses below it are not part of 第三章.
                chapter=chapter_match.group(1) if chapter_match else None
                # A heading states where the units below it live; it is not a citable
                # clause of its own.  Emitting it as a policy_clause made every chapter
                # heading an uncitable unit and a permanent clause_locator_missing gap.
                continue
            clause_match=CLAUSE_NUMBER_PATTERN.match(text)
            metadata={"paragraph":paragraph_index,
                      "clause_no":clause_match.group(0) if clause_match else None,
                      "chapter_no":chapter,
                      "chapter_heading":heading}
            units.append(KnowledgeUnitDraft("policy_clause",heading or f"段落 {paragraph_index}",text,source_heading=heading,metadata=metadata))
        elif element.tag.endswith("}tbl"):
            table_index+=1;table=Table(element,doc)
            for row_index,row in enumerate(table.rows,1):
                text=" | ".join(cell.text.strip() for cell in row.cells)
                if text.strip(" | "):units.append(KnowledgeUnitDraft("table_row",f"表格 {table_index} 第{row_index}行",text,source_heading=heading,source_cell_range=f"table:{table_index}:row:{row_index}",metadata={"locator":{"table_index":table_index,"row_no":row_index},"chapter_no":chapter}))
    return units,[]

def _pdf(content):
    """Parse a PDF; scanned pages go through OCR only when ocr_enabled is on."""
    if not getattr(get_settings(),"ocr_enabled",False):return _pdf_text_layer(content)
    return _pdf_with_ocr(content)


def _pdf_text_layer(content):
    from pypdf import PdfReader
    units=[];warnings=[]
    for page_no,page in enumerate(PdfReader(BytesIO(content)).pages,1):
        text=(page.extract_text() or "").strip()
        if text:units.append(KnowledgeUnitDraft("policy_clause",f"第 {page_no} 页",text,source_page_no=page_no))
        else:warnings.append(f"第 {page_no} 页无可提取文本，未启用 OCR")
    return units,warnings


def _pdf_with_ocr(content):
    """OCR fallback for pages without a text layer, truncated at ocr_max_pages."""
    from pypdf import PdfReader
    from .ocr import run_ocr
    settings=get_settings()
    try:pages=list(PdfReader(BytesIO(content)).pages)
    except Exception as exc:raise ValueError(f"PDF 无法读取：{type(exc).__name__}: {exc}") from exc
    raw_max_pages=getattr(settings,"ocr_max_pages",50)
    max_pages=50 if raw_max_pages is None else int(raw_max_pages)
    if max_pages<=0:return [],[f"ocr_max_pages={max_pages} 无效，未处理任何页面"]
    units=[];warnings=[]
    state={"engine":None,"loaded":False,"provider_reported":False}
    if len(pages)>max_pages:warnings.append(f"扫描件共 {len(pages)} 页，超过 ocr_max_pages={max_pages}，仅处理前 {max_pages} 页（已截断）")
    for page_no,page in enumerate(pages[:max_pages],1):
        text=(page.extract_text() or "").strip()
        if text:units.append(KnowledgeUnitDraft("policy_clause",f"第 {page_no} 页",text,source_page_no=page_no));continue
        units.extend(_ocr_pdf_page(page,page_no,settings,run_ocr,warnings,state))
    return units,warnings


def _ocr_pdf_page(page,page_no,settings,run_ocr,warnings,state):
    from .ocr import page_images, resolve_provider_name
    images=page_images(page)
    if not images:warnings.append(f"第 {page_no} 页无可提取文本且无可 OCR 图像，已跳过");return []
    engine,provider_error=_ocr_engine(settings,state)
    if engine is None:
        if provider_error and not state.get("provider_reported"):
            state["provider_reported"]=True;warnings.append(provider_error)
        warnings.append(f"第 {page_no} 页无可提取文本：OCR provider '{resolve_provider_name(getattr(settings,'ocr_provider',None))}' 不可用，已跳过该页")
        return []
    result=run_ocr(images,provider=getattr(settings,"ocr_provider",None),engine=engine)
    warnings.extend(result.warnings)
    if not result.available:warnings.append(f"第 {page_no} 页无可提取文本：OCR provider '{result.provider}' 不可用，已跳过该页");return []
    text="\n".join(part.strip() for part in result.texts if part and part.strip()).strip()
    if not text:warnings.append(f"第 {page_no} 页 OCR 未识别出文本");return []
    units=_text_units(text,f"第 {page_no} 页（OCR）","policy_clause")
    for unit in units:unit.source_page_no=page_no
    return units


def _ocr_engine(settings,state):
    """Load the OCR engine at most once per document, caching failures too."""
    if not state["loaded"]:
        from .ocr import load_engine_or_warn
        engine,error=load_engine_or_warn(getattr(settings,"ocr_provider",None))
        state["loaded"]=True;state["engine"]=engine;state["error"]=error
    return state["engine"],state.get("error")


def _image(content,suffix):
    """Parse a standalone scanned image; callers gate this on ocr_enabled."""
    from .ocr import image_from_bytes, run_ocr
    settings=get_settings()
    title=f"扫描图片 {suffix.lstrip('.')}"

    image=image_from_bytes(content)
    if image is None:return [],[f"{title} 无法解码为图像，OCR 已跳过"]
    result=run_ocr([image],provider=getattr(settings,"ocr_provider",None))
    warnings=list(result.warnings)
    if not result.available:warnings.append(f"{title} OCR provider '{result.provider}' 不可用，已跳过");return [],warnings
    text="\n".join(part.strip() for part in result.texts if part and part.strip()).strip()
    if not text:warnings.append(f"{title} OCR 未识别出文本");return [],warnings
    units=_text_units(text,title,"policy_clause")
    for unit in units:unit.source_page_no=1
    return units,warnings
