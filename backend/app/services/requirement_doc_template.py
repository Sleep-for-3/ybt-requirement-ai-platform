"""客户需求文档模板渲染器（P0-D）。

设计边界
--------
只支持「经适配与验收的客户 .docx 模板」，不做任意模板自动理解引擎：

1. 适配人员用纯数据配置（:class:`TemplateConfig` / dict）把模板中的
   ``{{...}}`` 锚点映射到固定的逻辑键（见 :data:`SUPPORTED_ANCHOR_KEYS`）。
2. 渲染器只做三件事：封面/固定章节的标量填充、字段级重复表格、审核记录重复表格。
   其余内容（页眉、页脚、分节、样式、编号定义、页面设置）一律原位保留。
3. 只消费冻结 JSON，不查数据库、不联网、不调模型、不执行宏。
4. 同一模板 + 同一冻结内容 + 同一 ``frozen_at`` 连续渲染两次，输出字节完全一致：
   ``core_properties`` 的 created/modified 固定为 ``frozen_at``，zip 条目时间戳与顺序
   由 :func:`_deterministic_package` 固定。

锚点唯一性
----------
每个声明的锚点在模板中必须恰好出现一次（段落内或某个表格样板行内）。锚点可被 Word
拆到多个 ``w:r`` 中（例如 ``{{co`` / ``ver`` / ``.name}}``），渲染前会先按 run 顺序
合并文本再定位，替换时把值写入起始 run，保留该 run 的字符格式。锚点缺失/重复/落在
超链接内都会抛出带中文诊断的异常，绝不静默产出残缺正式文档。
"""
from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass, field, asdict
from datetime import date, datetime, timezone
from io import BytesIO
import hashlib
import json
import re
from typing import Any, Iterable, Mapping, Sequence
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


# ---------------------------------------------------------------------------
# 常量：逻辑键词表（适配契约）与默认填充值
# ---------------------------------------------------------------------------

DOCUMENT_TITLE = "跑批脚本需求说明"
DEFAULT_FILL = "待业务补充"
DEFAULT_FROZEN_AT = datetime(2000, 1, 1, 0, 0, 0)
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

COVER_KEYS: tuple[str, ...] = (
    "cover.title",
    "cover.name",
    "cover.status",
    "cover.content_version",
    "cover.content_digest",
    "cover.summary",
    "cover.frozen_at",
)
SECTION_KEYS: tuple[str, ...] = (
    "section.scope",
    "section.fields_intro",
    "section.lineage",
    "section.script_basis",
    "section.policy",
    "section.gaps",
    "section.review",
)
TABLE_KEYS: tuple[str, ...] = ("table.fields", "table.review")
SUPPORTED_ANCHOR_KEYS: tuple[str, ...] = COVER_KEYS + SECTION_KEYS + TABLE_KEYS

FIELD_COLUMN_KEYS: tuple[str, ...] = (
    "seq",
    "name",
    "code",
    "business_definition",
    "business_content",
    "technical_content",
    "source_system",
    "processing_logic",
    "path_status",
    "owner_record",
    "ai_business_candidate",
    "ai_technical_candidate",
)
REVIEW_COLUMN_KEYS: tuple[str, ...] = ("step", "decision", "reviewer", "decided_at", "comment")
TABLE_COLUMN_KEYS: dict[str, tuple[str, ...]] = {
    "table.fields": FIELD_COLUMN_KEYS,
    "table.review": REVIEW_COLUMN_KEYS,
}

# 模板未显式给出列映射时按此默认列序填充（适配人员仍应显式配置）。
DEFAULT_TABLE_COLUMNS: dict[str, dict[str, int]] = {
    "table.fields": {
        "seq": 0,
        "name": 1,
        "code": 2,
        "business_definition": 3,
        "business_content": 4,
        "technical_content": 5,
        "source_system": 6,
        "processing_logic": 7,
        "path_status": 8,
    },
    "table.review": {"step": 0, "decision": 1, "reviewer": 2, "decided_at": 3, "comment": 4},
}

DEFAULT_CONFIG_VERSION = "p0d-1"

# 现有 Word 导出（backend/app/services/requirement_word.py）实际使用的一级章节，
# 顺序与文案保持一致；默认模板生成器复用它。
DEFAULT_SECTION_TITLES: tuple[str, ...] = (
    "需求范围与业务说明",
    "字段业务说明",
    "来源系统与分层路径",
    "加工规则与脚本证据",
    "制度依据与差异",
    "待确认事项",
    "审核与正式交付记录",
)


# ---------------------------------------------------------------------------
# 异常：中文消息，直接告诉适配人员缺哪个键
# ---------------------------------------------------------------------------


class RequirementDocTemplateError(RuntimeError):
    """模板渲染相关错误的基类。"""


class TemplateConfigError(RequirementDocTemplateError):
    """模板配置（注册/适配数据）本身不合法。"""


class TemplateHashMismatch(RequirementDocTemplateError):
    """传入模板字节与配置中固定的 template_hash 不一致。"""


class TemplateAnchorError(RequirementDocTemplateError):
    """锚点缺失、重复、位置非法或无法替换。"""


class FrozenContentScopeError(RequirementDocTemplateError):
    """冻结内容超出模板适配声明的 supported_scope。"""


class FrozenContentError(RequirementDocTemplateError):
    """冻结内容结构不满足渲染前置条件。"""


# ---------------------------------------------------------------------------
# 配置模型与校验（纯数据，不建数据库模型）
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class TemplateConfig:
    """客户模板的注册/适配配置。

    ``anchors``: 逻辑键 -> 模板中的锚点标识（``{{...}}``，可带中文标签）。
    ``required_keys``: 缺失即失败的必需映射（必须是 anchors 的子集）。
    ``optional_keys``: 允许模板缺失的锚点（缺失时记入诊断 skipped，不失败）。
    ``table_columns``: 重复表格的逻辑列键 -> 0 基列下标。
    ``supported_scope``: 适配验收范围，例如 ``{"min_fields": 0, "max_fields": 500,
        "sections": ["需求范围与业务说明", ...]}``；``min_fields``/``max_fields`` 会被强制执行。
    """

    template_hash: str
    config_version: str
    supported_scope: dict[str, Any]
    anchors: dict[str, str]
    required_keys: tuple[str, ...] = ()
    optional_keys: tuple[str, ...] = ()
    table_columns: dict[str, dict[str, int]] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["required_keys"] = list(self.required_keys)
        data["optional_keys"] = list(self.optional_keys)
        return data


@dataclass(frozen=True)
class TableProfile:
    """模板中的一张表（用于分析报告）。"""

    index: int
    rows: int
    columns: int
    header_text: str


@dataclass(frozen=True)
class TemplateAnalysis:
    """模板结构分析结果。

    缺失/重复锚点只报告不抛错；配置本身非法（未知逻辑键、必需映射缺失）仍抛
    :class:`TemplateConfigError`，模板哈希不符抛 :class:`TemplateHashMismatch`。
    """

    template_hash: str
    config_version: str
    found_anchors: dict[str, tuple[str, ...]]
    missing_keys: tuple[str, ...]
    duplicate_keys: dict[str, tuple[str, ...]]
    unsupported_keys: tuple[str, ...]
    section_count: int
    paragraph_count: int
    table_count: int
    tables: tuple[TableProfile, ...]
    header_text: tuple[str, ...]
    footer_text: tuple[str, ...]
    style_names: tuple[str, ...]
    numbering_abstract_ids: tuple[str, ...]
    numbering_num_ids: tuple[str, ...]
    max_numbering_levels: int
    core_modified: str | None

    @property
    def ok(self) -> bool:
        return not self.missing_keys and not self.duplicate_keys and not self.unsupported_keys

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["found_anchors"] = {k: list(v) for k, v in self.found_anchors.items()}
        data["duplicate_keys"] = {k: list(v) for k, v in self.duplicate_keys.items()}
        for key in ("missing_keys", "unsupported_keys", "header_text", "footer_text",
                    "style_names", "numbering_abstract_ids", "numbering_num_ids"):
            data[key] = list(data[key])
        data["tables"] = [asdict(t) for t in self.tables]
        data["ok"] = self.ok
        return data


@dataclass(frozen=True)
class RenderDiagnostics:
    """一次渲染的可追溯诊断。"""

    template_hash: str
    config_version: str
    content_digest: str
    frozen_at: str
    filled_anchors: dict[str, str]
    skipped_keys: tuple[str, ...]
    table_rows: dict[str, int]
    section_count: int
    header_text: tuple[str, ...]
    footer_text: tuple[str, ...]
    warnings: tuple[str, ...]
    output_sha256: str

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["skipped_keys"] = list(self.skipped_keys)
        data["header_text"] = list(self.header_text)
        data["footer_text"] = list(self.footer_text)
        data["warnings"] = list(self.warnings)
        return data


def template_hash(template_bytes: bytes) -> str:
    """模板原文件的 sha256（小写十六进制）。"""
    return hashlib.sha256(template_bytes).hexdigest()


def _normalize_token(value: Any, key: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise TemplateConfigError(f"锚点映射 {key!r} 的值必须是非空字符串，实际为 {value!r}")
    token = value.strip()
    if token.startswith("{{") and token.endswith("}}"):
        inner = token[2:-2].strip()
        if not inner:
            raise TemplateConfigError(f"锚点映射 {key!r} 的占位符为空：{value!r}")
        return "{{" + inner + "}}"
    if "{{" in token or "}}" in token:
        raise TemplateConfigError(
            f"锚点映射 {key!r} 的占位符 {value!r} 格式不完整，应为 {{{{标签}}}} 形式")
    return "{{" + token + "}}"


def parse_config(config: TemplateConfig | Mapping[str, Any]) -> TemplateConfig:
    """把 dict / TemplateConfig 归一化为经过校验的 :class:`TemplateConfig`。"""
    if isinstance(config, TemplateConfig):
        return _build_config({
            "template_hash": config.template_hash,
            "config_version": config.config_version,
            "supported_scope": config.supported_scope,
            "anchors": config.anchors,
            "required_keys": config.required_keys,
            "optional_keys": config.optional_keys,
            "table_columns": config.table_columns,
        })
    if not isinstance(config, Mapping):
        raise TemplateConfigError(
            f"模板配置必须是 dict 或 TemplateConfig，实际类型为 {type(config).__name__}")
    return _build_config(config)


def _build_config(config: Mapping[str, Any]) -> TemplateConfig:
    missing = [name for name in ("template_hash", "config_version", "supported_scope", "anchors")
               if name not in config]
    if missing:
        raise TemplateConfigError("模板配置缺少必需字段：" + "、".join(missing))

    raw_anchors = config["anchors"]
    if not isinstance(raw_anchors, Mapping) or not raw_anchors:
        raise TemplateConfigError("模板配置 anchors 必须是非空映射：逻辑键 -> {{锚点}}")

    anchors: dict[str, str] = {}
    for key, value in raw_anchors.items():
        if key not in SUPPORTED_ANCHOR_KEYS:
            raise TemplateConfigError(
                f"未知的锚点逻辑键 {key!r}；当前渲染器支持的逻辑键为：" + "、".join(SUPPORTED_ANCHOR_KEYS))
        anchors[key] = _normalize_token(value, key)

    seen_tokens: dict[str, str] = {}
    for key, token in anchors.items():
        if token in seen_tokens:
            raise TemplateConfigError(
                f"锚点 {token} 同时映射到 {seen_tokens[token]!r} 和 {key!r}，模板适配存在歧义")
        seen_tokens[token] = key

    required = _str_tuple(config.get("required_keys") or (), "required_keys")
    optional = _str_tuple(config.get("optional_keys") or (), "optional_keys")
    for key in required:
        if key not in anchors:
            raise TemplateConfigError(
                f"必需映射 {key!r} 未在 anchors 中声明，无法保证正式文档完整")
        if key not in SUPPORTED_ANCHOR_KEYS:
            raise TemplateConfigError(f"必需映射 {key!r} 不是渲染器支持的逻辑键")
    for key in optional:
        if key not in anchors:
            raise TemplateConfigError(f"可选映射 {key!r} 未在 anchors 中声明")
        if key in required:
            raise TemplateConfigError(f"锚点 {key!r} 不能同时声明为必需映射和可选映射")

    table_columns: dict[str, dict[str, int]] = {}
    raw_columns = config.get("table_columns") or {}
    if not isinstance(raw_columns, Mapping):
        raise TemplateConfigError("模板配置 table_columns 必须是映射：表逻辑键 -> {列键: 下标}")
    for table_key, columns in raw_columns.items():
        if table_key not in TABLE_KEYS:
            raise TemplateConfigError(f"table_columns 中的表键 {table_key!r} 不合法，应为 table.fields / table.review")
        if table_key not in anchors:
            raise TemplateConfigError(f"table_columns 声明了 {table_key!r}，但 anchors 未声明该表的重复行锚点")
        if not isinstance(columns, Mapping):
            raise TemplateConfigError(f"{table_key} 的列映射必须是 dict")
        normalized: dict[str, int] = {}
        for column_key, index in columns.items():
            if column_key not in TABLE_COLUMN_KEYS[table_key]:
                raise TemplateConfigError(
                    f"{table_key} 的列键 {column_key!r} 不受支持；支持：" + "、".join(TABLE_COLUMN_KEYS[table_key]))
            if isinstance(index, bool) or not isinstance(index, int) or index < 0:
                raise TemplateConfigError(f"{table_key}.{column_key} 的列下标必须是非负整数，实际为 {index!r}")
            normalized[column_key] = index
        index_to_column: dict[int, str] = {}
        for column_key, index in normalized.items():
            if index in index_to_column:
                raise TemplateConfigError(
                    f"{table_key} 的列下标 {index} 同时被 {index_to_column[index]!r} 和 {column_key!r} 使用")
            index_to_column[index] = column_key
        table_columns[table_key] = normalized

    scope = config["supported_scope"]
    if not isinstance(scope, Mapping):
        raise TemplateConfigError("模板配置 supported_scope 必须是 dict")
    scope_dict = dict(scope)
    for bound in ("min_fields", "max_fields"):
        if bound in scope_dict:
            value = scope_dict[bound]
            if isinstance(value, bool) or not isinstance(value, int) or value < 0:
                raise TemplateConfigError(f"supported_scope.{bound} 必须是非负整数，实际为 {value!r}")
    if "min_fields" in scope_dict and "max_fields" in scope_dict and scope_dict["min_fields"] > scope_dict["max_fields"]:
        raise TemplateConfigError("supported_scope.min_fields 不能大于 max_fields")
    if "sections" in scope_dict and not isinstance(scope_dict["sections"], (list, tuple)):
        raise TemplateConfigError("supported_scope.sections 必须是章节名称列表")

    hash_value = config["template_hash"]
    if not isinstance(hash_value, str) or not re.fullmatch(r"[0-9a-f]{64}", hash_value.strip().lower()):
        raise TemplateConfigError("template_hash 必须是 64 位小写十六进制字符串（原文件 sha256）")
    version = config["config_version"]
    if not isinstance(version, str) or not version.strip():
        raise TemplateConfigError("config_version 必须是非空字符串")

    return TemplateConfig(
        template_hash=hash_value.strip().lower(),
        config_version=version.strip(),
        supported_scope=scope_dict,
        anchors=anchors,
        required_keys=required,
        optional_keys=optional,
        table_columns=table_columns,
    )


def _str_tuple(values: Iterable[Any], name: str) -> tuple[str, ...]:
    result: list[str] = []
    for value in values:
        if not isinstance(value, str) or not value.strip():
            raise TemplateConfigError(f"{name} 只能包含非空字符串，实际为 {value!r}")
        if value not in result:
            result.append(value)
    return tuple(result)


def validate_config(config: TemplateConfig | Mapping[str, Any]) -> TemplateConfig:
    """校验配置；校验失败抛 :class:`TemplateConfigError`（中文消息含键名）。"""
    return parse_config(config)


# ---------------------------------------------------------------------------
# 模板加载与锚点扫描（含跨 run 合并）
# ---------------------------------------------------------------------------


@dataclass
class _AnchorHit:
    key: str
    token: str
    location: str
    paragraph: Paragraph
    table_el: Any | None
    row_index: int | None
    supported: bool = True


def _load_template(template_bytes: bytes, config: TemplateConfig) -> Any:
    if not isinstance(template_bytes, (bytes, bytearray)) or not template_bytes:
        raise RequirementDocTemplateError("模板字节为空，无法渲染客户需求文档")
    actual = template_hash(bytes(template_bytes))
    if actual != config.template_hash:
        raise TemplateHashMismatch(
            f"模板哈希不一致：配置声明 template_hash={config.template_hash}，"
            f"实际文件 sha256={actual}；请重新适配并固定模板版本，不要直接替换模板文件")
    try:
        return Document(BytesIO(bytes(template_bytes)))
    except Exception as exc:  # python-docx 抛 PackageNotFoundError 等
        raise RequirementDocTemplateError(
            f"无法打开模板文件（需要 .docx：{type(exc).__name__}: {exc}）；"
            "请确认模板是未加密的 .docx，而不是 .doc/.docm 改名文件") from exc


def _iter_block_paragraphs(container, location_prefix: str, container_label: str,
                           hits: list[_AnchorHit], stats: dict[str, int],
                           token_map: Mapping[str, str], seen_tc: set,
                           table_el=None, row_index: int | None = None) -> None:
    """遍历 container 直属的 w:p / w:tbl（递归表格单元格）。"""
    for child in container.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            stats["paragraphs"] += 1
            _scan_paragraph(child, location_prefix, hits, token_map, table_el, row_index)
        elif tag == qn("w:tbl"):
            stats["tables"] += 1
            rows = child.findall(qn("w:tr"))
            for r_idx, tr in enumerate(rows):
                for c_idx, tc in enumerate(tr.findall(qn("w:tc"))):
                    # 用元素对象本身去重（并借集合持有引用）。绝不能只存 id()：
                    # lxml 代理对象被回收后 id 会被复用，导致后续单元格被误跳过，
                    # 表现为锚点扫描随机丢表格（0 字段/长文本等场景曾偶发）。
                    if tc in seen_tc:
                        continue
                    seen_tc.add(tc)
                    _iter_block_paragraphs(
                        tc,
                        f"{location_prefix}表格第{r_idx + 1}行第{c_idx + 1}列",
                        container_label, hits, stats, token_map, seen_tc,
                        table_el=child, row_index=r_idx,
                    )


def _scan_paragraph(p_el, location: str, hits: list[_AnchorHit],
                    token_map: Mapping[str, str], table_el, row_index: int | None) -> None:
    paragraph = Paragraph(p_el, None)
    runs_text = "".join(run.text for run in paragraph.runs)
    full_text = p_el.text
    for key, token in token_map.items():
        count = runs_text.count(token)
        if count:
            for ordinal in range(count):
                suffix = "" if count == 1 else f"（同段第{ordinal + 1}处）"
                hits.append(_AnchorHit(key, token, location + suffix, paragraph,
                                       table_el, row_index, supported=True))
        elif token in full_text:
            hits.append(_AnchorHit(key, token, location + "（超链接/域内，当前渲染器不支持替换）",
                                   paragraph, table_el, row_index, supported=False))


def _header_footer_parts(doc) -> list[tuple[str, str, Any]]:
    """返回 (标签, 角色, 部件元素) 列表；同一部件的多个节引用只扫描一次。"""
    roots: list[tuple[str, str, Any]] = []
    seen: set[int] = set()
    roles = (("w:headerReference", "header", "页眉"), ("w:footerReference", "footer", "页脚"))
    for s_idx, section in enumerate(doc.sections):
        for tag_name, role, label in roles:
            for ref in section._sectPr.findall(qn(tag_name)):
                r_id = ref.get(qn("r:id"))
                part = section.part.related_parts.get(r_id) if r_id else None
                if part is None or id(part) in seen:
                    continue
                seen.add(id(part))
                kind = ref.get(qn("w:type")) or "default"
                roots.append((f"第{s_idx + 1}节{label}({kind})", role, part.element))
    return roots


def _scan_anchors(doc, config: TemplateConfig) -> list[_AnchorHit]:
    token_map = {key: config.anchors[key] for key in sorted(config.anchors)}
    hits: list[_AnchorHit] = []
    stats = {"paragraphs": 0, "tables": 0}
    seen_tc: set[int] = set()
    _iter_block_paragraphs(doc.element.body, "正文", "body", hits, stats, token_map, seen_tc)
    for label, role, root in _header_footer_parts(doc):
        _iter_block_paragraphs(root, label, role, hits, stats, token_map, seen_tc)
    return hits


def _group_hits(hits: Iterable[_AnchorHit]) -> dict[str, list[_AnchorHit]]:
    grouped: dict[str, list[_AnchorHit]] = {}
    for hit in hits:
        grouped.setdefault(hit.key, []).append(hit)
    return grouped


def _header_footer_text(doc) -> tuple[tuple[str, ...], tuple[str, ...]]:
    headers: list[str] = []
    footers: list[str] = []
    for label, role, root in _header_footer_parts(doc):
        texts = [p.text for p in root.iter(qn("w:p"))]
        joined = "\n".join(text for text in texts if text).strip()
        if role == "header" and joined:
            headers.append(joined)
        elif role == "footer" and joined:
            footers.append(joined)
    return tuple(headers), tuple(footers)


def _numbering_profile(doc) -> tuple[tuple[str, ...], tuple[str, ...], int]:
    try:
        numbering = doc.part.numbering_part.element
    except (KeyError, NotImplementedError, ValueError):
        return (), (), 0
    abstract_ids = tuple(el.get(qn("w:abstractNumId")) or "" for el in numbering.findall(qn("w:abstractNum")))
    num_ids = tuple(el.get(qn("w:numId")) or "" for el in numbering.findall(qn("w:num")))
    max_levels = 0
    for abstract in numbering.findall(qn("w:abstractNum")):
        max_levels = max(max_levels, len(abstract.findall(qn("w:lvl"))))
    return abstract_ids, num_ids, max_levels


def _table_profiles(doc) -> tuple[TableProfile, ...]:
    profiles: list[TableProfile] = []
    for idx, tbl in enumerate(doc.element.body.iterchildren(qn("w:tbl")), start=1):
        trs = tbl.findall(qn("w:tr"))
        columns = 0
        for tr in trs:
            columns = max(columns, len(tr.findall(qn("w:tc"))))
        header = ""
        if trs:
            header = " | ".join(
                "".join(t.text or "" for t in tc.iter(qn("w:t")))
                for tc in trs[0].findall(qn("w:tc"))
            )
        profiles.append(TableProfile(index=idx, rows=len(trs), columns=columns, header_text=header))
    return tuple(profiles)


def _core_modified(doc) -> str | None:
    value = doc.core_properties.modified
    return value.isoformat() if value is not None else None


def analyze_template(template_bytes: bytes, config: TemplateConfig | Mapping[str, Any]) -> TemplateAnalysis:
    """分析客户模板：锚点命中/缺失/重复、段落表格结构、页眉页脚、样式与编号定义。

    只读操作，不修改模板；模板哈希与配置不一致时抛 :class:`TemplateHashMismatch`。
    """
    cfg = parse_config(config)
    doc = _load_template(template_bytes, cfg)
    hits = _scan_anchors(doc, cfg)
    grouped = _group_hits(hits)
    found = {key: tuple(hit.location for hit in grouped[key]) for key in sorted(grouped)}
    declared = set(cfg.anchors)
    missing = tuple(sorted(declared - set(found)))
    duplicate = {key: tuple(hit.location for hit in grouped[key]) for key in sorted(grouped)
                 if len(grouped[key]) > 1}
    unsupported = tuple(sorted({key for key in grouped if any(not hit.supported for hit in grouped[key])}))
    headers, footers = _header_footer_text(doc)
    abstract_ids, num_ids, max_levels = _numbering_profile(doc)
    styles = tuple(sorted({style.name for style in doc.styles if style.name}))
    stats = {"paragraphs": 0, "tables": 0}
    seen_tc: set[int] = set()
    _iter_block_paragraphs(doc.element.body, "", "", [], stats, {}, seen_tc)
    return TemplateAnalysis(
        template_hash=template_hash(bytes(template_bytes)),
        config_version=cfg.config_version,
        found_anchors=found,
        missing_keys=missing,
        duplicate_keys=duplicate,
        unsupported_keys=unsupported,
        section_count=len(doc.sections),
        paragraph_count=stats["paragraphs"],
        table_count=stats["tables"],
        tables=_table_profiles(doc),
        header_text=headers,
        footer_text=footers,
        style_names=styles,
        numbering_abstract_ids=abstract_ids,
        numbering_num_ids=num_ids,
        max_numbering_levels=max_levels,
        core_modified=_core_modified(doc),
    )


# ---------------------------------------------------------------------------
# 冻结内容 -> 填充值（只读 content JSON）
# ---------------------------------------------------------------------------


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # XML 1.0 不允许的控制字符，与既有 Word 导出保持一致地剔除。
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)


def _value_or(value: Any, default: str) -> str:
    text = _clean_text(value).strip()
    return text if text else default


def _content_digest(content: Mapping[str, Any]) -> str:
    # 延迟导入：避免本模块在导入期拉起 models / SQLAlchemy，便于独立测试。
    from app.services.requirement_scope import content_digest

    return content_digest(dict(content))


def _require_content(content: Mapping[str, Any]) -> Mapping[str, Any]:
    if not isinstance(content, Mapping):
        raise FrozenContentError(f"冻结内容必须是 dict，实际类型为 {type(content).__name__}")
    if not isinstance(content.get("requirement"), Mapping):
        raise FrozenContentError("冻结内容缺少 requirement 对象，不能渲染正式需求文档")
    fields = content.get("fields")
    if not isinstance(fields, list):
        raise FrozenContentError("冻结内容缺少 fields 列表（0 个字段应为空列表），不能渲染正式需求文档")
    return content


def _parse_frozen_at(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime(value.year, value.month, value.day)
    elif isinstance(value, str) and value.strip():
        text = value.strip().replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(text)
        except ValueError as exc:
            raise FrozenContentError(f"frozen_at 不是合法的 ISO 时间：{value!r}") from exc
    else:
        raise FrozenContentError(f"frozen_at 必须是 datetime 或 ISO 字符串，实际为 {value!r}")
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed.replace(microsecond=0)


def _derive_frozen_at(content: Mapping[str, Any]) -> datetime:
    formal = content.get("formal_delivery") or {}
    candidates: list[Any] = [formal.get("approved_at")]
    for row in formal.get("review_record") or []:
        candidates.append(row.get("decided_at"))
    candidates.extend([content.get("frozen_at"), content.get("generated_at")])
    for candidate in candidates:
        if not candidate:
            continue
        try:
            return _parse_frozen_at(candidate)
        except FrozenContentError:
            continue
    return DEFAULT_FROZEN_AT


def _target_table_label(content: Mapping[str, Any]) -> str:
    target = content.get("target_table") or {}
    label = " ".join(str(target.get(key) or "") for key in ("table_name", "table_code")).strip()
    return label or DEFAULT_FILL


def _path_confirmed(record: Mapping[str, Any]) -> bool:
    path = record.get("confirmed_path") or {}
    if not (path.get("confirmed_by") and path.get("rule_ids") and path.get("rationale")):
        return False
    return path.get("technical_hash") == _content_digest(record.get("lineage") or {})


def _scope_block(content: Mapping[str, Any]) -> str:
    scope = content["requirement"]
    lines = [f"目标表：{_target_table_label(content)}"]
    for key, label in (("background", "业务背景"), ("objective", "业务目的"),
                       ("inclusion", "纳入范围"), ("exclusion", "排除范围"),
                       ("effective_date", "生效日期"), ("acceptance_criteria", "验收要求")):
        lines.append(f"{label}：{_value_or(scope.get(key), DEFAULT_FILL)}")
    return "\n".join(lines)


def _fields_intro_block(content: Mapping[str, Any]) -> str:
    fields = content.get("fields") or []
    return (f"本固定内容共 {len(fields)} 个字段，逐字段列示业务定义、业务正文、技术正文、"
            f"来源系统与加工规则。数据来源与口径以冻结版本为准。")


def _lineage_block(content: Mapping[str, Any]) -> str:
    basis = content.get("script_basis") or {}
    lines: list[str] = []
    if basis:
        for record in content.get("fields") or []:
            field = record.get("field") or {}
            code = field.get("field_code") or "未标注"
            state = "已确认" if _path_confirmed(record) else "待核验"
            lines.append(f"字段路径：{code}：{state}")
        for table in basis.get("metadata") or []:
            assignment = table.get("assignment") or {}
            lines.append("物理表：" + ".".join(
                str(table.get(key) or "未标注") for key in ("database_name", "schema_name", "table_name")))
            lines.append(f"数据层级：{_value_or(assignment.get('layer_name'), '未分类')}")
            lines.append(f"业务系统：{_value_or(assignment.get('business_system_name') or assignment.get('business_system_id'), '未关联')}")
    if not basis.get("metadata"):
        lines.append("尚无已固定的目录分层路径，需要核验上游元数据。脚本跨层读写不表示脚本只属于某一层。")
    return "\n".join(lines)


def _script_basis_block(content: Mapping[str, Any]) -> str:
    basis = content.get("script_basis") or {}
    lines: list[str] = []
    versions = {version.get("id"): version for version in basis.get("versions") or []}

    def ref(node: Any) -> str:
        node = node or {}
        return ".".join(str(node.get(key) or "未标注") for key in
                        ("database_name", "schema_name", "table_name", "column_name"))

    for rule in basis.get("rules") or []:
        script = versions.get(rule.get("script_version_id")) or {}
        lines.append(f"规则 {rule.get('rule_id') or '未标注'}：出处 {script.get('path', '未标注')} "
                     f"版本 {script.get('version_no', '未标注')} 行 {rule.get('source_line_start')} 至 "
                     f"{rule.get('source_line_end')} 语句 {rule.get('statement_id')}")
        lines.append(f"字段来源：{ref(rule.get('source'))}")
        lines.append(f"写入目标：{ref(rule.get('target'))}")
        for key, label in (("transformation_expression", "加工表达式"), ("join_condition", "关联条件"),
                           ("filter_condition", "过滤条件"), ("aggregation_rule", "聚合规则"),
                           ("code_mapping_rule", "码值转换")):
            if rule.get(key):
                lines.append(f"{label}：{_clean_text(rule[key])}")
    if not basis.get("rules"):
        lines.append("本版本尚未固定脚本规则。现有人工技术说明不能替代可追溯的脚本证据。")
    for script in basis.get("versions") or []:
        lines.append(f"脚本版本标识：{script.get('path')} v{script.get('version_no')} "
                     f"SHA256 {script.get('file_hash')}")
    template = basis.get("template")
    if template:
        lines.append(f"监管模板版本：编号 {template.get('id')} 版本 {template.get('version_no')} "
                     f"监管版本 {template.get('regulatory_version') or '未标注'}")
    return "\n".join(lines)


def _policy_block(content: Mapping[str, Any]) -> str:
    basis = content.get("script_basis") or {}
    decisions = (content.get("policy_comparisons") or {}).get("decisions") or {}
    labels = {"matched": "人工判断匹配", "conflict": "存在冲突", "missing_implementation": "缺少实现",
              "pending": "待确认", "not_compared": "未对照"}
    lines: list[str] = []
    for unit in (basis.get("policy_snapshot") or {}).get("evidence") or []:
        lines.append(f"制度资料：{unit.get('title') or ('制度资料片段 ' + str(unit.get('unit_id')))}")
        lines.append(f"制度版本：资料版本 {unit.get('document_version_id')} "
                     f"监管版本 {unit.get('regulatory_version') or '未标注'} "
                     f"来源类别 {unit.get('source_category')}")
        lines.append(f"条款原文：{_clean_text(unit.get('content'))}")
        row = decisions.get(str(unit.get("unit_id"))) or {}
        lines.append(f"人工对照结果：{labels.get(row.get('status'), '未确认')}")
        lines.append(f"关联实现规则：{_value_or('、'.join(row.get('rule_ids') or []), '未关联')}")
        lines.append(f"核验理由：{_value_or(row.get('rationale'), '待确认')}")
        lines.append(f"差异和待确认事项：{_value_or(row.get('difference'), '未填写差异说明')}")
        if row:
            lines.append(f"人工确认记录：确认人 {row.get('confirmed_by')} 时间 {row.get('confirmed_at')}")
    if not (basis.get("policy_snapshot") or {}).get("evidence"):
        lines.append("缺少制度依据。脚本现状不能自动升级为监管要求。")
    return "\n".join(lines)


def _gaps_block(content: Mapping[str, Any]) -> str:
    gaps = content.get("gaps") or []
    if not gaps:
        return "本固定版本未记录未解决事项，是否可正式交付仍以审核记录为准。"
    return "\n".join(f"事项：{_value_or(gap.get('message'), '待核验')}" for gap in gaps)


def _review_block(content: Mapping[str, Any]) -> str:
    formal = content.get("formal_delivery") or {}
    if not formal:
        return "本固定版本尚无正式交付审核记录，当前为草稿，不得作为正式交付。"
    lines = [f"正式交付版本：{_value_or(formal.get('version_no'), DEFAULT_FILL)}"]
    records = formal.get("review_record") or []
    if not records:
        lines.append("审核记录：尚未记录。")
    for row in records:
        lines.append(
            f"{row.get('step') or '审核步骤'}：{row.get('decision')} 审核人 {row.get('reviewer')} "
            f"时间 {row.get('decided_at')} 意见 {row.get('comment') or '未填写'}")
    return "\n".join(lines)


def _scalar_values(content: Mapping[str, Any], frozen_at: datetime) -> dict[str, str]:
    required = content["requirement"]
    formal = content.get("formal_delivery") or {}
    digest = _content_digest(content)
    version = formal.get("content_version") if formal else required.get("content_version", 0)
    frozen_date = frozen_at.strftime("%Y-%m-%d %H:%M:%S")
    return {
        "cover.title": DOCUMENT_TITLE,
        "cover.name": _value_or(required.get("name"), DEFAULT_FILL),
        "cover.status": "已审核正式交付" if formal else "草稿 待审核 不作为正式交付",
        "cover.content_version": _value_or(version, "0"),
        "cover.content_digest": digest,
        "cover.summary": ("本文记录本需求范围内的字段说明、脚本实现事实、制度依据和人工核验结果。"
                          "业务背景、目的与验收要求尚未填写的部分保留待补充标记。"),
        "cover.frozen_at": frozen_date,
        "section.scope": _scope_block(content),
        "section.fields_intro": _fields_intro_block(content),
        "section.lineage": _lineage_block(content),
        "section.script_basis": _script_basis_block(content),
        "section.policy": _policy_block(content),
        "section.gaps": _gaps_block(content),
        "section.review": _review_block(content),
    }


def _field_row_values(content: Mapping[str, Any], index: int, record: Mapping[str, Any]) -> dict[str, str]:
    field_data = record.get("field") or {}
    business = record.get("business") or {}
    lineage = record.get("lineage") or {}
    ownership = (content.get("manual_ownership") or {}).get(str(field_data.get("id"))) or {}
    owner_text = "；".join(
        f"{key} {value.get('kind', 'manual')} 确认人 {value.get('actor_id')}"
        for key, value in ownership.items())
    path_state = "待核验"
    if content.get("script_basis"):
        path_state = "已确认" if _path_confirmed(record) else "待核验"
    return {
        "seq": str(index + 1),
        "name": _value_or(field_data.get("field_name"), "未命名字段"),
        "code": _value_or(field_data.get("field_code"), ""),
        "business_definition": _value_or(business.get("business_definition"), DEFAULT_FILL),
        "business_content": _value_or(business.get("final_content"), "待人工补充或采用候选"),
        "technical_content": _value_or(lineage.get("final_content"), "待人工补充或采用候选"),
        "source_system": _value_or(lineage.get("source_system_name"), "待核验"),
        "processing_logic": _value_or(lineage.get("processing_logic"), "详见脚本事实，尚需人工核验"),
        "path_status": path_state,
        "owner_record": _value_or(owner_text, "尚无本需求人工维护记录"),
        "ai_business_candidate": _value_or(business.get("ai_generated_content"), ""),
        "ai_technical_candidate": _value_or(lineage.get("ai_generated_content"), ""),
    }


def _field_rows(content: Mapping[str, Any]) -> list[dict[str, str]]:
    return [_field_row_values(content, index, record)
            for index, record in enumerate(content.get("fields") or [])]


def _review_rows(content: Mapping[str, Any]) -> list[dict[str, str]]:
    formal = content.get("formal_delivery") or {}
    rows: list[dict[str, str]] = []
    for record in formal.get("review_record") or []:
        rows.append({
            "step": _value_or(record.get("step"), "审核步骤"),
            "decision": _value_or(record.get("decision"), DEFAULT_FILL),
            "reviewer": _value_or(record.get("reviewer"), DEFAULT_FILL),
            "decided_at": _value_or(record.get("decided_at"), DEFAULT_FILL),
            "comment": _value_or(record.get("comment"), "未填写"),
        })
    return rows


def _table_rows(content: Mapping[str, Any], table_key: str) -> list[dict[str, str]]:
    if table_key == "table.fields":
        return _field_rows(content)
    if table_key == "table.review":
        return _review_rows(content)
    raise TemplateConfigError(f"未知的重复表格键 {table_key!r}")


# ---------------------------------------------------------------------------
# 锚点替换原语（跨 run 合并）
# ---------------------------------------------------------------------------


def _run_offsets(runs: Sequence[Any]) -> list[tuple[int, int]]:
    offsets: list[tuple[int, int]] = []
    position = 0
    for run in runs:
        length = len(run.text)
        offsets.append((position, position + length))
        position += length
    return offsets


def _run_at(offsets: Sequence[tuple[int, int]], position: int) -> int:
    for index, (start, end) in enumerate(offsets):
        if start <= position < end:
            return index
    raise RequirementDocTemplateError(
        f"锚点定位失败：字符位置 {position} 不在任何 run 内（锚点文本为空 run 边界异常）")


def _continuation_paragraph(paragraph: Paragraph, lines: Sequence[str], run) -> None:
    """在锚点段落后插入续行，继承段落/字符格式但不继承编号（避免重复编号）。"""
    base_p = paragraph._p
    p_pr = base_p.find(qn("w:pPr"))
    r_pr = run._r.find(qn("w:rPr")) if run is not None else None
    cursor = base_p
    for line in lines:
        new_p = OxmlElement("w:p")
        if p_pr is not None:
            copied = deepcopy(p_pr)
            num_pr = copied.find(qn("w:numPr"))
            if num_pr is not None:
                copied.remove(num_pr)
            new_p.append(copied)
        if line:
            new_r = OxmlElement("w:r")
            if r_pr is not None:
                new_r.append(deepcopy(r_pr))
            text_el = OxmlElement("w:t")
            text_el.set(_XML_SPACE, "preserve")
            text_el.text = line
            new_r.append(text_el)
            new_p.append(new_r)
        cursor.addnext(new_p)
        cursor = new_p


def _replace_token_in_paragraph(paragraph: Paragraph, token: str, value: str) -> bool:
    runs = paragraph.runs
    merged = "".join(run.text for run in runs)
    start = merged.find(token)
    if start < 0:
        return False
    offset = start + len(token)
    if offset < len(merged) and merged.find(token, offset) >= 0:
        raise TemplateAnchorError(
            f"锚点 {token} 在同一段落 [{(paragraph.text or '')[:40]}] 中出现多次，无法确定填充位置")
    end = offset
    offsets = _run_offsets(runs)
    first = _run_at(offsets, start)
    last = _run_at(offsets, end - 1)
    lines = value.split("\n")
    head = lines[0]
    first_start, _ = offsets[first]
    first_text = runs[first].text
    if first == last:
        runs[first].text = first_text[:start - first_start] + head + first_text[end - first_start:]
    else:
        runs[first].text = first_text[:start - first_start] + head
        for middle in range(first + 1, last):
            runs[middle].text = ""
        last_start, _ = offsets[last]
        runs[last].text = runs[last].text[end - last_start:]
    if len(lines) > 1:
        _continuation_paragraph(paragraph, lines[1:], runs[first])
    return True


def _capture_cell_templates(sample_tr) -> list[tuple[Any | None, Any | None]]:
    templates: list[tuple[Any | None, Any | None]] = []
    for tc in sample_tr.findall(qn("w:tc")):
        paragraph = tc.find(qn("w:p"))
        p_pr = paragraph.find(qn("w:pPr")) if paragraph is not None else None
        run = paragraph.find(qn("w:r")) if paragraph is not None else None
        r_pr = run.find(qn("w:rPr")) if run is not None else None
        templates.append((deepcopy(p_pr) if p_pr is not None else None,
                          deepcopy(r_pr) if r_pr is not None else None))
    return templates


def _rewrite_row(new_tr, values: Mapping[int, str], templates: Sequence[tuple[Any, Any]]) -> None:
    cells = new_tr.findall(qn("w:tc"))
    for index, tc in enumerate(cells):
        tc_pr = tc.find(qn("w:tcPr"))
        for child in list(tc):
            if child is not tc_pr:
                tc.remove(child)
        text = _clean_text(values.get(index, ""))
        lines = text.split("\n") if text else [""]
        p_pr, r_pr = templates[index] if index < len(templates) else (None, None)
        for line in lines:
            new_p = OxmlElement("w:p")
            if p_pr is not None:
                copied = deepcopy(p_pr)
                num_pr = copied.find(qn("w:numPr"))
                if num_pr is not None:
                    copied.remove(num_pr)
                new_p.append(copied)
            if line:
                new_r = OxmlElement("w:r")
                if r_pr is not None:
                    new_r.append(deepcopy(r_pr))
                text_el = OxmlElement("w:t")
                text_el.set(_XML_SPACE, "preserve")
                text_el.text = line
                new_r.append(text_el)
                new_p.append(new_r)
            tc.append(new_p)


def _fill_repeat_table(hit: _AnchorHit, rows: Sequence[Mapping[str, str]],
                       columns: Mapping[str, int], token_map: Mapping[str, str]) -> int:
    table_el = hit.table_el
    if table_el is None or hit.row_index is None:
        raise TemplateAnchorError(
            f"重复表格锚点 {hit.token}（{hit.key}）必须位于表格行内，实际位置：{hit.location}")
    trs = table_el.findall(qn("w:tr"))
    if hit.row_index == 0:
        raise TemplateAnchorError(
            f"重复表格锚点 {hit.token}（{hit.key}）位于表头行（{hit.location}），"
            "请把它放在第一行数据样板行内")
    if hit.row_index >= len(trs):
        raise TemplateAnchorError(f"重复表格锚点 {hit.token}（{hit.key}）所在行已不存在")
    sample_tr = trs[hit.row_index]
    templates = _capture_cell_templates(sample_tr)
    width = len(templates)
    for column_key, index in columns.items():
        if index >= width:
            raise TemplateAnchorError(
                f"{hit.key} 的列映射 {column_key!r} 指向第 {index} 列，但样板行只有 {width} 列；"
                "请修正配置 table_columns")
    # 样板行里不允许再放其他锚点，否则复制出来的行会包含未替换占位符。
    other_tokens = [token for token in token_map.values() if token != hit.token]
    for tc in sample_tr.findall(qn("w:tc")):
        cell_text = "".join(el.text or "" for el in tc.iter(qn("w:t")))
        for token in other_tokens:
            if token in cell_text:
                raise TemplateAnchorError(
                    f"重复表格样板行（{hit.location}）内还包含其他锚点 {token}，会把未替换占位符复制到每一行；"
                    "请把该锚点移出样板行")
    index_columns = {index: column_key for column_key, index in columns.items()}
    ordered_rows: list[dict[int, str]] = []
    for row in rows:
        ordered = {index: _clean_text(row.get(column_key, "")) for index, column_key in index_columns.items()}
        ordered_rows.append(ordered)
    for values in ordered_rows:
        new_tr = deepcopy(sample_tr)
        _rewrite_row(new_tr, values, templates)
        sample_tr.addprevious(new_tr)
    table_el.remove(sample_tr)
    return len(ordered_rows)


# ---------------------------------------------------------------------------
# 确定性保存
# ---------------------------------------------------------------------------


def _deterministic_package(data: bytes) -> bytes:
    """重排 zip：固定条目时间戳/属性，保持原条目顺序，保证两次渲染字节一致。"""
    source = zipfile.ZipFile(BytesIO(data))
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    source.close()
    return output.getvalue()


def _save_document(doc, frozen_at: datetime) -> bytes:
    properties = doc.core_properties
    properties.created = frozen_at
    properties.modified = frozen_at
    buffer = BytesIO()
    doc.save(buffer)
    return _deterministic_package(buffer.getvalue())


# ---------------------------------------------------------------------------
# 渲染入口
# ---------------------------------------------------------------------------


def _check_scope(content: Mapping[str, Any], config: TemplateConfig) -> None:
    field_count = len(content.get("fields") or [])
    minimum = config.supported_scope.get("min_fields")
    maximum = config.supported_scope.get("max_fields")
    if minimum is not None and field_count < minimum:
        raise FrozenContentScopeError(
            f"字段数 {field_count} 小于模板适配验收范围 min_fields={minimum}，"
            "请先补充该模板适配或改用受支持的模板")
    if maximum is not None and field_count > maximum:
        raise FrozenContentScopeError(
            f"字段数 {field_count} 超过模板适配验收范围 max_fields={maximum}，"
            "请先完成大字段量模板适配后再正式导出")


def render(template_bytes: bytes, frozen_content: Mapping[str, Any],
           config: TemplateConfig | Mapping[str, Any], *,
           frozen_at: datetime | str | None = None) -> tuple[bytes, RenderDiagnostics]:
    """基于客户 .docx 骨架原位填充，返回 ``(新文档字节, 诊断)``。

    只消费冻结 JSON；不查库、不联网、不调模型、不执行宏。参数 ``frozen_at`` 用于固定
    文档属性时间（缺省时依次取 formal_delivery.approved_at / review_record.decided_at /
    content.frozen_at / content.generated_at，仍无则用 :data:`DEFAULT_FROZEN_AT`）。
    """
    cfg = parse_config(config)
    content = _require_content(frozen_content)
    _check_scope(content, cfg)
    doc = _load_template(template_bytes, cfg)

    hits = _scan_anchors(doc, cfg)
    grouped = _group_hits(hits)
    declared = set(cfg.anchors)
    found = set(grouped)
    missing = sorted(declared - found)
    missing_required = [key for key in cfg.required_keys if key in missing]
    missing_optional = [key for key in missing if key in cfg.optional_keys]
    missing_fatal = [key for key in missing if key not in cfg.optional_keys]

    duplicate = {key: [hit.location for hit in grouped[key]] for key in sorted(grouped)
                 if len(grouped[key]) > 1}
    unsupported = sorted({key for key in grouped if any(not hit.supported for hit in grouped[key])})
    problems: list[str] = []
    if missing_required:
        problems.append("缺失必需锚点：" + "、".join(
            f"{key}（{cfg.anchors[key]}）" for key in missing_required))
    if missing_fatal:
        problems.append("模板中未找到声明锚点：" + "、".join(
            f"{key}（{cfg.anchors[key]}）" for key in missing_fatal))
    if duplicate:
        problems.append("锚点重复：" + "；".join(
            f"{key}（{cfg.anchors[key]}）出现 {len(locations)} 次：" + " -> ".join(locations)
            for key, locations in duplicate.items()))
    if unsupported:
        problems.append("锚点位于超链接/域内暂不支持：" + "、".join(
            f"{key}（{cfg.anchors[key]}）" for key in unsupported))
    for key in cfg.table_columns:
        if key in grouped and grouped[key][0].table_el is None:
            problems.append(f"{key} 的重复行锚点不在表格内：" + " -> ".join(
                hit.location for hit in grouped[key]))
    if problems:
        raise TemplateAnchorError("客户模板锚点校验失败，已拒绝生成不完整正式文档。" + " | ".join(problems))

    frozen_moment = _parse_frozen_at(frozen_at) if frozen_at is not None else _derive_frozen_at(content)
    scalar_values = _scalar_values(content, frozen_moment)
    table_rows = {key: _table_rows(content, key) for key in TABLE_KEYS if key in cfg.anchors}
    table_columns = {key: dict(cfg.table_columns.get(key) or DEFAULT_TABLE_COLUMNS[key])
                     for key in table_rows}

    warnings: list[str] = []
    for key in sorted(declared - set(SUPPORTED_ANCHOR_KEYS)):
        warnings.append(f"忽略未知逻辑键 {key}")
    for key in missing_optional:
        warnings.append(f"可选锚点 {key}（{cfg.anchors[key]}）在模板中不存在，已跳过")
    skipped = set(SUPPORTED_ANCHOR_KEYS) - declared | set(missing_optional)

    filled: dict[str, str] = {}
    table_counts: dict[str, int] = {}
    token_map = {key: cfg.anchors[key] for key in sorted(cfg.anchors)}

    for key in sorted(cfg.anchors):
        if key in TABLE_KEYS:
            continue
        if key not in grouped:
            # 仅当该键声明为 optional 且模板确实缺失时才会走到这里；已记入 skipped。
            continue
        hit = grouped[key][0]
        if not _replace_token_in_paragraph(hit.paragraph, hit.token, scalar_values[key]):
            raise TemplateAnchorError(
                f"锚点 {hit.token}（{key}）在 {hit.location} 定位失败，模板可能被并发修改")
        filled[key] = hit.location

    for key in TABLE_KEYS:
        if key not in cfg.anchors or key not in grouped:
            continue
        hit = grouped[key][0]
        rows = table_rows[key]
        columns = table_columns[key]
        count = _fill_repeat_table(hit, rows, columns, token_map)
        filled[key] = hit.location
        table_counts[key] = count

    headers, footers = _header_footer_text(doc)
    if not headers:
        warnings.append("模板未定义页眉，输出同样无页眉")
    if not footers:
        warnings.append("模板未定义页脚，输出同样无页脚")

    output = _save_document(doc, frozen_moment)
    reopened = Document(BytesIO(output))
    residue = _find_residue(reopened, cfg)
    if residue:
        raise TemplateAnchorError("渲染后仍存在未替换占位符，已拒绝输出：" + "、".join(residue))

    diagnostics = RenderDiagnostics(
        template_hash=cfg.template_hash,
        config_version=cfg.config_version,
        content_digest=scalar_values["cover.content_digest"],
        frozen_at=frozen_moment.isoformat(sep=" "),
        filled_anchors=filled,
        skipped_keys=tuple(sorted(skipped)),
        table_rows=table_counts,
        section_count=len(reopened.sections),
        header_text=headers,
        footer_text=footers,
        warnings=tuple(warnings),
        output_sha256=hashlib.sha256(output).hexdigest(),
    )
    return output, diagnostics


def _find_residue(doc, config: TemplateConfig) -> list[str]:
    """渲染后全文搜索 ``{{...}}``；任何残留都视为失败（不静默出正式文档）。"""
    texts: list[str] = []
    for paragraph in doc.element.body.iter(qn("w:p")):
        texts.append("".join(el.text or "" for el in paragraph.iter(qn("w:t"))))
    for _label, _role, root in _header_footer_parts(doc):
        for paragraph in root.iter(qn("w:p")):
            texts.append("".join(el.text or "" for el in paragraph.iter(qn("w:t"))))
    joined = "\n".join(texts)
    residue = {match.group(0) for match in re.finditer(r"\{\{[^{}]{1,160}\}\}", joined)}
    for token in config.anchors.values():
        if token in joined:
            residue.add(token)
    return sorted(residue)


# ---------------------------------------------------------------------------
# 内置默认模板（与现有 export_word 骨架等价）
# ---------------------------------------------------------------------------


def _configure_default_styles(document) -> None:
    """与 backend/app/services/requirement_word.py 中既有版式保持一致。"""
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name, size in (("Normal", 11), ("Title", 21), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.paragraph_format.space_after = Pt(7)
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
    title_properties = document.styles["Title"].element.get_or_add_pPr()
    border = title_properties.find(qn("w:pBdr"))
    if border is not None:
        title_properties.remove(border)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15


def _default_anchors() -> dict[str, str]:
    keys = ("cover.name", "cover.status", "cover.content_version", "cover.content_digest",
            "cover.summary", "section.scope", "section.fields_intro", "section.lineage",
            "section.script_basis", "section.policy", "section.gaps", "section.review",
            "table.fields", "table.review")
    return {key: "{{" + key + "}}" for key in keys}


def build_default_template() -> bytes:
    """用代码生成与现有 ``export_word`` 等价的内置默认模板（含 7 个一级章节文案锚点）。

    该模板是不可变的 ``bytes``；同一 Python 环境连续调用输出字节一致。
    """
    document = Document()
    _configure_default_styles(document)
    document.add_paragraph(DOCUMENT_TITLE, "Title")

    def labeled(label: str, token: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.add_run(label + "：").bold = True
        paragraph.add_run(token)

    labeled("需求名称", "{{cover.name}}")
    labeled("文档状态", "{{cover.status}}")
    labeled("固定内容版本", "{{cover.content_version}}")
    labeled("固定内容标识", "{{cover.content_digest}}")
    document.add_paragraph("{{cover.summary}}")

    section_tokens = {
        "需求范围与业务说明": "{{section.scope}}",
        "字段业务说明": "{{section.fields_intro}}",
        "来源系统与分层路径": "{{section.lineage}}",
        "加工规则与脚本证据": "{{section.script_basis}}",
        "制度依据与差异": "{{section.policy}}",
        "待确认事项": "{{section.gaps}}",
        "审核与正式交付记录": "{{section.review}}",
    }
    for title, token in section_tokens.items():
        document.add_heading(title, level=1)
        document.add_paragraph(token)
        if title == "字段业务说明":
            field_table = document.add_table(rows=2, cols=9)
            field_table.style = "Table Grid"
            headers = ("序号", "字段名称", "字段代码", "业务定义", "业务正文",
                       "技术正文", "来源系统", "加工规则", "路径核验")
            for index, text in enumerate(headers):
                cell = field_table.rows[0].cells[index]
                cell.text = ""
                cell.paragraphs[0].add_run(text).bold = True
            field_table.rows[1].cells[0].text = "{{table.fields}}"
        if title == "审核与正式交付记录":
            review_table = document.add_table(rows=2, cols=5)
            review_table.style = "Table Grid"
            for index, text in enumerate(("审核步骤", "结论", "审核人", "时间", "意见")):
                cell = review_table.rows[0].cells[index]
                cell.text = ""
                cell.paragraphs[0].add_run(text).bold = True
            review_table.rows[1].cells[0].text = "{{table.review}}"
    properties = document.core_properties
    properties.created = DEFAULT_FROZEN_AT
    properties.modified = DEFAULT_FROZEN_AT
    buffer = BytesIO()
    document.save(buffer)
    return _deterministic_package(buffer.getvalue())


def default_config(template_bytes: bytes | None = None) -> TemplateConfig:
    """内置默认模板的默认配置（template_hash 由模板字节现算，避免手工维护哈希）。"""
    if template_bytes is None:
        template_bytes = build_default_template()
    anchors = _default_anchors()
    return TemplateConfig(
        template_hash=template_hash(bytes(template_bytes)),
        config_version=DEFAULT_CONFIG_VERSION,
        supported_scope={"min_fields": 0, "max_fields": 500,
                         "sections": list(DEFAULT_SECTION_TITLES)},
        anchors=anchors,
        required_keys=tuple(anchors),
        optional_keys=(),
        table_columns={key: dict(value) for key, value in DEFAULT_TABLE_COLUMNS.items()},
    )


__all__ = [
    "COVER_KEYS",
    "DEFAULT_CONFIG_VERSION",
    "DEFAULT_FILL",
    "DEFAULT_FROZEN_AT",
    "DEFAULT_SECTION_TITLES",
    "DEFAULT_TABLE_COLUMNS",
    "DOCUMENT_TITLE",
    "FrozenContentError",
    "FrozenContentScopeError",
    "RenderDiagnostics",
    "RequirementDocTemplateError",
    "SUPPORTED_ANCHOR_KEYS",
    "SECTION_KEYS",
    "TABLE_KEYS",
    "TableProfile",
    "TemplateAnalysis",
    "TemplateAnchorError",
    "TemplateConfig",
    "TemplateConfigError",
    "TemplateHashMismatch",
    "analyze_template",
    "build_default_template",
    "default_config",
    "parse_config",
    "render",
    "template_hash",
    "validate_config",
]
