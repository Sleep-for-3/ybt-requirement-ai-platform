"""Narrative Word export from the same immutable JSON used by Excel.

No database, live retrieval, model call, or file execution occurs here.
"""
from io import BytesIO
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def export_word(content):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name, size in (("Normal", 11), ("Title", 21), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.paragraph_format.space_after = Pt(7)
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
    # Office's built-in Title style adds a blue rule; this document uses spacing only.
    title_properties = document.styles["Title"].element.get_or_add_pPr()
    border = title_properties.find(qn("w:pBdr"))
    if border is not None:
        title_properties.remove(border)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    def text(value, default="待业务补充"):
        return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value if value not in (None, "") else default))
    def paragraph(label, value, default="待业务补充"):
        p = document.add_paragraph()
        p.add_run(label + "：").bold = True
        p.add_run(text(value, default))
        return p
    def heading(value, level=1):
        document.add_heading(text(value), level=level)
    def ref(node):
        return ".".join(str(node.get(k) or "未标注") for k in ("database_name", "schema_name", "table_name", "column_name"))

    scope = content["requirement"]
    formal = content.get("formal_delivery")
    document.add_paragraph("跑批脚本需求说明", "Title")
    paragraph("需求名称", scope.get("name"))
    paragraph("文档状态", "已审核正式交付" if formal else "草稿 待审核 不作为正式交付")
    version = formal.get("content_version") if formal else scope.get("content_version", 0)
    paragraph("固定内容版本", version)
    # A common immutable content identifier is useful when comparing exports.
    from app.services.requirement_scope import content_digest
    paragraph("固定内容标识", content_digest(content))
    document.add_paragraph("本文记录本需求范围内的字段说明、脚本实现事实、制度依据和人工核验结果。业务背景、目的与验收要求尚未填写的部分保留待补充标记。")

    heading("需求范围与业务说明")
    target = content.get("target_table") or {}
    paragraph("目标表", " ".join(str(target.get(k) or "") for k in ("table_name", "table_code")).strip())
    for key, label in (("background", "业务背景"), ("objective", "业务目的"), ("inclusion", "纳入范围"),
                       ("exclusion", "排除范围"), ("effective_date", "生效日期")):
        paragraph(label, scope.get(key))
    paragraph("验收要求", scope.get("acceptance_criteria"))

    heading("字段业务说明")
    for record in content.get("fields", []):
        field = record["field"]
        heading(f"{field.get('field_name') or '未命名字段'} {field.get('field_code') or ''}", 2)
        business, lineage = record.get("business") or {}, record.get("lineage") or {}
        paragraph("业务定义", business.get("business_definition"))
        paragraph("业务正文", business.get("final_content"), "待人工补充或采用候选")
        paragraph("技术正文", lineage.get("final_content"), "待人工补充或采用候选")
        if business.get("ai_generated_content"):
            paragraph("AI 业务解释候选", business["ai_generated_content"])
        if lineage.get("ai_generated_content"):
            paragraph("AI 技术解释候选", lineage["ai_generated_content"])
        paragraph("来源系统", lineage.get("source_system_name"), "待核验")
        paragraph("加工规则", lineage.get("processing_logic"), "详见脚本事实，尚需人工核验")
        ownership = content.get("manual_ownership", {}).get(str(field.get("id")), {})
        paragraph("内容维护记录", "；".join(f"{k} {v.get('kind', 'manual')} 确认人 {v.get('actor_id')}" for k, v in ownership.items()), "尚无本需求人工维护记录")

    basis = content.get("script_basis") or {}
    heading("来源系统与分层路径")
    if basis:
        from app.services.requirement_paths import path_issues
        blocked = {i["field_id"] for i in path_issues(content)}
        for record in content.get("fields", []):
            path = record.get("confirmed_path") or {}
            paragraph("字段路径", f"{record['field']['field_code']}：{'待核验' if record['field']['id'] in blocked else '已确认'}")
            paragraph("路径规则", "、".join(path.get("rule_ids", [])), "未确认")
            paragraph("路径核验依据", path.get("rationale"), "未确认")
    for table in basis.get("metadata", []):
        assignment = table.get("assignment") or {}
        paragraph("物理表", ".".join(str(table.get(k) or "未标注") for k in ("database_name", "schema_name", "table_name")))
        paragraph("数据层级", assignment.get("layer_name"), "未分类")
        paragraph("业务系统", assignment.get("business_system_name") or assignment.get("business_system_id"), "未关联")
    if not basis.get("metadata"):
        document.add_paragraph("尚无已固定的目录分层路径，需要核验上游元数据。脚本跨层读写不表示脚本只属于某一层。")

    heading("加工规则与脚本证据")
    versions = {v["id"]: v for v in basis.get("versions", [])}
    for rule in basis.get("rules", []):
        heading(rule["rule_id"], 2)
        script = versions.get(rule["script_version_id"], {})
        paragraph("出处", f"{script.get('path', '未标注')} 版本 {script.get('version_no', '未标注')} 行 {rule.get('source_line_start')} 至 {rule.get('source_line_end')} 语句 {rule.get('statement_id')}")
        paragraph("字段来源", ref(rule["source"]))
        paragraph("写入目标", ref(rule["target"]))
        for key, label in (("transformation_expression", "加工表达式"), ("join_condition", "关联条件"),
                           ("filter_condition", "过滤条件"), ("aggregation_rule", "聚合规则"), ("code_mapping_rule", "码值转换")):
            if rule.get(key):
                paragraph(label, rule[key])
    if not basis.get("rules"):
        document.add_paragraph("本版本尚未固定脚本规则。现有人工技术说明不能替代可追溯的脚本证据。")
    for script in basis.get("versions", []):
        paragraph("脚本版本标识", f"{script['path']} v{script['version_no']} SHA256 {script['file_hash']}")
    template = basis.get("template")
    if template:
        paragraph("监管模板版本", f"编号 {template['id']} 版本 {template['version_no']} 监管版本 {template.get('regulatory_version') or '未标注'}")

    heading("制度依据与差异")
    decisions = content.get("policy_comparisons", {}).get("decisions", {})
    labels = {"matched": "人工判断匹配", "conflict": "存在冲突", "missing_implementation": "缺少实现", "pending": "待确认"}
    for unit in basis.get("policy_snapshot", {}).get("evidence", []):
        heading(unit.get("title") or f"制度资料片段 {unit['unit_id']}", 2)
        paragraph("制度版本", f"资料版本 {unit['document_version_id']} 监管版本 {unit.get('regulatory_version') or '未标注'} 来源类别 {unit.get('source_category')}")
        paragraph("条款原文", unit["content"])
        row = decisions.get(str(unit["unit_id"])) or {}
        paragraph("人工对照结果", labels.get(row.get("status"), "未确认"))
        paragraph("关联实现规则", "、".join(row.get("rule_ids", [])), "未关联")
        paragraph("核验理由", row.get("rationale"), "待确认")
        paragraph("差异和待确认事项", row.get("difference"), "未填写差异说明")
        if row:
            paragraph("人工确认记录", f"确认人 {row.get('confirmed_by')} 时间 {row.get('confirmed_at')}")
    if not basis.get("policy_snapshot", {}).get("evidence"):
        document.add_paragraph("缺少制度依据。脚本现状不能自动升级为监管要求。")

    heading("待确认事项")
    for gap in content.get("gaps", []):
        paragraph("事项", gap.get("message"), "待核验")
    if not content.get("gaps"):
        document.add_paragraph("本固定版本未记录未解决事项，是否可正式交付仍以审核记录为准。")
    if formal:
        heading("审核与正式交付记录")
        paragraph("正式交付版本", formal.get("version_no"))
        for row in formal.get("review_record", []):
            paragraph(row.get("step") or "审核步骤", f"{row.get('decision')} 审核人 {row.get('reviewer')} 时间 {row.get('decided_at')} 意见 {row.get('comment') or '未填写'}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
