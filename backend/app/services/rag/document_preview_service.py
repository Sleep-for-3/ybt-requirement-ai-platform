"""Authorized version selection, safe preview metadata, and original bytes."""
from pathlib import PurePosixPath
from urllib.parse import quote

from fastapi import HTTPException, Response
from sqlalchemy import select

from app.models import KnowledgeDocument, KnowledgeDocumentVersion, KnowledgeUnit
from app.services.auth.permission_service import PermissionService
from app.services.knowledge_evidence import unit_locator
from app.services.storage import get_storage_service


def authorized_version(db, principal, project_id, document_id, version_id=None):
    try:
        project = PermissionService(db, principal).require_project_permission(project_id, "knowledge.search")
    except HTTPException:
        raise HTTPException(404, "Resource not found") from None
    doc = db.get(KnowledgeDocument, document_id)
    visible = doc and (
        doc.project_id == project_id
        or (doc.knowledge_scope == "global" and doc.confidentiality_level != "restricted")
    )
    if not visible or doc.document_status == "archived":
        raise HTTPException(404, "Resource not found")
    statement = select(KnowledgeDocumentVersion).where(
        KnowledgeDocumentVersion.document_id == doc.id,
        KnowledgeDocumentVersion.project_id == doc.project_id,
    )
    if version_id is not None:
        version = db.scalar(statement.where(KnowledgeDocumentVersion.id == version_id))
    else:
        version = db.scalar(statement.where(KnowledgeDocumentVersion.id == doc.current_version_id)) if doc.current_version_id else None
        # Old integrations updated only current_version_no. Keep that read contract while
        # governed activation writes both pointers atomically.
        if version is None or version.version_no != doc.current_version_no:
            version = db.scalar(statement.where(KnowledgeDocumentVersion.version_no == doc.current_version_no))
    if version is None:
        raise HTTPException(404, "Resource not found")
    return doc, version


def safe_filename(name):
    name = PurePosixPath(str(name).replace("\\", "/")).name
    return "".join(c for c in name if ord(c) >= 32 and ord(c) != 127)[:255] or "document"


def preview_document(db, doc, version):
    units = db.scalars(select(KnowledgeUnit).where(
        KnowledgeUnit.document_id == doc.id,
        KnowledgeUnit.document_version_id == version.id,
        KnowledgeUnit.project_id == doc.project_id,
    ).order_by(KnowledgeUnit.id).limit(2001)).all()
    warnings = []
    # Old versions have no persisted warning list: never borrow current warnings.
    if (version.warnings_json or (version.version_no == doc.current_version_no and doc.warnings_json)):
        warnings.append("文档解析存在警告，部分内容可能未提取。")
    file_type = PurePosixPath(version.file_name).suffix.lower().lstrip(".")
    if file_type == "pdf" and not units:
        warnings.append("PDF 无可提取文本，可能需要 OCR；本服务未启用 OCR。")
    if version.parse_status == "failed":
        warnings.append("此版本解析失败，可尝试查看原文。")
    if len(units) > 2000:
        warnings.append("预览仅显示前 2000 个块，请查看原文获取完整内容。")
    result = {
        "document_id": doc.id, "project_id": doc.project_id,
        "document_version_id": version.id, "version_no": version.version_no,
        "file_name": safe_filename(version.file_name), "file_type": file_type,
        "parse_status": version.parse_status, "warnings": warnings,
        "truncated": len(units) > 2000,
        "blocks": [{"block_id": f"knowledge-unit-{unit.id}", "block_type": unit.unit_type,
                    "knowledge_unit_id": unit.id, "text": unit.content, "locator": unit_locator(unit)}
                   for unit in units[:2000]],
    }

    if file_type == "xlsx":
        from io import BytesIO
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(BytesIO(get_storage_service().read(version.storage_path)), read_only=True, data_only=True)
            try:
                result["sheets"] = [{"name": sheet.title,
                    "rows": [["" if value is None else str(value)[:2000] for value in row]
                             for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 500), max_col=min(sheet.max_column or 1, 50), values_only=True)],
                    "truncated": (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 50}
                    for sheet in workbook.worksheets[:20]]
                if len(workbook.worksheets) > 20 or any(sheet["truncated"] for sheet in result["sheets"]):
                    warnings.append("网格预览有大小限制，请下载原文查看全部内容。")
            finally:
                workbook.close()
        except Exception:
            warnings.append("网格预览暂不可用，请查看解析内容或下载原文。")
    return result


def original_content(version):
    try:
        data = get_storage_service().read(version.storage_path)
    except Exception:
        # Neither local paths nor object-store errors belong in an HTTP response.
        raise HTTPException(404, "Resource not found") from None
    name = safe_filename(version.file_name)
    mime = {
        ".pdf": "application/pdf",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain", ".md": "text/plain", ".sql": "text/plain",
    }.get(PurePosixPath(name).suffix.lower(), "application/octet-stream")
    return Response(data, media_type=mime, headers={
        "Content-Disposition": f"inline; filename=\"document{PurePosixPath(name).suffix if PurePosixPath(name).suffix.lower() in {'.pdf','.xlsx','.docx','.txt','.md','.sql'} else ''}\"; filename*=UTF-8''{quote(name, safe='')}",
        "Cache-Control": "no-store", "X-Content-Type-Options": "nosniff",
        "Content-Security-Policy": "sandbox",
    })
