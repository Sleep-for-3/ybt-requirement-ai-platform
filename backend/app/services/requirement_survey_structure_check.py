"""一表通调研表式导出产物的结构自检：只读、只判断"这份文件是否还像标准模板"。

为什么需要它：真实浏览器里我们已经验证过"能导出、字节可复现"，但评审真正要确认的是
"这份文件的结构是否仍然符合监管模板（表级信息 6 行、按来源系统分组、11 个基础列、
每个字段一行、完整度与来源说明在尾部）"。这些结构一旦被后续改动破坏，光看文件能打开
是发现不了的，所以把验收口径写成可执行的检查。
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.requirement_survey_export import (
    _BASE_COLUMNS,
    _BUSINESS_COLUMNS,
    _SOURCE_COLUMNS,
    _TABLE_HEADERS,
)

STRUCTURE_CHECK_VERSION = "survey-structure-check-v1"
STRUCTURE_CONTRACT_VERSION = "survey-structure-contract-v1"

# 结构契约要按"渲染时的结构"判定，否则历史交付会被新规则误判为不合格。
# 真实证据：产物 #20 渲染于"完整度统计/来源说明"加入导出器之前，用当前契约自检
# 会得到 failed_keys=['completeness','source_notes']——那份文件在当时是合规的。
_SURVEY_RENDERER_PREFIX = "survey-requirement-doc-"
_GENERATION_PATTERN = re.compile(r"v(\d+)")
_LEGACY_KEYS = ("cover", "table_meta", "groups", "header", "rows")
_COMPLETENESS_KEYS = _LEGACY_KEYS + ("completeness",)
_FULL_KEYS = _COMPLETENESS_KEYS + ("source_notes",)
_CONTRACTS = {
    "legacy": {
        "label": "v7 及更早契约（表级信息 / 分组 / 列名 / 数据行）",
        "required_keys": _LEGACY_KEYS,
    },
    "completeness": {
        "label": "v8 契约（另有交付完整度统计，尚无来源说明）",
        "required_keys": _COMPLETENESS_KEYS,
    },
    "full": {
        "label": "v9+ 契约（完整度统计 + 来源说明）",
        "required_keys": _FULL_KEYS,
    },
}

_FIELD_CODE = re.compile(r"^[A-Za-z][A-Za-z0-9_]{1,31}$")
_PLACEHOLDER = "待业务补充"
# 导出里的完整度标题是「交付完整度（自动统计，不替代业务确认）」，按前缀写死会漏判，
# 所以用"包含完整度"作为锚点。
_COMPLETENESS_MARK = "完整度"
# 完整度统计与来源说明整块都在数据行之后；判定"数据行结束"不能只靠一个标题，
# 否则标题被清空时统计行会被误当成数据行（本轮自检在缺失场景里就是这么误报的）。
_FOOTER_MARKS = (
    "完整度",
    "字段总数",
    "有业务口径正文",
    "有集市映射",
    "多来源字段",
    "仅跑批脚本直接产出",
    "未登记来源字段",
    "来源说明",
)
_SOURCE_NOTES_PREFIX = "来源说明"
_BASE_TITLE = "一表通数据基础信息"
_EAST_TITLE = "EAST5.0同源映射列"


class StructureCheckError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def _cell(sheet, row: int, column: int) -> Any:
    return sheet.cell(row=row, column=column).value


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _find_row(sheet, predicate, *, limit: int = 80) -> int | None:
    for row in range(1, min(sheet.max_row, limit) + 1):
        if predicate(row):
            return row
    return None


def _upgrade_hint(contract: dict[str, Any]) -> str | None:
    """历史产物按当时的契约判定为通过，但要说明"通过"是按旧版式通过的。"""
    generation = contract.get("generation")
    if generation is None or generation >= 9:
        return None
    missing = "完整度统计与来源说明" if generation < 8 else "来源说明"
    return (
        f"该产物用 v{generation} 渲染器生成，按当时的契约判定通过；它没有当前的{missing}。"
        "如需当前版式，请按最新内容重新冻结并导出（历史文件不会被覆盖）。"
    )


def contract_for_renderer(renderer_version: str | None) -> dict[str, Any] | None:
    """按渲染器版本返回当时的结构契约；不是一表通调研表式则返回 None。"""
    value = (renderer_version or "").strip()
    if not value.startswith(_SURVEY_RENDERER_PREFIX):
        return None
    match = _GENERATION_PATTERN.search(value[len(_SURVEY_RENDERER_PREFIX):])
    generation = int(match.group(1)) if match else 0
    if generation >= 9:
        key = "full"
    elif generation == 8:
        key = "completeness"
    else:
        key = "legacy"
    spec = _CONTRACTS[key]
    return {
        "version": STRUCTURE_CONTRACT_VERSION,
        "generation": generation,
        "renderer_version": value,
        "label": spec["label"],
        "required_keys": list(spec["required_keys"]),
    }


def check_survey_structure(
    payload: bytes, *, renderer_version: str | None = None
) -> dict[str, Any]:
    if not payload:
        raise StructureCheckError("artifact_empty", "产物文件为空，无法自检。")
    contract = contract_for_renderer(renderer_version)
    if renderer_version and contract is None:
        raise StructureCheckError(
            "artifact_not_survey",
            "这份产物不是一表通调研表式（xlsx），没有可自检的表级信息与列组结构；"
            "Word 与标准版式产物请人工核对，或用结构自检以外的验收方式。",
        )
    if contract is None:
        contract = {
            "version": STRUCTURE_CONTRACT_VERSION,
            "generation": None,
            "renderer_version": None,
            "label": _CONTRACTS["full"]["label"],
            "required_keys": list(_FULL_KEYS),
        }
    try:
        workbook = load_workbook(BytesIO(payload), data_only=True)
    except Exception as error:  # noqa: BLE001 - 任何解析失败都按"不可读"处理
        raise StructureCheckError("artifact_not_readable", f"产物不是可读的 xlsx：{error}") from error
    sheet = workbook.active
    checks: list[dict[str, Any]] = []

    def add(key: str, label: str, ok: bool, detail: str, *, severity: str = "blocking") -> None:
        checks.append(
            {"key": key, "label": label, "ok": bool(ok), "detail": detail, "severity": severity}
        )

    title_ok = _text(_cell(sheet, 1, 1)) == "返回目录"
    add("cover", "首行目录锚点", title_ok, f"A1={_text(_cell(sheet, 1, 1)) or '空'}")

    labels = [label for _, label, _, _ in _TABLE_HEADERS]
    meta_rows = [_find_row(sheet, lambda row, text=label: _text(_cell(sheet, row, 1)) == text) for label in labels]
    meta_ok = all(row is not None for row in meta_rows) and meta_rows == sorted(meta_rows)
    meta_values = [
        _text(_cell(sheet, row, 3)) if row else ""
        for row in meta_rows
    ]
    add(
        "table_meta",
        "表级信息 6 行",
        meta_ok,
        "行号=" + ",".join(str(row) for row in meta_rows) + "；取值=" + " | ".join(meta_values),
    )

    title_row = _find_row(sheet, lambda row: _text(_cell(sheet, row, 1)) == _BASE_TITLE)
    group_titles: list[str] = []
    if title_row:
        for column in range(1, sheet.max_column + 1):
            value = _text(_cell(sheet, title_row, column))
            if value:
                group_titles.append(value)
    business_groups = [value[len("业务口径-"):] for value in group_titles if value.startswith("业务口径-")]
    source_groups = [value[len("溯源-"):] for value in group_titles if value.startswith("溯源-")]
    groups_ok = (
        title_row is not None
        and _BASE_TITLE in group_titles
        and _EAST_TITLE in group_titles
        and business_groups == source_groups
        and len(business_groups) >= 1
    )
    add(
        "groups",
        "按来源系统分组（业务口径/溯源成对）",
        groups_ok,
        f"第 {title_row or '?'} 行；业务口径组={business_groups}；溯源组={source_groups}",
    )

    header_row = _find_row(sheet, lambda row: _text(_cell(sheet, row, 1)) == "数据项编码")
    header_labels: list[str] = []
    if header_row:
        for column in range(1, sheet.max_column + 1):
            header_labels.append(_text(_cell(sheet, header_row, column)))
    base_labels = [label for _, label, _ in _BASE_COLUMNS]
    business_labels = [label for _, label, _ in _BUSINESS_COLUMNS]
    source_labels = [label for _, label, _ in _SOURCE_COLUMNS]
    expected_header = list(base_labels)
    for _ in business_groups:
        expected_header.extend(business_labels)
        expected_header.extend(source_labels)
    actual_header = [label for label in header_labels if label]
    header_ok = bool(header_row) and actual_header == expected_header
    add(
        "header",
        f"列名行（{len(base_labels)} 基础列 + 每组 {len(business_labels)}+{len(source_labels)} 列）",
        header_ok,
        f"第 {header_row or '?'} 行；实际 {len(actual_header)} 列；期望 {len(expected_header)} 列",
    )

    data_rows: list[int] = []
    completeness_row = None
    notes_row = None
    if header_row:
        for row in range(header_row + 1, sheet.max_row + 1):
            first = _text(_cell(sheet, row, 1))
            if any(mark in first for mark in _FOOTER_MARKS):
                break
            if first:
                data_rows.append(row)
        # 尾部两块单独定位：即使完整度标题被清空，也不能把"来源说明也缺了"当成结论
        for row in range(header_row + 1, sheet.max_row + 1):
            first = _text(_cell(sheet, row, 1))
            if completeness_row is None and _COMPLETENESS_MARK in first:
                completeness_row = row
            if notes_row is None and first.startswith(_SOURCE_NOTES_PREFIX):
                notes_row = row

    bad_codes = [
        (_text(_cell(sheet, row, 1)), row)
        for row in data_rows
        if not _FIELD_CODE.match(_text(_cell(sheet, row, 1)))
    ]
    empty_names = [row for row in data_rows if not _text(_cell(sheet, row, 2))]
    rows_ok = bool(data_rows) and not bad_codes and not empty_names
    add(
        "rows",
        "每个数据项一行（编码合法、名称非空）",
        rows_ok,
        f"{len(data_rows)} 行；非法编码 {bad_codes[:3]}；缺名称 {empty_names[:3]}",
    )

    add(
        "completeness",
        "交付完整度统计在数据行之后",
        completeness_row is not None,
        f"第 {completeness_row or '?'} 行",
    )
    add(
        "source_notes",
        "来源说明在完整度之后",
        notes_row is not None and (completeness_row is None or notes_row > completeness_row),
        f"第 {notes_row or '?'} 行",
    )

    placeholders = 0
    for row in data_rows:
        for column in range(1, sheet.max_column + 1):
            if _text(_cell(sheet, row, column)) == _PLACEHOLDER:
                placeholders += 1
    add(
        "placeholders",
        "占位单元格比例（信息项，不判定失败）",
        True,
        f"{placeholders} 个「{_PLACEHOLDER}」单元格（数据 {len(data_rows)} 行 × {sheet.max_column} 列）",
        severity="info",
    )

    required = set(contract["required_keys"])
    skipped_keys: list[str] = []
    for item in checks:
        if item["severity"] != "blocking" or item["key"] in required:
            continue
        # 渲染时还没有这项尾注：按"渲染时结构"判定，降为信息项并写明原因，
        # 既不假装通过，也不把当时合规的文件判成不合格。
        item["severity"] = "info"
        item["detail"] = f"{item['detail']}（该版本渲染时还没有这项尾注，按渲染时结构判定不记为失败）"
        skipped_keys.append(item["key"])

    blocking = [item for item in checks if item["severity"] == "blocking"]
    failed = [item for item in blocking if not item["ok"]]
    return {
        "check_version": STRUCTURE_CHECK_VERSION,
        "sheet_name": sheet.title,
        "field_rows": len(data_rows),
        "column_count": sheet.max_column,
        "row_count": sheet.max_row,
        "groups": business_groups,
        "placeholder_cells": placeholders,
        "checks": checks,
        "contract": contract,
        "skipped_keys": skipped_keys,
        "upgrade_hint": _upgrade_hint(contract),
        "summary": {
            "blocking_total": len(blocking),
            "blocking_failed": len(failed),
            "info_total": len(checks) - len(blocking),
        },
        "verdict": "pass" if not failed else "fail",
        "failed_keys": [item["key"] for item in failed],
    }
