from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from build_deep_tutorial import (
    BLUE_DARK, MUTED, TEAL,
    add_bullets, add_callout, add_chapter_intro, add_code,
    add_numbers, add_para, add_self_check, add_table, configure_book,
)

OUT_DIR = Path(__file__).resolve().parent
OUT_FILE = OUT_DIR / "10_从离线批处理到Flink实时开发_工作实战与求职项目.docx"


def cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BATCH TO STREAM · PROJECT-BASED TEXTBOOK")
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("从离线批处理到 Flink 实时开发")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("用银行实时交易风控与指标平台，讲透时间、状态、可靠性、生产排障与求职表达")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("— 面向会 SQL / ETL / 数仓，但第一次接触流处理的开发者 —")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    for label, value in (
        ("实操基线", "Apache Flink 2.2.1 · Kafka Connector 5.0.0 · Flink CDC 3.6.0"),
        ("贯穿项目", "银行实时交易风控与实时指标平台"),
        ("讲解顺序", "批处理类比 → SQL 快速入门 → Java 状态编程 → 故障演练 → 简历面试"),
        ("资料校验日期", date.today().isoformat()),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(label + "：")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(BLUE_DARK)
        p.add_run(value)
    doc.add_page_break()


def sections(doc: Document, entries: Iterable[tuple[str, str]]) -> None:
    for title, body in entries:
        doc.add_heading(title, level=2)
        add_para(doc, body)


def chapter(doc: Document, number: int, title: str, lead: str,
            outcomes: list[str], entries: list[tuple[str, str]],
            questions: list[str], exercise: str) -> None:
    add_chapter_intro(doc, f"第 {number} 章　{title}", lead, outcomes)
    sections(doc, entries)
    add_self_check(doc, questions, exercise)


def build() -> Path:
    doc = Document()
    configure_book(doc)
    doc.sections[0].header.paragraphs[0].text = "从批处理到实时开发  ·  Flink 工作实战教材"
    doc.core_properties.title = "从离线批处理到 Flink 实时开发：工作实战与求职项目"
    doc.core_properties.subject = "Flink、Kafka、CDC、状态计算、可靠性、生产排障与项目包装"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Flink, Kafka, CDC, 实时数仓, Exactly-once, 银行风控"
    cover(doc)

    doc.add_heading("先说明：这不是一本术语大纲", level=1)
    add_para(doc, """你已经做过批量数据开发，这是一项优势，不是包袱。离线开发里的源表、分区、调度、重跑、幂等、数据质量、维表和汇总层，在实时系统中仍然存在；只是“一天跑一次”变成“每来一条就推进一点”，调度依赖变成事件驱动，临时中间表变成长期状态，重跑日期分区变成从 Kafka Offset 或 Flink Savepoint 恢复。

这本书不会只写“Watermark 用来处理乱序”然后结束。你会看到交易为什么乱序、Watermark 到底承诺什么、空闲分区为什么卡住窗口、迟到交易怎样补偿、SQL 和 Java 怎样实现、上线后看哪些指标，以及面试官继续追问时怎样回答。所有核心概念都放进同一个银行项目，直到它成为一条能运行、能恢复、能解释、能排障的实时链路。""")
    add_callout(doc, "版本选择不是越新越好",
        "资料校验时 Flink 2.3.0 是最新核心版本，但官方 Kafka Connector 5.0.0 与 Flink CDC 3.6.0 的兼容范围落在 Flink 2.2.x，所以实操统一使用 Flink 2.2.1。生产工作先看组件兼容矩阵，再看“最新”二字。",
        "warning")

    doc.add_heading("全书路线", level=1)
    add_numbers(doc, [
        "批处理与流处理的心智转换", "Kafka 事件账本", "Flink 架构、并行度与 Slot",
        "本地可复现实验环境", "第一条 Flink SQL", "Event Time 与 Watermark",
        "窗口与动态表", "State、Timer 与 TTL", "Checkpoint、Savepoint 与 Exactly-once",
        "去重、幂等和 Upsert", "Join、维表与 CDC", "Java DataStream 风控规则",
        "CEP 行为序列", "实时数仓与流批校准", "反压、倾斜与性能调优",
        "生产部署、升级和监控", "十一类生产故障", "测试、安全与发布门槛",
        "完整作品集", "简历、面试与 45 天路线",
    ])

    chapter(doc, 1, "先换脑子：实时不是把批任务每分钟跑一次",
        "离线作业像每天关门后盘库存；实时作业像收银台每扫一件商品就更新库存。二者都在算账，但时间、状态和恢复方式完全不同。",
        ["用批处理语言解释流处理", "区分有界与无界数据", "判断需求是否值得实时化"],
        [
            ("1.1 把旧经验翻译成新术语",
             """Hive 业务日期对应事件时间窗口，调度依赖对应事件持续到达，中间表对应 Flink Managed State，重跑某天分区对应从 Kafka Offset 或 Savepoint 回放，全量加增量抽取对应 CDC 的 Snapshot 加 Change Log。你不是从零开始，只是要重新理解“什么时候算完”和“记忆放在哪里”。

批量表是有界数据：2026-08-03 分区最终会结束，所以 COUNT(*) 可以等最后一行。交易 Topic 是无界数据：只要系统运行，事件就继续来。对无限流问“总共有多少条”没有天然答案，只能问“截至现在累计多少”或“每五分钟多少”。窗口、业务主键和触发器，就是给无限数据划出可结算边界。"""),
            ("1.2 状态为什么取代中间表",
             """离线 SQL 可以把中间结果写入一张表，下一节点读取。实时程序不能每来一条就扫描历史表，它把“同一账户十分钟内已经转出几次”保存在 State 中。这个 State 长期存在，要跟随 Key 分片、参与 Checkpoint、故障时恢复、扩容时重新分配。普通 Java Map 做不到这些。

实时化也不是越多越好。欺诈告警、设备监控、实时推荐特征适合秒级或分钟级；监管月报最终报送、历史全量重算、模型训练通常适合批量。常见正确方案是流批并存：实时链路负责快，离线链路负责最终核对与大范围回补。"""),
            ("1.3 开工前先问业务",
             """先问允许多大延迟、错误结果代价、迟到数据怎样补、峰值吞吐、状态保存多久、结果需要更新还是只追加。若这些问题没有答案，直接讨论“上 Flink 还是 Spark”只是技术选型表演。业务 SLA 决定 Watermark，恢复目标决定 Checkpoint，审计要求决定事件与规则版本，外部系统承载力决定 Sink。"""),
        ],
        ["为什么无界流上的全局 COUNT 没有自然终点？", "离线重跑与 Kafka 回放最大的外部副作用是什么？", "哪些需求不值得实时化？"],
        "把你手头一个 T+1 作业画成“源→清洗→关联→汇总→输出”，为每段标出实时化后新增的时间、状态和恢复问题。")

    chapter(doc, 2, "Kafka：可回放的事件账本",
        "把 Kafka 想成按主题分类、可以从书签继续阅读的流水账。Flink 是持续读账并计算的人。",
        ["解释 Topic、Partition、Offset 与 Consumer Group", "设计消息 Key", "理解保留与回放"],
        [
            ("2.1 六个词讲透 Kafka",
             """Topic 是一类事件的日志本，例如 transaction_events；Partition 是同一本日志拆成的有序分册；Offset 是分区内页码；Producer 写事件；Consumer 读事件；Consumer Group 让多个实例协作分担分区。分区提供并行度，也规定顺序边界：Kafka 只保证单个分区内部顺序，不保证整个 Topic 全局有序。

若同一账户交易顺序影响风控，应以 account_id 作为 Key，让相同账户进入同一分区。代价是超级活跃账户可能变成热分区。因此 Key 设计永远在业务顺序、数据均匀和并行吞吐之间取舍。"""),
            ("2.2 Offset 不是业务去重键",
             """生产者重试可能把同一个 event_id 写成两条 Kafka 记录，它们拥有不同 Offset；所以 Offset 唯一不能证明业务唯一。反过来，从旧位置回放时同一条记录会再次经过下游，即使 Offset 没变，非幂等 Sink 仍可能插入第二行。业务去重依赖稳定 event_id，落地幂等依赖稳定结果主键。

Retention 决定日志能回放多久，不代表 Flink 状态也必须保存同样久。Kafka 可以保存七天用于灾难恢复，而在线精确去重只保存二十四小时，再用数据库唯一键和离线对账兜底。"""),
            ("2.3 一个合格事件契约",
             """交易事件至少包含 event_id、account_id、customer_id、txn_type、amount、currency、device_id、ip、event_time、ingest_time 和 schema_version。event_time 是业务发生时刻，ingest_time 是平台收到时刻，两者差值反映入口延迟。金额使用 DECIMAL 或“分”为单位的 long，不用二进制浮点数。生产者重试时 event_id 不变，消费者才能识别重复。

事件契约还要写清 Key、可空字段、枚举、时区、兼容策略、敏感字段、错误事件去向和保留期。没有契约的 JSON 只是“碰巧能解析”的字符串。"""),
        ],
        ["为什么同一账户需要相同 Key？", "Offset 为什么不能替代 event_id？", "消费者数多于分区会怎样？"],
        "为 transaction_events 写一份完整事件契约，并制作正常、重复、破损、迟到四类样本。")

    doc.add_heading("交易事件样例", level=2)
    add_code(doc, """{
  "event_id": "txn-20260804-000001",
  "account_id": "A10086",
  "customer_id": "C9001",
  "txn_type": "TRANSFER_OUT",
  "amount": 18888.00,
  "currency": "CNY",
  "device_id": "D-71",
  "ip": "10.20.30.40",
  "event_time": "2026-08-04T09:30:01.123+08:00",
  "ingest_time": "2026-08-04T09:30:02.004+08:00",
  "schema_version": 1
}""")

    chapter(doc, 3, "Flink 集群、作业图、并行度与 Slot",
        "写 SQL 只是入口。排障时必须知道谁协调、谁执行、数据为何卡在某一个算子。",
        ["区分 JobManager 与 TaskManager", "解释并行度与 Slot", "理解 Operator UID"],
        [
            ("3.1 总调度室与生产车间",
             """JobManager 接收作业、生成执行图、申请资源、协调 Checkpoint 和故障恢复，像总调度室。TaskManager 真正执行 Source、Map、Join、Window、Sink，像生产车间。Task Slot 是资源调度单位，不等于线程，也不意味着每个算子独占一套 CPU。

生产中 JobManager 要高可用；TaskManager 可以横向扩展。故障恢复不是全量从头跑，而是让状态回到最近成功 Checkpoint，让可回放 Source 回到同一逻辑位置。"""),
            ("3.2 SQL 怎样变成 Subtask",
             """SQL 或 DataStream 先形成逻辑算子，优化器做过滤下推、投影裁剪和 Join 规划，可链接的算子组成 Operator Chain 以减少网络与序列化，随后每个算子按 Parallelism 展开为多个 Subtask，最后调度到 Slot。

Parallelism 是当前实例数；Max Parallelism 决定 Keyed State 将来可扩展的上限。状态按 Key Group 分片，扩缩容时重新分配。生产状态算子应设置稳定 UID 和显式 Max Parallelism，否则改代码后从 Savepoint 恢复可能无法匹配旧状态。"""),
            ("3.3 为什么加并行度不一定有用",
             """Kafka 只有两个分区时，Source 并行度八意味着六个实例没数据；单个热 Key 仍只能落到一个 Keyed Subtask；同步 JDBC Sink 被一个慢数据库限制时，Source 扩容只会更快制造反压。调优必须沿 Source 到 Sink 找到第一处真正瓶颈，而不是看 CPU 低就把所有并行度翻倍。"""),
        ],
        ["JobManager 和 TaskManager 各负责什么？", "Kafka 两分区、Source 并行度八会怎样？", "状态算子为何要固定 UID？"],
        "打开 Flink Web UI 的 Job Graph，写出每个算子的输入、输出、并行度和可能瓶颈。")

    chapter(doc, 4, "本地环境：可启动不等于可用",
        "本地实验的目标不是模拟生产规模，而是制造可重复证据：能造数据、能看结果、能故障恢复、能观察迟到与重复。",
        ["选出兼容组件", "理解容器网络", "建立验收清单"],
        [
            ("4.1 推荐版本基线",
             """教程采用 Flink 2.2.1、Kafka Connector 5.0.0、Flink CDC 3.6.0。校验日 Flink 2.3.0 虽为最新核心版本，但其 Kafka Connector 页面仍提示没有对应连接器；Connector 5.0.0 与 CDC 3.6.0 明确支持 2.2.x。搭环境要把 Flink、Connector、CDC、JDK、数据库版本放在一张矩阵中核对。

Kafka、PostgreSQL、Flink JobManager/TaskManager、Prometheus、Grafana 可用 Docker Compose 组合。容器访问 Kafka 使用 kafka:9092，宿主机客户端使用外部 Listener；容器里的 localhost 指向容器自己，这是最常见的“端口明明开着却连不上”。"""),
            ("4.2 七项验收",
             """容器健康只是第一项。还要验证 Topic 可生产消费、SQL Client 能加载 Connector JAR、Flink Job 为 RUNNING 且记录计数变化、合法与脏数据分别到正确出口、重复 event_id 不产生第二条结果、乱序与极晚数据按策略处理、kill TaskManager 后能从 Checkpoint 恢复。

Checkpoint、Savepoint、Kafka 和 PostgreSQL 数据需要持久卷。停止环境时不要随手删除卷，否则你测试到的不是恢复而是全新初始化。Secret 用环境变量或 Secret 注入，不写进 Git。"""),
            ("4.3 README 应成为操作合同",
             """README 要给出版本、端口、启动、健康检查、造数、查看结果、停止和常见故障。每个“成功”都配可观察证据：Flink UI 中的 Job ID、Kafka 输出样例、PostgreSQL 行数、Checkpoint ID。别人从零照 README 能跑通，才算环境真正交付。"""),
        ],
        ["容器内为何不能用 localhost 访问 Kafka？", "为什么核心最新不代表组合可用？", "证明环境可用需要哪些证据？"],
        "把七项验收做成 verify.ps1；脚本返回非零即代表某条业务链未通过。")

    chapter(doc, 5, "第一条 Flink SQL：持续查询",
        "DDL 定义如何解释一条流，INSERT SELECT 提交后持续消费，而不是执行完退出。",
        ["创建 Kafka Source/Sink", "理解持续作业", "设计脏数据出口"],
        [
            ("5.1 DDL 没有搬运数据",
             """CREATE TABLE 只是声明去哪里读、字节怎样解释成列、业务时间是哪一列、愿意等待多少乱序。METADATA 暴露 partition 和 offset 后，错误才能回到原始记录。scan.startup.mode 决定首次从最早、最新、时间戳还是 Consumer Group 已提交位置开始，改错会造成漏读或大范围重放。

json.ignore-parse-errors 只能避免作业被单条坏消息打死，不能代表错误可以消失。生产链路应把 Parse Error、Missing Field、Unknown Schema 分流到 dirty_transaction_events，带 raw_payload、error_code、partition、offset、发现时间，并有人监控和回补。"""),
            ("5.2 持续运行意味着持续治理",
             """批量 INSERT SELECT 的源会读完，实时源不会结束，因此作业长期 RUNNING。修改 SQL 不是保存文件后自动生效，而是一次有状态发布：保存旧状态、停旧作业、发布新版本、从 Savepoint 恢复或明确放弃状态。作业名字、版本、配置、规则和 Connector 必须可追踪。"""),
        ],
        ["Flink DDL 是否真的创建物理表？", "为什么要保留 partition/offset？", "实时 INSERT SELECT 何时结束？"],
        "发送合法、金额为负、JSON 破损三条消息，证明它们分别进入正确出口。")

    doc.add_heading("Kafka Source SQL", level=2)
    add_code(doc, """CREATE TABLE transaction_events (
  event_id STRING,
  account_id STRING,
  customer_id STRING,
  txn_type STRING,
  amount DECIMAL(18, 2),
  currency STRING,
  event_time TIMESTAMP_LTZ(3),
  ingest_time TIMESTAMP_LTZ(3),
  kafka_partition INT METADATA FROM 'partition' VIRTUAL,
  kafka_offset BIGINT METADATA FROM 'offset' VIRTUAL,
  WATERMARK FOR event_time AS event_time - INTERVAL '5' SECOND
) WITH (
  'connector' = 'kafka',
  'topic' = 'transaction_events',
  'properties.bootstrap.servers' = 'kafka:9092',
  'properties.group.id' = 'bank-risk-v1',
  'scan.startup.mode' = 'earliest-offset',
  'format' = 'json',
  'json.ignore-parse-errors' = 'true'
);""")

    chapter(doc, 6, "Event Time 与 Watermark：时间听谁的",
        "移动网络、重试、跨机房和分区负载会制造乱序。Watermark 是系统对业务时间进度的估计。",
        ["区分三种时间", "解释 Watermark", "制定迟到策略"],
        [
            ("6.1 三种时间",
             """Event Time 是事件真实发生时刻，适合重放与稳定结果；Ingestion Time 是进入 Flink Source 的时刻；Processing Time 是算子机器处理时的当前时钟。风控和交易指标通常以 Event Time 为主，因为历史回放时仍能得到相同窗口归属。Processing Time 简单且低延迟，但机器变慢或重跑会改变结果。

Watermark=09:30:00 等价于“我认为 09:30:00 之前的数据基本到齐，可以推进对应计算”。它不是绝对事实，而是延迟与完整性的约定。等待五秒可容忍常见乱序，但窗口至少多等约五秒；等十分钟更完整，却可能失去告警价值。"""),
            ("6.2 并行分区的最小值陷阱",
             """每个 Kafka 分区独立生成 Watermark，下游使用最小值。某分区长期没有数据时，它的 Watermark 不推进，整个窗口可能不出结果。table.exec.source.idle-timeout 可把空闲分区标记为 Idle，但要监控误判：真正网络故障的慢分区不能被悄悄忽略。

迟到要分三档：Watermark 前到达的轻度乱序正常计算；窗口触发后但仍在允许迟到范围的事件更新旧结果；超过范围的极晚事件进入 Side Output 或补偿 Topic。银行风控可先发及时告警再补充更新，监管最终数字由离线对账兜底。"""),
            ("6.3 参数必须来自分布而非感觉",
             """记录 ingest_time-event_time 的延迟直方图，分时段观察 P95、P99 和极端值，再结合业务 SLA 选择 Watermark。若大多数事件 800ms 内到达、P99 为 4 秒，五秒可能合理；若月末上游会补发两小时历史数据，不能把 Watermark 改成两小时，应单独走补偿链路。"""),
        ],
        ["Event Time 为何适合重跑？", "空闲分区为何卡住窗口？", "Watermark 越保守越好吗？"],
        "构造十条乱序事件，画出事件时间、到达时间、Watermark，并判定正常、更新或极晚。")

    chapter(doc, 7, "窗口、动态表与 Changelog",
        "窗口给无限流划账期；动态表解释为什么聚合结果会更新、撤回或覆盖。",
        ["选择窗口", "理解 Changelog", "匹配 Sink 语义"],
        [
            ("7.1 窗口先看业务边界",
             """Tumbling Window 固定且不重叠，适合每五分钟指标；Sliding/Hopping Window 固定长度并按步长滑动，适合“过去十分钟、每分钟刷新”；Session Window 按一段无活动时间闭合，适合会话。规则“十分钟内五次转出”若使用十分钟滚动窗口，会漏掉跨 09:30 边界的行为；滑动窗口更符合业务，但步长越小，状态和计算越多。

count、sum、min、max 优先使用增量聚合，不要把窗口内全部交易塞进 ListState 到期遍历。只有需要完整明细、排序或复杂模式时才保存事件，并设置 TTL、上限和异常大 Key 保护。"""),
            ("7.2 动态表为什么会撤回",
             """账户 A 当前累计 100 元时输出 +I(A,100)，新交易到达后旧结果失效，可能输出 -U(A,100) 与 +U(A,120)，或以 Upsert 形式用主键覆盖为 120。若下游只会 Append，就会留下 100 与 120 两行，产生重复含义。

普通 Kafka Topic 适合追加事实；Upsert Kafka 和 JDBC Upsert 需要主键；文件系统更偏 Append 和后续 Compact。PRIMARY KEY NOT ENFORCED 表示 Flink 信任主键语义但不替你检查，真实数据库仍应有唯一约束。"""),
            ("7.3 Window TVF",
             """Flink SQL 的 HOP、TUMBLE、CUMULATE 等 Table-Valued Function 会显式产生 window_start 与 window_end，便于下游把窗口边界作为幂等结果键。业务告警不能只存账户和金额，还应存窗口、Watermark/处理时刻、规则版本与证据。"""),
        ],
        ["为什么滚动窗口会漏跨边界行为？", "聚合更新为何不能盲写 Append Sink？", "NOT ENFORCED 是否代表数据库无需唯一键？"],
        "用 TUMBLE 与 HOP 计算同一组跨边界交易，对比结果和状态成本。")

    doc.add_heading("十分钟滑动聚合 SQL", level=2)
    add_code(doc, """SELECT account_id, window_start, window_end,
       COUNT(*) AS txn_count, SUM(amount) AS total_amount
FROM TABLE(
  HOP(TABLE transaction_events, DESCRIPTOR(event_time),
      INTERVAL '1' MINUTE, INTERVAL '10' MINUTE)
)
WHERE txn_type = 'TRANSFER_OUT'
GROUP BY account_id, window_start, window_end;""")

    chapter(doc, 8, "State、Timer 与 TTL：实时开发的分水岭",
        "状态是流处理的记忆；Timer 是未来给自己设置的闹钟；TTL 是防止记忆无限长大的兜底。",
        ["选择状态类型", "使用事件时间 Timer", "控制状态规模"],
        [
            ("8.1 State 不是普通 HashMap",
             """普通 Map 不会随 Key 分区，不会自动参与 Checkpoint，TaskManager 崩溃可能全丢，扩缩容也不会重分配。Managed State 由 Flink 管理，跟随 Keyed Stream 分片并恢复。ValueState 保存一个值，ListState 保存集合，MapState 保存小映射，AggregatingState 保存增量聚合，BroadcastState 把小规则表复制到所有并行实例。

状态使用之前通常先 keyBy(account_id)。此后每个 Key 看到自己的逻辑状态，同一个函数实例服务很多 Key。把 account_id 忘记 keyBy 或 Key 选错，规则就会把不同账户混在一起。"""),
            ("8.2 Timer 不是 sleep",
             """十分钟后清理历史交易时注册 Event-time Timer，让 Watermark 到达指定时刻后回调 onTimer。定时器与状态一起快照和恢复；开线程 sleep 会阻塞、不可恢复且无法随 Key 管理。Processing-time Timer 依赖机器时钟，历史回放可能得出不同结果。

TTL 不保证在过期毫秒立即物理删除，所以不能替代精确业务定时器。正确做法常是 Timer 精确移除十分钟前事件，TTL 设成三十分钟作为异常兜底。TTL 太短会误删业务记忆，太长会扩大磁盘、Checkpoint 和恢复压力。"""),
            ("8.3 状态容量要先估算",
             """状态大小约等于 Key 数量×每 Key 条目×序列化后字节×开销系数。峰值 Key、最长保留期、迟到与恢复时间都要进入估算。先用紧凑结构和增量聚合降低条目，再选择 HashMap、EmbeddedRocksDB 或其他后端；不要指望换 RocksDB 自动修复无限增长的设计。"""),
        ],
        ["普通 Map 为什么不可靠？", "Event-time Timer 与 Processing-time Timer 差异是什么？", "TTL 为什么不能代替精准删除？"],
        "设计“十分钟五次转出”的 Key、State、Timer、清理、迟到和恢复行为。")

    chapter(doc, 9, "Checkpoint、Savepoint 与端到端 Exactly-once",
        "Checkpoint 是自动恢复快照，Savepoint 是人为迁移档案；Exactly-once 是整条链路的协作结果。",
        ["解释 Barrier", "区分两类快照", "识别一致性边界"],
        [
            ("9.1 Barrier 怎样拍一致画面",
             """Checkpoint Coordinator 向 Source 注入 Barrier。Barrier 随记录向下游流动，多输入算子在对齐模式下等待同编号 Barrier 到齐，再快照状态并继续传递。Kafka Offset、窗口累计、去重状态和 Sink 事务位置因此对应同一逻辑时刻。故障后它们一起回退并重放。

严重反压时 Barrier 排在大量 Buffer 之后，Start Delay 与 Alignment Time 升高。Unaligned Checkpoint 把在途数据也纳入快照，使 Barrier 越过积压，但会增加状态存储 I/O。它能改善反压下的快照时间，不能修复慢 Sink 的业务根因。"""),
            ("9.2 Checkpoint 与 Savepoint",
             """Checkpoint 周期自动触发，面向故障恢复和系统管理；Savepoint 由人主动触发，面向版本升级、迁移、扩缩容和回滚。升级前 Stop-with-Savepoint，保持所有状态算子 UID 稳定，新版本从该路径恢复。更改状态数据类型、删除有状态算子或改变 UID 都可能破坏兼容。

端到端 Exactly-once 还要求 Source 可回放、Sink 支持事务或幂等。Kafka 到 Flink 状态即使完全一致，若用普通 INSERT 写没有唯一键的 PostgreSQL 表，恢复重放仍会产生第二行。对 HTTP 外部调用要带稳定 idempotency_key；对 JDBC 用业务主键 Upsert；对 Kafka Exactly-once Sink 使用事务并与 Checkpoint 提交。"""),
            ("9.3 参数来自恢复目标",
             """Checkpoint Interval 表示故障时可能重放的时间范围和 Sink 事务可见节奏，不是越短越安全。频率太高会争抢 CPU、网络和存储。应同时看成功率、duration、start delay、alignment、size 与恢复时间，用压测找到满足 RPO/RTO 的区间。"""),
        ],
        ["Barrier 为什么需要对齐？", "Unaligned Checkpoint 解决与不解决什么？", "无主键 JDBC INSERT 为什么破坏端到端一致性？"],
        "kill 一个 TaskManager，记录 Checkpoint ID、恢复时间、Lag 和结果是否重复。")

    chapter(doc, 10, "去重、幂等与可回放结果键",
        "重复来自生产重试、Kafka 写入、Flink 回放、Sink 超时和人工补数，没有一个开关能消灭全部重复。",
        ["设计业务去重键", "分层防重复", "估算去重状态"],
        [
            ("10.1 三层防线",
             """入口层要求生产者重试使用同一 event_id；计算层按 event_id KeyBy，以 ValueState+TTL 过滤短期重复；落地层使用稳定 alert_id 唯一约束与 Upsert，抵抗恢复、回放和超时重试。每层解决不同来源，不能只在 Flink 做一次 distinct 就宣布完成。

alert_id 应可重算，例如 hash(rule_code、account_id、window_start、rule_version)。随机 UUID 每次回放都会变化，数据库无法识别同一业务结果。窗口更新时主键保持相同，Value 覆盖为新结果。"""),
            ("10.2 去重状态不是免费午餐",
             """每天十亿 event_id 保存七天会非常庞大。TTL 由最大在线重复窗口和回放策略决定，不必机械等于 Kafka Retention。可采用短期精确状态、数据库唯一键、离线对账的分层方案。Bloom Filter 省空间但有误判，在资金类精确过滤中不能未经评估直接使用。"""),
            ("10.3 回放必须先审计外部副作用",
             """从 earliest-offset 重放前列出所有 Sink：Kafka Topic 是否会追加第二份、数据库是否 Upsert、告警短信是否再次发送、外部接口是否幂等。高风险外部动作应从“计算结果 Topic”由独立服务消费，消费服务记录 alert_id 和发送状态，避免 Flink 恢复时直接重复通知客户。"""),
        ],
        ["列出四种重复来源。", "为什么随机 UUID 不适合告警幂等键？", "回放前为什么要盘点外部副作用？"],
        "注入相同 event_id 三次、故障恢复一次、Sink 超时一次，断言最终只有一条业务结果。")

    chapter(doc, 11, "Join、维表与 Flink CDC",
        "交易事实要关联客户等级、账户状态和动态规则；真正的难点是拿哪个时刻的维度，以及如何避免外部数据库拖慢作业。",
        ["选择 Join 方式", "理解 CDC Snapshot+Log", "管理 Schema Evolution"],
        [
            ("11.1 四种关联方式",
             """Stream-Stream Interval Join 适合两边都是事件流，必须限制时间范围，否则状态无限增长。Temporal Table Join 让事实按业务时刻关联当时版本的维表。Lookup Join 每条按 Key 查询外部库，必须考虑缓存、异步、超时和限流。BroadcastState 适合很小且要快速推送的规则集。

每秒两万笔交易同步 SELECT PostgreSQL，会把 Flink 吞吐变成数据库往返速度，也可能打垮业务库。Lookup Cache 和 Async I/O 是改进，稳定方案常是通过 CDC 将维表变化同步进 Flink State 或低延迟 KV 存储。"""),
            ("11.2 CDC 不是定时查 update_time",
             """CDC 读取数据库变更日志中的 INSERT、UPDATE、DELETE。首次启动先做一致性 Snapshot 建基线，再从记录的日志位置持续消费。PostgreSQL 需要逻辑复制权限与稳定 slot.name；复制槽会阻止未消费 WAL 清理，所以 CDC 作业长期停机可能撑爆数据库磁盘。

Flink CDC 3.6 支持 PostgreSQL Schema Evolution 等能力，但“连接器能识别 DDL”不等于所有下游都能无损接受。新增可空列通常容易；删除、重命名、缩窄类型、改变主键要双写、灰度和回滚计划。"""),
            ("11.3 规则必须版本化",
             """risk_rule_config 至少包含 rule_code、rule_version、threshold、effective_from、enabled、updated_at。告警保存命中的 rule_version 和实际阈值。若只保存“当前规则”，一个月后没人能解释历史告警为何触发。处理时最新与事件时有效是两种语义，需求必须明确。"""),
        ],
        ["无边界 Stream-Stream Join 为什么危险？", "同步 Lookup 怎样制造反压？", "复制槽为何可能撑满 PostgreSQL？"],
        "比较 JDBC Lookup、CDC Temporal Table、BroadcastState 的延迟、压力、一致性与复杂度。")

    doc.add_heading("Temporal Join 示意", level=2)
    add_code(doc, """SELECT t.event_id, t.account_id, t.amount, t.event_time,
       a.account_status, a.risk_level
FROM transaction_events AS t
LEFT JOIN account_profile
FOR SYSTEM_TIME AS OF t.event_time AS a
ON t.account_id = a.account_id;""")

    chapter(doc, 12, "DataStream Java：有状态风控规则",
        "SQL 覆盖常规 ETL；精细清理、动态定时器、侧输出和复杂状态逻辑需要 DataStream API。",
        ["使用 KeyedProcessFunction", "维护时间状态", "输出可解释证据"],
        [
            ("12.1 规则设计先于代码",
             """规则：同一账户十分钟内转出至少五笔且总额不低于五万元。Key 是 account_id；状态保存时间范围内的轻量交易或分桶聚合；Event-time Timer 清理过期数据；TTL 作为兜底；重复在上游先按 event_id 过滤；极晚数据进补偿流。

最容易写出的实现是每条读取整个 ListState、过滤、再整体写回，教学可用，流量大时成本很高。生产可使用 MapState 按时间桶保存 count/sum，或维护紧凑队列与 AggregateState。金额使用分为单位的 long 或 BigDecimal，不能用 double。"""),
            ("12.2 告警必须有证据",
             """输出 alert_id、account_id、rule_code、rule_version、window_start/end、阈值、实际次数、实际金额、相关 event_id、维表版本、event_time 和 processing_time。这样业务能复核、开发能回放、审计能解释。只有“高风险=0.92”的结果无法说明来自哪些事实。

Late Event、Parse Error、Unknown Schema、Rule Error 可以用 Side Output 分流，但每条侧输出都要有消费者、保留、告警和回补。没人看的 DLQ 只是延迟的数据丢失。"""),
            ("12.3 稳定 UID 与状态兼容",
             """每个有状态算子设置 uid，例如 rapid-transfer-rule-v1。升级时可以改判断逻辑，但不能随意改变状态类型；需要变更时设计状态迁移、双跑或从新状态启动。代码 Review 除了业务条件，还要检查 Key、时间语义、状态上限、Timer 数量、序列化和异常路径。"""),
        ],
        ["规则前为何 keyBy(account_id)？", "ListState 全量读写的性能问题是什么？", "告警证据为何是业务功能？"],
        "实现规则并覆盖不足次数、金额不足、同时满足、跨边界、重复、迟到六组测试。")

    doc.add_heading("KeyedProcessFunction 骨架", level=2)
    add_code(doc, """public final class RapidTransferRule
    extends KeyedProcessFunction<String, Transaction, RiskAlert> {

  private transient MapState<Long, BucketStat> minuteBuckets;

  @Override
  public void processElement(
      Transaction tx, Context ctx, Collector<RiskAlert> out) throws Exception {
    long minute = tx.eventTimeMillis() / 60_000L;
    BucketStat stat = minuteBuckets.get(minute);
    if (stat == null) stat = new BucketStat();
    stat.add(tx.eventId(), tx.amountFen());
    minuteBuckets.put(minute, stat);

    long cleanupAt = (minute + 11) * 60_000L;
    ctx.timerService().registerEventTimeTimer(cleanupAt);

    WindowEvidence evidence = aggregateLastTenMinutes(minuteBuckets, minute);
    if (evidence.count() >= 5 && evidence.sumFen() >= 5_000_000L) {
      out.collect(RiskAlert.from(tx, evidence, "RAPID_TRANSFER_V1"));
    }
  }

  @Override
  public void onTimer(
      long timestamp, OnTimerContext ctx, Collector<RiskAlert> out) throws Exception {
    minuteBuckets.remove(timestamp / 60_000L - 11);
  }
}
// 教学骨架：生产实现需补去重、迟到、状态 TTL、稳定 UID、异常和测试。""")

    chapter(doc, 13, "CEP：从单点阈值到行为序列",
        "复杂事件处理识别“试探→多笔转出→快速提现”这样的顺序模式。",
        ["理解 Pattern", "控制模式状态", "判断何时不用 CEP"],
        [
            ("13.1 Pattern 的业务含义",
             """begin 定义起点，next 要求严格紧邻，followedBy 允许中间有其他事件，timesOrMore 会制造多个匹配分支，within 限制整条模式必须在指定时间内完成。每个词都对应业务假设，不是 API 花样。乱序、重复和迟到会改变匹配，因此仍需 Event Time、Watermark 和去重。

模式可能重叠命中，需设置跳过策略和告警合并，否则同一事件序列产生告警风暴。超时的部分匹配也可能有业务价值，例如试探后未提现，可输出到超时侧流用于观察。"""),
            ("13.2 什么时候不要用 CEP",
             """简单次数、金额、唯一设备数用 SQL 或 State 更清晰。规则极多且业务频繁配置时，可考虑规则引擎或将规则 DSL 编译成算子。不要为了简历出现 CEP，把一个 count>=5 的规则写成难以理解和测试的模式。技术价值来自解决问题，不来自英文缩写数量。"""),
        ],
        ["next 与 followedBy 差异是什么？", "within 为什么必需？", "什么规则用普通状态更合适？"],
        "为“异地登录后五分钟内大额转账”画序列，写明乱序、重复、缺失和超时处理。")

    chapter(doc, 14, "实时数仓：分层、口径与流批校准",
        "实时数仓仍需要 ODS、DWD、DWS、ADS、口径、血缘与回补，否则只会产生更多 Topic。",
        ["设计实时分层", "解释流批差异", "建立对账"],
        [
            ("14.1 分层不是照搬名字",
             """ODS 尽量保留原始事件和 Kafka 元数据；DWD 做解析、标准化、去重、维度补齐；DWS 形成账户五分钟交易等可复用主题汇总；ADS 面向告警和看板。每层要有明确消费方、Key、Changelog 模式、保留和 SLA，不是每过一个 SELECT 就新建一层。

核心指标口径记录业务定义、维度、Event Time 列、窗口、Watermark、迟到策略、去重键、维表时态、规则版本和负责团队。否则“交易金额”在不同作业中可能一个含冲正、一个不含。"""),
            ("14.2 为什么实时与离线对不上",
             """实时可能截断极晚数据，离线 T+1 包含；实时关联当前维表，离线按历史拉链；去重范围与空值处理不同；上游修复后离线重跑，实时未回补；实时规则升级后历史结果仍按旧版产生。发现差异时应下钻到 event_id，而不是只看总额差了多少。

每天用离线权威结果校准实时结果，分别比较条数与金额，按机构、产品和时间桶定位差异。差异超阈值自动生成待处理项，并明确是实时补偿、离线更正还是口径问题。"""),
        ],
        ["ODS 为何保留 partition/offset？", "列出三种流批不一致原因。", "口径文档至少记录什么？"],
        "选择熟悉的离线指标，写一版实时口径和 T+1 校准方案。")

    chapter(doc, 15, "反压、倾斜与性能调优",
        "调优像查水管：数据向下游流，压力从慢点向上游传。先找瓶颈，再动参数。",
        ["定位反压", "解决热 Key", "选择状态后端"],
        [
            ("15.1 反压阅读顺序",
             """先看 Kafka Lag 与端到端延迟是否持续上升；在 Web UI 找 backPressuredTimeMsPerSecond 高的上游；再向下游找 busyTimeMsPerSecond 高、吞吐低的第一处瓶颈；最后判断 CPU、I/O、GC、网络、序列化、外部系统或热 Key。Source 高反压通常不代表 Source 自己慢，而是下游消费不过来。

同步外部查询改 Async I/O，并设置容量、超时、重试和熔断；JDBC Sink 用合理批量、并发和 Upsert；减少重复 JSON 解析、对象创建和大记录；聚合用增量状态。Unaligned Checkpoint 只改善反压下 Barrier 传播，不能让慢数据库变快。"""),
            ("15.2 热 Key",
             """一个账户或机构占 80% 流量时，一个 Subtask 满载，其余很闲。增加并行度不能拆开同一个 Key。若业务允许，可增加 Key 粒度；聚合可用随机盐值做两阶段汇总；热点名单可单独分流；时间桶可以拆大 Key。每种方案都要检查是否破坏顺序、窗口和精确性。

Kafka 分区数决定 Source 最大有效并行度，Sink 并发受下游承载力限制。资源规划不能只看平均吞吐，要压峰值、突发、恢复追赶和热点。"""),
            ("15.3 状态后端",
             """HashMapStateBackend 工作状态在 JVM Heap，访问快但受堆与 GC 限制；EmbeddedRocksDBStateBackend 使用本地磁盘和内存缓存，支持更大状态及增量快照，但有序列化与磁盘成本；ForSt 等能力要按目标版本和部署环境验证。无论选什么，状态无 TTL 和大对象设计都会出问题。"""),
        ],
        ["Source 高反压应向哪里找？", "为何加并行度不能解决单个热 Key？", "HashMap 与 RocksDB 怎样取舍？"],
        "制造 80% 流量集中于一个账户，对比普通 keyBy 与两阶段聚合的 Subtask Busy。")

    chapter(doc, 16, "生产部署、升级、监控与安全",
        "本地 RUNNING 只是开始。生产作业必须可升级、可回滚、可观测并隔离集群执行入口。",
        ["选择部署模式", "执行 Savepoint 升级", "建立监控与安全基线"],
        [
            ("16.1 部署与升级",
             """共享开发可用 Standalone Session Cluster，企业大数据平台常见 YARN，云原生团队常用 Kubernetes Application Mode。官方生产建议倾向 Application Mode 的应用隔离。选择要匹配团队现有运维能力，不为时髦增加一个没人会维护的平台。

安全升级顺序：核对 Flink/Connector/JDK/状态兼容；Stop-with-Savepoint；记录 Job ID、路径、规则版本；发布新 JAR/镜像并保持 UID；从 Savepoint 启动；验证恢复、Lag 追平、结果和 Checkpoint；异常则回滚版本从原 Savepoint 恢复。"""),
            ("16.2 监控不是只看 RUNNING",
             """流量看 recordsIn/Out、bytes 和 Kafka Lag；延迟看入口、处理、端到端 P95/P99；负载看 busy/backpressured/idle；Checkpoint 看成功率、duration、start delay、alignment、size；状态看 state size、RocksDB 和 TTL；质量看脏数据、迟到、重复和规则命中；外部系统看 Sink 延迟、连接池、数据库锁和 CDC WAL。

阈值要对应行动。例如 Lag 超阈值先确认流量突增还是吞吐下降；Checkpoint 连续失败要在可恢复窗口耗尽前升级告警；迟到率突然变化可能是上游时钟或批量补发。"""),
            ("16.3 安全与治理",
             """Flink 可执行用户代码，Web UI 和 REST API 不应暴露公网。使用内网、TLS、认证、RBAC 与网络策略；Kafka 使用 ACL 和加密；数据库账号最小权限；Secret 不进 Git 或日志。账户、设备、IP 属于敏感数据，日志与告警证据要脱敏、授权和限期保存。"""),
        ],
        ["Application Mode 的优势是什么？", "升级前为何固定 UID？", "Checkpoint 成功是否代表业务健康？"],
        "写一份发布 Runbook：前置检查、Savepoint、启动、观察、回滚条件、责任人。")

    chapter(doc, 17, "十一类真实故障：先证据，后动作",
        "生产排障最怕先重启。重启可能暂时清掉症状，也会丢掉最关键的指标和现场。",
        ["建立排查树", "避免危险操作", "完成复盘"],
        [
            ("17.1 症状到证据",
             """Kafka Lag 涨：看算子吞吐、busy/backpressure，常见为慢 Sink、分区不足、热 Key。Checkpoint 超时：看 start delay、alignment、state size 与存储 I/O。结果重复：查 event_id、恢复时间、Sink 主键。窗口不出数：看 Watermark 与空闲分区。迟到暴增：看入口延迟分布、上游发布和时钟。

单 Subtask 100%：看 Key 分布。TaskManager OOM：区分 Heap、Direct、Managed、RocksDB 和网络内存，检查大对象与无 TTL 状态。PostgreSQL 压力高：查 Lookup QPS、慢 SQL、连接池、JDBC 批次。CDC WAL 撑盘：查复制槽消费位置。升级不能恢复：查 UID 和状态类型。流批不一致：下钻 event_id、维表时态与口径版本。"""),
            ("17.2 一个完整反压案例",
             """现象：上午十点 Lag 上升，Source 高反压。沿图向下发现 JDBC Sink busy 接近 100%，批次只有二十条，数据库延迟从 8ms 升到 60ms。临时限流非关键流量并增大安全批次；永久修复为主键 Upsert、索引优化、合理并发，必要时先写 Kafka 再异步落库。修复后验证 Lag 能追平、Checkpoint 恢复、数据库无锁等待，并补峰值压测。

危险操作包括：未留证就反复重启；随意改 group.id 或删除 Consumer Group；未确认数据位置就删除复制槽；从旧 Savepoint 恢复到非幂等 Sink；为追求快照速度直接关闭 Exactly-once；为让窗口出数把 Watermark 调得过激。"""),
            ("17.3 复盘写什么",
             """复盘包括影响范围、发现方式、时间线、直接原因、系统性根因、临时止血、永久修复、监控缺口、测试缺口、责任人与完成时间。不要写“某同事操作失误”就结束，要问为什么权限、校验、回滚和告警没有挡住。"""),
        ],
        ["Source 反压高为何不能直接判定 Source 慢？", "CDC 停机为何数据库磁盘仍增长？", "重启前保存哪些证据？"],
        "选一个批处理故障，改写为实时故障复盘：影响、证据、根因、修复和预防。")

    chapter(doc, 18, "测试、质量与上线门槛",
        "实时作业永远不结束，不能靠等它跑完再核对。测试必须控制时间、乱序、状态与故障。",
        ["设计四层测试", "构造边界数据", "定义上线门槛"],
        [
            ("18.1 四层测试",
             """第一层是纯函数单测：解析、金额、规则判断，不启动 Flink。第二层用算子 Harness 控制 Event Time、Watermark、Timer 和 State。第三层用 MiniCluster 或 Testcontainers 验证真实序列化、Kafka、PostgreSQL 与 Checkpoint。第四层做端到端故障演练：kill TaskManager、重启、回放、升级、回滚。

必测数据包括正常、边界、空值、非法枚举、金额精度；轻度乱序、允许迟到、极晚；重复 event_id、同业务不同 event_id；热 Key、空闲分区、峰值突发；维表更新和 Schema 变更；Checkpoint 中断与 Sink 超时。"""),
            ("18.2 上线门槛",
             """至少准备兼容矩阵、事件契约、质量规则、状态估算、Checkpoint 压测、容量余量、告警、Dashboard、Savepoint 演练、回滚方案、安全评审和值班 Runbook。个人作品集可把部分写成“生产化设计”，但必须区分已实现、已测试和计划，不能声称在真实银行生产上线。"""),
            ("18.3 数据质量不是脏数据计数",
             """质量规则要可行动：schema_invalid_rate、required_field_null_rate、duplicate_rate、late_rate、unknown_currency、amount_outlier、source_to_sink_count_gap。每个指标有阈值、负责人、原始样本定位和处理策略。只把错误打日志、没人响应，不叫质量治理。"""),
        ],
        ["为什么 Processing Time 测试不稳定？", "算子测试怎样推进 Event Time？", "作品集哪些能力只能写设计？"],
        "为 RapidTransferRule 写 Given-When-Then 表，覆盖至少八种输入序列。")

    add_chapter_intro(doc, "第 19 章　完整作品集：银行实时交易风控与指标平台",
        "这一章把所有零件装成一个可展示项目。目标不是伪造生产规模，而是用可复现实验展示实时系统关键难题。",
        ["搭建端到端链路", "形成可验证证据", "分里程碑交付"])
    add_callout(doc, "项目一句话",
        "以 Kafka 接收银行交易事件，Flink SQL 完成清洗、事件时间聚合与维表关联，DataStream 实现去重、状态定时器和风控规则，Flink CDC 同步 PostgreSQL 动态配置，结果写入 Kafka 与 PostgreSQL，并用 Prometheus/Grafana 监控延迟、反压和 Checkpoint。",
        "teal")
    doc.add_heading("19.1 业务验收", level=2)
    add_numbers(doc, [
        "合法交易进入 DWD，非法记录进入 DLQ 且能定位原始 Offset。",
        "相同 event_id 重复三次只影响结果一次。",
        "乱序五秒内进入正确窗口，极晚事件进入补偿 Topic。",
        "PostgreSQL 更新阈值后经 CDC 生效，告警带 rule_version。",
        "规则输出次数、金额、窗口和 event_id 证据。",
        "TaskManager 故障后从 Checkpoint 恢复且结果不重复。",
        "Dashboard 显示 Lag、P95 延迟、反压、Checkpoint 和脏数据率。",
    ])
    doc.add_heading("19.2 架构图（文字版）", level=2)
    add_code(doc, """交易系统 / 数据生成器
    → Kafka transaction_events
    → Flink ODS/DWD ──错误──→ dirty_transaction_events
    → 去重 + Watermark + 标准化
       ├→ Flink SQL 窗口指标
       └→ DataStream State/Timer/CEP 风控
PostgreSQL risk_rule_config ──Flink CDC──→ 动态规则状态
    → Kafka risk_alerts + PostgreSQL Upsert + OLAP 指标
    → Prometheus / Grafana / 告警处理页面""")
    doc.add_heading("19.3 推荐仓库结构", level=2)
    add_code(doc, """bank-realtime-risk-platform/
├─ README.md                 # 架构、启动、验收、真实边界
├─ docker-compose.yml        # Kafka/Postgres/Flink/监控
├─ sql/                      # Source、DWD、窗口、Sink
├─ flink-job/                # Java State/Timer/CEP
├─ cdc/risk-rule.yaml
├─ generator/                # 正常、重复、乱序、热点
├─ monitoring/               # Prometheus/Grafana
├─ tests/                    # 单元、集成、故障演练
├─ docs/
│  ├─ event-contract.md
│  ├─ metric-definitions.md
│  ├─ architecture-decisions.md
│  └─ incident-runbook.md
└─ scripts/start.ps1 stop.ps1 verify.ps1""")
    doc.add_heading("19.4 三条规则与证据", level=2)
    add_table(doc, ["规则", "实现", "输出证据"], [
        ["十分钟快速转出", "Keyed State + Event-time Timer", "次数、金额、event_id、窗口"],
        ["五分钟大额累计", "Flink SQL HOP Window", "阈值、实际金额、Watermark"],
        ["试探→转出→提现", "CEP Pattern", "匹配序列、规则版本、超时"],
    ], [2.1, 2.2, 2.2])
    doc.add_heading("19.5 七个里程碑", level=2)
    add_numbers(doc, [
        "M1：Kafka→Flink SQL→Kafka，跑通合法与脏数据。",
        "M2：Event Time、Watermark、窗口、极晚侧输出。",
        "M3：Java 去重与快速转出规则，补 Harness 测试。",
        "M4：PostgreSQL CDC 动态规则与 Temporal/Broadcast 关联。",
        "M5：JDBC Upsert、Checkpoint 故障恢复与幂等验证。",
        "M6：Prometheus/Grafana、慢 Sink 与热点 Key 演练。",
        "M7：README、架构决策、压测报告、演示视频和讲稿。",
    ])
    doc.add_heading("19.6 数字必须测量", level=2)
    add_para(doc, """不要虚构“每秒十万条、延迟五十毫秒”。记录机器配置、事件大小、Kafka 分区、Flink 并行度、状态规模、Checkpoint 间隔和 Sink，再测稳定吞吐、P50/P95/P99 端到端延迟、最大 Lag、Checkpoint P95、故障恢复、重复率和迟到率。简历数字来自这张报告，并标注本地压测。

作品集真正加分的是证据链：一键启动、自动验收、Job Graph、Dashboard、故障前后对比、测试报告、架构决策和限制说明。一个诚实可复现的小项目，比无法解释的“亿级平台”更可信。""")
    add_self_check(doc, ["项目真实难点是什么？", "怎样证明故障恢复正确？", "哪些数字能写进简历？"],
        "先完成 M1 并提交代码、README 与验收输出；下一步写 Issue，不要一次复制所有组件。")

    add_chapter_intro(doc, "第 20 章　简历、面试与 45 天执行路线",
        "包装不是夸大，而是把业务、职责、技术选择、难题、证据与边界组织清楚。",
        ["写项目经历", "完成三分钟讲解", "按阶段达到投递门槛"])
    doc.add_heading("20.1 简历模板", level=2)
    add_code(doc, """项目：银行实时交易风控与实时指标平台（个人作品集 / 本地压测）
技术：Flink 2.2.1、Kafka、Flink CDC 3.6、PostgreSQL、Prometheus、Grafana

• 设计 Kafka 事件契约与 account_id 分区策略，保留 event_id、event_time、
  partition/offset，实现可追踪和可回放入口。
• 使用 Flink SQL 构建 ODS/DWD/DWS 链路，完成 Watermark、滑动窗口、
  Changelog/Upsert，并将极晚和异常数据分流到补偿主题。
• 使用 KeyedProcessFunction、Managed State、Event-time Timer 实现
  十分钟快速转出规则，以稳定 alert_id + 数据库唯一键验证回放幂等。
• 通过 Flink CDC 同步动态规则，保留 rule_version 解释历史告警。
• 建立 Lag、反压、Checkpoint 和质量监控，完成 TaskManager 故障、
  热点 Key 与慢 Sink 演练。
• 本地压测：峰值 __ 条/秒，P95 延迟 __ ms，Checkpoint P95 __ 秒，
  故障恢复 __ 秒（环境和脚本见仓库）。""")
    doc.add_heading("20.2 三分钟讲解", level=2)
    add_numbers(doc, [
        "30 秒：批量风控为何不够快，业务 SLA 是什么。",
        "40 秒：Kafka、Flink SQL/DataStream、CDC、Sink、监控架构。",
        "60 秒：乱序、重复、动态维表、Exactly-once、反压如何解决。",
        "30 秒：测试、故障演练与本地实测数字。",
        "20 秒：真实边界和下一步，而不是冒充生产项目。",
    ])
    doc.add_heading("20.3 高频追问", level=2)
    add_table(doc, ["问题", "回答骨架"], [
        ["为什么 Flink？", "事件时间、状态、低延迟、生态与现有技术栈；不贬低其他方案"],
        ["Exactly-once？", "可回放 Source + Checkpoint + 状态一致 + 事务/幂等 Sink"],
        ["Watermark 五秒依据？", "入口延迟分布 + 业务 SLA + 极晚补偿"],
        ["状态无限大？", "Timer 清理 + TTL 兜底 + 紧凑结构 + 状态监控"],
        ["动态规则？", "CDC + 版本化 Temporal/Broadcast State"],
        ["反压怎么查？", "Lag→backpressure→下游第一瓶颈→CPU/I/O/倾斜/外部系统"],
        ["如何回补？", "侧输出留证、受控回放、幂等主键、离线校准"],
        ["项目规模？", "如实说环境、数据生成器、实测数字与未覆盖范围"],
    ], [2.4, 4.1], font_size=8.7)
    doc.add_heading("20.4 45 天路线", level=2)
    add_table(doc, ["阶段", "学习与实现", "交付物"], [
        ["1～7 天", "流处理心智、Kafka Key/Partition/Offset、事件契约", "Kafka 演示 + README"],
        ["8～15 天", "Flink SQL、Watermark、窗口、Changelog", "M1/M2 + 测试数据"],
        ["16～25 天", "State、Timer、TTL、去重、Java 测试", "M3 + 状态设计"],
        ["26～33 天", "CDC、Checkpoint、幂等、Savepoint", "M4/M5 + 故障报告"],
        ["34～40 天", "监控、反压、热点、压测", "M6 + Dashboard"],
        ["41～45 天", "验收、文档、压测数字、讲稿", "M7 + 演示视频"],
    ], [1.3, 3.2, 2.0])
    add_callout(doc, "可投递门槛",
        "能独立解释时间、状态、Checkpoint、Changelog、CDC 和反压；仓库能一键启动；有重复/乱序/故障测试；简历数字可复现；能诚实说清已实现与生产化设计。",
        "success")
    add_self_check(doc, ["怎样证明简历数字？", "面试官真正想听 API 还是取舍？", "未做过的生产能力怎样回答？"],
        "录一次三分钟讲解，删除所有无法用代码、测试、截图或报告证明的句子。")

    doc.add_page_break()
    doc.add_heading("附录 A　术语速查", level=1)
    add_table(doc, ["英文", "中文理解", "工作解释"], [
        ["Event", "事件", "已经发生、应可追踪的业务事实"],
        ["Stream", "流", "持续到达、通常无界的事件序列"],
        ["Watermark", "事件时间进度", "此前数据基本到齐的估计"],
        ["State", "状态", "算子跨事件保存的业务记忆"],
        ["Checkpoint", "自动恢复快照", "恢复状态与 Source 位置"],
        ["Savepoint", "人工迁移快照", "升级、扩缩容与回滚"],
        ["Backpressure", "反压", "下游慢迫使上游降速"],
        ["Changelog", "变更日志", "动态表的 INSERT/UPDATE/DELETE"],
        ["Upsert", "插入或更新", "按主键覆盖结果"],
        ["CDC", "变更数据捕获", "从数据库日志读取行级变化"],
        ["Idempotency", "幂等", "重复执行的业务效果等同一次"],
        ["Replay", "回放", "从历史 Offset 或时间重新处理"],
        ["Lag", "积压", "消费者落后生产者多少"],
        ["Throughput", "吞吐", "单位时间处理记录或字节"],
        ["Latency", "延迟", "事件发生到结果可见的时间"],
    ], [1.5, 2.0, 3.0], font_size=8.6)

    doc.add_heading("附录 B　官方资料与版本边界", level=1)
    add_para(doc, """以下资料在 2026-08-04 校验。版本会变化，搭建前应重新查看 Downloads、Connector 兼容范围和 Release Notes。教程代码强调机制与工程方法，不替代目标版本官方文档。""")
    add_bullets(doc, [
        "Flink Downloads：https://flink.apache.org/downloads/",
        "Flink Stable Docs：https://nightlies.apache.org/flink/flink-docs-stable/",
        "Flink 2.2 Kafka Connector：https://nightlies.apache.org/flink/flink-docs-release-2.2/docs/connectors/table/kafka/",
        "Production Readiness：https://nightlies.apache.org/flink/flink-docs-stable/docs/ops/production_ready/",
        "Event Time：https://nightlies.apache.org/flink/flink-docs-stable/docs/concepts/time/",
        "Stateful Processing：https://nightlies.apache.org/flink/flink-docs-release-2.3/docs/concepts/stateful-stream-processing/",
        "Checkpoint under Backpressure：https://nightlies.apache.org/flink/flink-docs-stable/docs/ops/state/checkpointing_under_backpressure/",
        "Flink CDC 3.6：https://flink.apache.org/2026/03/30/apache-flink-cdc-3.6.0-release-announcement/",
        "Kafka Downloads：https://kafka.apache.org/community/downloads/",
    ])
    add_callout(doc, "最后一句",
        "实时开发不是把 SQL 跑得更频繁，而是在没有终点的数据上管理时间、状态、变化和故障。把这四件事讲清楚、做出来、测出来，你的批处理经验就会成为优势。",
        "success")

    doc.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print(build())

