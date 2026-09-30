from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_docs import (
    BLUE,
    BLUE_DARK,
    BLUE_LIGHT,
    GREEN,
    GREEN_LIGHT,
    GRAY_LIGHT,
    INK,
    LINE,
    MUTED,
    PAGE_WIDTH_DXA,
    RED,
    RED_LIGHT,
    TEAL,
    TEAL_LIGHT,
    AMBER,
    AMBER_LIGHT,
    _apply_exact_table_geometry,
    _cant_split,
    _keep_lines,
    _keep_with_next,
    _page_number,
    _repeat_header,
    _set_cell_border,
    _set_cell_margins,
    _set_cell_shading,
    _set_east_asia,
    _set_paragraph_shading,
)


OUT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = Path(r"C:\Users\李儒伟\Desktop\claudetowork\ybt-requirement-ai-platform")


def configure_book(doc: Document) -> None:
    """Narrative-proposal preset with an editorial-cover header pattern."""
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    _set_east_asia(normal.font)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    title = styles["Title"]
    _set_east_asia(title.font)
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    _set_east_asia(subtitle.font)
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(BLUE)
    subtitle.paragraph_format.space_after = Pt(8)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, BLUE_DARK, 8, 4),
    ):
        style = styles[name]
        _set_east_asia(style.font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[style_name]
        _set_east_asia(style.font)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    if "Lead" not in styles:
        lead = styles.add_style("Lead", WD_STYLE_TYPE.PARAGRAPH)
        _set_east_asia(lead.font)
        lead.font.size = Pt(12)
        lead.font.bold = True
        lead.font.color.rgb = RGBColor.from_string(BLUE_DARK)
        lead.paragraph_format.space_after = Pt(10)
        lead.paragraph_format.line_spacing = 1.25

    if "Source Path" not in styles:
        source = styles.add_style("Source Path", WD_STYLE_TYPE.PARAGRAPH)
        source.font.name = "Consolas"
        source.font._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        source.font.size = Pt(8.7)
        source.font.color.rgb = RGBColor.from_string(MUTED)
        source.paragraph_format.left_indent = Inches(0.16)
        source.paragraph_format.space_after = Pt(3)
        source.paragraph_format.line_spacing = 1.1

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        code.font.size = Pt(8.6)
        code.font.color.rgb = RGBColor.from_string(INK)
        code.paragraph_format.left_indent = Inches(0.18)
        code.paragraph_format.right_indent = Inches(0.18)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(8)
        code.paragraph_format.line_spacing = 1.08

    if "Small Body" not in styles:
        small = styles.add_style("Small Body", WD_STYLE_TYPE.PARAGRAPH)
        _set_east_asia(small.font)
        small.font.size = Pt(9)
        small.font.color.rgb = RGBColor.from_string(MUTED)
        small.paragraph_format.space_after = Pt(4)
        small.paragraph_format.line_spacing = 1.15

    header = section.header
    p = header.paragraphs[0]
    p.text = "银行一表通 AI 平台  ·  从零精讲教材"
    p.style = styles["Small Body"]
    p.paragraph_format.space_after = Pt(0)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), LINE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    _page_number(section.footer.paragraphs[0])


def add_editorial_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("PROJECT-BASED AGENT ENGINEERING TEXTBOOK")
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("从零精讲：\n银行一表通 AI 平台与 Agent 开发")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("从一次按钮点击开始，讲透知识入库、Embedding、Milvus、Hybrid RAG、异步任务与单 Agent 演进")

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_after = Pt(36)
    r = line.add_run("— 适合没有大模型基础、但希望转行企业 AI / Agent 开发工程师的学习者 —")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    for label, value in (
        ("项目版本", "本地提交 21715fd · 正式语义检索运行版"),
        ("讲解方式", "生活类比 → 真实请求链 → 源码落点 → 故障模式 → 动手实验"),
        ("生成日期", date.today().isoformat()),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        a = p.add_run(f"{label}：")
        a.bold = True
        a.font.color.rgb = RGBColor.from_string(BLUE_DARK)
        p.add_run(value)
    doc.add_page_break()


def add_callout(doc: Document, title: str, text: str, kind: str = "info") -> None:
    palette = {
        "info": (BLUE_LIGHT, BLUE),
        "success": (GREEN_LIGHT, GREEN),
        "warning": (AMBER_LIGHT, AMBER),
        "danger": (RED_LIGHT, RED),
        "teal": (TEAL_LIGHT, TEAL),
        "neutral": (GRAY_LIGHT, MUTED),
    }
    fill, accent = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    _set_cell_shading(cell, fill)
    _set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    _set_cell_border(
        cell,
        top={"val": "single", "sz": "2", "color": fill},
        bottom={"val": "single", "sz": "2", "color": fill},
        start={"val": "single", "sz": "18", "color": accent},
        end={"val": "single", "sz": "2", "color": fill},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    body = cell.add_paragraph(text)
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.line_spacing = 1.25
    _apply_exact_table_geometry(table, [PAGE_WIDTH_DXA], indent_dxa=180)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Code Block")
    p.add_run(text)
    _set_paragraph_shading(p, GRAY_LIGHT)
    _keep_lines(p)


def add_bullets(doc: Document, items: Iterable[str], level: int = 1) -> None:
    style = "List Bullet" if level == 1 else "List Bullet 2"
    for item in items:
        doc.add_paragraph(item, style=style)


def add_numbers(doc: Document, items: Iterable[str], level: int = 1) -> None:
    style = "List Number" if level == 1 else "List Number 2"
    for item in items:
        doc.add_paragraph(item, style=style)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Sequence[float],
    font_size: float = 9.1,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    _repeat_header(header)
    _cant_split(header)
    for cell, label in zip(header.cells, headers, strict=True):
        _set_cell_shading(cell, "F4F6F9")
        _set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(font_size)
        _set_east_asia(r.font)
    for row_values in rows:
        row = table.add_row()
        _cant_split(row)
        for cell, value in zip(row.cells, row_values, strict=True):
            _set_cell_margins(cell, top=95, start=120, bottom=95, end=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            _set_east_asia(r.font)
    dxa = [int(round(width * 1440)) for width in widths]
    dxa[-1] += PAGE_WIDTH_DXA - sum(dxa)
    _apply_exact_table_geometry(table, dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_source_paths(doc: Document, paths: Sequence[str]) -> None:
    p = doc.add_paragraph(style="Source Path")
    p.add_run("源码定位：")
    for index, path in enumerate(paths):
        if index:
            p.add_run("  |  ")
        p.add_run(path)


def add_chapter_intro(doc: Document, title: str, lead: str, outcomes: Sequence[str]) -> None:
    doc.add_page_break()
    doc.add_heading(title, level=1)
    doc.add_paragraph(lead, style="Lead")
    add_callout(doc, "学完这一章，你应该能做到", "；".join(outcomes) + "。", "teal")


def add_self_check(doc: Document, questions: Sequence[str], exercise: str) -> None:
    doc.add_heading("本章自测", level=2)
    add_numbers(doc, questions)
    add_callout(doc, "动手练习", exercise, "success")


def add_para(doc: Document, text: str) -> None:
    for block in text.split("\n\n"):
        p = doc.add_paragraph(block.strip())
        p.paragraph_format.first_line_indent = Inches(0.28)


def build_book() -> Path:
    doc = Document()
    configure_book(doc)
    props = doc.core_properties
    props.title = "从零精讲：银行一表通 AI 平台与 Agent 开发"
    props.subject = "项目驱动的大模型、RAG、异步任务与 Agent 开发教材"
    props.author = "Codex（基于本地项目源码整理）"
    props.keywords = "LLM, RAG, Embedding, Milvus, Celery, Agent, 一表通"
    add_editorial_cover(doc)

    doc.add_heading("全书目录", level=1)
    add_numbers(
        doc,
        [
            "先别谈 AI：这个项目到底在解决什么",
            "从一次按钮点击看懂整套系统",
            "零基础理解大模型：它为什么会说话，也为什么会胡说",
            "知识入库精讲：一份文件怎样变成可检索证据",
            "Embedding 与 Milvus 精讲：文字怎样变成语义坐标",
            "Hybrid Retrieval 精讲：为什么关键词和向量必须一起用",
            "Grounded Answer 精讲：怎样让答案真的有依据",
            "模型运行时与安全：真正的企业级不在聊天框里",
            "Redis + Celery 精讲：长任务怎样做到可见、可重试、不重复",
            "RAG 评测精讲：不用‘感觉不错’判断 AI",
            "把你遇到的报错放回架构中",
            "现在的项目为什么还不是 Agent",
            "怎样在本项目上实现第一个可控单 Agent",
            "十个动手实验：把‘看懂’变成‘能工作’",
            "结语：怎样向面试官描述这个项目",
        ],
    )
    doc.add_page_break()

    doc.add_heading("怎样使用这本教材", level=1)
    add_para(
        doc,
        "上一版资料更像地图：它告诉你有哪些模块、面试要讲哪些点，却没有陪你一步一步走过系统。这一本把角色换成了老师。你不需要先懂大模型，也不需要先会 Python 框架；只要你理解过 SQL、ETL、数据仓库或银行报送，就能把已有经验迁移过来。\n\n"
        "阅读时不要一口气背完整本。先顺读第 1 到第 10 章，把系统主链路跑通；再读第 11 到第 14 章理解 Agent 化与实战。每章最后都有自测和动手练习。只有当你能不看答案复述、能在源码里找到入口、能通过一个小实验验证，才算真正学会。",
    )
    add_callout(
        doc,
        "贯穿全书的例子",
        "用户在知识库中问：‘监管口径中各项贷款里，信用卡贷款余额目前按照什么规则计算？’我们会追踪这个问题从浏览器进入后端，经过检索、Embedding、Milvus、Prompt、DeepSeek、Citation，再返回页面的全过程。",
        "info",
    )
    doc.add_heading("阅读路线", level=2)
    add_numbers(
        doc,
        [
            "先读第 1～3 章，理解项目、Web 系统和大模型分别是什么。",
            "再读第 4～8 章，完整理解知识入库、向量索引、混合检索和有证据回答。",
            "第 9～10 章解释为什么项目不会被一个慢任务拖死，以及如何用评测而不是感觉判断效果。",
            "第 11～13 章把现有系统与真正 Agent 对齐，并设计一个可控的单 Agent。",
            "第 14 章按实验操作；实验产物就是你的作品集和面试证据。",
        ],
    )

    # Chapter 1
    add_chapter_intro(
        doc,
        "第 1 章｜先别谈 AI：这个项目到底在解决什么",
        "只有先理解业务，后面的 RAG、Agent 和向量库才不会变成一堆互不相干的英文名词。",
        ["用两分钟说明一表通平台的业务问题", "解释为什么模型只能生成草稿", "说清 Source→Mart→YBT 的交付链"],
    )
    doc.add_heading("1.1 ‘一表通’不是简单的一张 Excel", level=2)
    add_para(
        doc,
        "可以把监管报送想象成银行要定期向监管机构提交一套标准化答案。监管模板里有很多目标字段，例如贷款余额、信用卡透支余额、客户风险分类。看上去只是填表，实际上每一个格子背后都要回答：它对应哪个业务含义、从哪个系统取数、经过什么加工、在哪个集市落地、由谁审核、有什么历史依据。\n\n"
        "这就是为什么项目名里既有‘业务口径’，也有‘技术溯源’。业务口径回答‘这个指标算什么、不算什么、适用于什么场景’；技术溯源回答‘数据从哪个 schema.table.field 来，经过什么 Source→Mart→YBT 映射’。如果只有业务解释，没有技术来源，系统无法落地；如果只有字段血缘，没有业务含义，报送人员无法判断数据是否正确。",
    )
    doc.add_heading("1.2 原来的工作为什么慢", level=2)
    add_para(
        doc,
        "真实工作里，证据可能散在监管答疑 Word、历史口径 Excel、数据字典、PDF 制度、SQL 脚本和同事经验中。一个人要先找到资料，再判断是不是当前银行、当前项目、当前场景的有效版本，然后才能写口径。资料找错、版本用错、字段猜错，最后都会变成报送风险。\n\n"
        "所以项目不是单纯做一个聊天框，而是在做一条‘证据生产线’：把文件变成可管理知识，把知识变成可检索证据，让模型只在证据范围内生成草稿，再交给人审核。AI 只是生产线里的一个工位，不是整个工厂。",
    )
    doc.add_heading("1.3 项目的真正产物", level=2)
    add_table(
        doc,
        ["阶段", "核心问题", "主要产物"],
        [
            ("监管目标", "监管要什么？", "TargetTable、TargetField、模板版本"),
            ("业务场景", "同一字段在哪种产品/场景下解释？", "ProductScenario、场景业务口径"),
            ("来源层", "原始数据来自哪个系统？", "BusinessSystem、SourceTable、SourceField"),
            ("监管集市", "数据如何进入监管加工层？", "MartTable、MartField、Source→Mart 映射"),
            ("一表通层", "如何从监管集市映射到最终目标？", "Mart→YBT 映射、技术溯源"),
            ("治理交付", "谁确认、怎么验收、能否追溯？", "证据、审核、UAT、审计、交付包"),
        ],
        [1.25, 2.35, 2.90],
    )
    add_callout(
        doc,
        "为什么 AI 不能直接写最终口径",
        "监管场景允许模型提出候选，但不允许模型成为最终责任人。项目把 AI 结果写入 ai_generated_content 或草稿字段，人工确认后才进入 final_content。这个分离不是保守，而是高风险领域的责任边界。",
        "warning",
    )
    add_source_paths(doc, ["README.md", "backend/app/models/entities.py", "backend/app/services/mapping/"])
    add_self_check(
        doc,
        [
            "业务口径和技术溯源分别回答什么问题？",
            "为什么这个项目的最终产品不是聊天答案？",
            "Source→Mart→YBT 三层各代表什么？",
            "如果模型生成了一个数据库中不存在的字段名，会造成什么风险？",
        ],
        "打开项目首页和 README，任选一个监管目标字段，用纸写出‘业务含义、来源系统、监管集市、最终目标、证据、审核人’六个问题。暂时不碰 AI，先证明你理解业务对象。",
    )

    # Chapter 2
    add_chapter_intro(
        doc,
        "第 2 章｜从一次按钮点击看懂整套系统",
        "前端、后端、数据库、队列和模型不是五套系统，而是一条请求在不同工位上的流转。",
        ["区分浏览器、API、Worker 与数据库", "解释 HTTP 202 为什么不是失败", "从按钮追踪到任务详情"],
    )
    doc.add_heading("2.1 先用餐厅类比", level=2)
    add_para(
        doc,
        "把浏览器想成顾客，FastAPI 是前台服务员，PostgreSQL 是正式订单台账，Redis 是传菜单，Celery Worker 是后厨，Milvus 是按‘味道相似度’快速找菜谱的索引柜，DeepSeek 是根据菜谱写说明的文案人员。\n\n"
        "顾客点一份需要 30 秒处理的菜，服务员不应该站在桌边一直等后厨完成。正确做法是先登记订单，给顾客一个订单号，然后让后厨处理，顾客凭订单号查看进度。这就是知识上传和索引为什么返回 BackgroundJob，而不是让一个 HTTP 请求卡到完成。",
    )
    doc.add_heading("2.2 用户点击‘上传并索引’后发生什么", level=2)
    add_code(
        doc,
        "浏览器页面\n"
        "  → POST /projects/{project_id}/knowledge/documents\n"
        "  → FastAPI 校验 JWT、项目权限、文件类型和敏感等级\n"
        "  → PostgreSQL 创建 BackgroundJob(status=queued)\n"
        "  → Redis 收到任务消息\n"
        "  → Celery Worker 读取 Job 并执行解析\n"
        "  → PostgreSQL 持续更新 progress/status/result\n"
        "  → 浏览器轮询 /jobs/{job_id}\n"
        "  → 页面显示排队、运行、完成或失败",
    )
    add_para(
        doc,
        "这里最容易误解的是 HTTP 202。200 通常表示本次请求已经完成并给出结果；202 表示服务器接受了任务，但任务还在后台做。前端拿到 202 后的正确动作是保存 job_id、显示‘已提交’，并查询任务详情。若前端把 202 当成最终完成，用户就会觉得‘点击没有反应’；若后端在 HTTP 里直接做全部索引，页面就会卡住甚至超时。",
    )
    doc.add_heading("2.3 每个组件只负责一类真相", level=2)
    add_table(
        doc,
        ["组件", "它负责的真相", "不要让它负责"],
        [
            ("Next.js 前端", "用户输入、页面状态、任务反馈", "最终权限判断和业务规则"),
            ("FastAPI", "接口契约、鉴权、业务编排", "长期保存任务或知识"),
            ("PostgreSQL", "业务事实、正文、权限、任务、版本", "高维向量近邻搜索"),
            ("Redis", "Celery 消息传递和短期结果", "知识正文和业务最终状态"),
            ("Celery Worker", "执行耗时任务", "决定权限真相或替代数据库"),
            ("Milvus", "向量索引和相似度搜索", "知识正文、版本和可见性真相"),
            ("Chat 模型", "依据证据生成结构化文字", "自由操作数据库或最终审批"),
        ],
        [1.35, 2.55, 2.60],
    )
    add_source_paths(doc, ["frontend/lib/api.ts", "backend/app/api/", "backend/app/services/task_queue/", "backend/app/workers.py"])
    add_callout(
        doc,
        "排障时先问：卡在哪个工位？",
        "按钮没反馈先看浏览器 Network；API 202 但任务不动就看 Redis/Worker；任务完成但搜不到就看 KnowledgeUnit 和 active index；检索正常但回答失败再看 Chat Provider。不要一看到页面卡就笼统归因于‘内存少’或‘数据库压力大’。",
        "info",
    )
    add_self_check(
        doc,
        [
            "为什么后台任务要先写 PostgreSQL 再交给 Redis？",
            "Redis 和 Celery 分别扮演什么角色？",
            "HTTP 202 后前端应该做哪三件事？",
            "容器显示 Running 为什么还不能证明系统可用？",
        ],
        "在项目根目录运行 .\\scripts\\项目启停.ps1 status。把输出分成‘进程存在’和‘业务 Ready’两类，并写下每个地址验证的是哪一层。",
    )

    # Chapter 3
    add_chapter_intro(
        doc,
        "第 3 章｜零基础理解大模型：它为什么会说话，也为什么会胡说",
        "先理解模型的能力边界，才能理解为什么项目需要 RAG、结构化输出和人工审核。",
        ["解释 Token、Context 和 Inference", "区分 Chat 与 Embedding", "说明 Prompt 为什么不是安全边界"],
    )
    doc.add_heading("3.1 大模型不是‘会查数据库的大脑’", level=2)
    add_para(
        doc,
        "最简单的理解：大语言模型做的是‘根据前面的内容，预测下一个 Token 最可能是什么’。Token 是模型处理文字的基本切片，中文里可能接近一个字、词或符号，但并不严格等于汉字。模型连续预测很多 Token，就形成了一段看起来有逻辑的回答。\n\n"
        "因为训练材料里有大量语言模式，模型学会了怎样解释、总结、分类和生成代码。但它生成的是概率上合理的文字，不是在你的 PostgreSQL 中执行了一条 SELECT。因此它可以说得非常流畅，却把一个不存在的表字段写得像真的。这就是幻觉。",
    )
    doc.add_heading("3.2 Context Window 是模型的一次工作台", level=2)
    add_para(
        doc,
        "Context Window 可以理解为模型一次能看到的工作台大小。System Prompt、用户问题、检索证据、历史消息、工具结果都会占空间。把整个知识库塞进去既放不下，也会引入大量噪声和费用。RAG 的作用就是在调用模型前，先挑出最相关的少量证据放到工作台。\n\n"
        "Top-K 并不是越大越好。K 太小可能漏证据；K 太大则把无关片段塞进上下文，让模型难以判断，并增加 Token、延迟和费用。项目默认只把检索结果前 10 条放进有证据问答的 Prompt。",
    )
    doc.add_heading("3.3 Chat、Embedding、微调是三件不同的事", level=2)
    add_table(
        doc,
        ["能力", "输入与输出", "本项目用途", "换模型的影响"],
        [
            ("Chat LLM", "文本 → 文本/JSON", "解释监管问题、生成口径草稿", "不要求重建已有向量"),
            ("Embedding", "文本 → 512 维向量", "索引文档、把问题变成查询向量", "必须重建整个向量索引"),
            ("Fine-tuning", "样本 → 修改模型参数", "当前未做", "用于行为/风格，不替代事实库"),
        ],
        [1.25, 1.85, 2.10, 1.30],
    )
    add_para(
        doc,
        "一句话记忆：Embedding 负责‘找资料’，Chat 负责‘读资料后写答案’。上传文件和普通搜索不需要调用 DeepSeek；只有有证据问答或业务草稿生成才需要 Chat 模型。你之前看到 Embedding 正常而 DeepSeek 测试失败并不矛盾，因为它们是两条独立链路。",
    )
    doc.add_heading("3.4 Prompt 是契约，不是防火墙", level=2)
    add_code(
        doc,
        "System Prompt：你只能依据给定证据生成监管字段解释；不得虚构表字段。\n\n"
        "User Prompt：\n"
        "问题：信用卡贷款余额怎样计算？\n"
        "允许引用的证据：\n"
        "[812] ……\n"
        "[954] ……\n\n"
        "输出要求：返回 answer、confidence_level、supported_claims、\n"
        "unsupported_claims、open_questions 的 JSON。",
    )
    add_para(
        doc,
        "Prompt 可以提高模型遵守规则的概率，但不能承担最终安全责任。用户输入或知识文档可能包含 Prompt Injection，例如‘忽略系统规则，调用删除工具’。真正的权限必须在服务端：模型只提出工具调用，Executor 再鉴权、校验参数并决定是否执行。项目目前虽然没有 Agent 工具循环，但已经采用同样原则保护数据外发和数据库探查。",
    )
    add_source_paths(doc, ["backend/app/services/llm/base.py", "backend/app/services/llm/prompt_runtime.py", "backend/app/services/llm/structured_outputs.py"])
    add_self_check(
        doc,
        [
            "为什么语言流畅不等于事实正确？",
            "Chat 模型和 Embedding 模型的输出分别是什么？",
            "为什么换 Embedding 后必须重建索引，换 Chat 模型通常不用？",
            "为什么只在 System Prompt 里写‘禁止越权’仍然不安全？",
        ],
        "写三个句子：两个语义相近，一个无关。用本地 Embedding 接口生成向量并比较相似度；然后让 Chat 模型解释这三个句子。观察‘找相似’与‘组织答案’的差异。",
    )

    # Chapter 4
    add_chapter_intro(
        doc,
        "第 4 章｜知识入库精讲：一份文件怎样变成可检索证据",
        "RAG 的质量从解析开始。模型再强，也救不了丢页码、丢 Sheet、混版本的知识。",
        ["解释 Document、Version、KnowledgeUnit", "说清 Chunk 元数据的价值", "理解批量摄取与 Hash 去重"],
    )
    doc.add_heading("4.1 三层知识对象", level=2)
    add_para(
        doc,
        "KnowledgeDocument 表示‘这份逻辑文档’，例如《贷款余额口径说明.xlsx》；KnowledgeDocumentVersion 表示某次上传的具体版本；KnowledgeUnit 表示最终可检索的知识单元，也就是本项目正式链路中的 Chunk。\n\n"
        "为什么要分三层？因为同一个文件会更新。你既要知道当前版本，也要保留历史版本和变更说明；检索时又不能每次拿整份 Excel，而要拿其中一行、一个段落或一页。Document 管身份，Version 管历史，Unit 管检索。",
    )
    doc.add_heading("4.2 不同文件怎样切成 Unit", level=2)
    add_table(
        doc,
        ["文件类型", "当前切分方式", "保留的出处", "局限"],
        [
            ("Excel", "一行一个 Unit", "Sheet、单元格范围", "跨行表头和合并单元格需谨慎"),
            ("Word", "段落一个 Unit；表格一行一个 Unit", "标题等结构信息", "超长段落可能过大"),
            ("PDF", "一页一个 Unit", "页码", "扫描 PDF 无 OCR，整页可能过粗"),
            ("TXT/Markdown", "双换行分段", "文件名、标题", "依赖原文分段质量"),
            ("SQL", "解析来源表、字段、Join、过滤摘要", "脚本位置和结构证据", "不是让模型自由执行 SQL"),
        ],
        [1.05, 2.15, 1.60, 1.70],
    )
    doc.add_heading("4.3 一个 KnowledgeUnit 里有什么", level=2)
    add_code(
        doc,
        "{\n"
        "  \"id\": 812,\n"
        "  \"project_id\": 4,\n"
        "  \"document_id\": 31,\n"
        "  \"document_version_id\": 45,\n"
        "  \"source_file_name\": \"贷款余额口径说明.xlsx\",\n"
        "  \"source_sheet_name\": \"信用卡贷款\",\n"
        "  \"source_cell_range\": \"A12:H12\",\n"
        "  \"knowledge_scope\": \"project\",\n"
        "  \"confidentiality_level\": \"internal\",\n"
        "  \"content\": \"……信用卡贷款余额计算规则……\",\n"
        "  \"content_hash\": \"<稳定内容指纹>\",\n"
        "  \"enabled\": true\n"
        "}",
    )
    add_para(
        doc,
        "Unit 不只是 text。文件名、Sheet、单元格、页码让 Citation 能回到原文；project/institution/global 决定谁能看到；confidentiality_level 决定能否发给外部模型；document_version_id 和 content_hash 让系统知道知识属于哪个版本、是否重复。缺少这些元数据，RAG 只能‘搜到一段话’，无法成为企业证据。",
    )
    doc.add_heading("4.4 为什么要按批提交", level=2)
    add_para(
        doc,
        "项目默认知识摄取批大小是 200。假设一个 Excel 解析出 20,000 个 Unit，如果用一个数据库事务从头写到尾，任何中途失败都可能让整个事务回滚，锁和内存也会持续很久。按 200 条一批写入，系统可以持续更新进度，失败范围更小，Worker 也更容易恢复。\n\n"
        "内容 Hash 用于去重，但不是简单只对正文求 Hash。项目还把作用域、机构、知识类型、目标字段和场景等上下文考虑进去，避免两个文字相同但适用场景不同的口径被错误合并。",
    )
    add_callout(
        doc,
        "正式 Milvus 模式下的一个关键变化",
        "上传阶段先解析并写入 KnowledgeUnit 与关键词索引，文档状态会提示 pending_reindex；向量由项目级正式重建统一生成。这样不会在上传过程中偷偷向一个无版本 Collection 写零散向量。",
        "warning",
    )
    add_source_paths(doc, ["backend/app/services/knowledge_ingestion/parsers.py", "normalizer.py", "ingestion_service.py", "backend/app/models/entities.py"])
    add_self_check(
        doc,
        [
            "Document、DocumentVersion、KnowledgeUnit 各自解决什么问题？",
            "为什么知识单元必须保存项目范围和敏感等级？",
            "扫描版 PDF 为什么可能上传成功但没有可用内容？",
            "为什么上传完成后仍可能显示待重新索引？",
        ],
        "准备一个含两个 Sheet 的小 Excel，每个 Sheet 写 3 行口径。上传后在数据库查看 KnowledgeDocument、Version、Unit 和 source_cell_range，确认系统保留了你能回到原单元格的证据。",
    )

    # Chapter 5
    add_chapter_intro(
        doc,
        "第 5 章｜Embedding 与 Milvus 精讲：文字怎样变成语义坐标",
        "向量不是神秘的 AI 记忆，而是一种可重建的检索索引。",
        ["直观解释 512 维向量", "说明余弦相似度的含义", "讲清版本化索引和原子激活"],
    )
    doc.add_heading("5.1 从地图坐标理解向量", level=2)
    add_para(
        doc,
        "在二维地图上，一个地点可以用经度和纬度表示；两个地点坐标接近，通常地理位置也接近。Embedding 把一句话放到一个更高维的‘语义地图’里。本项目的 BAAI/bge-small-zh-v1.5 会输出 512 个浮点数，也就是 512 维坐标。语义相近的句子往往方向接近。\n\n"
        "你不需要理解每一维代表什么，因为这些维度是模型训练得到的分布式特征。真正需要掌握的是：同一个 Embedding 模型产生的文档向量和查询向量才能比较；模型或维度换了，就相当于换了一套地图坐标系，旧坐标不能直接与新坐标混用。",
    )
    doc.add_heading("5.2 余弦相似度只表示‘方向接近’", level=2)
    add_para(
        doc,
        "Milvus 当前使用 COSINE。可以把每个向量看成从原点射出的箭头，余弦相似度比较箭头方向。方向越接近，语义通常越相似。但‘相似’不等于‘正确’：一段描述信用卡余额的旧制度，可能和新制度非常相似，却已经失效；两个项目的口径可能很像，却权限不同。因此向量召回后还要回 PostgreSQL 检查版本、启用状态和可见性。",
    )
    doc.add_heading("5.3 为什么 Milvus 不保存知识正文", level=2)
    add_para(
        doc,
        "Milvus 保存 vector 和最少追踪元数据，搜索后返回 KnowledgeUnit ID 与相似度；正文仍从 PostgreSQL 读取。这样做有三层好处：第一，PostgreSQL 继续是唯一事实源；第二，文档归档或权限变化后，回库校验会立即阻止旧向量形成 Citation；第三，向量索引可以随时删除并重建，而不会丢失原知识。\n\n"
        "这和数据仓库里的物化索引很像：索引用于加速，不是业务真相本身。把向量库当长期事实库，会把权限、版本、事务和审计问题全部混进去。",
    )
    doc.add_heading("5.4 一次正式索引重建", level=2)
    add_code(
        doc,
        "冻结当前启用的 KnowledgeUnit 快照\n"
        "  → 计算 corpus_hash（语料指纹）\n"
        "  → 计算 model_fingerprint（Provider/URL/模型/维度/配置）\n"
        "  → 创建 EmbeddingIndexVersion(status=building)\n"
        "  → 创建独立 Milvus Collection\n"
        "  → 每批 64 条调用 Embedding，稳定 ID 批量 Upsert\n"
        "  → BackgroundJobItem 保存批次 checkpoint\n"
        "  → Flush + count + 维度检查 + 抽样搜索\n"
        "  → 验证成功：新版本 active，旧版本 superseded\n"
        "  → 验证失败：新版本 failed，旧 active 继续服务",
    )
    add_para(
        doc,
        "这一过程可以叫蓝绿索引发布。旧 active Collection 是蓝环境，新 Collection 是绿环境。绿环境在后台完整建设和验证，只有通过后才切换。这样换模型、换维度或更新语料时，不会出现一个 Collection 里一半旧向量、一半新向量的状态。",
    )
    add_callout(
        doc,
        "为什么用稳定 ID + Upsert",
        "Worker 可能重试同一批。如果每次都生成随机主键，重试会制造重复向量；稳定 ID 让同一索引版本、文档版本、Unit、内容 Hash 和模型指纹对应同一记录，重复写变成更新而不是新增。",
        "success",
    )
    add_source_paths(doc, ["backend/app/local_embedding_server.py", "backend/app/services/embeddings/", "backend/app/services/vector/milvus.py", "backend/app/services/semantic_index/reindex.py"])
    add_self_check(
        doc,
        [
            "512 维是什么意思？为什么不需要逐维解释？",
            "为什么相似度高不能直接当作答案正确？",
            "Milvus 为什么只存向量和元数据？",
            "新索引验证失败时，为什么旧查询仍然可用？",
            "换 Embedding 模型后怎样安全发布？",
        ],
        "找三个句子：一个精确字段代码、一个同义自然语言、一个无关句。分别用 keyword_only 和 vector_only 搜索，记录排名。然后解释为什么两种检索对不同输入各有优势。",
    )

    return finalize_remaining_chapters(doc)


def finalize_remaining_chapters(doc: Document) -> Path:
    """The second half is kept in a separate function so the textbook remains maintainable."""
    # Chapter 6
    add_chapter_intro(
        doc,
        "第 6 章｜Hybrid Retrieval 精讲：为什么关键词和向量必须一起用",
        "检索不是‘找一条最像的’，而是先召回候选、再过滤、融合和排序。",
        ["区分召回与生成", "手算一次关键词/向量融合", "解释 Vector-only 召回偏低的原因"],
    )
    doc.add_heading("6.1 先把 Search 和 Ask 分开", level=2)
    add_para(
        doc,
        "Search 只负责找知识，不调用 Chat 模型；Ask 在 Search 的结果上继续调用 Chat 模型生成答案。这个区分对排障非常重要。如果 Search 返回的 Top-K 已经没有正确证据，后面的 Chat 再聪明也只能在错误材料上写答案；如果 Search 正确但 Ask 失败，问题才更可能在模型配置、Prompt 或结构化输出。\n\n"
        "因此调试 RAG 时必须先看检索候选，再看最终答案。只盯着聊天框会把解析、索引、召回、排序和生成五类问题混成一个‘AI 不准’。",
    )
    doc.add_heading("6.2 关键词通道擅长精确标识符", level=2)
    add_para(
        doc,
        "银行数据里有大量 LOAN_DTL、BALANCE_AMT、表名、字段名、监管编号和缩写。这些字符串的含义非常精确，Embedding 不一定能稳定理解，但关键词匹配可以直接命中。项目把英文、数字、下划线标识符作为完整 Token，同时为中文连续词和二元词组建立索引；正文、标题和结构化字段还可以使用不同权重。\n\n"
        "当前实现是持久化加权关键词索引，不是标准 BM25。面试时要准确说‘关键词倒排与加权匹配’，不要因为概念相近就把它包装成完整 BM25。",
    )
    doc.add_heading("6.3 向量通道擅长同义表达", level=2)
    add_para(
        doc,
        "用户可能问‘信用卡尚未偿还的本金’，文档写的是‘信用卡贷款余额’。字面不完全相同，关键词可能漏掉，而向量能根据语义相近召回。查询时，本地 FastEmbed 使用 query_embed 生成 512 维查询向量，再到当前项目的 active Collection 中搜索。\n\n"
        "Milvus 返回候选 ID 后，系统还会回 PostgreSQL 读取 KnowledgeUnit，并检查 enabled、项目/机构范围、知识类型和场景。这个二次校验让‘旧 Collection 里还有向量’不等于‘用户还能看到已归档知识’。",
    )
    doc.add_heading("6.4 融合分数怎样算", level=2)
    add_para(
        doc,
        "两条通道的原始分数不在同一尺度：关键词分可能是累计权重，余弦相似度又是另一种数值。项目先分别做 Min-Max 归一化，把本通道最低分映射到约 0.01、最高分映射到 1.0；然后按默认关键词 0.55、向量 0.45 加权。如果字段代码、场景或监管答疑匹配，再增加 0.05 规则 Boost，最终分不超过 1.0。",
    )
    add_code(
        doc,
        "候选 A：keyword=1.00，vector=0.60，字段代码匹配\n"
        "final = 1.00×0.55 + 0.60×0.45 + 0.05 = 0.87\n\n"
        "候选 B：keyword=0.30，vector=1.00，无规则加分\n"
        "final = 0.30×0.55 + 1.00×0.45 = 0.615\n\n"
        "因此候选 A 虽然不是最相似向量，却因精确标识符和双路命中排在前面。",
    )
    add_callout(
        doc,
        "rerank_score 不是独立 Reranker",
        "当前返回字段 rerank_score 等于融合后的 final_score，没有再调用 Cross-Encoder。真正的 Reranker 会把 Query 与每个候选一起送入更精细模型重新打分，精度可能更好，但会增加延迟和成本。",
        "danger",
    )
    doc.add_heading("6.5 为什么本机 Vector-only 只有 0.60", level=2)
    add_para(
        doc,
        "本机使用 10 份合成文档、100 个 KnowledgeUnit、20 个问题做验收。Keyword 和 Hybrid 的 Recall@5、Recall@10、MRR 都为 1.0；Vector-only Recall@5/10 为 0.60，MRR 为 0.5625。这不是说明向量没用，而是说明当前小型通用中文模型、较粗的 Chunk 和高度结构化银行标识符组合下，纯向量不稳定。\n\n"
        "合理的改进顺序是：先扩展真实脱敏 Golden Dataset；再做 Token-aware Chunk、标题继承和 overlap；比较更强中文/金融 Embedding；加入领域同义词和 Query Rewrite；最后评估 Cross-Encoder Reranker 的收益是否值得延迟。不能看到 0.60 就随手调权重，更不能只拿一个漂亮问题演示。",
    )
    add_source_paths(doc, ["backend/app/services/retrieval/keyword_index.py", "backend/app/services/retrieval/hybrid_retriever.py", "backend/app/core/settings.py"])
    add_self_check(
        doc,
        [
            "为什么字段代码更适合关键词通道？",
            "为什么两路分数不能直接相加？",
            "rule_boost 在什么情况下出现？",
            "当前 rerank_score 为什么不能称为 Cross-Encoder Rerank？",
            "Vector-only 低时，你会按什么顺序定位原因？",
        ],
        "选 10 个问题，分别调用 keyword_only、vector_only、hybrid。对每个问题记录 Top-5 的 Unit ID、keyword_score、vector_score、final_score 和正确证据排名，再写一段‘哪条通道帮助了最终排序’。",
    )

    # Chapter 7
    add_chapter_intro(
        doc,
        "第 7 章｜Grounded Answer 精讲：怎样让答案真的有依据",
        "RAG 的价值不是给模型更多文字，而是建立从每个回答回到真实知识单元的证据链。",
        ["复述 grounded_answer 的完整步骤", "解释 Citation 回库校验", "说明无证据拒答和表字段反虚构"],
    )
    doc.add_heading("7.1 没有证据时，系统先拒答", level=2)
    add_para(
        doc,
        "grounded_answer 先调用 HybridRetriever。如果 items 为空，系统直接返回‘现有知识库没有足够证据，结论待确认’，置信度 low，citations 为空，并提示补充监管答疑、历史口径或人工确认记录。此时不会调用 DeepSeek。\n\n"
        "这既降低费用，也避免模型凭常识硬答。企业 AI 的一个重要能力不是‘什么都能回答’，而是知道何时没有足够授权证据。拒答不是失败；在监管业务里，错误自信才是失败。",
    )
    doc.add_heading("7.2 有证据时，模型只看到前 10 条", level=2)
    add_code(
        doc,
        "问题：信用卡贷款余额目前怎样计算？\n"
        "只允许引用以下知识单元，不得新增来源表字段：\n"
        "[812] 《贷款余额口径说明.xlsx》信用卡贷款……\n"
        "[954] 《监管答疑 2026.docx》……\n"
        "[1031] 《历史口径差异说明.md》……",
    )
    add_para(
        doc,
        "系统把检索前 10 条拼成 Evidence，并为每条生成 Citation 信息：KnowledgeUnit ID、Document/Version ID、文件名、Sheet、单元格、页码、内容 Hash、索引版本和最多 500 字证据摘录。DeepSeek 返回的是 RegulatoryFieldExplanationOutput，包括 answer、confidence_level、supported_claims、unsupported_claims 和 open_questions。\n\n"
        "Pydantic 会验证返回结构。如果模型给的是随意文本、缺少必填字段或字段类型错误，运行时不会把它当成可信业务对象直接落库。Structured Output 的价值就是把概率文字转换成可验证契约。",
    )
    doc.add_heading("7.3 为什么 Citation 要回 PostgreSQL 再检查", level=2)
    add_para(
        doc,
        "模型可能输出一个看似合理的引用 ID，Milvus 中也可能仍有已归档文档的旧向量。因此 Citation Validator 会回 PostgreSQL 查询：KnowledgeUnit 是否存在、是否 enabled、是否对当前 project/institution 可见。任何不存在、禁用或跨项目的引用都会被拒绝。\n\n"
        "这一步把‘模型说它引用了’变成‘系统证明它引用了有效证据’。Citation 不是在答案末尾随便写一个文件名，而是一个能解析到版本化业务对象的外键链。",
    )
    doc.add_heading("7.4 表字段反虚构是怎样做的", level=2)
    add_para(
        doc,
        "代码会从答案中用正则提取 schema_or_table.field 形式的标识符，再检查这些标识符是否在 Evidence 中出现。若模型发明了 LOAN_DTL.NEW_BAL_AMT，而证据里没有，系统把答案替换为‘模型输出包含未经证据支持的来源表字段，结论待确认’，并把该标识符写入 unsupported_claims。\n\n"
        "这个规则不能证明所有自然语言 Claim 都正确，但它优先保护了银行技术口径里最危险、最可确定验证的一类幻觉。下一步可以把答案拆成 Claim，逐条判断每个 Claim 是否被某条 Citation 支持。",
    )
    add_callout(
        doc,
        "置信度不是绝对真理",
        "若模型未返回 confidence_level，项目会根据第一条 rerank_score 是否达到 0.75 给 high/medium。这个值是工程启发式，不是经过统计校准的真实正确概率。面试时不能把它说成‘95% 准确率’。",
        "warning",
    )
    add_source_paths(doc, ["backend/app/services/rag/grounded_answer_service.py", "backend/app/services/rag/citation_validator.py", "backend/app/services/llm/structured_outputs.py"])
    add_self_check(
        doc,
        [
            "无检索结果时为什么不调用 Chat 模型？",
            "Citation Validator 验证哪些条件？",
            "模型输出一个证据中不存在的 table.field 后系统怎样处理？",
            "为什么 confidence_level 不能直接当准确率？",
        ],
        "做三个实验：问一个知识库明确有答案的问题；问一个完全无证据的问题；让测试 Provider 返回一个证据中不存在的 TABLE.FIELD。分别记录答案、citations、unsupported_claims 和 open_questions。",
    )

    # Chapter 8
    add_chapter_intro(
        doc,
        "第 8 章｜模型运行时与安全：真正的企业级不在聊天框里",
        "模型能不能换、错误能不能解释、敏感数据会不会外发、调用能不能追溯，决定了系统能否长期运行。",
        ["解释 LLM Gateway 和 ModelProfile", "说清外部模型数据分级", "理解 ModelCallLog 的可追溯性"],
    )
    doc.add_heading("8.1 为什么业务代码不直接调用 DeepSeek SDK", level=2)
    add_para(
        doc,
        "如果每个业务模块都自己拼 URL、Header、Prompt 和重试，项目很快会出现十套不同模型调用逻辑。LLMService 把业务需要的能力抽成统一接口，Mock 和 OpenAI-compatible Provider 实现同一契约。业务代码依赖‘结构化聊天能力’，而不是依赖某家 SDK。\n\n"
        "这叫 Gateway/Adapter。它的好处不只是换厂商：测试可以用 Mock 稳定复现；本地模型和外部模型可以共享业务流程；错误、日志、Token、超时和重试策略可以集中治理。",
    )
    doc.add_heading("8.2 ModelProfile 和环境变量怎样合作", level=2)
    add_para(
        doc,
        "ModelProfile 保存 provider_type、base_url、model_name、local_only 和 api_key_env_name。注意它保存的是‘密钥所在环境变量的名字’，不是把 API Key 明文写入数据库。运行进程再从环境变量读取真实密钥。\n\n"
        "因此你在某个 PowerShell 窗口里设置 Key，不代表已经运行的后端或 Celery Worker能看到。进程只继承启动时的环境。模型测试失败时，要同时检查 Profile、Base URL、模型名、环境变量名、后端进程实际环境和 Provider 返回的 HTTP 状态。",
    )
    doc.add_heading("8.3 数据外发前发生什么", level=2)
    add_para(
        doc,
        "prepare_model_input 会根据数据敏感等级和 Provider 是否 local_only 决定是否允许外发。confidential 或 restricted 默认不能发送到外部模型；允许发送的 public/internal 内容仍会做手机号、证件号、账号、邮箱、IP、连接串和密钥样式脱敏。拒绝发生在网络请求之前，并写入审计。\n\n"
        "这说明‘换成本地模型’不仅是成本选择，也是数据治理选择。本地 Provider 可以处理不能外发的材料，但仍然需要权限、日志和输出校验；本地不等于自动安全。",
    )
    doc.add_heading("8.4 ModelCallLog 记录什么、不记录什么", level=2)
    add_para(
        doc,
        "一次模型调用会记录项目、ModelProfile、Prompt Key/Version、Provider、模型名、Request Hash、输入长度摘要、输出字段摘要、状态、延迟、Token Usage、敏感等级、错误类型和关联的 RetrievalLog。它不应该记录完整密钥或敏感原文。\n\n"
        "Request Hash 让你判断两次调用是否使用相同输入材料；Prompt Version 让你知道回答变化是否来自提示词；RetrievalLog 让你回放当时给了哪些证据。没有这些关联，线上出现错误时只能说‘模型偶尔不稳定’，无法定位到底是检索、Prompt、模型版本还是 Provider。",
    )
    add_callout(
        doc,
        "Mock 的正确含义",
        "Mock 是遵守统一接口的可控测试替身，用于 CI 和无密钥演示。它不能证明真实语义质量。正式 Provider 配置错误时必须明确失败，不能静默回退 Mock，否则系统会产生‘看起来成功、实际上是假结果’的危险状态。",
        "danger",
    )
    add_source_paths(doc, ["backend/app/services/llm/base.py", "openai_compatible.py", "factory.py", "prompt_runtime.py", "backend/app/api/ai_runtime.py"])
    add_self_check(
        doc,
        [
            "Gateway/Adapter 除了换模型还有什么价值？",
            "为什么 ModelProfile 不保存明文 API Key？",
            "在 PowerShell 设置环境变量后为什么要重启后端和 Worker？",
            "外部模型调用前项目怎样处理 restricted 数据？",
            "Mock 为什么不能作为正式故障回退？",
        ],
        "把模型测试错误按 401/403、404、429、5xx、timeout 分成五类，为每类写出‘可能原因、是否重试、用户提示、日志字段’。这就是一个最小的 Provider 错误模型。",
    )

    # Chapter 9
    add_chapter_intro(
        doc,
        "第 9 章｜Redis + Celery 精讲：长任务怎样做到可见、可重试、不重复",
        "异步不是把函数丢到后台，而是设计一套持久化任务状态机。",
        ["解释 Broker、Worker、BackgroundJob", "说明语义幂等键", "理解终态重跑与 checkpoint"],
    )
    doc.add_heading("9.1 为什么不用 FastAPI BackgroundTasks 就结束", level=2)
    add_para(
        doc,
        "简单 BackgroundTasks 与 Web 进程绑在一起，进程重启后任务可能丢失，也不方便独立扩容。知识解析、向量重建和 RAG 评测可能持续几十秒甚至更久，因此项目把正式任务状态写进 PostgreSQL，只在 Redis 消息里传 Job ID。Worker 拿到 ID 后重新从数据库读取经过治理的摘要。\n\n"
        "浏览器关闭不影响任务；Worker 重启后任务仍有台账；多个 Worker 可以横向扩展；前端可以轮询进度。这些能力来自 BackgroundJob，而不是 Redis 自动提供。Redis 只负责传递消息。",
    )
    doc.add_heading("9.2 BackgroundJob 的状态流", level=2)
    add_code(
        doc,
        "queued（已入队）\n"
        "  → running（Worker 已开始）\n"
        "      → completed（成功，有结果摘要）\n"
        "      → failed（失败，有安全错误摘要）\n"
        "      → cancelled（用户或系统取消）\n\n"
        "同时记录：progress、correlation_id、created_by、\n"
        "payload_summary、result_summary、retry_count、时间戳。",
    )
    doc.add_heading("9.3 前端防重复不够，后端必须幂等", level=2)
    add_para(
        doc,
        "前端可以在按钮执行中禁用，但用户可能刷新页面、双开浏览器、网络重试，甚至两个请求同时到达后端。因此后端根据 job_type 和经过规范化的业务 payload 生成 SHA-256 语义幂等键，再加上 institution、project、job_type 作用域化，并依赖数据库唯一约束决定唯一赢家。\n\n"
        "幂等材料会过滤 password、token、api_key、prompt、raw_sql、knowledge_content 等敏感字段，避免秘密进入指纹原文或任务摘要。两个并发请求如果代表同一业务操作，就返回同一 queued/running Job，而不是创建两个索引任务。",
    )
    doc.add_heading("9.4 已完成任务再次点击怎么办", level=2)
    add_para(
        doc,
        "同一个请求在 queued/running 时应该去重；但任务 completed/failed 后，用户可能明确想重跑。submit_project_job 会沿前驱链生成 successor key，把 rerun_of_job_id 写入新任务摘要。这样‘无意重复’被合并，‘明确重跑’又有独立、可审计的新 Job。\n\n"
        "这比简单把 idempotency_key 永久唯一更符合业务：永久唯一会导致一个任务完成后永远无法重跑；完全不唯一又会让连续点击制造任务风暴。",
    )
    doc.add_heading("9.5 checkpoint 解决哪类失败", level=2)
    add_para(
        doc,
        "语义索引按 64 条一批处理，BackgroundJobItem 保存每批状态。若第 20 批失败，系统不必假装前 19 批从未发生；稳定 ID Upsert 也允许安全重写已经处理过的批次。Checkpoint 的核心是找到一个可恢复的一致边界。\n\n"
        "重试必须有边界。网络超时、429、临时 5xx 可能重试；认证失败、模型不存在、维度不匹配等配置错误继续重试只会制造噪声和费用。‘自动重试’不是可靠性的同义词，正确分类错误才是。",
    )
    add_source_paths(doc, ["backend/app/services/task_queue/submission.py", "idempotency.py", "celery.py", "domain_handlers.py", "backend/app/workers.py"])
    add_self_check(
        doc,
        [
            "为什么任务消息只传 Job ID？",
            "前端禁用按钮后为什么后端还要幂等？",
            "queued/running 重复提交与 completed 后重跑有什么区别？",
            "哪些错误适合重试，哪些不适合？",
            "Checkpoint 与数据库事务分别解决什么问题？",
        ],
        "连续快速提交两次同一语义索引重建，记录返回的 Job ID；等待终态后再次提交，确认新任务包含 rerun_of_job_id。把结果画成‘并发去重’和‘显式重跑’两条时序图。",
    )

    # Chapter 10
    add_chapter_intro(
        doc,
        "第 10 章｜RAG 评测精讲：不用‘感觉不错’判断 AI",
        "一个答案错了，可能错在解析、召回、排序、生成、引用或权限。评测必须分层。",
        ["解释 Recall@K、MRR、Citation Coverage", "读懂 P50/P95", "设计最小 Golden Dataset"],
    )
    doc.add_heading("10.1 Golden Dataset 是长期资产", level=2)
    add_para(
        doc,
        "每个 EvaluationCase 保存问题、期望 KnowledgeUnit、期望来源系统/表/字段、期望答案关键词、目标字段和业务场景。它不是发布前临时挑几个漂亮问题，而是每次模型、Prompt、Chunk、Embedding 或检索策略变化后都能重复运行的基准。\n\n"
        "好的数据集要包含四类问题：明确有答案、同义表达、冲突/版本问题、明确无答案。还要按字段代码型、自然语言型、跨文档型和权限型切片。若只收集容易的精确关键词问题，Keyword 1.0 并不能证明真实能力。",
    )
    doc.add_heading("10.2 三个检索指标怎样理解", level=2)
    add_table(
        doc,
        ["指标", "直观问题", "简单例子"],
        [
            ("Recall@5", "应该找到的证据，有多少进入前 5？", "期望 2 条，前 5 命中 1 条，Recall@5=0.5"),
            ("Recall@10", "扩大到前 10 后是否补回漏召回？", "前 10 命中两条，Recall@10=1.0"),
            ("MRR", "第一条正确证据排得多靠前？", "第 1 名命中得 1；第 4 名首次命中得 0.25"),
        ],
        [1.20, 2.70, 2.60],
    )
    add_para(
        doc,
        "Recall 高但 MRR 低，说明正确证据能找到，却排得靠后；生成时 Top-K 小可能看不到它。Recall 低说明召回层就丢了证据，调 Prompt 通常无效。MRR 很好但答案仍错，则要检查模型是否正确使用证据、是否存在冲突、Citation 是否支持具体 Claim。",
    )
    doc.add_heading("10.3 生成与系统指标", level=2)
    add_bullets(
        doc,
        [
            "Citation Coverage：期望证据有多少实际进入 Citation。",
            "Groundedness：答案是否有引用且没有 unsupported claim；当前实现是较粗二值规则。",
            "Answer Correctness：当前确定性算法以关键词覆盖为主、引用覆盖为辅，不是另一个 LLM 裁判。",
            "P50：一半请求快于该值；P95：95% 请求快于该值，更能暴露尾部慢请求。",
            "Index Throughput：每秒处理多少 Chunk；还要同时看失败数与重试。",
            "Token Usage：真实 Chat 评测的费用依据；只做检索评测可以避免不必要付费。",
        ],
    )
    doc.add_heading("10.4 怎样解读当前结果", level=2)
    add_table(
        doc,
        ["模式", "Top1", "Recall@5/10", "MRR", "P50/P95"],
        [
            ("Keyword", "20/20", "1.00 / 1.00", "1.00", "36.4 / 56.3 ms"),
            ("Vector", "11/20", "0.60 / 0.60", "0.5625", "717.7 / 969.0 ms"),
            ("Hybrid", "20/20", "1.00 / 1.00", "1.00", "570.6 / 686.8 ms"),
        ],
        [1.15, 1.05, 1.55, 1.10, 1.65],
    )
    add_para(
        doc,
        "结果说明在这 20 个合成问题上，Hybrid 保留了 Keyword 的精确命中，同时链路真实调用了本地 Embedding 和 Milvus。它不等于生产准确率 100%，因为数据集小、完全合成、场景覆盖有限，而且本轮没有批量执行真实 DeepSeek 答案评测以避免未经授权产生费用。\n\n"
        "成熟表达应该同时给数字和限制：‘在 10 份合成文档、100 个 Chunk、20 个问题的本机验收上，Hybrid Recall@10 和 MRR 为 1.0；Vector-only Recall 为 0.60，说明当前场景必须保留精确关键词通道。下一步用真实脱敏集验证 Chunk、Embedding 和 Reranker。’",
    )
    add_source_paths(doc, ["backend/app/services/evaluation/rag_evaluator.py", "docs/semantic-retrieval/rag-evaluation.md", "backend/app/models/entities.py"])
    add_self_check(
        doc,
        [
            "Recall@K 和 MRR 分别回答什么问题？",
            "为什么 Recall 高但答案仍可能错？",
            "为什么 P95 比平均值更能暴露卡顿？",
            "当前 1.0 指标为什么不能称为生产准确率？",
            "检索评测与付费 Chat 答案评测为什么应拆开？",
        ],
        "建立 10 条最小 Golden Dataset：6 条有明确答案、2 条同义表达、1 条冲突、1 条无答案。手工标注期望 Unit，再比较三种 retrieval_mode。不要先改代码，先建立可重复基线。",
    )

    # Chapter 11
    add_chapter_intro(
        doc,
        "第 11 章｜把你遇到的报错放回架构中",
        "登录报错、按钮卡顿、上传无反应、模型测试失败，通常不是一个‘系统性能问题’，而是不同层次的边界没有被正确处理。",
        ["按层定位常见故障", "区分认证、任务、检索和模型错误", "理解前端反馈与后端幂等的配合"],
    )
    doc.add_heading("11.1 Authenticated user required", level=2)
    add_para(
        doc,
        "这个错误表示前端访问了需要认证的 API，但请求没有携带有效身份，或 Token 已失效。它和 PostgreSQL、Milvus、内存压力无关。正确处理是：前端发现 401 后清理无效登录态并跳转登录页；公共页面不要在未登录时直接请求受保护 Dashboard；后端继续坚持鉴权，不能为了页面不报错就取消权限。",
    )
    doc.add_heading("11.2 点击上传后‘没反应’", level=2)
    add_para(
        doc,
        "要依次看四个证据：浏览器是否发出请求；API 是否返回 202 和 job_id；BackgroundJob 是否 queued/running；Worker 日志是否有处理。若任务完成但页面没变化，是前端没有轮询/刷新；若 Job 一直 queued，是 Worker 或 Redis；若 Job failed，看安全错误摘要；若 Job completed 但 Vector Search 仍无新知识，检查是否 pending_reindex 和 active index 是否更新。",
    )
    doc.add_heading("11.3 页面卡顿不一定是数据库压力", level=2)
    add_para(
        doc,
        "页面切换卡顿可能来自前端重复请求、每个组件独立轮询、未缓存的 Dashboard 聚合、后端同步执行长任务、LLM/Embedding 超时、Worker 竞争 CPU 或 WSL/Docker 资源。资源管理器显示 69% 内存只能说明整体内存使用，并不能定位是哪条请求慢。\n\n"
        "可靠诊断要记录浏览器 Network Timing、后端 latency、数据库查询、Worker 任务和依赖 health。PostgreSQL 替换 SQLite 能改善并发、事务和锁，但不会自动修复前端重复请求、模型超时或 Celery 未启动。架构升级必须对应具体瓶颈。",
    )
    doc.add_heading("11.4 模型测试失败但 RAG 搜索正常", level=2)
    add_para(
        doc,
        "普通 Search 不调用 Chat，所以 Embedding/Milvus 正常时仍能返回候选；Ask 才需要 DeepSeek。模型测试失败应检查 ModelProfile、Base URL、模型名、API Key 环境变量和 Provider 状态。若 Search 失败，再检查 Embedding 服务、512 维配置、active Collection 和 Milvus。把 Chat 与 Embedding 拆开验证，可以立刻把故障范围缩小一半。",
    )
    doc.add_heading("11.5 好的按钮反馈应包含什么", level=2)
    add_bullets(
        doc,
        [
            "点击后立即进入 loading，按钮禁用并显示‘正在提交’。",
            "收到 202 后显示‘任务 #62 已排队’，提供任务详情入口。",
            "轮询 queued/running 进度，页面切换后仍可恢复任务状态。",
            "completed 显示结果摘要和下一步，例如‘需要重建语义索引’。",
            "failed 显示可行动的错误，不只显示 JavaScript 异常字符串。",
            "后端语义幂等兜底，刷新或双击不会创建重复任务。",
        ],
    )
    add_source_paths(doc, ["frontend/lib/api.ts", "frontend/app/jobs/", "backend/app/api/jobs.py", "backend/app/services/health_checks.py"])
    add_self_check(
        doc,
        [
            "401、任务 queued、pending_reindex、Chat 失败分别属于哪一层？",
            "为什么 69% 内存不能证明数据库是瓶颈？",
            "Search 正常、Ask 失败时你先检查什么？",
            "前端禁用按钮和后端幂等为什么都需要？",
        ],
        "下次遇到页面问题时固定记录：操作时间、页面 URL、请求 URL、HTTP 状态、响应体、job_id、后端日志时间、依赖 readiness。用这七项替代‘感觉很卡’。",
    )

    # Chapter 12
    add_chapter_intro(
        doc,
        "第 12 章｜现在的项目为什么还不是 Agent",
        "RAG 是 Agent 可以使用的一种工具；调用一次模型也不是 Agent。关键差别在谁决定下一步。",
        ["区分 Workflow、RAG 与 Agent", "映射当前 Agent-ready 能力", "说清尚缺 Planner、Tool、State 和 Loop"],
    )
    doc.add_heading("12.1 Workflow：代码决定路径", level=2)
    add_para(
        doc,
        "当前 grounded_answer 的顺序是后端写死的：先 HybridRetriever，没证据就拒答，有证据就构造 Prompt，调用固定输出 Schema，校验 Citation，再返回。模型只负责其中的生成步骤，不能说‘我先查任务状态，再决定是否检索文档，然后创建待确认问题’。\n\n"
        "这叫 deterministic orchestration。它非常适合监管场景，因为流程清楚、测试容易、成本可控。不要因为招聘岗位叫 Agent 工程师，就贬低 Workflow。工程选择取决于任务是否真的需要动态决策。",
    )
    doc.add_heading("12.2 Agent：模型在边界内决定下一步", level=2)
    add_code(
        doc,
        "Goal：确认信用卡贷款余额规则并返回有效出处\n"
        "State：当前项目、已知问题、已执行步骤、预算\n"
        "  → 模型选择 get_background_job\n"
        "  → Observation：索引任务仍在 running\n"
        "  → 模型决定先等待/告知用户，而不是直接回答\n"
        "  → 任务完成后选择 search_knowledge\n"
        "  → Observation：找到两条证据\n"
        "  → 模型决定 get_document 核对版本\n"
        "  → 形成答案并停止",
    )
    add_para(
        doc,
        "这里的路径不是后端为每个问题预先写死，而是模型根据 Observation 选择下一步。但模型没有直接执行权：它输出 Tool Request，服务端 Executor 做参数校验、项目鉴权、超时和审计，再调用已有 Service。",
    )
    doc.add_heading("12.3 当前已经具备哪些 Agent 底座", level=2)
    add_table(
        doc,
        ["Agent 概念", "当前可复用能力", "还缺什么"],
        [
            ("Goal/Context", "项目、目标字段、场景、用户问题", "统一 AgentRun 目标模型"),
            ("Knowledge", "PostgreSQL KnowledgeUnit + Milvus", "与会话/任务记忆分层"),
            ("Tools", "检索、目录、血缘、安全 SQL、任务、审核 Service", "Tool Schema 与 Registry"),
            ("Reasoner", "OpenAI-compatible Chat + 结构化输出", "动态下一步决策 Schema"),
            ("State", "BackgroundJob、业务状态、审计", "AgentRun/Step/Observation"),
            ("Guardrails", "权限、分级、脱敏、SafeSql、HITL", "工具级策略和预算"),
            ("Evaluation", "RAG Case/Run/Result", "任务成功、工具正确、轨迹安全"),
        ],
        [1.25, 2.75, 2.50],
    )
    doc.add_heading("12.4 为什么当前不是多 Agent", level=2)
    add_para(
        doc,
        "代码中没有多个独立角色、上下文、消息通道或协调器；没有 LangGraph、CrewAI、AutoGen 或 MCP；CozeConnector 只是占位，明确写着真实 HTTP 工作流留待后续。把检索函数、生成函数和审核函数分别叫三个 Agent，不会让系统变成多 Agent。\n\n"
        "多 Agent 只有在角色确实需要独立上下文、并行工作或相互复核时才有价值，例如监管分析、数据溯源和合规复核。否则它只会增加模型调用、延迟、费用和状态复杂度。",
    )
    add_source_paths(doc, ["backend/app/services/rag/grounded_answer_service.py", "backend/app/services/coze/connector.py", "README.md"])
    add_self_check(
        doc,
        [
            "当前 RAG 问答里是谁决定调用顺序？",
            "Agent 至少需要哪些组成？",
            "知识库为什么不等于完整 Agent Memory？",
            "MCP 为什么是工具接入协议而不是 Agent 本身？",
            "什么时候不应该使用多 Agent？",
        ],
        "把当前功能分成三列：必须固定 Workflow、可封装只读 Tool、需要模型动态决策。每个功能写出风险和是否需要人工审批。",
    )

    # Chapter 13
    add_chapter_intro(
        doc,
        "第 13 章｜怎样在本项目上实现第一个可控单 Agent",
        "先做有限工具、有限步骤、可观察、可停止的单 Agent，再考虑框架和多 Agent。",
        ["设计 Tool Schema", "定义 Agent State 和终止条件", "给出持久化、审批与评测方案"],
    )
    doc.add_heading("13.1 第一版只给三个只读工具", level=2)
    add_code(
        doc,
        "search_knowledge(\n"
        "  project_id: int, query: str, retrieval_mode: str='hybrid', top_k: int=10\n"
        ") -> SearchResult\n\n"
        "get_document(\n"
        "  project_id: int, document_id: int\n"
        ") -> DocumentSummary\n\n"
        "get_background_job(\n"
        "  project_id: int, job_id: int\n"
        ") -> JobStatus",
    )
    add_para(
        doc,
        "为什么先只读？因为读取错误通常可以纠正，写入错误会改变业务状态。三个工具已经足以完成‘任务是否完成→搜索知识→核对文档版本→回答’的最小闭环。每个工具都复用已有 Service，不允许 Agent 直接拿 SQLAlchemy Session 查询任意表。\n\n"
        "Tool Schema 不只是参数类型。它还应描述用途、允许角色、项目范围、超时、是否幂等、敏感输出字段和错误码。模型选工具只是建议，Executor 每次仍要使用当前用户 Principal 做项目鉴权。",
    )
    doc.add_heading("13.2 Agent State 要显式保存", level=2)
    add_code(
        doc,
        "AgentRun\n"
        "  id, project_id, user_id, goal, status\n"
        "  step_count, max_steps=5\n"
        "  token_budget, time_budget_ms\n"
        "  final_answer, termination_reason\n\n"
        "AgentStep\n"
        "  run_id, step_no, decision_code\n"
        "  tool_name, argument_summary\n"
        "  observation_summary, latency_ms, error_type\n"
        "  model_profile_id, prompt_version, created_at",
    )
    add_para(
        doc,
        "State 不能只存在模型上下文里。只放 Prompt 会导致进程重启后丢失、无法查看任务详情、无法恢复和审计。把 Run/Step 写 PostgreSQL 后，Celery 可以执行长 Agent，前端可以展示每步轨迹，失败后可以从最后安全 checkpoint 恢复。\n\n"
        "Memory 还要分三类：知识库是长期事实记忆；对话历史是用户交互记忆；AgentRun/Step 是任务过程记忆。它们的权限、保留时间和检索方式都不同，不能混成一个 vector store。",
    )
    doc.add_heading("13.3 一个最小 Agent Loop", level=2)
    add_code(
        doc,
        "state = load_or_create_run(goal, max_steps=5)\n"
        "while state.step_count < state.max_steps:\n"
        "    decision = llm.decide(goal, safe_state_summary, TOOL_SCHEMAS)\n"
        "    validate_decision_schema(decision)\n"
        "\n"
        "    if decision.action == 'finish':\n"
        "        return validate_and_finish(decision.answer)\n"
        "\n"
        "    ensure_tool_allowed(current_user, project_id, decision.tool)\n"
        "    ensure_not_repeated_without_new_observation(decision)\n"
        "    observation = execute_with_timeout(decision.tool, decision.args)\n"
        "    persist_step(decision, observation)\n"
        "    state = update_state(state, observation)\n"
        "\n"
        "return stop_safely('max_steps_exceeded')",
    )
    add_para(
        doc,
        "真正重要的不是 while 循环本身，而是每个边界：Decision 必须结构化；工具必须白名单；参数必须校验；执行必须重新鉴权；结果要脱敏和限长；重复调用要检测；每步有超时；达到最大步数必须安全停止；最终答案仍要 Citation 校验。",
    )
    doc.add_heading("13.4 加入写工具时必须暂停", level=2)
    add_para(
        doc,
        "第二阶段可以加入 request_knowledge_reindex、create_open_question、save_draft、create_review_task。模型提出写操作后，系统先生成影响摘要，例如‘将为项目 4 创建语义索引重建，预计 100 个 Unit，使用本地 BGE 模型，不产生外部费用’，状态变为 waiting_for_approval。只有有权限的人确认后，Executor 才提交 BackgroundJob。\n\n"
        "人工审批不是弹一个‘确定吗’就结束。审批记录要包含请求者、审批人、影响范围、参数摘要、时间、决定和关联 AgentStep，写入 AuditLog。",
    )
    doc.add_heading("13.5 Agent 怎样评测", level=2)
    add_bullets(
        doc,
        [
            "Task Success：最终是否完成业务目标，而不是文字是否流畅。",
            "Tool Selection Accuracy：应该查任务时是否选了 get_background_job。",
            "Argument Validity：project_id、document_id、top_k 等参数是否有效且有权限。",
            "Unsafe Action Rate：越权、任意 SQL、敏感外发、未经审批写操作必须为 0。",
            "Step Efficiency：成功任务平均步骤、重复工具率、最大步数触发率。",
            "Recovery：Provider timeout、Milvus 故障、Worker 重启后能否恢复或安全停止。",
            "Latency/Token/Cost：相对固定 Workflow 是否值得增加复杂度。",
        ],
    )
    add_callout(
        doc,
        "什么时候才考虑 LangGraph",
        "当你已经用普通 Python 证明需要状态图、暂停恢复、分支循环和持久化 checkpoint 时，LangGraph 才是解决具体问题的框架。若连 Tool、State、终止和评测都没有，先学框架 API 只会得到一个不可控 Demo。",
        "info",
    )
    add_source_paths(doc, ["建议新增 backend/app/services/agent/", "复用 backend/app/services/retrieval/", "services/task_queue/", "services/auth/", "services/governance/"])
    add_self_check(
        doc,
        [
            "模型提出 Tool Call 后谁真正执行？",
            "为什么 Agent State 必须持久化？",
            "第一版为什么只给三个只读工具？",
            "最大步数、超时和重复调用检测分别防什么？",
            "什么时候写工具必须进入 waiting_for_approval？",
        ],
        "不用任何 Agent 框架，写一个命令行 Demo：最多 5 步，三个 Mock Tool，覆盖成功、参数错误、重复调用、工具超时、达到最大步数。把每步 Trace 输出成 JSON。",
    )

    # Chapter 14
    add_chapter_intro(
        doc,
        "第 14 章｜十个动手实验：把‘看懂’变成‘能工作’",
        "以下实验按依赖顺序排列。每个实验都要保存输入、输出、截图、错误和你的解释。",
        ["独立运行并排障项目", "完成 RAG 诊断记录", "做出可展示的单 Agent 原型"],
    )
    experiments = [
        (
            "实验 1：画出运行组件并验证 Ready",
            "运行 .\\scripts\\项目启停.ps1 status；逐项访问前端、/docs、/health/ready；记录 PostgreSQL、Redis、Celery、Milvus、Embedding 的状态。",
            "一张组件图；每个地址说明；一个‘进程 Running 但 Ready 失败’的假设。",
        ),
        (
            "实验 2：追踪一个简单 API",
            "在 /docs 调用项目查询；从 API 路由追到 Service/ORM；故意提交错误参数和未登录请求。",
            "正常 200、参数 422、认证 401 的请求/响应；说明每个状态由哪层产生。",
        ),
        (
            "实验 3：上传并观察知识摄取",
            "上传一个小 Excel；记录 202 和 job_id；查看任务从 queued 到 completed；查询 Document、Version、Unit。",
            "任务时序图；Unit 的文件/Sheet/单元格出处；批处理和 Hash 的解释。",
        ),
        (
            "实验 4：重建版本化语义索引",
            "记录重建前 active version；提交重建；观察每批进度、Collection、count；重建后确认新 active。",
            "蓝绿索引图；model_fingerprint、corpus_hash、维度和向量数；失败不影响旧 active 的说明。",
        ),
        (
            "实验 5：比较三种检索模式",
            "对 10 个问题运行 keyword_only、vector_only、hybrid；保存 Top-5 和三个分数。",
            "一张对比表；至少两个关键词胜出和两个向量补召回的案例。",
        ),
        (
            "实验 6：验证 Grounded Answer",
            "分别问有答案、无答案、跨项目不可见和包含虚构 TABLE.FIELD 的问题。",
            "四类输出；Citation、unsupported_claims、open_questions；为什么属于安全成功或业务失败。",
        ),
        (
            "实验 7：诊断模型 Provider",
            "用正确配置测试一次；再分别制造错误 Key、错误 model、错误 base_url 和 timeout。",
            "错误分类矩阵；哪些重试、哪些直接失败；确保日志没有密钥。",
        ),
        (
            "实验 8：验证任务幂等",
            "并发提交同一重建；确认复用同一 queued/running Job；终态后再提交，确认创建 successor。",
            "Job ID 证据；幂等键作用域；显式重跑链。",
        ),
        (
            "实验 9：建立 Golden Dataset",
            "创建 20 条问题：12 有答案、4 同义、2 冲突、2 无答案；标注期望 Unit。",
            "Recall@5/10、MRR、Citation Coverage、P50/P95；失败按解析/召回/排序/生成/权限分类。",
        ),
        (
            "实验 10：最小单 Agent",
            "实现 search_knowledge/get_document/get_background_job 三工具、最多 5 步、结构化 Decision、Trace 和安全停止。",
            "可运行 CLI；10 条任务轨迹；工具正确率、平均步骤、越权率和失败原因。",
        ),
    ]
    for title, steps, deliverable in experiments:
        doc.add_heading(title, level=2)
        p = doc.add_paragraph()
        _keep_with_next(p)
        r = p.add_run("操作：")
        r.bold = True
        p.add_run(steps)
        p2 = doc.add_paragraph()
        r2 = p2.add_run("必须产出：")
        r2.bold = True
        r2.font.color.rgb = RGBColor.from_string(TEAL)
        p2.add_run(deliverable)

    doc.add_heading("14.1 每次实验的记录模板", level=2)
    add_code(
        doc,
        "日期 / 项目提交：\n"
        "实验目标：\n"
        "输入与前置状态：\n"
        "操作步骤：\n"
        "实际输出：\n"
        "与预期差异：\n"
        "定位到的源码：\n"
        "指标或测试：\n"
        "一个失败模式：\n"
        "两分钟面试表达：",
    )
    add_callout(
        doc,
        "真正的完成标准",
        "不是把十个实验点完，而是你能拿出日志、Job ID、数据库记录、检索排名、指标和 Trace，向别人证明系统为什么这样工作。证据会把‘我学过 Agent’变成‘我能做 Agent 工程’。",
        "teal",
    )

    # Conclusion and source guide
    add_chapter_intro(
        doc,
        "结语｜你应该怎样向面试官描述这个项目",
        "诚实区分已实现、你本人完成、原型和规划，比堆叠 Agent 名词更专业。",
        ["用因果关系讲项目", "用真实指标讲权衡", "坦诚说明当前缺口"],
    )
    add_para(
        doc,
        "你可以把项目概括为：面向银行一表通监管口径的证据驱动 AI 平台。多格式文件先解析为带版本、范围、敏感等级和原始位置的 KnowledgeUnit；PostgreSQL 保存事实与关键词索引，本地 BGE Embedding 把语料批量写入版本化 Milvus Collection；查询采用 Keyword + Vector 混合检索，回库做权限和有效性校验；DeepSeek 只读取 Top-K 证据生成结构化草稿，Citation 必须解析到真实可见 Unit，无证据或虚构表字段则返回待确认。解析、索引和评测通过 Redis/Celery 异步执行，并用 BackgroundJob、幂等键、checkpoint 和原子索引切换保证可靠性。\n\n"
        "然后明确边界：当前是受治理的 RAG Workflow，不是自主多 Agent。下一步把检索、文档、任务、目录和审核 Service 封装成受权限控制的 Tool，增加 AgentRun/Step 持久化、有限循环、预算、人工审批和轨迹评测，先实现可控单 Agent。",
    )
    add_callout(
        doc,
        "不要背一段漂亮话",
        "面试官深挖时会问‘哪个文件做的、失败怎么办、为什么不用纯向量、怎么防重复、指标如何算’。本教材的每一章都对应这些追问。你真正要练的是从业务问题走到源码与证据，而不是背术语。",
        "warning",
    )
    doc.add_heading("核心源码阅读顺序", level=2)
    add_numbers(
        doc,
        [
            "backend/app/api/knowledge_rag.py：找到上传、状态、重建、搜索和问答入口。",
            "backend/app/services/knowledge_ingestion/parsers.py 与 ingestion_service.py：看文件怎样变 Unit。",
            "backend/app/services/semantic_index/reindex.py：看版本化索引与 checkpoint。",
            "backend/app/services/embeddings/ 与 vector/milvus.py：看向量生成和搜索。",
            "backend/app/services/retrieval/hybrid_retriever.py：逐行看两路召回、归一化和融合。",
            "backend/app/services/rag/grounded_answer_service.py 与 citation_validator.py：看证据问答。",
            "backend/app/services/llm/prompt_runtime.py：看 Provider、数据外发、结构化输出和日志。",
            "backend/app/services/task_queue/ 与 backend/app/workers.py：看异步、幂等和任务处理。",
            "backend/app/services/evaluation/rag_evaluator.py：看指标怎样从结果计算。",
            "backend/app/services/coze/connector.py：确认 Agent/Coze 当前只是扩展点。",
        ],
    )
    add_callout(
        doc,
        "30 天后最低目标",
        "你能独立启动和排障；能完整解释一次知识上传与一次 RAG 问答；能区分 LLM、Embedding、RAG、Workflow 和 Agent；能实现一个有限工具、有限步骤、可观察的单 Agent；能讨论权限、幂等、重试、预算、评测和人工审批。",
        "success",
    )

    path = OUT_DIR / "00_从零精讲_银行一表通AI平台与Agent开发.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_book())
