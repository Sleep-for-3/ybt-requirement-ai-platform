# -*- coding: utf-8 -*-
"""生成《银行智能平台 · 10 项落地建议与现有能力对照（增量评估）》桌面 .docx。

运行（本机）：
  C:\\Users\\admin\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe gen_platform_gap_report.py
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = os.path.join(
    os.environ["USERPROFILE"],
    "Desktop",
    "银行智能平台-10项落地建议与现有能力对照-2026-10-17.docx",
)

AS_OF = "2026-10-17"
AVAIL_CM = 17.0

doc = Document()

_sec = doc.sections[0]
_sec.left_margin = Cm(2.0)
_sec.right_margin = Cm(2.0)
_sec.top_margin = Cm(2.0)
_sec.bottom_margin = Cm(2.0)

_normal = doc.styles["Normal"]
_normal.font.name = "Microsoft YaHei"
_normal.font.size = Pt(10.5)
_normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

RED = (0xC0, 0x00, 0x00)
GREEN = (0x1F, 0x6F, 0x3C)
GREY = (0x59, 0x59, 0x59)
ORANGE = (0xB5, 0x6A, 0x00)


def _fix_run(run, size=10.5, bold=None, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    return run


def h(text, level=1):
    p = doc.add_heading("", level=level)
    _fix_run(p.add_run(text), size={0: 20, 1: 15, 2: 12.5, 3: 11.5}[level], bold=True)
    return p


def para(text, size=10.5, bold=False, color=None, indent=0, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    _fix_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def bullet(text, size=10.5, color=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    _fix_run(p.add_run(text), size=size, color=color)
    return p


def table(headers, rows, widths=None, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, htxt in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        _fix_run(cell.paragraphs[0].add_run(htxt), size=font_size, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _fix_run(p.add_run(str(val)), size=font_size)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def card(title, verdict, verdict_color, evidence, gain, caveat=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(1)
    _fix_run(p.add_run(title), size=11, bold=True)
    _fix_run(p.add_run("　【" + verdict + "】"), size=10.5, bold=True, color=verdict_color)
    for label, val in (("平台现状", evidence), ("加了的收益", gain)) + ((("注意事项", caveat),) if caveat else ()):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.left_indent = Cm(0.6)
        _fix_run(q.add_run(label + "："), size=9.5, bold=True)
        _fix_run(q.add_run(val), size=9.5)


# ============================== 封面 ==============================
h("银行智能平台 · 10 项落地建议与现有能力对照（增量评估）", 0)
para(
    "评估时点：%s（Asia/Shanghai）　｜　评估对象：ybt-requirement-ai-platform 现有代码库　｜　方法：逐项对照真实实现，只读代码核对，未运行"
    % AS_OF,
    size=10,
    color=GREY,
)
para(
    "结论先行：10 项建议中，5 项是真缺口（建议加，按此顺序），3 项是小增量（有条件地加），2 项建议不加——"
    "其中「Langfuse 自托管」与平台已有能力高度重叠，且会引入第二套审核状态真源。",
    size=10.5,
    bold=True,
)

h("对照总表", 1)
table(
    ["建议项", "判定", "平台现状（代码证据）", "加上去能提升什么"],
    [
        [
            "1. PaddleOCR PP-StructureV3",
            "强烈建议",
            "parsers.py:95-103 的 _pdf() 只用 pypdf 抽文本，抽不到就直接告警「未启用 OCR」——扫描件产出为 0",
            "解锁能力：扫描版制度/报送模板从「完全不可入库」变为可入库；单元格坐标让引用高亮可落到表格格子",
        ],
        [
            "2. bge-reranker-base（FastEmbed）",
            "强烈建议",
            "hybrid_retriever.py:287 的 \"rerank_score\" 直接等于 final_score——字段名是假的，检索链路里没有任何交叉编码器",
            "把 rerank_score 变成真的重排分；长文档里精确条文不再被高覆盖长文挤出前 5",
        ],
        [
            "3. bm25s + jieba",
            "强烈建议",
            "keyword_index.py:22 正则切分 + 2 字 bigram；:48 按 title 1.25 / structured 1.5 取最大权重。无 IDF、无长度归一",
            "补上词项区分度（「1104 报表」比「管理」重）与文档长度归一；hybrid 召回的准确侧提升",
        ],
        [
            "4. PG 哈希链 + Merkle 证明",
            "强烈建议",
            "governance/audit.py 的 record_audit() 只写 AuditLog 行，含脱敏与 correlation_id，但无 prev_hash、无链式校验",
            "审计日志从「约定不可改」变为「改了一定被发现」，可向监管做外部可验证证明",
        ],
        [
            "5. RefChecker 思路移植",
            "建议",
            "citation_validator.py 仅 48 行，只校验引用「存在且可见」；grounded_answer_service.py:130/135 是规则式拒绝，不是蕴含判定",
            "引用校验从「编号合法」升级为「被引原文确实支持该结论」，产出引用精确率",
        ],
        [
            "6. chonkie 切分层",
            "可加",
            "text_processing.py:10 是固定窗 chunk_text(1000, overlap=150)；但入库解析层已按页/表/行/单元格产出结构化单元",
            "仅对长篇幅制度条文（连续散文）有实质收益；表格类已有结构，收益接近 0",
        ],
        [
            "7. RRF 融合",
            "可加",
            "hybrid_retriever.py:258-262 是分数级加权和 keyword*kw + vector*vw + rule_boost，依赖 settings.hybrid_*_weight 标定",
            "免调权重；但对本平台更关键的是 AUTHORITY_RANK（:19/:254）现在是 0.05 的 rule_boost，须提为独立 tie-break 阶段",
        ],
        [
            "8. RAGAS",
            "部分可加",
            "rag_evaluator.py 已有 recall@5/@10、MRR、source/table/field hit、citation_coverage、groundedness，且显式给了分母",
            "只值 LLM-judge 那一部分：把「代理指标」换成裁判式 faithfulness；文件自己写明 not_expert_accuracy",
        ],
        [
            "9. Langfuse 自托管",
            "建议不加",
            "governance/workflow.py 已有 10+ 条多角色门流程；ai_skills.py:342 已有 independent_review_required 独立性校验；RagEvaluationCase 支持 feedback→用例提升",
            "会重复建设，且引入 Postgres+ClickHouse+Redis+S3 四个新服务、第二套审核状态真源；改为扩展 RagEvaluationCase 加双人复核字段",
        ],
        [
            "10. instructor",
            "建议不加",
            "openai_compatible.py:144-162 已实现 validate → repair（\"You repair JSON formatting\" 修复轮）→ revalidate；structured_outputs.py 本身就是纯 pydantic",
            "instructor 的核心行为已被手写实现，边际价值接近 0；Presidio 降为纵深防御（content_redactor.py 已覆盖 7 类正则）",
        ],
    ],
    widths=[3.4, 1.7, 6.2, 5.7],
    font_size=8,
)

doc.add_page_break()

# ============================== 一 ==============================
h("一、五项真缺口：建议按此顺序加", 1)
para(
    "排序依据是「解锁程度 > 收益/成本 > 合规风险」。前三项都在检索与文档链路上，改动集中、可独立验收；"
    "先做 1、2 两项的原因见下。",
    size=10,
)

card(
    "① PaddleOCR PP-StructureV3（Apache-2.0）",
    "强烈建议 · 最高优先",
    GREEN,
    "knowledge_ingestion/parsers.py 的 _pdf() 用 pypdf 逐页 extract_text()，取不到文本时写一条告警直接跳过（原文："
    "「第 N 页无可提取文本，未启用 OCR」）。也就是说扫描版监管文件目前一键入库后不会产生任何知识单元。",
    "这是唯一一项「从 0 到 1」的建议。它解锁两类此前完全无法处理的输入：扫描版制度文件、影印版报送模板。"
    "PP-StructureV3 同时给出表格单元格坐标与文本坐标，使平台已有的 locator / source_cell_range 字段能真正落到格子级，"
    "引用高亮不再只到页。",
    "PP-OCRv6 tiny 仅 1.5M 参数，CPU 可跑；Windows 官方支持。注意 PaddleOCR-VL 权重许可未核实，先只用 PP-StructureV3 的代码部分。",
)

card(
    "② bge-reranker-base（MIT，走 FastEmbed TextCrossEncoder）",
    "强烈建议 · 最高优先",
    GREEN,
    "hybrid_retriever.py:287 写的是 \"rerank_score\": round(final_score, 4)——重排分就是融合分的副本，"
    "整条知识检索链路里不存在交叉编码器。需要说明：平台确实有一个「模型重排」，但它是 ai_skills/field_rerank.py 的"
    "LLM 重排，只作用于字段候选（api/ai_skills.py:105、agent 工具 rerank_field_candidates），与段落重排是两件事。",
    "把 rerank_score 变成真的重排分，并让顶层结果由交叉编码器重新排序。对当前最痛的场景——"
    "一条精确的监管条文淹没在长篇文件的高覆盖命中里——效果最直接。",
    "FastEmbed 已内置 TextCrossEncoder，理论上零新依赖；C-MTEB T2Reranking map 67.28。上线后应把 hybrid_retriever 的"
    "top-k 放大再重排，否则重排无候选可用。",
)

card(
    "③ bm25s + jieba 重构 keyword_index.py",
    "强烈建议",
    GREEN,
    "keyword_index.py:22 tokenize() 英文按标识符边界、中文取 2 字 bigram；:48 weighted_tokens() 对每个 token 取"
    "content 1.0 / title 1.25 / structured 1.5 的最大权重，同义词按 0.85 折扣。hybrid_retriever.py:351-357 的打分是"
    "「命中权重 / 总权重 × 0.7」的覆盖率，既无 IDF 也无文档长度归一。",
    "补齐两件关键词检索的必备件：词项区分度与长度归一。现状下一份长文档只要覆盖了全部查询词，"
    "得分与一条短精确条文相同；引入 IDF 后低频监管术语的权重才合理。现有同义词表与三级字段权重应保留为 BM25 的词项 boost，"
    "而不是丢掉。",
    "jieba 为 MIT、bm25s 为 MIT，风险低。改造集中在单文件，可用已有 rag_evaluator 的 recall@5/@10 直接做前后对照。",
)

card(
    "④ PostgreSQL 哈希链 + Merkle 包含证明（自研）",
    "强烈建议",
    GREEN,
    "services/governance/audit.py 的 record_audit() 写入 AuditLog，含 before/after 摘要、correlation_id、"
    "IP 掩码与敏感键过滤（SENSITIVE_KEYS），质量不低；但表中没有 prev_hash，也没有任何链式或树式校验字段。",
    "审计日志目前的不可篡改性只是「约定」：任何能连库的人都能改一行历史而不留痕。加上 prev_hash 链后，"
    "改一行会让后续全部校验失败；再用 Merkle 包含证明即可对外部（监管/审计方）证明「某条记录在第 N 版日志中存在」，"
    "而不必交出全表。",
    "算法本身简单，自研成本低于引入 immudb（服务端许可未核实）或 pymerkle（GPL-3.0）的合规成本。"
    "需要同时设计周期性外部锚定（对象存储/文件摘要）才闭环。",
)

card(
    "⑤ RefChecker 思路移植到 citation_validator.py",
    "建议",
    GREEN,
    "citation_validator.py 全文仅 48 行，做的是「引用编号是否解析到可见且启用的 KnowledgeUnit」——存在性与可见性校验。"
    "grounded_answer_service.py 已在做规则式拒绝：:130 拒绝未经证据支持的表字段、:135 拒绝超出本次证据范围的知识单元编号。"
    "也就是说，有伪造拦截，但没有「引用的原文是否真的蕴含这句结论」的判定。",
    "把引用校验从「编号合法」升级为「支持性成立」，从而产出 citation precision。这是评估报告里"
    "「无依据率」分母做准的前提——现在该分母只能靠规则命中，无法回答「引对地方了吗」。"
    "建议移植判定思路与提示词，而非直接依赖库（该库研究级、有停更风险）。",
    "需要与 citation_validator 的 fail-closed 语义协调：蕴含判定引入模型后，如何避免把「判定失败」误当成「引用造假」。",
)

# ============================== 二 ==============================
h("二、三项小增量：有条件地加", 1)

card(
    "⑥ chonkie 切分层",
    "可加 · 收益受限于已结构化",
    ORANGE,
    "text_processing.py:10 的 chunk_text(text, chunk_size=1000, overlap=150) 是固定窗口切分，无语义或 late chunking。"
    "但入库侧 knowledge_ingestion/parsers.py 早已按语义结构产出单元：PDF 按页、表格按行、单元格带 "
    "source_cell_range 与 locator。",
    "实质收益只在「连续散文形态的长制度条文」上——那里当前是 1000 字硬切，可能切断一条完整规则。"
    "对表格与已结构化内容收益接近 0。",
    "建议先跑一次统计：现有知识单元中由固定窗口产生的占比，若很低则优先级继续下压。",
)

card(
    "⑦ RRF 融合",
    "可加 · 但要处理 AUTHORITY_RANK",
    ORANGE,
    "hybrid_retriever.py:212-220 做权重归一后，:258-262 用 keyword_score*kw + vector_score*vw + rule_boost 融合。"
    "AUTHORITY_RANK（:19 定义，:254 使用）目前只贡献两件事：写入 authority_rank 字段、以及命中时给 0.05 的 rule_boost。",
    "RRF 免去对 settings.hybrid_vector_weight / hybrid_keyword_weight 的标定，在关键词分与向量分尺度不同时更稳健。",
    "关键约束：本平台的权威等级（regulatory_formal 6 / regulatory_qa 5 / internal_policy 4 …）是合规语义，"
    "不能在 RRF 里被当成一路稀疏信号压平。正确做法是 RRF 融合后，把 AUTHORITY_RANK 提升为独立的 tie-break 阶段。",
)

card(
    "⑧ RAGAS（只取其 LLM-judge 部分）",
    "部分可加",
    ORANGE,
    "rag_evaluator.py 已有 recall_at_5/10、MRR、source/table/field_hit_rate、citation_coverage、groundedness，"
    "并且显式给出分母：successful_query_count、answer_correctness_denominator、generation_coverage、status_counts，"
    "metric_notes 里还主动声明「以上均为代理指标，不等于银行专家判定的业务正确率」。",
    "唯一真正新增的信号是裁判式 faithfulness：用语义蕴含替代现在的关键词覆盖代理。"
    "这能回答「回答是否被证据支持」，而现在只能回答「关键词是否出现」。",
    "必须固定裁判模型并为其单独留档（平台已有 get_prompt_runtime() + ModelCallLog 可承接），"
    "否则裁判本身成为不可审计的黑箱；同时需关闭其匿名遥测。",
)

doc.add_page_break()

# ============================== 三 ==============================
h("三、两项建议不加：会重复建设", 1)
para(
    "这一节是本次评估与早前那份选型报告的主要差异。早前报告的「Langfuse 自托管」与「instructor」两项，"
    "在逐项核对真实实现后判定为重复建设。",
    size=10,
    color=RED,
    bold=True,
)

h("⑨ Langfuse 自托管 —— 建议不加，改为扩展自有评测表", 2)
para("早前报告把它列为「唯一能一步到位补上人工标注+双人复核闭环」。核对后这个判断需要修正：", size=10)

table(
    ["早前认为缺的能力", "平台实际上已有的实现", "位置"],
    [
        ["多角色审核流程", "10+ 条工作流定义，含 business_review → technical_review → final_review 多步角色门，强制校验指派人是否持有该角色",
         "governance/workflow.py:46-85、:188-197、:444-456"],
        ["双人 / 独立性复核", "已有独立复核校验：人工复核者不得是该测试运行创建者，冲突直接 409",
         "api/ai_skills.py:342 independent_review_required"],
        ["业务反馈采集", "AIUserFeedback 表，记录 correct_source_system / table / field、rating、comment、output_hash",
         "models/entities.py:597"],
        ["反馈升级为评测用例", "由人工纠正值直接构造 RagEvaluationCase，并回链 source_feedback_id",
         "api/knowledge_rag.py:361-370"],
        ["显式分母的评测", "answer_coverage_denominator、answer_correctness_denominator、generation_coverage、status_counts",
         "evaluation/rag_evaluator.py:288-342"],
    ],
    widths=[3.4, 8.8, 4.8],
    font_size=8,
)

para(
    "结论：所谓「P1 人工标注缺口」缺的不是标注工具，而是两样具体的东西——① 成规模的、经业务确认的期望标签集；"
    "② 评测用例级的双人复核字段（现在的独立性校验只在技能人工复核上，用例本身没有第二审核人）。"
    "引入 Langfuse 需要新增 Postgres + ClickHouse + Redis + S3/MinIO 四个服务，并产生第二套审核状态真源，"
    "属于典型的重复所有者问题。"
    "建议改为：在 RagEvaluationCase 上增加 reviewer_a / reviewer_b / reviewed_at 字段，"
    "复用现有 ReviewTask 与角色门，把双人复核做成用例状态机的一个状态。",
    size=10,
)

h("⑩ instructor —— 建议不加（核心行为已手写实现）", 2)
bullet(
    "services/llm/openai_compatible.py:144-162 已经实现了 instructor 的核心循环：先 model_validate，"
    "捕获 ValidationError 后用「You repair JSON formatting. Return only the corrected JSON object.」发起一次修复轮，"
    "再次 model_validate，仍失败才降级。"
)
bullet(
    "services/llm/structured_outputs.py 本身就是一组纯 pydantic 模型（BaseModel + model_validator），"
    "没有非结构化解析环节需要 instructor 接管。"
)
bullet(
    "因此 instructor 的能力已被覆盖，引入它只会增加一层依赖与一处行为差异。"
    "早前报告的「零边际成本」判断在依赖关系上成立（它确实是 RAGAS 的依赖），但在能力上是零收益。"
)
bullet(
    "同批的 Presidio：content_redactor.py 已覆盖 7 类正则（手机号、证件号、账号、邮箱、IP、数据库连接串、密码/密钥），"
    "确实缺实体级识别（客户姓名、地址）。但平台既定姿态是「不自动连接银行数据库、仅使用授权脱敏样本」"
    "（requirement_uat.py:88、:96），真实 PII 流量有界，故 Presidio 降级为纵深防御，不必进第一批。",
    color=GREY,
)

# ============================== 四 ==============================
h("四、一句话汇总", 1)
para(
    "把 10 项建议按「平台是否已有等价物」重新切分后，真正值得投入的顺序是："
    "先做 ① PaddleOCR（唯一的能力从 0 到 1）与 ② 真重排（修掉一个名不副实的字段），"
    "再做 ③ BM25 与 ④ 审计哈希链（补齐两个基础件的空缺），"
    "然后 ⑤ 引用蕴含判定（把分母做准）。"
    "⑥⑦⑧ 视评测数据再定。⑨⑩ 不做。"
    "其中 ③ 与 ⑦ 的价值可以直接用现有 rag_evaluator 的 recall@5/@10 做前后对照来验证，"
    "建议在动工前先把一组对照基线跑出来。",
    size=10.5,
)

para(
    "—— 本报告全部判断来自对 ai-platform/backend/app 现有代码的只读核对，逐条标注了文件与行号，可复核。"
    "未运行任何测试或基准，性能与准确率类结论均待实测。 ——",
    size=9,
    color=GREY,
)

doc.save(OUT)
print("SAVED:", OUT)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
