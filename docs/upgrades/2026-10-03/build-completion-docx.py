# -*- coding: utf-8 -*-
"""Generate the Desktop completion report (DOCX) for the DSH upgrade task."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\admin\Desktop\银行智能平台-DSH升级任务-完成情况说明-2026-10-04.docx"
FONT = "Microsoft YaHei"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text, size=10.5, bold=False, color=None, style=None, space_after=4):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        set_font(r, size=size)
        p.paragraph_format.space_after = Pt(2)


def table(doc, header, rows, widths=None, size=9.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_font(r, size=size, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            set_font(r, size=size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


doc = Document()
for s in doc.sections:
    s.left_margin = Cm(2.2)
    s.right_margin = Cm(2.2)

title = doc.add_paragraph()
tr = title.add_run("银行智能平台（ybt-requirement-ai-platform）\nDSH 升级任务 — 完成情况说明")
set_font(tr, size=18, bold=True)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

meta = doc.add_paragraph()
mr = meta.add_run(
    "生成日期：2026-10-04    分支：dsh/banking-semantic-agent-v2    当前提交：8d76006（已推送 GitHub）\n"
    "基线：17486f1（《银行智能平台评估报告-2026-10-03》审查基线）    "
    "工作目录：C:\\Users\\admin\\Downloads\\ai-platform-dsh-20260930-150943\\ai-platform"
)
set_font(mr, size=9, color=(0x55, 0x55, 0x55))
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------------- 1 ----------------
doc.add_heading("一、任务概要", level=1)
para(doc, "依据《银行智能平台评估报告-2026-10-03》与《DSH升级任务书-2026-10-03》，在现有代码上完成安全修复、正式业务闭环和可验证的工程发布，分三个阶段：")
bullets(doc, [
    "第一阶段 W01–W06：供应链与发布基线、机构熔断与刷新令牌原子轮换、审核内容不可变、SQL 限额/脱敏/连接器契约、后台任务幂等与恢复、人工编辑与模型配置完整性，并修复 B23 过期测试契约。",
    "第二阶段 W07/W08/W11：候选历史与真实业务状态、AI 评测与外发治理、代表性 UAT 包。",
    "第三阶段 W09/W10：可恢复备份、真实依赖集成与压力故障验收。",
])
para(doc, "约束：每个工作包独立提交并留存实施/验收记录；保护现有业务数据、权限与外发守卫；隔离环境测试；不删除失败测试、不关闭 lint、不扩大权限来换取全绿。", size=10)

# ---------------- 2 ----------------
doc.add_heading("二、总体进度一览（截至本次提交）", level=1)
table(doc,
      ["工作包", "对应问题", "状态", "关键证据 / 提交"],
      [
          ["W01 供应链与发布基线", "B01/B17/B18/过期迁移 head", "已修复 / 已验证（含运行实例升级 + 浏览器回归）", "1782bac 0950bcd 909f0f8 7b9bc46 7640d97 7714227 caf60bd f452b29 8d76006"],
          ["W02 机构熔断 + 令牌原子轮换", "B02/BA05、B05/BA04", "已修复 / 已验证", "16767d7、2193915；真实 PG 20 并发仅 1 次成功"],
          ["W03 审核内容与产物不可变", "B03/BA01", "已修复 / 已验证", "行锁补测 1d6cb6e；真实 PG 8 并发仅 1 个开启新修订"],
          ["W04 SQL 限额/脱敏/连接器契约", "B04/B06/B09", "已修复 / 已验证", "真实 PG 验证子查询/CTE/字符串无法绕过限额"],
          ["W05 后台任务幂等与恢复", "B07/BA06", "已修复 / 已验证（迁移已在业务库执行）", "f4fba42、888e0ea；真实 PG 12 并发仅 1 个领取"],
          ["W06 人工编辑与模型配置", "B10–B13、B15–B16", "已修复 / 已验证（真实浏览器交互回归仍缺）", "ac6ffbd db1cc57 6ea5779 f6e90ec 4f6cd3d"],
          ["B23 过期测试契约", "迁移 head 硬编码 / 退役 token 误伤", "已修复 / 已验证", "1782bac"],
          ["W07 候选历史与真实状态", "B14、进度误导、假空", "大部分完成（剩五类状态汇总条）", "c75311d b79ddbb fae2f74 8652686 d942718 a38c676 52b0779 9ed2377 8bd1d7e"],
          ["W08 AI 外发治理", "外发判定分散", "网关已收敛且不可绕过；黄金集/阈值待银行", "5a910f2；项目 11 仍不在白名单"],
          ["W09 可恢复备份", "B08、备份命名误导", "命名已修正、工具与可恢复性已验证；系统级备份待批准", "45e5cc1 cfc2775 64cbf6f"],
          ["W10 真实依赖与压力", "真实依赖未集成", "真实 PG 层验收已完成；多 worker/对象存储/压力待授权", "w02/w03/w04/w10 系列脚本"],
          ["W11 代表性 UAT 包", "缺业务验收包", "模板已备；执行待四类角色与真实样本", "646aff2 5223471"],
      ],
      widths=[3.6, 3.0, 4.6, 5.4])

# ---------------- 3 ----------------
doc.add_heading("三、本次（窗口卡住后续做）完成的工作", level=1)

doc.add_heading("3.1 补齐 B18 第三处落地缺口：前端从未展示发布标识差异", level=2)
bullets(doc, [
    "问题：frontend/lib/build-info.ts 的 frontendIdentity()/identityMismatches() 早已存在且有 6 例单测，但没有任何界面调用；旧前端配新后端时界面看起来完全正常。",
    "改动：新增 frontend/components/ReleaseIdentityNotice.tsx（仅在不一致或任一侧为 unknown 时告警，一致时静默），挂载于 AppShell 顶栏；/api/version 为无鉴权端点，未登录也能核对。",
    "提交：f452b29（代码）、9b5d00c（记录）。验证：前端 node --test 253 passed、tsc exit 0、lint exit 0、构建 exit 0。",
])

doc.add_heading("3.2 W01/B01：前端安全升级在运行实例上落地（Next 15.5.27）", level=2)
bullets(doc, [
    "依赖：npm ci → next 15.5.27（维护 LTS）、react/react-dom 19.3.0、eslint-config-next 15.5.27；package.json / package-lock.json 未变。",
    "构建：启动器 local-deploy.ps1 build → exit 0；生成 frontend/.next/BUILD_ID。",
    "升级后仓库级验证：node --test 253 passed、tsc --noEmit exit 0、next lint exit 0、生产构建 exit 0。",
    "运行实例：前端以 Next 15.5.27 生产构建在 127.0.0.1:3000 提供服务；按本机约束未启动 Celery worker/beat。",
    "附带说明：frontend/next-env.d.ts 由 Next 15 重新生成（新增 .next/types/routes.d.ts 引用），已随本次提交。",
])

doc.add_heading("3.3 发布标识真正注入运行实例（发现并修复一个真实坑）", level=2)
bullets(doc, [
    "现象：GET /api/version 长期返回 app_commit=unknown，即使 backend/build-info.json 存在。",
    "根因：PowerShell 5.1 的 Set-Content -Encoding utf8 会写入 UTF-8 BOM（EF BB BF），导致 Python json.loads 失败 → 回退为空 → unknown。",
    "修复：改用无 BOM 写入（UTF8Encoding($false)），并给前端构建注入 NEXT_PUBLIC_APP_COMMIT / NEXT_PUBLIC_BUILD_TIME。",
    "结果：/api/version 报出 app_commit=9b5d00c…、build_time=2026-10-04T13:29:05Z；前端内联标识与之一致，新告警条保持静默。",
])

doc.add_heading("3.4 W05 迁移 202610030001 在业务库执行（经您批准，先转储）", level=2)
bullets(doc, [
    "触发证据：升级后实例上 GET /api/projects/5/requirement-workspace → 500，日志明确「字段 background_jobs.lease_owner 不存在」（代码已要求新列，库未迁移）。",
    "迁移前全库转储：.local-run/backups/pre-migration-20261004-214454/ybt_dsh_handoff_v2-pre-202610030001.dump，11.8 MB，SHA-256 4C9B8E47…4B60。",
    "执行：alembic upgrade head → 202610020051 → 202610030001（exit 0）；alembic current = 202610030001 (head)。",
    "迁移后核对：新列 lease_owner / lease_expires_at 与索引 ix_background_jobs_lease_expires_at 均存在；业务数据行数前后一致（projects 11 / users 29 / target_fields 22 / background_jobs 123 / mapping_versions 3）；/api/health/ready 的 alembic_revision 转为 healthy。",
])

doc.add_heading("3.5 升级后的真实浏览器回归（真实 Chrome）", level=2)
table(doc,
      ["检查项", "结果"],
      [
          ["登录页渲染与标题", "标题「一表通字段口径智能辅助平台」；登录表单完整，含「仅限脱敏模拟数据环境」提示"],
          ["登录", "本地管理员登录成功 → 重定向 /projects"],
          ["会话持久（续期基础）", "sessionStorage 的 access/refresh token 刷新后仍生效，恢复已登录态"],
          ["发布标识一致", "浏览器内 GET /api/version 200；新告警条不出现（一致时不误报）"],
          ["项目切换", "顶栏下拉 11 → 5；URL 变为 /projects?projectId=5"],
          ["需求（工作台）", "迁移前 500 / 迁移后 200；页面无「无法连接服务」，显示 17/17 个字段"],
          ["需求进度不误导（W07）", "步骤条显示 AI 口径分析 9/17、人工校核与导出 2/17（部分完成不显示全部完成）"],
          ["语义目录", "/semantics?projectId=5 渲染「共 14 个语义概念」，筛选齐全"],
          ["交付", "/deliverables?projectId=5 正常渲染（0 个交付包，空态而非报错）"],
      ],
      widths=[4.2, 12.4])

# ---------------- 4 ----------------
doc.add_heading("四、当前运行实例状态", level=1)
table(doc,
      ["组件", "地址/端口", "状态"],
      [
          ["PostgreSQL 18.4（便携）", "127.0.0.1:5432", "运行中（库 ybt_dsh_handoff_v2，已迁移到 202610030001）"],
          ["Redis（便携）", "127.0.0.1:6379", "运行中（PING→PONG）"],
          ["本地 Embedding（FastEmbed）", "127.0.0.1:11434", "运行中"],
          ["Milvus 向量栈（WSL Docker）", "127.0.0.1:19530", "运行中（standalone/etcd/minio 均 healthy）"],
          ["后端 FastAPI", "127.0.0.1:8000", "运行中（/health/live 200）"],
          ["前端 Next", "127.0.0.1:3000", "运行中（Next 15.5.27 生产构建）"],
          ["Celery Worker / Beat", "—", "未启动（遵循本机约束：进程操作仅限 app.main:app）"],
      ],
      widths=[5.0, 4.0, 7.6])
para(doc, "就绪检查 /api/health/ready = not_ready：application / database / alembic_revision / storage / redis / vector_store / llm_provider / embedding_provider / semantic_index / disk_space 共 10 项 healthy；仅 task_queue unhealthy。", size=10)
para(doc, "说明：task_queue 不健康是预期且如实记录——按本机约束未启动 Celery worker，队列面无消费者；数据库层的原子领取与 owner 围栏已在隔离真实 PostgreSQL 单独验证。", size=10, color=(0x99, 0x33, 0x00))

# ---------------- 5 ----------------
doc.add_heading("五、如何独立复跑（主审可用）", level=1)
para(doc, "前置（测试套默认语义）：$env:TASK_QUEUE_PROVIDER='inline'; $env:AUTH_MODE='optional'", size=10)
bullets(doc, [
    "后端全量回归：cd backend; & \".venv\\Scripts\\python.exe\" -m pytest tests -q  →  期望 1433 passed（exit 0）",
    "前端逻辑/规则：cd frontend; <node> --test tests/*.test.mjs  →  期望 253 passed",
    "前端类型/构建：<node> node_modules\\typescript\\bin\\tsc --noEmit --incremental false；<node> node_modules\\next\\dist\\bin\\next build  →  exit 0",
    "隔离真实依赖（只写自己的隔离库）：docs/upgrades/2026-10-03/w02_postgres_concurrency.py、w10_postgres_job_claim.py、w03_postgres_mapping_guard.py、w04_postgres_safe_query.py、w10_postgres_lease_fencing.py、w05_postgres_migration_rehearsal.py、w09_backup_plan.py",
    "逐条主张的复跑方式见：docs/upgrades/2026-10-03/VERIFICATION-INDEX.md",
    "发布/回滚步骤见：docs/upgrades/2026-10-03/RELEASE-RUNBOOK.md；本次升级/迁移记录见：docs/upgrades/2026-10-03/W01-R2B-DEPLOY-AND-MIGRATION.md",
])

# ---------------- 6 ----------------
doc.add_heading("六、未完成 / 受阻项（不得当作通过）", level=1)
table(doc,
      ["类别", "事项", "缺什么"],
      [
          ["需银行专家", "≥30 正/负/歧义例黄金标注集与试行业务准确率阈值", "银行专家标注；不得由我方编造真值"],
          ["需银行安全", "各素材/提示路径应申报的数据分级", "业务/安全策略判定（现有网关已收敛且不可绕过）"],
          ["需银行运维", "系统级备份策略：加密、权限、保留周期、异机副本、一致性点", "运维策略与独立恢复环境"],
          ["需产品/银行", "容量假设（并发用户、后台任务、资产规模）与压测", "先给容量数字，再定阈值"],
          ["需授权", "真实 Redis/Celery 多 worker 重投递与 >900s 租约续期", "授权启动 worker 进程"],
          ["需授权", "银行身份/只读库接入与 Oracle/SQL Server/DB2 连接器真实验收", "真实只读账号、驱动、网络"],
          ["需样本", "UAT 四类角色业务主流程、三档断点（390/768/桌面）折叠恢复", "测试账号 + 真实样本"],
          ["需确认", "业务准确率阈值须在试点前明确，不得事后用成功率替换", "银行专家"],
      ],
      widths=[3.0, 8.6, 5.0])

# ---------------- 7 ----------------
doc.add_heading("七、风险与回滚", level=1)
bullets(doc, [
    "业务库迁移：本次仅新增可空列与索引（向后兼容），执行前已全库转储；如需回退，可先 alembic downgrade -1 回到 202610020051，或直接用转储恢复。",
    "前端升级：如需回退，回退 frontend/package.json、package-lock.json 与两个页面文件到升级前提交 → npm ci → 启动器 build → start（见 Runbook §5）。",
    "已识别的发布流程缺口：启动器 build 不注入 NEXT_PUBLIC_APP_COMMIT / NEXT_PUBLIC_BUILD_TIME，常规发布路径下前端标识会是 unknown（本次已手工注入，建议固化进 Runbook 或启动器）。",
    "写 backend/build-info.json 必须使用无 BOM 的 UTF-8，否则 /api/version 会静默返回 unknown。",
])

# ---------------- 8 ----------------
doc.add_heading("八、审阅入口", level=1)
bullets(doc, [
    "实施与验收总记录：docs/upgrades/2026-10-03/README.md",
    "本次升级与迁移记录：docs/upgrades/2026-10-03/W01-R2B-DEPLOY-AND-MIGRATION.md",
    "逐条主张复跑索引：docs/upgrades/2026-10-03/VERIFICATION-INDEX.md",
    "发布与回滚 Runbook：docs/upgrades/2026-10-03/RELEASE-RUNBOOK.md",
    "W11 UAT 包：docs/upgrades/2026-10-03/W11-UAT-PACKAGE.md、w11/",
    "审查复现证据：docs/reviews/2026-10-03/",
    "GitHub：https://github.com/Sleep-for-3/ybt-requirement-ai-platform（分支 dsh/banking-semantic-agent-v2）",
])

doc.save(OUT)
print("SAVED:", OUT)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
