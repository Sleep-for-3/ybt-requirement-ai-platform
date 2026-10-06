"""Generate the Desktop supplement report (.docx) for the R01-R08 round."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(r"C:\Users\admin\Desktop\银行智能平台-2026-10-06补交R01-R08交付说明.docx")
CN_FONT = "微软雅黑"


def set_cjk(run) -> None:
    run.font.name = CN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)


def para(document, text: str, *, size: int = 10.5, bold: bool = False):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_cjk(run)
    return p


def bullet(document, text: str, *, size: int = 10.5):
    p = document.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(size)
    set_cjk(run)
    return p


def heading(document, text: str, level: int = 1):
    h = document.add_heading(level=level)
    run = h.add_run(text)
    run.font.size = Pt(15 if level == 1 else 12.5)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    set_cjk(run)
    return h


def table(document, headers: list[str], rows: list[list[str]]):
    t = document.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
        set_cjk(run)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(9)
            set_cjk(run)
    return t


def main() -> int:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = CN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    title = document.add_heading(level=0)
    run = title.add_run("银行智能平台 2026-10-06 补交交付说明（R01–R08）")
    run.font.size = Pt(20)
    set_cjk(run)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("分支 dsh/banking-semantic-agent-v2　·　复核基线 3db4db6 → 最终提交 9427cb6")
    r.font.size = Pt(10)
    set_cjk(r)

    heading(document, "一、结论", 1)
    para(document, "《银行智能平台-3db4db6交付复核与补交任务-2026-10-06》的 8 项补交（R01–R08）"
                   "已全部完成源码修复：先 P1（R01、R02、R04、R05、R06、R07），再 P2（R03、R08）。"
                   "每项独立提交，均给出问题、修改文件、commit、修复前证据、修复后结果与验证边界。")
    para(document, "两个阻断项已解决：R01 的正常续期后原请求重放（此前 refresh 200 后仍返回 401）、"
                   "R02 的失租后领域提交（此前失租尝试仍提交文件/通知/审计，接管后重复）。", bold=True)
    para(document, "此外补交了复核指出的两类“原任务范围内”接线：C07 的评测结果页、"
                   "C11 的公共进度区（不再另行排期）。")

    heading(document, "二、最终验证（真实计数与退出码）", 1)
    table(document, ["检查", "结果"], [
        ["后端全量回归", "1519 passed / 0 failed / 0 skipped（21:11，退出码 0）"],
        ["前端单元测试", "281 passed / 0 failed"],
        ["前端 tsc / lint / build", "全部退出码 0"],
        ["R06 发布回滚门禁", "14/14 检查通过（退出码 0）"],
        ["R07 真实队列验收", "13/13 步骤 ok（退出码 0）"],
    ])
    para(document, "计数对照：后端 3db4db6 为 1507，本轮新增 12 条 → 1519；"
                   "前端 267 → 281（新增 R01 6、R04 4、R08 4）。")

    heading(document, "三、逐项交付", 1)
    table(document, ["编号", "优先级", "主题", "Commit", "修复前 → 修复后"], [
        ["R01", "P1", "正常续期后原请求仍失败 / 旧响应归属", "4e99ce1", "8 passed/5 failed → 13/0"],
        ["R02", "P1", "失租后领域提交与重复结果", "1e5b837", "3 failed/1 passed → 4/0"],
        ["R03", "P2", "默认模型摘要保真", "0881b7e", "（与 R04/R05 同批）"],
        ["R04", "P1", "回答评分与结果页接线", "0881b7e+090b7aa", "后端 8 failed → 25/0"],
        ["R05", "P1", "执行与评分使用不可变快照", "0881b7e", "同上"],
        ["R06", "P1/P2", "发布与回滚门禁", "54d2c81", "8 failed/6 passed → 14/14"],
        ["R07", "P1", "真实队列验收有效性", "229c6e4", "空断言通过 → 13/13 退出码 0"],
        ["R08", "P2", "持续恢复、按任务隔离、公共进度区", "9427cb6", "10 passed/4 failed → 14/0"],
    ])

    heading(document, "四、各修复的实质", 1)
    bullet(document, "R01：把“身份转换”与“同身份令牌轮换”分开。续期只条件更新令牌、不改变身份代次，"
                     "原请求因此能真正重放；续期前先核验请求归属，旧账号的晚到 401 不会为新会话启动续期；"
                     "飞行记录只清理自己。")
    bullet(document, "R02：新增每次尝试的执行权协议（不可变 owner token + 失租信号），"
                     "把校验与持久化放进同一事务；领域提交边界在失租时整体回滚，"
                     "已写入的对象存储对象被回收，接管后只保留一份有效结果。")
    bullet(document, "R03：运行开始时解析并固定真正采用的模型档案，摘要、逐条结果与模型调用三者一致；"
                     "运行中改动默认配置不会悄悄换模型。")
    bullet(document, "R04：回答代理与证据代理分开计算（仅引文命中不再取得回答满分）；"
                     "摘要给出生成覆盖率、分母、逐条状态计数与实际模型；结果页显示这些信息、"
                     "逐条状态与原因，无生成样本时显示“不适用”；合并元数据时保留已有的降级原因。")
    bullet(document, "R05：执行前把用例的输入与真值冻结成不可变快照，执行与评分只读快照；"
                     "运行中/运行后编辑都不影响该次评测，且可从快照重算数据集版本。")
    bullet(document, "R06：发布时“上一版”改为从配置文件读取（不再被新标签覆盖）；"
                     "回滚的就绪检查耗尽、compose 失败、前端无响应都明确非零退出，"
                     "并如实说明配置文件已被修改、保留改动前备份；身份核验扩展到提交+构建时间与实际组件。")
    bullet(document, "R07：验收改用真正会成功、且留下可观测领域结果的作业；"
                     "断言首次必须成功且领域计数增量恰好为 1，重投后必须零新增效果；"
                     "Redis 守卫按键类型区分队列与结果键（不再 WRONGTYPE）；重置既有隔离库必须先显式授权。")
    bullet(document, "R08：轮询恢复成功后清掉停轮状态并继续排程；故障状态按任务归属（不再全局串扰），"
                     "恢复/终态/解除订阅/退出都会清理；公共进度区统一接入断连提示与“立即重试”，"
                     "9 个使用该组件的页面自动获得。")

    heading(document, "五、如实记录的与复核结论不同之处", 1)
    bullet(document, "R06：复核指出“失败不改配置文件”的笼统描述不成立 —— 本轮确认并修正："
                     "镜像校验阶段确实不改，但 compose 之后失败确实会留下已修改的配置；"
                     "脚本与文档均已明确写清。")
    bullet(document, "R07：复核指出夹具必然失败却算通过 —— 本轮实证：换用有效夹具后首次投递真的经过 broker"
                     "并且成功（领域计数增量 1）。此前“通过”源于夹具必失败 + 断言接受 failed + 复用旧幂等键。")
    bullet(document, "R08：复核的 visibility 剩余问题先被我复现为失败（4 条新用例修复前全红），再修复。")

    heading(document, "六、未验证 / 待验收（不得当作通过）", 1)
    for text in [
        "R08 的真实浏览器验证未执行（需访问正在运行的 3000/8000 与真实业务作业，本轮不操作既有服务）。",
        "R02 中 6 个在领域服务内部自行提交的 handler 未接入执行权（已逐项列出）。",
        "R06 未在真实 Docker/Compose 执行；构建清单尚未接入 CI。",
        "R07 未多机部署，未验证 broker 高可用。",
        "R05 未验证大用例集的 JSON 体积，未验证“从快照恢复并重跑”。",
        "R01 未真实浏览器验证；身份代次为每标签页进程内。",
        "未在 CI（Linux runner）执行；未做容器镜像层扫描；安全扫描沿用上一轮 0 漏洞结论。",
        "真实银行资料与签署人仍缺失（W11 银行验收未完成）。",
    ]:
        bullet(document, text)

    heading(document, "七、边界与保护", 1)
    for text in [
        "未操作业务库、未执行其迁移（只对隔离库做迁移验证）。",
        "未改动原有 3000/8000 服务。",
        "未清共享 Redis：验收全程在专用库运行，共享库的哨兵消息前后完全未变。",
        "未调用未批准的真实模型（模型边界为替身或 mock）。",
        "未部署：源码补交完成不等于现有运行实例已生效。",
    ]:
        bullet(document, text)

    heading(document, "八、交付物位置", 1)
    para(document, r"ai-platform\docs\upgrades\2026-10-06-supplement\：台账 README.md、"
                   r"R01–R08 逐项记录、final-verification.md（最终计数）与 R06 回归脚本。", size=10)

    document.save(OUT)
    print(OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
