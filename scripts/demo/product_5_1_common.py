from __future__ import annotations

import hashlib
import json
import random
import sqlite3
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEMO_SEED = 20260830
COLLECTION_DATE = date(2026, 8, 31)
JULY_COLLECTION_DATE = date(2026, 7, 31)
DEMO_PROJECT_NAME = "一表通产品数据智能演示"
DEMO_PROJECT_CODE = "PRODUCT_5_1_DEMO"
TARGET_TABLE_CODE = "YBT_5_1_PRODUCT_BUSINESS_BASIC"
TARGET_TABLE_NAME = "表5.1 产品业务基本信息"

ROOT = Path(__file__).resolve().parents[2]
DEMO_ROOT = ROOT / "demo" / "product_5_1"
DEFAULT_RUNTIME_ROOT = ROOT / ".demo_runtime" / "product_5_1"

SOURCE_DATABASES = {
    "product_center": "src_product_center.db",
    "core_deposit": "src_core_deposit.db",
    "credit_loan": "src_credit_loan.db",
    "wealth_agency": "src_wealth_agency.db",
    "treasury_market": "src_treasury_market.db",
    "reference": "src_reference.db",
    "regulatory_mart": "mart_regulatory.db",
}

TARGET_COLUMN_BY_CODE = {
    "E010001": "product_id",
    "E010002": "institution_id",
    "E010003": "product_name",
    "E010004": "product_number",
    "E010005": "accounting_subject_type",
    "E010007": "product_category",
    "E010008": "proprietary_flag",
    "E010009": "currency",
    "E010010": "product_term",
    "E010011": "establishment_date",
    "E010012": "maturity_date",
    "E010013": "product_issue",
    "E010014": "interest_rate_type",
    "E010015": "product_status_code",
    "E010018": "agency_institution_name",
    "E010016": "remark",
    "E010017": "collection_date",
}

SCENARIOS = [
    ("CORPORATE_DEPOSIT", "单位存款", "deposit"),
    ("PERSONAL_DEPOSIT", "个人存款", "deposit"),
    ("CORPORATE_LOAN", "公司贷款", "loan"),
    ("PERSONAL_LOAN", "个人贷款", "loan"),
    ("CORPORATE_OVERDRAFT", "法人账户透支", "loan"),
    ("WEALTH", "理财", "wealth"),
    ("FUND", "基金", "agency"),
    ("INSURANCE", "保险", "agency"),
    ("BOND", "债券", "treasury"),
    ("BOND_LENDING", "债券借贷", "treasury"),
    ("BANK_CARD", "银行卡", "card"),
    ("CENTRAL_BANK_BUSINESS", "中央银行业务", "treasury"),
    ("AGENCY_BUSINESS", "代客业务", "agency"),
]

SEMANTIC_CONCEPTS = [
    ("PRODUCT", "产品", "business_term", "满足金融消费者最小定制、最细颗粒度且可唯一识别的产品或业务。"),
    ("PRODUCT_IDENTIFIER", "产品标识", "business_term", "用于跨系统唯一识别产品或业务的稳定标识。"),
    ("PRODUCT_CATEGORY", "产品类别", "code_set", "监管产品类别代码，可按业务实质多值填报。"),
    ("ACCOUNTING_SUBJECT_TYPE", "科目类型", "code_set", "产品或业务对应的会计科目大类。"),
    ("PROPRIETARY_PRODUCT_FLAG", "自营标识", "code_set", "区分自营、代客及混合产品。"),
    ("CURRENCY", "币种", "dimension", "产品业务涉及的标准币种集合。"),
    ("PRODUCT_TERM", "产品期限", "business_rule", "持续性产品为0，一次性产品为成立日至到期日的持续天数。"),
    ("PRODUCT_LIFECYCLE", "产品生命周期", "business_term", "产品成立、存续、停用和到期的生命周期。"),
    ("PRODUCT_STATUS", "产品状态", "code_set", "监管状态01正常、02停用。"),
    ("INTEREST_RATE_TYPE", "利率类型", "code_set", "固定、浮动或固定/浮动利率。"),
    ("AGENCY_INSTITUTION", "代客产品所属机构", "business_term", "仅代客产品填报的所属机构名称。"),
    ("REPORTING_INSTITUTION", "报送机构", "business_term", "实际业务经办最小单位或监管允许的共同上级机构。"),
    ("PRODUCT_ISSUE", "产品期次", "business_term", "同一产品品牌或系列的发行期次。"),
    ("BUSINESS_VS_PRODUCT", "业务与产品边界", "regulatory_rule", "表5.1同时覆盖产品类和业务类数据。"),
]


@dataclass(frozen=True)
class Product:
    product_id: str
    product_code: str
    product_name: str
    domain: str
    category: str
    subject_type: str
    proprietary_flag: str
    currency: str | None
    lifecycle_type: str
    launch_date: date
    maturity_date: date | None
    stop_date: date | None
    status: str
    rate_type: str | None
    issue_no: int | None
    institution_id: str
    issuer_name: str | None
    department_id: str
    product_line: str
    product_class: str
    product_group: str
    base_product: str
    customer_segment: str
    channel: str
    region: str
    amount_limit: int
    description: str
    source_status: str


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_regulatory_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    first_six: list[list[str | None]] = []
    for row in worksheet.iter_rows(min_col=1, max_col=6, values_only=True):
        values = [None if value is None else str(value).strip() for value in row]
        if any(value not in (None, "") for value in values):
            first_six.append(values)
    table_row = next((row for row in first_six if row[0] == "数据表名称"), None)
    header_index = next((index for index, row in enumerate(first_six) if row[0] == "数据项编码"), None)
    if table_row is None or table_row[2] != TARGET_TABLE_NAME or header_index is None:
        raise ValueError(f"工作簿不是预期的{TARGET_TABLE_NAME}: {path}")
    headers = first_six[header_index]
    expected_headers = [
        "数据项编码", "数据项名称", "数据类别", "数据格式",
        "字段业务定义（监管原始口径）", "字段业务定义（监管定义细化）",
    ]
    if headers != expected_headers:
        raise ValueError(f"监管表头不匹配: {headers}")
    fields: list[dict[str, Any]] = []
    supplemental_fields: list[dict[str, Any]] = []
    for row in first_six[header_index + 1 :]:
        code, name, category, data_format, original, refined = row
        if code and code.startswith("E"):
            fields.append({
                "field_code": code, "field_name": name, "data_category": category,
                "data_format": data_format, "regulatory_original_definition": original,
                "regulatory_refined_definition": refined,
                "physical_column": TARGET_COLUMN_BY_CODE.get(code),
            })
        elif code == "补充字段" or (fields and name and not any(row[2:])):
            supplemental_fields.append({"field_name": name, "rule": category})
    missing = sorted(set(TARGET_COLUMN_BY_CODE) - {item["field_code"] for item in fields})
    if missing:
        raise ValueError(f"监管工作簿缺少字段: {missing}")
    canonical = json.dumps(first_six, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return {
        "source_path": str(path.resolve()),
        "source_sha256": sha256_file(path),
        "regulatory_columns_sha256": hashlib.sha256(canonical).hexdigest(),
        "sheet_name": worksheet.title,
        "table_name": TARGET_TABLE_NAME,
        "collection_scope": next(row[2] for row in first_six if row[0] == "采集范围"),
        "reporting_scope_detail": next(row[2] for row in first_six if row[0] == "一表通报送范围确认"),
        "headers": expected_headers,
        "fields": fields,
        "supplemental_fields": supplemental_fields,
    }


def connect_fresh(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()
    connection = sqlite3.connect(path)
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def execute_many(connection: sqlite3.Connection, sql: str, rows: Iterable[tuple[Any, ...]]) -> None:
    connection.executemany(sql, list(rows))
    connection.commit()


def generate_products(seed: int = DEMO_SEED) -> tuple[list[Product], list[dict[str, Any]]]:
    rng = random.Random(seed)
    specs = [
        ("deposit", 340, ["0101", "0102", "0103", "0106", "0107", "0201", "0202", "0204", "0208"], "02"),
        ("loan", 370, ["1101", "1102", "1103", "1202", "1501", "1601", "1801", "1802", "1901"], "01"),
        ("wealth", 160, ["3201"], "06"),
        ("fund", 90, ["3208"], "06"),
        ("insurance", 60, ["3208"], "06"),
        ("treasury", 180, ["2101", "2201", "2202", "2301", "2401", "2501", "2701", "2801", "3401", "3402"], "01"),
        ("card", 80, ["0201", "1802", "3103"], "00"),
    ]
    fixed = {
        ("deposit", 0): "单位活期存款", ("deposit", 1): "单位一年期定期存款",
        ("deposit", 2): "个人三年期定期存款", ("loan", 0): "流动资金贷款",
        ("loan", 1): "法人账户透支", ("loan", 2): "个人住房贷款",
        ("loan", 3): "普惠经营贷款", ("wealth", 0): "净值型理财产品",
        ("fund", 0): "A类基金", ("fund", 1): "C类基金",
        ("insurance", 0): "代销保险", ("card", 0): "借记卡产品",
        ("treasury", 0): "企业债券", ("treasury", 1): "债券借贷",
        ("treasury", 2): "存放央行款项", ("treasury", 3): "向中央银行借款",
    }
    domain_prefix = {"deposit": "DEP", "loan": "LOAN", "wealth": "WLT", "fund": "FUND", "insurance": "INS", "treasury": "TRY", "card": "CARD"}
    domain_cn = {"deposit": "存款", "loan": "贷款", "wealth": "理财", "fund": "基金", "insurance": "保险", "treasury": "金融市场", "card": "银行卡"}
    products: list[Product] = []
    issues: list[dict[str, Any]] = []
    global_index = 0
    for domain, count, categories, subject in specs:
        for index in range(count):
            global_index += 1
            category = categories[index % len(categories)]
            prefix = domain_prefix[domain]
            product_id = f"YBT{prefix}{global_index:08d}"
            product_code = f"{prefix}{index + 1:06d}"
            product_name = fixed.get((domain, index), f"{domain_cn[domain]}产品{index + 1:04d}")
            lifecycle_type = "continuous" if domain in {"deposit", "loan", "card"} or category in {"2201", "2202"} else "one_off"
            launch_date = date(2018, 1, 1) + timedelta(days=rng.randint(0, 2800))
            maturity_date = None if lifecycle_type == "continuous" else launch_date + timedelta(days=rng.choice([90, 180, 365, 730, 1095]))
            stopped = index % 17 == 0 and index > 3
            status = "inactive" if stopped else "active"
            stop_date = (launch_date + timedelta(days=rng.randint(120, 1200))) if stopped else None
            proprietary_flag = "02" if domain in {"fund", "insurance"} else ("03" if domain == "card" and index % 11 == 0 else "01")
            issuer_name = f"{rng.choice(['华夏','海岳','东洲','中原'])}{'基金管理有限公司' if domain == 'fund' else '保险股份有限公司'}" if proprietary_flag == "02" else None
            rate_type = None if domain in {"fund", "insurance", "treasury"} and category not in {"2801"} else ("02" if domain in {"loan", "wealth"} else "01")
            currency = rng.choice(["CNY", "CNY", "CNY", "USD", "EUR"])
            source_status = status
            product = Product(
                product_id=product_id, product_code=product_code, product_name=product_name, domain=domain,
                category=category, subject_type=subject, proprietary_flag=proprietary_flag, currency=currency,
                lifecycle_type=lifecycle_type, launch_date=launch_date, maturity_date=maturity_date,
                stop_date=stop_date, status=status, rate_type=rate_type, issue_no=(index + 1 if domain in {"wealth", "treasury"} else None),
                institution_id=f"LIC00000001{(index % 24) + 1:04d}", issuer_name=issuer_name,
                department_id=f"D{(index % 12) + 1:03d}", product_line=f"{domain_cn[domain]}产品线",
                product_class=f"{domain_cn[domain]}产品类{index % 4 + 1}", product_group=f"{domain_cn[domain]}产品组{index % 8 + 1}",
                base_product=f"{domain_cn[domain]}基础产品{index % 16 + 1}", customer_segment=rng.choice(["CORP", "RETAIL", "SME", "FI"]),
                channel=rng.choice(["BRANCH", "MOBILE", "ONLINE", "AGENCY"]), region=rng.choice(["NATIONAL", "EAST", "SOUTH", "WEST"]),
                amount_limit=rng.choice([100000, 500000, 1000000, 5000000]),
                description=f"用于监管演示的{product_name}，包含完整产品层级与生命周期属性。", source_status=source_status,
            )
            products.append(product)

    def replace(index: int, **changes: Any) -> None:
        current = products[index]
        products[index] = Product(**{**current.__dict__, **changes})

    # Controlled exceptions are deterministic and remain below five percent of records.
    for index in range(340 + 370 + 160, 340 + 370 + 160 + 8):
        replace(index, issuer_name=None)
        issues.append(_issue(products[index], "MISSING_AGENCY_ISSUER", "warning", "代客产品缺少所属机构"))
    for index in range(20, 28):
        replace(index, status="inactive", stop_date=None)
        issues.append(_issue(products[index], "STOPPED_WITHOUT_STOP_DATE", "warning", "停用产品未维护停用日期，目标层应回填采集日期"))
    for index in range(40, 47):
        replace(index, status="active", stop_date=date(2026, 5, 31))
        issues.append(_issue(products[index], "ACTIVE_WITH_STOP_DATE", "error", "正常产品存在停用日期"))
    for index in range(340 + 370, 340 + 370 + 5):
        replace(index, maturity_date=None)
        issues.append(_issue(products[index], "ONE_OFF_WITHOUT_MATURITY", "error", "一次性产品缺少到期日期"))
    for index in range(100, 106):
        replace(index, category="9999")
        issues.append(_issue(products[index], "INVALID_PRODUCT_CATEGORY", "error", "产品类别不在监管代码表"))
    for index in range(150, 156):
        replace(index, currency=None)
        issues.append(_issue(products[index], "MISSING_CURRENCY", "error", "产品币种缺失"))
    for index in range(200, 204):
        replace(index, product_code=products[199].product_code)
        issues.append(_issue(products[index], "DUPLICATE_PRODUCT_NUMBER", "warning", "内部产品编号重复"))
    for index in range(360, 365):
        replace(index, source_status="inactive" if products[index].status == "active" else "active")
        issues.append(_issue(products[index], "CROSS_SYSTEM_STATUS_MISMATCH", "warning", "产品中心与业务系统状态不一致"))
    overdraft_index = next(i for i, item in enumerate(products) if item.product_name == "法人账户透支")
    replace(overdraft_index, subject_type="00")
    issues.append(_issue(products[overdraft_index], "OVERDRAFT_SUBJECT_AMBIGUITY", "warning", "法人账户透支存在资产/负债/其他科目解释冲突，应形成开放问题"))
    return products, issues


def _issue(product: Product, code: str, severity: str, description: str) -> dict[str, Any]:
    return {"product_id": product.product_id, "product_name": product.product_name, "issue_code": code, "severity": severity, "expected_behavior": description}


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
