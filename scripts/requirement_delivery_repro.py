"""Reproduce R1/R2/R3 for the requirement-delivery upgrade round.

R1: an XLSX whose data cell holds a formula that was never recalculated (no cached
    value) loses that cell silently, because the knowledge parser reads
    ``data_only=True`` and reports nothing.
R2: a candidate whose citations are all authorized but whose body says the
    opposite of the evidence passes ``check_candidate_compliance`` with no
    violation, because the contract only checks reference membership.
R3: exporting the same frozen content twice yields different bytes, because the
    DOCX writer stamps a fresh timestamp into the package.

The script only observes and prints; it never asserts the broken behaviour as
desired.  Run from ``backend`` so the app package is importable:

    .\\.venv\\Scripts\\python.exe ..\\scripts\\requirement_delivery_repro.py --json out.json
"""
from __future__ import annotations

import argparse
from hashlib import sha256
from io import BytesIO
import json
import os
import sys
import time

from openpyxl import Workbook

# Keep the run isolated: importing the app must never reach a real provider.
for _name, _value in {
    "LLM_PROVIDER": "mock",
    "LLM_BASE_URL": "",
    "LLM_MODEL": "",
    "LLM_API_KEY_ENV_NAME": "",
    "EMBEDDING_PROVIDER": "mock",
    "EMBEDDING_BASE_URL": "",
    "EMBEDDING_MODEL": "",
    "EMBEDDING_DIMENSION": "0",
    "EMBEDDING_API_KEY_ENV_NAME": "",
    "VECTOR_STORE_PROVIDER": "mock",
    "MILVUS_URI": "",
}.items():
    os.environ.setdefault(_name, _value)

from app.services.knowledge_ingestion.parsers import parse_document  # noqa: E402
from app.services.requirement_candidate_contract import check_candidate_compliance  # noqa: E402
from app.services.requirement_word import export_word  # noqa: E402


def r1_formula_without_cache() -> dict:
    """Build a header-recognisable XLSX where the 口径 column is an uncached formula."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "字段口径"
    sheet.append(["字段代码", "字段名称", "口径"])
    # Written by a program that never opened the file in Excel: openpyxl stores the
    # formula without a cached result, so data_only=True reads the cell as None.
    sheet.append(["CUST_NO", "客户号", "=IF(1=1,\"只报送正常贷款余额\",\"\")"])
    buffer = BytesIO()
    workbook.save(buffer)
    raw = buffer.getvalue()

    units, warnings = parse_document("r1-formula.xlsx", raw, "regulatory_qa")
    texts = [unit.content for unit in units]

    # Read the same bytes twice: with and without data_only, to show the formula
    # text exists in the file while the "value" read by ingestion is empty.
    from openpyxl import load_workbook

    formula_cell = load_workbook(BytesIO(raw), data_only=False)["字段口径"]["C2"].value
    value_cell = load_workbook(BytesIO(raw), data_only=True)["字段口径"]["C2"].value
    return {
        "case": "R1",
        "title": "无缓存公式静默丢值",
        "file_bytes": len(raw),
        "formula_cell_data_only_false": formula_cell,
        "formula_cell_data_only_true": value_cell,
        "unit_count": len(units),
        "unit_texts": texts,
        "unit_type": [unit.unit_type for unit in units],
        "warnings": warnings,
        "observed_loss": (
            "口径" not in "".join(texts)
        ),
        "warnings_mention_formula": any("公式" in str(item) or "缓存" in str(item) for item in warnings),
    }


def r2_authorized_but_wrong() -> dict:
    """Authorized references, body asserts the opposite of the evidence."""
    context = {
        "physical_sources": [],
        "evidence": [
            {
                "unit_id": 11,
                "title": "个人贷款统计制度",
                "regulatory_version": "2026-v1",
                "source_category": "policy",
                "content": "只报送正常贷款余额。",
            }
        ],
        "script_basis": {
            "rules": [
                {
                    "rule_id": "script-edge-1",
                    "filter_condition": "LOAN_STATUS = 'NORMAL'",
                    "join_condition": None,
                    "aggregation_rule": "SUM(LOAN_BALANCE)",
                    "code_mapping_rule": None,
                    "transformation_expression": "LOAN_BALANCE",
                }
            ]
        },
    }
    candidate = {
        "business_definition": "",
        "processing_logic": "取贷款余额",
        "final_content": "报送所有逾期贷款余额，无需筛选正常状态。",
        "physical_references": [],
        "evidence_unit_ids": [11],
        "script_rule_ids": ["script-edge-1"],
        "policy_comparisons": [],
        "gaps": [],
    }
    violations = check_candidate_compliance(context, candidate)
    return {
        "case": "R2",
        "title": "合法引用配错误正文",
        "evidence_text": context["evidence"][0]["content"],
        "candidate_body": candidate["final_content"],
        "violations": violations,
        "contradiction_undetected": violations == [],
    }


def r3_word_bytes_unstable() -> dict:
    """Render the same frozen content twice and compare file hashes."""
    content = {
        "requirement": {
            "name": "个人贷款余额报送需求",
            "content_version": 1,
            "target_table_id": 1,
            "field_ids": [],
            "document_ids": [],
            "source_table_ids": [],
            "mart_table_ids": [],
        },
        "target_table": {"table_name": "监管客户表", "table_code": "EAST_CUSTOMER"},
        "fields": [],
        "gaps": [],
        "formal_delivery": {
            "content_version": 1,
            "version_no": "v1",
            "review_record": [
                {"step": "业务审核", "decision": "通过", "reviewer": "p0c-biz",
                 "decided_at": "2026-10-23T10:00:00+08:00", "comment": "同意"}
            ],
        },
        "script_basis": {},
        "policy_comparisons": {},
        "manual_ownership": {},
    }
    first = export_word(content)
    time.sleep(1.2)
    second = export_word(content)
    return {
        "case": "R3",
        "title": "同快照 Word 重复导出字节不一致",
        "first_sha256": sha256(first).hexdigest(),
        "second_sha256": sha256(second).hexdigest(),
        "byte_identical": first == second,
        "first_size": len(first),
        "second_size": len(second),
    }


def r1_after() -> dict:
    """R1 修复后：无缓存公式必须变成可定位的阻断项。"""
    from app.services.knowledge_ingestion.parsers import parse_document_with_report

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "字段口径"
    sheet.append(["字段代码", "字段名称", "口径"])
    sheet.append(["CUST_NO", "客户号", "=IF(1=1,\"只报送正常贷款余额\",\"\")"])
    buffer = BytesIO()
    workbook.save(buffer)
    raw = buffer.getvalue()

    units, warnings, report = parse_document_with_report("r1-formula.xlsx", raw, "regulatory_qa")
    payload = report.to_dict() if report is not None else {}
    return {
        "case": "R1-after",
        "title": "无缓存公式→阻断项可见",
        "unit_texts": [unit.content for unit in units],
        "warnings": warnings,
        "formula_cells": payload.get("formula_cells"),
        "formula_uncached": payload.get("formula_uncached"),
        "blocking_count": payload.get("blocking_count"),
        "has_blocking": bool(report.has_blocking) if report is not None else None,
        "warnings_mention_formula": any("公式" in str(item) for item in warnings),
    }


def r2_after() -> dict:
    """R2 修复后：引用合法但正文反转固定过滤条件时必须判为冲突。"""
    from app.services.requirement_claim_support import evaluate_claim_support

    context = {
        "physical_sources": [],
        "evidence": [
            {
                "unit_id": 11,
                "title": "个人贷款统计制度",
                "regulatory_version": "2026-v1",
                "source_category": "policy",
                "content": "只报送正常贷款余额。",
            }
        ],
        "script_basis": {
            "rules": [
                {
                    "rule_id": "script-edge-1",
                    "filter_condition": "LOAN_STATUS = 'NORMAL'",
                    "join_condition": None,
                    "aggregation_rule": "SUM(LOAN_BALANCE)",
                    "code_mapping_rule": None,
                    "transformation_expression": "LOAN_BALANCE",
                }
            ]
        },
    }
    candidate = {
        "business_definition": "",
        "processing_logic": "取贷款余额",
        "final_content": "报送所有逾期贷款余额，无需筛选正常状态。",
        "physical_references": [],
        "evidence_unit_ids": [11],
        "script_rule_ids": ["script-edge-1"],
        "policy_comparisons": [],
        "gaps": [],
        "claim_support_version": "claim-support-v1",
        "claims": [
            {
                "claim_id": "claim-1",
                "section": "business",
                "text": "报送所有逾期贷款余额，无需筛选正常状态。",
                "assertion_kind": "filter",
                "support": [
                    {
                        "source": "script_rule",
                        "ref_id": "script-edge-1",
                        "status": "supported",
                        "method": "deterministic_fact_match",
                        "method_version": "claim-support-v1",
                        "detail": "",
                        "ref_version": "",
                    }
                ],
            }
        ],
    }
    report = evaluate_claim_support(context, candidate)
    return {
        "case": "R2-after",
        "title": "合法引用配错误正文→判定冲突",
        "status": report.status,
        "conflict_count": report.conflict_count,
        "unknown_count": report.unknown_count,
        "codes": [item.code for item in report.findings],
        "violations": report.violations,
        "contradiction_detected": report.status == "conflict",
    }


def r3_after() -> dict:
    """R3 修复后：同一冻结内容重复渲染必须字节一致。"""
    from app.services.requirement_word import export_word_with_template

    content = {
        "requirement": {
            "name": "个人贷款余额报送需求",
            "content_version": 1,
            "target_table_id": 1,
            "field_ids": [],
            "document_ids": [],
            "source_table_ids": [],
            "mart_table_ids": [],
        },
        "target_table": {"table_name": "监管客户表", "table_code": "EAST_CUSTOMER"},
        "fields": [],
        "gaps": [],
        "formal_delivery": {"content_version": 1, "version_no": "v1", "review_record": []},
        "script_basis": {},
        "policy_comparisons": {},
        "manual_ownership": {},
    }
    first, _ = export_word_with_template(content, frozen_at="2026-10-23T00:00:00+08:00")
    time.sleep(1.2)
    second, _ = export_word_with_template(content, frozen_at="2026-10-23T00:00:00+08:00")
    return {
        "case": "R3-after",
        "title": "同快照重复渲染→字节一致",
        "first_sha256": sha256(first).hexdigest(),
        "second_sha256": sha256(second).hexdigest(),
        "byte_identical": first == second,
        "first_size": len(first),
        "second_size": len(second),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json", default="")
    args = parser.parse_args()
    results = [r1_formula_without_cache(), r2_authorized_but_wrong(), r3_word_bytes_unstable()]
    fixed = [r1_after(), r2_after(), r3_after()]
    payload = {
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "python": sys.version.split()[0],
        "cases": results,
        "after": fixed,
    }
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    print(text)
    if args.json:
        with open(args.json, "w", encoding="utf-8") as handle:
            handle.write(text + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
